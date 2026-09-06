"""G131 — Raporlama: şablon CRUD (`report_templates`), Excel/CSV export, koşu logu
(`report_runs`), çıktı saklama + temizlik, `/runs` + `/runs/{id}/download`.

Sözleşme: docs/plan/raporlama-plani-2026-09-06.md §2.4-2.5, §2.7.

Düzen (G130 `env` reçetesi): süreç içi sqlite (StaticPool), `routes.reports.SessionLocal`
sqlite fabrikasına bağlanır; kimlik `get_current_user` override'ı + `ADMIN_EMAILS` env'i
(GERÇEK `require_admin` koşar). Çıktı dizini `settings.rapor_cikti_dizini` → `tmp_path`.
"""
import csv
import datetime as dt
import hashlib
import io
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import models
from config.settings import Settings, settings
from database import _MIGRATIONS, Base
from schemas_rapor import KolonBasligi
from services.rapor import cikti, kosu_logu

ADMIN = "yonetici@hanyaloglu-acar.av.tr"
ADMIN2 = "ikinci.yonetici@hanyaloglu-acar.av.tr"
USER = "avukat@hanyaloglu-acar.av.tr"
T1 = "tenant-hanyaloglu"
T2 = "tenant-baska"

TEMPLATES = "/api/reports/templates"
EXPORT = "/api/reports/export"
RUNS = "/api/reports/runs"
PREVIEW = "/api/reports/preview"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

KOLONLAR = ["tracking_no", "subject", "opening_date", "created_at", "maddi_tazminat", "active"]


def _veri_yukle(db):
    db.add_all([
        models.Case(tracking_no="HA.G131.1", tenant_id=T1, status="DERDEST", subject="=CMD()|calc",
                    opening_date=dt.date(2025, 3, 1), active=True, maddi_tazminat=Decimal("1234.50"),
                    created_at=dt.datetime(2026, 1, 5, 10, 30)),
        models.Case(tracking_no="HA.G131.2", tenant_id=None, status="KARAR", subject="+artı ile başlar",
                    opening_date=dt.date(2024, 6, 15), active=False, maddi_tazminat=Decimal("-250"),
                    created_at=dt.datetime(2026, 1, 6, 9, 0)),
        models.Case(tracking_no="HA.G131.3", tenant_id=None, status="DERDEST", subject="Normal konu",
                    opening_date=None, active=True, created_at=dt.datetime(2026, 1, 7, 9, 0)),
        models.Case(tracking_no="HA.G131.4", tenant_id=T2, status="DERDEST", subject="Baska tenant",
                    created_at=dt.datetime(2026, 1, 8, 9, 0)),
        models.Case(tracking_no="HA.G131.5", tenant_id=T1, status="DERDEST", subject="Silinmis",
                    deleted_at=dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc), created_at=dt.datetime(2026, 1, 9)),
    ])
    db.commit()


@pytest.fixture()
def env(monkeypatch, tmp_path):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from dependencies import get_current_user
    from routes import reports as route_mod

    monkeypatch.setenv("ADMIN_EMAILS", f"{ADMIN},{ADMIN2}")
    cikti_dizini = tmp_path / "rapor_ciktilari"
    monkeypatch.setattr(settings, "rapor_cikti_dizini", str(cikti_dizini))
    monkeypatch.setattr(settings, "rapor_max_satir", 50_000)
    monkeypatch.setattr(settings, "rapor_cikti_saklama_gun", 30)

    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    maker = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    monkeypatch.setattr(route_mod, "SessionLocal", maker)

    db = maker()
    try:
        _veri_yukle(db)
    finally:
        db.close()

    def _client(email=ADMIN, tid=T1):
        app = FastAPI()
        app.include_router(route_mod.router)
        app.dependency_overrides[get_current_user] = lambda: {"preferred_username": email, "tid": tid}
        return TestClient(app, raise_server_exceptions=False)

    yield SimpleNamespace(db=maker, client=_client, dizin=cikti_dizini, tmp=tmp_path)
    engine.dispose()


def _tanim(kolonlar=KOLONLAR, filtreler=(), siralama=(("tracking_no", "asc"),), kaynak="davalar"):
    return {"veri_kaynagi": kaynak, "kolonlar": list(kolonlar), "filtreler": list(filtreler),
            "siralama": [{"alan": a, "yon": y} for a, y in siralama]}


def _sablon_govdesi(ad="Derdest davalar", paylasimli=False, **ek):
    govde = {"ad": ad, "aciklama": "test", "tanim": _tanim(), "paylasimli": paylasimli}
    govde.update(ek)
    return govde


def _export(client, format="xlsx", **ek):
    govde = {"tanim": _tanim(), "format": format, "kaynak": "manuel"}
    govde.update(ek)
    return client.post(EXPORT, json=govde)


def _runs(env):
    db = env.db()
    try:
        return db.query(models.ReportRun).order_by(models.ReportRun.id).all()
    finally:
        db.close()


def _sha(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


# ═══════════════════════════════════════════════════════════════════════════
# 1. Şablonlar
# ═══════════════════════════════════════════════════════════════════════════

def test_sablon_olustur_listele_guncelle_sil(env):
    """CRUD + sahiplik: PUT/DELETE yalnız `olusturan`; soft delete listeden düşer."""
    c = env.client()
    r = c.post(TEMPLATES, json=_sablon_govdesi())
    assert r.status_code == 201, r.text
    sablon = r.json()
    assert set(sablon) == {"id", "ad", "aciklama", "tanim", "olusturan", "paylasimli", "created_at", "updated_at"}
    assert sablon["olusturan"] == ADMIN and sablon["paylasimli"] is False
    assert sablon["tanim"]["kolonlar"] == KOLONLAR

    r = c.put(f"{TEMPLATES}/{sablon['id']}", json=_sablon_govdesi(ad="Yeni ad", paylasimli=True))
    assert r.status_code == 200, r.text
    assert r.json()["ad"] == "Yeni ad" and r.json()["paylasimli"] is True

    liste = c.get(TEMPLATES).json()
    assert [s["ad"] for s in liste] == ["Yeni ad"]

    assert c.delete(f"{TEMPLATES}/{sablon['id']}").status_code == 204
    assert c.get(TEMPLATES).json() == []
    db = env.db()
    try:
        satir = db.get(models.ReportTemplate, sablon["id"])
        assert satir is not None and satir.deleted_at is not None and satir.deleted_by == ADMIN
    finally:
        db.close()
    assert c.put(f"{TEMPLATES}/{sablon['id']}", json=_sablon_govdesi()).status_code == 404


def test_sablon_baskasinin_kaydi_403_ve_paylasimli_gorunurluk(env):
    """Kabul: başkasının şablonunu düzenleme/silme 403; paylaşımlı olmayan şablon başkasına listelenmez."""
    sahip, diger = env.client(ADMIN), env.client(ADMIN2)
    ozel = sahip.post(TEMPLATES, json=_sablon_govdesi(ad="Özel")).json()
    ortak = sahip.post(TEMPLATES, json=_sablon_govdesi(ad="Ortak", paylasimli=True)).json()
    digerin = diger.post(TEMPLATES, json=_sablon_govdesi(ad="Diğerinin")).json()

    assert [s["ad"] for s in sahip.get(TEMPLATES).json()] == ["Ortak", "Özel"]
    assert [s["ad"] for s in diger.get(TEMPLATES).json()] == ["Diğerinin", "Ortak"]

    for sid in (ozel["id"], ortak["id"]):
        assert diger.put(f"{TEMPLATES}/{sid}", json=_sablon_govdesi(ad="Ele geçir")).status_code == 403
        assert diger.delete(f"{TEMPLATES}/{sid}").status_code == 403
    assert sahip.delete(f"{TEMPLATES}/{digerin['id']}").status_code == 403
    # 403 sonrası kayıtlar değişmedi
    assert [s["ad"] for s in sahip.get(TEMPLATES).json()] == ["Ortak", "Özel"]


def test_sablon_gecersiz_tanim_422(env):
    c = env.client()
    r = c.post(TEMPLATES, json=_sablon_govdesi(tanim=_tanim(kolonlar=["yok_boyle_kolon"])))
    assert r.status_code == 422
    assert set(r.json()["detail"]) == {"alan", "sebep"}
    r = c.post(TEMPLATES, json={"ad": "", "tanim": _tanim()})
    assert r.status_code == 422 and r.json()["detail"]["alan"] == "ad"
    r = c.post(TEMPLATES, json=_sablon_govdesi(fazla="alan"))
    assert r.status_code == 422


def test_sablon_tenant_sizintisi_yok(env):
    """Başka tenant'ın paylaşımlı şablonu bile listelenmez; NULL tenant havuzu görünür."""
    env.client(ADMIN, tid=T2).post(TEMPLATES, json=_sablon_govdesi(ad="T2 ortak", paylasimli=True))
    db = env.db()
    try:
        db.add(models.ReportTemplate(ad="Havuz", tanim=_tanim(), olusturan=ADMIN2, paylasimli=True, tenant_id=None))
        db.commit()
    finally:
        db.close()
    assert [s["ad"] for s in env.client(ADMIN, tid=T1).get(TEMPLATES).json()] == ["Havuz"]


@pytest.mark.parametrize("metot,yol", [
    ("get", TEMPLATES), ("post", TEMPLATES), ("put", f"{TEMPLATES}/1"), ("delete", f"{TEMPLATES}/1"),
    ("post", EXPORT), ("get", RUNS), ("get", f"{RUNS}/1/download"),
])
def test_yonetici_olmayan_403(env, metot, yol):
    c = env.client(USER)
    r = getattr(c, metot)(yol, json={}) if metot in ("post", "put") else getattr(c, metot)(yol)
    assert r.status_code == 403 and r.json()["detail"] == "Yönetici yetkisi gerekli"


# ═══════════════════════════════════════════════════════════════════════════
# 2. Export — xlsx
# ═══════════════════════════════════════════════════════════════════════════

def test_export_xlsx_icerik_ve_kosu_satiri(env):
    """Kabul: koşu satırı tüm alanlarıyla dolu; indirilen baytların sha256'sı log ile eşit;
    xlsx geri okununca başlık/satır/tarih tipi doğru."""
    from openpyxl import load_workbook

    r = _export(env.client(), "xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(XLSX_MIME)
    assert 'attachment; filename="hukdok-rapor-manuel-' in r.headers["content-disposition"]
    assert r.headers["content-disposition"].endswith('.xlsx"')
    run_id = int(r.headers["x-rapor-kosu-id"])
    govde = r.content

    ws = load_workbook(io.BytesIO(govde))["Rapor"]
    basliklar = [c.value for c in ws[1]]
    assert basliklar[0] and len(basliklar) == len(KOLONLAR)
    satirlar = list(ws.iter_rows(min_row=2, values_only=False))
    assert [s[0].value for s in satirlar] == ["HA.G131.1", "HA.G131.2", "HA.G131.3"]      # T1 + NULL, silinmiş/T2 yok
    assert satirlar[0][1].value == "=CMD()|calc"                                          # xlsx'te önek yok
    assert isinstance(satirlar[0][2].value, (dt.date, dt.datetime))                       # opening_date gerçek tarih
    assert satirlar[0][2].number_format == "DD.MM.YYYY"
    assert isinstance(satirlar[0][3].value, dt.datetime) and satirlar[0][3].number_format == "DD.MM.YYYY HH:MM"
    assert satirlar[0][4].value == 1234.5 and satirlar[0][4].number_format == "#,##0.00"
    assert satirlar[0][5].value == "Evet" and satirlar[1][5].value == "Hayır"
    assert satirlar[2][2].value is None
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:F4"
    assert ws["A1"].fill.start_color.rgb.endswith("4A1530") and ws["A1"].font.bold

    kosular = _runs(env)
    assert len(kosular) == 1
    k = kosular[0]
    assert k.id == run_id and k.kullanici == ADMIN and k.tenant_id == T1
    assert k.baslangic is not None and k.format == "xlsx" and k.kaynak == "manuel" and k.veri_kaynagi == "davalar"
    assert k.satir_sayisi == 3 and k.kolon_sayisi == len(KOLONLAR)
    assert k.dosya_adi.startswith("hukdok-rapor-manuel-") and k.dosya_adi.endswith(".xlsx")
    assert k.dosya_boyutu == len(govde) and k.sha256 == _sha(govde)
    assert k.tanim == _tanim() and k.sablon_id is None and k.hata is None and k.sure_ms is not None
    assert k.dosya_yolu == str(env.dizin / f"{run_id}-{k.dosya_adi[:-5]}.xlsx")
    assert Path(k.dosya_yolu).read_bytes() == govde                                       # saklanan = indirilen


def test_export_csv_bom_ayrac_enjeksiyon(env):
    """CSV: utf-8-sig BOM + `;` + `= + - @` öneki (`'`); tarih GG.AA.YYYY; negatif sayı dokunulmaz."""
    r = _export(env.client(), "csv")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/csv")
    assert r.headers["content-disposition"].endswith('.csv"')
    govde = r.content
    assert govde.startswith(b"\xef\xbb\xbf")
    metin = govde.decode("utf-8-sig")
    satirlar = list(csv.reader(io.StringIO(metin), delimiter=";"))
    assert len(satirlar) == 4 and len(satirlar[0]) == len(KOLONLAR)
    assert satirlar[1][1] == "'=CMD()|calc"
    assert satirlar[2][1] == "'+artı ile başlar"
    assert satirlar[1][2] == "01.03.2025" and satirlar[1][3] == "05.01.2026 10:30:00"
    assert satirlar[2][4] == "-250.0"                        # sayı: önek YOK
    assert satirlar[1][5] == "Evet" and satirlar[3][2] == ""
    k = _runs(env)[0]
    assert k.format == "csv" and k.sha256 == _sha(govde) and k.dosya_boyutu == len(govde)
    assert k.dosya_yolu.endswith(".csv")


def test_export_413_satir_tavani_kosu_yazmaz(env, monkeypatch):
    monkeypatch.setattr(settings, "rapor_max_satir", 2)
    r = _export(env.client(), "xlsx")
    assert r.status_code == 413
    assert r.json()["detail"] == {"sebep": "satir_limiti", "toplam": 3, "limit": 2}
    assert _runs(env) == []
    assert not env.dizin.exists() or not any(env.dizin.iterdir())
    monkeypatch.setattr(settings, "rapor_max_satir", 3)
    assert _export(env.client(), "csv").status_code == 200


def test_export_422_bilinmeyen_kolon_ve_format(env):
    c = env.client()
    r = _export(c, "xlsx", tanim=_tanim(kolonlar=["tenant_id"]))
    assert r.status_code == 422 and r.json()["detail"]["alan"] == "kolonlar[0]"
    r = _export(c, "pdf")
    assert r.status_code == 422 and r.json()["detail"]["alan"] == "format"
    r = c.post(EXPORT, json={"tanim": _tanim(), "kaynak": "robot"})
    assert r.status_code == 422 and r.json()["detail"]["alan"] == "kaynak"
    assert _runs(env) == []


def test_export_sablon_id_ve_kaynak_asistan(env):
    c = env.client()
    sablon = c.post(TEMPLATES, json=_sablon_govdesi(ad="Favori")).json()
    r = _export(c, "csv", sablon_id=sablon["id"], kaynak="asistan")
    assert r.status_code == 200
    assert 'filename="hukdok-rapor-asistan-' in r.headers["content-disposition"]
    k = _runs(env)[0]
    assert k.sablon_id == sablon["id"] and k.kaynak == "asistan"
    kosu = c.get(RUNS).json()["kosular"][0]
    assert kosu["sablon_adi"] == "Favori" and kosu["kaynak"] == "asistan"
    assert _export(c, "csv", sablon_id=9999).status_code == 404


def test_onizleme_kosu_yazmaz(env):
    """Kabul: `/preview` koşu satırı YAZMAZ (K4)."""
    r = env.client().post(PREVIEW, json={"tanim": _tanim(), "sayfa": 1, "sayfa_boyu": 10})
    assert r.status_code == 200 and r.json()["toplam"] == 3
    assert _runs(env) == []
    assert not env.dizin.exists() or not any(env.dizin.iterdir())


def test_export_tenant_kisiti(env):
    """T2 kullanıcısı yalnız T2 + NULL satırları alır; koşu satırı T2 tenant'ıyla yazılır."""
    r = _export(env.client(ADMIN, tid=T2), "csv", tanim=_tanim(kolonlar=["tracking_no"]))
    satirlar = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert [s[0] for s in satirlar[1:]] == ["HA.G131.2", "HA.G131.3", "HA.G131.4"]
    assert _runs(env)[0].tenant_id == T2
    # T1'in /runs listesi T2 koşusunu görmez
    assert env.client(ADMIN, tid=T1).get(RUNS).json() == {"toplam": 0, "kosular": []}


# ═══════════════════════════════════════════════════════════════════════════
# 3. Koşular + indirme
# ═══════════════════════════════════════════════════════════════════════════

def test_runs_listesi_sekli_ve_sirasi(env):
    c = env.client()
    ilk = int(_export(c, "csv").headers["x-rapor-kosu-id"])
    ikinci = int(_export(env.client(ADMIN2), "xlsx").headers["x-rapor-kosu-id"])
    r = c.get(RUNS)
    assert r.status_code == 200
    govde = r.json()
    assert govde["toplam"] == 2
    assert [k["id"] for k in govde["kosular"]] == [ikinci, ilk]                # yeni → eski
    k = govde["kosular"][0]
    assert set(k) == {"id", "sablon_id", "sablon_adi", "format", "kaynak", "veri_kaynagi", "kolon_sayisi",
                      "satir_sayisi", "kullanici", "baslangic", "sure_ms", "dosya_adi", "dosya_boyutu", "sha256",
                      "dosya_mevcut", "hata", "tanim"}
    assert k["kullanici"] == ADMIN2 and k["dosya_mevcut"] is True and k["format"] == "xlsx"
    assert k["tanim"] == _tanim() and len(k["sha256"]) == 64
    assert c.get(RUNS, params={"limit": 1, "offset": 1}).json()["kosular"][0]["id"] == ilk
    assert c.get(RUNS, params={"limit": 0}).status_code == 422


def test_download_ayni_dosyayi_verir(env):
    """Kabul: `/runs/{id}/download` aynı dosyayı verir; sha256 eşit."""
    c = env.client()
    r = _export(c, "xlsx")
    run_id = int(r.headers["x-rapor-kosu-id"])
    d = env.client(ADMIN2).get(f"{RUNS}/{run_id}/download")             # başka yönetici de indirir
    assert d.status_code == 200
    assert d.content == r.content and _sha(d.content) == _runs(env)[0].sha256
    assert d.headers["content-type"].startswith(XLSX_MIME)
    assert d.headers["content-disposition"] == r.headers["content-disposition"]
    assert d.headers["x-rapor-kosu-id"] == str(run_id)
    assert c.get(f"{RUNS}/9999/download").status_code == 404
    assert env.client(ADMIN, tid=T2).get(f"{RUNS}/{run_id}/download").status_code == 404   # tenant


def test_download_traversal_reddi(env):
    """`dosya_yolu` dizin dışını ya da izinsiz uzantıyı gösterirse 404 — içerik sızmaz."""
    c = env.client()
    run_id = int(_export(c, "csv").headers["x-rapor-kosu-id"])
    disari = env.tmp / "gizli.csv"
    disari.write_text("sir", encoding="utf-8")
    icerde_txt = env.dizin / "not.txt"
    icerde_txt.write_text("sir", encoding="utf-8")
    kacis = env.dizin / ".." / "gizli.csv"
    for yol in (str(disari), str(icerde_txt), str(kacis)):
        db = env.db()
        try:
            run = db.get(models.ReportRun, run_id)
            run.dosya_yolu = yol
            db.commit()
        finally:
            db.close()
        d = c.get(f"{RUNS}/{run_id}/download")
        assert d.status_code == 404, yol
        assert b"sir" not in d.content
        assert c.get(RUNS).json()["kosular"][0]["dosya_mevcut"] is False


def test_temizlik_410_ve_dosya_mevcut_false(env):
    """Kabul: saklama süresi dolmuş dosyada 410; `dosya_mevcut=false`; satır (sha256 dahil) kalır."""
    c = env.client()
    r = _export(c, "xlsx")
    run_id = int(r.headers["x-rapor-kosu-id"])
    yol = Path(_runs(env)[0].dosya_yolu)
    assert yol.is_file()

    db = env.db()
    try:
        assert kosu_logu.temizle(db, simdi=dt.datetime.now(dt.timezone.utc) + dt.timedelta(days=29)) == 0
        assert yol.is_file()
        assert kosu_logu.temizle(db, simdi=dt.datetime.now(dt.timezone.utc) + dt.timedelta(days=31)) == 1
        assert kosu_logu.temizle(db, simdi=dt.datetime.now(dt.timezone.utc) + dt.timedelta(days=31)) == 0
    finally:
        db.close()
    assert not yol.exists()
    k = _runs(env)[0]
    assert k.dosya_yolu is None and k.sha256 == _sha(r.content) and k.satir_sayisi == 3
    assert c.get(f"{RUNS}/{run_id}/download").status_code == 410
    kosu = c.get(RUNS).json()["kosular"][0]
    assert kosu["dosya_mevcut"] is False and kosu["sha256"] == k.sha256


def test_temizlik_her_exportta_tembel_kosar(env):
    """Eski koşunun dosyası yeni bir export sırasında temizlenir; yeni dosya kalır."""
    c = env.client()
    eski_id = int(_export(c, "csv").headers["x-rapor-kosu-id"])
    eski_yol = Path(_runs(env)[0].dosya_yolu)
    db = env.db()
    try:
        run = db.get(models.ReportRun, eski_id)
        run.baslangic = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=31)
        db.commit()
    finally:
        db.close()
    yeni_id = int(_export(c, "xlsx").headers["x-rapor-kosu-id"])
    kosular = {k.id: k for k in _runs(env)}
    assert not eski_yol.exists() and kosular[eski_id].dosya_yolu is None
    assert Path(kosular[yeni_id].dosya_yolu).is_file()
    assert c.get(f"{RUNS}/{eski_id}/download").status_code == 410
    assert c.get(f"{RUNS}/{yeni_id}/download").status_code == 200


def test_temizlik_hatasi_exportu_dusurmez(env, monkeypatch, caplog):
    def _patla(db):
        raise RuntimeError("disk okunamadı")
    monkeypatch.setattr(kosu_logu, "temizle", _patla)
    with caplog.at_level("WARNING", logger="services.rapor.kosu_logu"):
        r = _export(env.client(), "csv")
    assert r.status_code == 200
    assert any("temizligi atlandi" in m for m in caplog.messages)
    assert all(rec.levelname != "ERROR" for rec in caplog.records)


def test_hata_yolunda_kosu_satiri_hata_ile_kalir(env, monkeypatch, caplog):
    """Kabul: üretim patlarsa 500; koşu satırı `hata` dolu kalır; yarım dosya silinir; TEK ERROR."""
    def _patla(format, kolonlar, satir_iter, hedef):
        hedef.write_bytes(b"yarim")
        raise RuntimeError("openpyxl patladı")
    monkeypatch.setattr(cikti, "ciktiyi_yaz", _patla)
    with caplog.at_level("WARNING"):
        r = _export(env.client(), "xlsx")
    assert r.status_code == 500 and r.json()["detail"] == "Rapor üretilemedi"
    kosular = _runs(env)
    assert len(kosular) == 1
    k = kosular[0]
    assert k.hata == "RuntimeError: openpyxl patladı" and k.dosya_yolu is None and k.sha256 is None
    assert k.kullanici == ADMIN and k.satir_sayisi == 3 and k.sure_ms is not None
    assert not any(env.dizin.glob("*"))
    assert sum(1 for rec in caplog.records if rec.levelname == "ERROR") == 1
    kosu = env.client().get(RUNS).json()["kosular"][0]
    assert kosu["hata"] == k.hata and kosu["dosya_mevcut"] is False
    assert env.client().get(f"{RUNS}/{k.id}/download").status_code == 410


# ═══════════════════════════════════════════════════════════════════════════
# 4. Servis birimleri
# ═══════════════════════════════════════════════════════════════════════════

@pytest.mark.parametrize("deger,beklenen", [
    ("=SUM(A1)", "'=SUM(A1)"), ("+1", "'+1"), ("-x", "'-x"), ("@cmd", "'@cmd"),
    ("normal", "normal"), ("", ""), (-5, -5), (1.5, 1.5), (None, None), (True, True),
])
def test_csv_hucresi_koru(deger, beklenen):
    assert cikti.csv_hucresi_koru(deger) == beklenen


def test_xlsx_ve_csv_uret_dosyaya_yazar(tmp_path):
    kolonlar = [KolonBasligi(anahtar="a", etiket="Ad", tip="metin"),
                KolonBasligi(anahtar="t", etiket="Tarih", tip="tarih"),
                KolonBasligi(anahtar="p", etiket="Tutar", tip="para")]
    satirlar = iter([{"a": "x", "t": "2025-01-02", "p": 10.0}, {"a": None, "t": None, "p": None}])
    ozet = cikti.ciktiyi_yaz("xlsx", kolonlar, satirlar, tmp_path / "a" / "r.xlsx")
    assert ozet.satir_sayisi == 2 and ozet.dosya_boyutu > 0 and len(ozet.sha256) == 64
    assert ozet.sha256 == _sha((tmp_path / "a" / "r.xlsx").read_bytes())
    ozet2 = cikti.ciktiyi_yaz("csv", kolonlar, iter([{"a": "y", "t": "2025-01-02T10:11:12", "p": 1}]),
                              tmp_path / "r.csv")
    assert ozet2.satir_sayisi == 1
    ham = (tmp_path / "r.csv").read_bytes()
    assert ham.startswith(b"\xef\xbb\xbf")
    assert ham.decode("utf-8-sig") == "Ad;Tarih;Tutar\r\ny;02.01.2025 10:11:12;1\r\n"
    with pytest.raises(ValueError):
        cikti.ciktiyi_yaz("pdf", kolonlar, iter([]), tmp_path / "r.pdf")


def test_cikti_yolu_ve_guvenlik(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "rapor_cikti_dizini", str(tmp_path / "c"))
    yol = kosu_logu.cikti_yolu(7, "hukdok-rapor-manuel-20260906-2310.xlsx")
    assert yol == tmp_path / "c" / "7-hukdok-rapor-manuel-20260906-2310.xlsx"
    assert (tmp_path / "c").is_dir()
    assert kosu_logu.cikti_yolu(8, "../../etc/passwd.csv").name == "8-passwd.csv"
    assert kosu_logu.cikti_yolu(9, "a b(1).xlsx").name == "9-a-b-1.xlsx"

    ic = tmp_path / "c" / "1-r.csv"
    ic.write_text("x")
    assert kosu_logu.yol_guvenli_mi(str(ic)) is True
    assert kosu_logu.yol_guvenli_mi(str(tmp_path / "c" / ".." / "1-r.csv")) is False
    assert kosu_logu.yol_guvenli_mi(str(tmp_path / "c" / "1-r.txt")) is False
    assert kosu_logu.yol_guvenli_mi(str(tmp_path / "1-r.csv")) is False
    assert kosu_logu.yol_guvenli_mi(None) is False and kosu_logu.yol_guvenli_mi("") is False


def test_cikti_dizini_varsayilani_backend_data(monkeypatch):
    monkeypatch.setattr(settings, "rapor_cikti_dizini", "")
    monkeypatch.setattr(Path, "mkdir", lambda self, *a, **k: None)
    beklenen = Path(kosu_logu.__file__).resolve().parent.parent.parent / "data" / "rapor_ciktilari"
    assert kosu_logu.cikti_dizini() == beklenen


# ═══════════════════════════════════════════════════════════════════════════
# 5. Env + migrasyon kuralı
# ═══════════════════════════════════════════════════════════════════════════

def test_settings_varsayilanlari_ve_env(monkeypatch):
    for ad in ("RAPOR_MAX_SATIR", "RAPOR_CIKTI_DIZINI", "RAPOR_CIKTI_SAKLAMA_GUN"):
        monkeypatch.delenv(ad, raising=False)
    s = Settings()
    assert (s.rapor_max_satir, s.rapor_cikti_dizini, s.rapor_cikti_saklama_gun) == (50_000, "", 30)
    monkeypatch.setenv("RAPOR_MAX_SATIR", "10")
    monkeypatch.setenv("RAPOR_CIKTI_DIZINI", "/app/data/rapor_ciktilari")
    monkeypatch.setenv("RAPOR_CIKTI_SAKLAMA_GUN", "bozuk")
    s = Settings()
    assert (s.rapor_max_satir, s.rapor_cikti_dizini, s.rapor_cikti_saklama_gun) == (10, "/app/data/rapor_ciktilari", 30)


def test_migrasyon_kurali_tablo_modelde_table_opu_yok():
    """Kabul: iki tablo modelde; `("table", ...)` op'u YAZILMADI (create_all yaratır, G041)."""
    assert models.ReportTemplate.__tablename__ == "report_templates"
    assert models.ReportRun.__tablename__ == "report_runs"
    for op in _MIGRATIONS:
        assert not (op[0] == "table" and op[1] in ("report_templates", "report_runs")), op[:2]
    fk = next(iter(models.ReportRun.__table__.c.sablon_id.foreign_keys))
    assert fk.column.table.name == "report_templates" and fk.ondelete == "SET NULL"
    # Modelde index YOK; FK index'i (G043 kuralı) koşulsuz ("index", ...) op'unda, IF NOT EXISTS'li
    assert not models.ReportRun.__table__.indexes and not models.ReportTemplate.__table__.indexes
    fk_indexleri = [
        ddl for op in _MIGRATIONS if op[0] == "index" and op[1] == "report_runs" for ddl in op[2]
        if "sablon_id" in ddl
    ]
    assert len(fk_indexleri) == 1 and "IF NOT EXISTS" in fk_indexleri[0]
    assert not [op for op in _MIGRATIONS if op[0] == "index" and op[1] == "report_templates"]


def test_export_akisi_yield_per_kullanir(env, monkeypatch):
    """Kabul (bellek): export motorun `yield_per` iteratörünü tüketir, listeye ALMAZ."""
    from services.rapor import motor
    cagrilar = []
    orijinal = motor.satirlari_akit

    def _izle(db, tanim, tenant_id, parca=None):
        it = orijinal(db, tanim, tenant_id, parca)
        cagrilar.append(it)
        return it
    monkeypatch.setattr(motor, "satirlari_akit", _izle)
    assert _export(env.client(), "xlsx").status_code == 200
    assert len(cagrilar) == 1 and hasattr(cagrilar[0], "__next__")
