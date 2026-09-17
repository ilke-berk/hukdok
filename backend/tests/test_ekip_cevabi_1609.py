"""Veri ekibinin 16.09 cevabındaki (Ek-5) kart düzeltmeleri script'i
(`scripts/ekip_cevabi_1609.py`): kart 14393 onarımı, klasör temizliği, 13.09 föy dönüşü,
mükerrer kart, kapatma, ayırma atlama, karar künyesi, son durum, ilişki.

**TEST VERİSİ KURALI:** ekibin gerçek Ek-5'i repoya girmez; sayfalar openpyxl ile
SENTETİK üretilir (aynı başlıklarla). Sabit tablolar prod kart id'leri taşıdığı için
testte monkeypatch ile bu DB'nin id'lerine çevrilir.
"""
from datetime import date

import openpyxl
import pytest

import models
from managers import case_manager, stage_decisions
from scripts import ekip_cevabi_1609 as ec
from tests.test_g064_aktarim_cekirdek import _kart
from tests.test_g179_ekip_cevabi_1209 import _belge, _foy, _taraf, db_env  # noqa: F401


@pytest.fixture()
def sabitsiz(monkeypatch):
    monkeypatch.setattr(ec, "CIFT_ISTISNALARI", {})
    monkeypatch.setattr(ec, "CIFT_YERINE_ILISKI", ())
    monkeypatch.setattr(ec, "AYIRMA_ATLA", {})
    monkeypatch.setattr(ec, "G179_KAPATMALAR", ())


def _ek5_yaz(yol, *, foy_donus=(), kapatma=(), ayirma=(), mukerrer=(), yeni_kart=(), ayni_dava=(),
             kunye=(), son_durum=(), deneme=(), arabuluculuk_metni=""):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ec.SAYFA_KUNYE
    ws.append(["Föy (SistemNo)", "Dosya No", "Kart", "Alan", "16.09 listenizde kartın bu alanı", "Cevabımız",
               "Neden", "Doğru yeri (bizdeki aşama zinciri)"])
    for foy, kart, zincir in kunye:
        ws.append([foy, "1.1", kart, "Yerel Karar No", "", "", "", zincir])

    ws = wb.create_sheet(ec.SAYFA_SON_DURUM)
    ws.append(["Kart", "Ofis dosya no", "Föy", "Sizdeki 'Dosya Son Durumu'", "Sizdeki esas"])
    for kart in son_durum:
        ws.append([kart, "", "", "İstinafda", ""])

    ws = wb.create_sheet(ec.SAYFA_AYNI_DAVA)
    ws.append(["Sınıf", "Mahkeme", "Esas", "Kart", "Ofis dosya no", "Ana Tür"])
    for sinif, esas, kart in ayni_dava:
        ws.append([sinif, "X Mahkemesi", esas, kart, "", ""])

    ws = wb.create_sheet(ec.SAYFA_MUKERRER)
    ws.append(["Kalması önerilen kart", "Künyesi", "Mükerrer kart", "Künyesi"])
    for kalan, muk in mukerrer:
        ws.append([kalan, "", muk, ""])

    ws = wb.create_sheet(ec.SAYFA_YENI_KART)
    ws.append(["Yeni kart", "Künyesi", "Bizde föyü olan kart", "Künyesi", "Bizim föyümüz"])
    for yeni, foylu in yeni_kart:
        ws.append([yeni, "", foylu, "", ""])

    ws = wb.create_sheet(ec.SAYFA_DENEME)
    ws.append(["Kart", "Ofis dosya no", "Mahkeme"])
    for kart in deneme:
        ws.append([kart, "", ""])

    ws = wb.create_sheet(ec.SAYFA_ACIK)
    ws.append(["Konu", "Bizdeki karşılığı", "Durum"])
    ws.append(["Arabuluculuk föyü ile dava föyü aynı kartta", arabuluculuk_metni, ""])

    ws = wb.create_sheet(ec.SAYFA_1309)
    ws.append(["Sınıf", "Kart", "Ofis dosya no", "13.09 listenizde", "16.09 listenizde", "Notumuz"])
    for foy, kart in foy_donus:
        ws.append(["Düzenleme canlıda görünmüyor", kart, "", f"föy {foy} · klasör [1.1]", "föy yok", ""])
    for kart in kapatma:
        ws.append(["Kapatma canlıda görünmüyor", kart, "", "kart listede yok", "kart açık", ""])
    for kart in ayirma:
        ws.append(["Ayırma canlıda görünmüyor", kart, "", "ayırma bildirildi", "", ""])
    wb.save(yol)
    return yol


def _kos(db_env, tmp_path, *, apply=False, kart_14393=False, **sayfalar):  # noqa: F811
    return ec.kos(db_env, ek5=_ek5_yaz(tmp_path / "ek5.xlsx", **sayfalar), apply=apply, kim="test",
                  kart_14393=kart_14393)


# ─── okuma ───────────────────────────────────────────────────────────────────

def test_ek5_oku_siniflari_ayirir(tmp_path):
    yol = _ek5_yaz(tmp_path / "ek5.xlsx", foy_donus=[("H-16708", 4674)], kapatma=[14335], ayirma=[2468],
                   mukerrer=[(5051, 5626)], yeni_kart=[(14395, 13161)], deneme=[14335, 14411],
                   ayni_dava=[("A · Aynı tür", "2019/83", 2668), ("A · Aynı tür", "2026/646;2019/83", 2669),
                              ("B · Arabuluculuk", "2023/503", 4235), ("B · Arabuluculuk", "2023/503", 13773),
                              ("D · Tür farklı", "2023/2393", 13644), ("D · Tür farklı", "2023/2393", 13645)],
                   kunye=[("H-14157", 13478, "Yerel Konya · esas 2020/284 · karar 2023/386 · 2023-07-05  →  "
                                             "İstinaf · esas 2023/1590 · karar 2025/12"),
                          ("i-12288", 9630, "")],
                   arabuluculuk_metni="kart 2468 (ARB-1 + H-2) · kart 4370 (ARB-3 + H-4)")
    veri = ec.ek5_oku(yol)
    assert veri["foy_donus"] == [("H-16708", 4674)]
    assert [k for k, _ in veri["kapatma"]] == [14335, 14411]          # 07'deki tekrar eklenmez
    assert veri["ayirma"] == [2468, 4370]
    assert [(a, b) for a, b, _ in veri["cift"]] == [(5051, 5626), (13161, 14395), ec.YOZGAT]
    assert [(a, b) for a, b, _ in veri["iliski"]] == [(4235, 13773)]  # D sınıfına dokunulmaz
    assert veri["kunye"] == [(13478, "H-14157", "2020/284", "2023/386")]


# ─── föy dönüşü ──────────────────────────────────────────────────────────────

def test_foy_donusu_tek_foylu_kaynak_kart_birlesir_bolunen_kaynak_kapanir(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        eski = _kart(db, "S4.QUICK......0469.HUKUK.00000", "2.500", esas_no="2026/1")
        yeni = _kart(db, "S4.QUICK......0590.HUKUK.00000", "2.500.00", esas_no="2026/1")
        for k in (eski, yeni):
            _taraf(db, k, "Quick Sigorta A.Ş.")
        _foy(db, yeni, "H-1")
        _belge(db, yeni, "dilekce.pdf")
        hukuk = _kart(db, "S4.QUICK......0494.HUKUK.00000", "2.506")
        arabu = _kart(db, "S4.QUICK......0492.ARABU.00000", "2.506")
        ortak = _kart(db, "S4.QUICK......0600.HUKUK.00000", "2.506.00")
        _foy(db, ortak, "H-2")
        _foy(db, ortak, "ARB-2")
        db.commit()
        ids = eski.id, yeni.id, hukuk.id, arabu.id, ortak.id
    finally:
        db.close()
    eski_id, yeni_id, hukuk_id, arabu_id, ortak_id = ids

    sonuc = _kos(db_env, tmp_path, apply=True,
                 foy_donus=[("H-1", eski_id), ("H-2", hukuk_id), ("ARB-2", arabu_id)])
    assert sonuc.sayim("foy", "RET") == 0

    db = db_env()
    try:
        assert db.query(models.CaseFoy).filter_by(sistem_no="H-1").one().case_id == eski_id
        assert db.get(models.Case, yeni_id).deleted_at is not None                 # birleşti
        assert db.query(models.CaseDocument).one().case_id == eski_id              # belge taşındı
        assert db.query(models.CaseFoy).filter_by(sistem_no="H-2").one().case_id == hukuk_id
        assert db.query(models.CaseFoy).filter_by(sistem_no="ARB-2").one().case_id == arabu_id
        assert db.get(models.Case, ortak_id).deleted_at is not None                # föysüz + belgesiz
    finally:
        db.close()

    ikinci = _kos(db_env, tmp_path, apply=True,
                  foy_donus=[("H-1", eski_id), ("H-2", hukuk_id), ("ARB-2", arabu_id)])
    assert ikinci.sayim("foy", "YAPILDI") == 0 and ikinci.sayim("foy", "ATLANDI") == 3


def test_foy_donusu_belgeli_kaynak_kart_kapanmaz(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        a = _kart(db, "D1.A..........0001.HUKUK.00000", "1.1")
        b = _kart(db, "D1.A..........0002.ARABU.00000", "1.1")
        ortak = _kart(db, "D1.A..........0003.HUKUK.00000", "1.1.00")
        _foy(db, ortak, "H-9")
        _foy(db, ortak, "ARB-9")
        _belge(db, ortak, "x.pdf")
        db.commit()
        a_id, b_id, ortak_id = a.id, b.id, ortak.id
    finally:
        db.close()

    sonuc = _kos(db_env, tmp_path, apply=True, foy_donus=[("H-9", a_id), ("ARB-9", b_id)])
    assert sonuc.sayim("foy", "RET") == 1
    db = db_env()
    try:
        assert db.get(models.Case, ortak_id).deleted_at is None
    finally:
        db.close()


# ─── mükerrer · kapatma · ayırma ─────────────────────────────────────────────

def test_cift_istisnasi_yonu_cevirir_yerine_iliski_birlestirmez(db_env, tmp_path, monkeypatch):  # noqa: F811
    db = db_env()
    try:
        oneri = _kart(db, "S7.SOMPO......0045.HUKUK.00000", "6.6036", esas_no="2020/911", court="Bakırköy 4. ATM")
        guncel = _kart(db, "S7.SOMPO......0046.HUKUK.00000", "6.6036", esas_no="2020/911;2026/639",
                       court="İstanbul 4. ATM")
        axa = _kart(db, "S3.C_KOSE.....0002.IDARE.00000", "", esas_no="2022/1034")
        nippon = _kart(db, "S0.T_AS.......0001.IDARI.10000", "", esas_no="2022/1034")
        _taraf(db, axa, "Axa Sigorta A.Ş.")
        _taraf(db, nippon, "Türk Nippon Sigorta A.Ş.")
        for k in (oneri, guncel):
            _taraf(db, k, "Sompo Sigorta A.Ş.")
        db.commit()
        oneri_id, guncel_id, axa_id, nippon_id = oneri.id, guncel.id, axa.id, nippon.id
    finally:
        db.close()
    monkeypatch.setattr(ec, "CIFT_ISTISNALARI", {(guncel_id, oneri_id): {"mahkeme_kontrolu": False,
                                                                        "esas_kontrolu": False}})
    monkeypatch.setattr(ec, "CIFT_YERINE_ILISKI", ((axa_id, nippon_id, "ayrı müvekkil"),))
    monkeypatch.setattr(ec, "AYIRMA_ATLA", {})
    monkeypatch.setattr(ec, "G179_KAPATMALAR", ())

    sonuc = _kos(db_env, tmp_path, apply=True, mukerrer=[(oneri_id, guncel_id)], yeni_kart=[(nippon_id, axa_id)],
                 kapatma=[guncel_id])
    assert sonuc.sayim("cift", "YAPILDI") == 1 and sonuc.sayim("cift", "ATLANDI") == 1
    assert sonuc.sayim("kapatma", "ATLANDI") == 1                    # birleşmenin kalanı kapatılmaz
    db = db_env()
    try:
        assert db.get(models.Case, oneri_id).deleted_at is not None
        assert db.get(models.Case, guncel_id).deleted_at is None
        assert db.get(models.Case, nippon_id).deleted_at is None
        assert db.query(models.CaseRelation).count() == 1
    finally:
        db.close()


def test_kapatma_foylu_karti_reddeder_ayirma_kapali_karti_atlar(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        deneme = _kart(db, "X1.I_KUTLUK...0005.HUKUK.00000", "", esas_no="9/9")
        _belge(db, deneme, "test.pdf")
        foylu = _kart(db, "HK-1", "")
        _foy(db, foylu, "H-5")
        db.commit()
        deneme_id, foylu_id = deneme.id, foylu.id
    finally:
        db.close()

    sonuc = _kos(db_env, tmp_path, apply=True, kapatma=[foylu_id], deneme=[deneme_id], ayirma=[deneme_id])
    assert sonuc.sayim("kapatma", "YAPILDI") == 1 and sonuc.sayim("kapatma", "RET") == 1
    assert sonuc.sayim("ayirma", "ATLANDI") == 1 and sonuc.sayim("ayirma", "RET") == 0


# ─── karar künyesi · son durum ───────────────────────────────────────────────

def test_kunye_eski_turun_karari_guncel_esasta_ise_duzeltilir(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        kart = _kart(db, "D1.N_DEMIREL..0001.HUKUK.00000", "1386.001.00", court="Konya 2. Tüketici Mahkemesi")
        case_manager.sync_current_esas(db, kart, "2025/69", court=kart.court)
        stage_decisions.add_stage_decision(db, kart, stage="YEREL", mahkeme=kart.court, esas_no="2025/69",
                                           karar_no="2023/386", karar_tarihi=date(2023, 7, 5))
        _foy(db, kart, "H-14157")
        db.commit()
        kart_id = kart.id
        assert kart.karar_no == "2023/386"
    finally:
        db.close()
    zincir = "Yerel Konya · esas 2020/284 · karar 2023/386 · 2023-07-05  →  İstinaf"

    sonuc = _kos(db_env, tmp_path, apply=True, kunye=[("H-14157", kart_id, zincir)], son_durum=[kart_id])
    assert sonuc.sayim("kunye", "YAPILDI") == 1
    db = db_env()
    try:
        kart = db.get(models.Case, kart_id)
        assert kart.karar_no is None and kart.karar_tarihi is None
        satirlar = (db.query(models.CaseStageDecision).filter_by(case_id=kart_id)
                    .order_by(models.CaseStageDecision.sira_no).all())
        assert [(s.esas_no, s.karar_no) for s in satirlar] == [("2020/284", "2023/386"), ("2025/69", None)]
        assert "2020/284" in {e.esas_no for e in db.query(models.CaseEsasNumber).filter_by(case_id=kart_id)}
    finally:
        db.close()

    ikinci = _kos(db_env, tmp_path, apply=True, kunye=[("H-14157", kart_id, zincir)])
    assert ikinci.sayim("kunye", "YAPILDI") == 0 and ikinci.sayim("kunye", "ATLANDI") == 1


def test_son_durum_yalniz_istinafda_ise_degisir(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        a = _kart(db, "S2.A..0833", "", dosya_son_durumu="İstinafda")
        b = _kart(db, "S2.A..0834", "", dosya_son_durumu="Bilirkişide")
        db.commit()
        a_id, b_id = a.id, b.id
    finally:
        db.close()
    sonuc = _kos(db_env, tmp_path, apply=True, son_durum=[a_id, b_id])
    assert sonuc.sayim("son_durum", "YAPILDI") == 1 and sonuc.sayim("son_durum", "ATLANDI") == 1
    db = db_env()
    try:
        assert db.get(models.Case, a_id).dosya_son_durumu == ec.SON_DURUM_YENI
    finally:
        db.close()


# ─── kart 14393 ──────────────────────────────────────────────────────────────

def test_kart_14393_paketin_yazdiklari_geri_alinir_foy_dogru_karta(db_env, tmp_path, sabitsiz, monkeypatch):  # noqa: F811
    db = db_env()
    try:
        dogru = _kart(db, "S3.AXA........0803.HUKUK.00000", "1.20788.00;27.001.00")
        kart = _kart(db, "S4.QUICK......0504.HUKUK.00100", "2.554.00", court="İstanbul Anadolu 6. Tüketici",
                     manevi_tazminat=2000000)
        kart.status = "MAHZEN"
        case_manager.sync_current_esas(db, kart, "2015/1367", court=kart.court)
        korunan = _taraf(db, kart, "Quick Sigorta A.Ş.")
        eklenen = _taraf(db, kart, "Corpus Sigorta Anonim Şirketi")
        eklenen.role = "Davalı"
        _belge(db, kart, "gercek.pdf", korunan)
        db.add(models.CaseLawyer(case_id=kart.id, name="Reyhan Duygun"))
        _foy(db, kart, "H-5441", eklenen)
        paket = "HUKDOK_TESLIM_HUKDOK_TESLIM_PAKETI_2026-09-15.xlsx"
        for alan, eski, yeni in (("court", "Kayseri 3. Tüketici", "İstanbul Anadolu 6. Tüketici"),
                                 ("esas_no", "2026/371", "2015/1367"), ("status", "DERDEST", "MAHZEN"),
                                 ("manevi_tazminat", "20000000.00", "2000000"),
                                 ("taraf", None, "Corpus Sigorta Anonim Şirketi (Davalı)"),
                                 ("avukat", None, "Reyhan Duygun")):
            db.add(models.CaseHistory(case_id=kart.id, field_name=alan, old_value=eski, new_value=yeni,
                                      source=paket))
        cikar = _kart(db, "S4.CORPUS.....0002.HUKUK.00000", "27.002.00;2.555.00")
        db.commit()
        dogru_id, kart_id, cikar_id = dogru.id, kart.id, cikar.id
    finally:
        db.close()
    monkeypatch.setattr(ec, "KART_14393", kart_id)
    monkeypatch.setattr(ec, "KART_14393_FOY", ("H-5441", dogru_id))
    monkeypatch.setattr(ec, "KLASOR_CIKAR", ((cikar_id, "2.555.00", "27.002.00"),))

    sonuc = _kos(db_env, tmp_path, apply=True, kart_14393=True)
    assert sonuc.sayim("k14393", "YAPILDI") == 1 and sonuc.sayim("klasor", "YAPILDI") == 1
    db = db_env()
    try:
        kart = db.get(models.Case, kart_id)
        assert (kart.court, kart.esas_no, kart.status) == ("Kayseri 3. Tüketici", "2026/371", "DERDEST")
        assert kart.manevi_tazminat == 20000000
        assert kart.klasor_no_2 == "2.554.00"
        assert [p.name for p in kart.parties] == ["Quick Sigorta A.Ş."]
        assert kart.lawyers == [] and len(kart.documents) == 1
        assert {e.esas_no for e in db.query(models.CaseEsasNumber).filter_by(case_id=kart_id)} == {"2026/371"}
        assert db.query(models.CaseFoy).filter_by(sistem_no="H-5441").one().case_id == dogru_id
        assert db.get(models.Case, cikar_id).klasor_no_2 == "27.002.00"
    finally:
        db.close()

    ikinci = _kos(db_env, tmp_path, apply=True, kart_14393=True)
    assert ikinci.sayim("k14393", "ATLANDI") == 1 and ikinci.sayim("klasor", "ATLANDI") == 1
