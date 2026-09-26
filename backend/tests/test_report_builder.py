"""G210 — takvim raporu: duruşma satırının Açıklama'sı davanın son durumu.

Kilitlenen davranışlar:
  1. davanın `dosya_son_durumu` dolu (strip sonrası boş değil) → `title` = o değer,
  2. boş/NULL/yalnız boşluk → eski davranış: `note`, o da yoksa "Duruşma",
  3. elle eklenen işaretler (`CalendarEvent`) değişmez,
  4. sütun başlıkları ve satır alan adları değişmez; Excel'de Açıklama hücresi son durumdur.

Testler süreç içi sqlite (StaticPool) üzerinde GERÇEK sorgu koşar.
"""
import io
from datetime import date

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import models
from database import Base

TENANT = "tenant-x"
START = date(2026, 10, 1)
END = date(2026, 10, 31)


@pytest.fixture()
def db():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def _dava(db, tracking_no, son_durum=None):
    case = models.Case(
        tracking_no=tracking_no,
        esas_no="2026/10",
        court="İstanbul 1. Asliye Hukuk",
        responsible_lawyer_name="Av. Deneme",
        dosya_son_durumu=son_durum,
    )
    db.add(case)
    db.flush()
    db.add(models.CaseParty(case_id=case.id, name="Müvekkil A", role="Davacı", party_type="CLIENT"))
    db.add(models.CaseParty(case_id=case.id, name="Karşı B", role="Davalı", party_type="COUNTER"))
    return case


def _durusma(db, case, gun, note=None, saat="10:00"):
    db.add(models.HearingDate(case_id=case.id, hearing_date=date(2026, 10, gun), hearing_time=saat, note=note))


def _satirlar(db):
    from report_builder import build_report_rows

    db.commit()
    return build_report_rows(db, TENANT, START, END)


def _tek_durusma(db, son_durum, note):
    case = _dava(db, "T-1", son_durum)
    _durusma(db, case, 5, note=note)
    rows = _satirlar(db)
    assert len(rows) == 1
    assert rows[0]["type"] == "Duruşma"
    return rows[0]


def test_son_durumlu_durusma_aciklamasi_son_durum(db):
    row = _tek_durusma(db, "  Bilirkişi raporu bekleniyor  ", "Ön inceleme")
    assert row["title"] == "Bilirkişi raporu bekleniyor"
    # Diğer alanlar dokunulmadan gelir
    assert row["esas_no"] == "2026/10"
    assert row["client"] == "Müvekkil A"
    assert row["counter"] == "Karşı B"


def test_son_durumsuz_notlu_durusma_eski_davranis(db):
    row = _tek_durusma(db, None, "Ön inceleme")
    assert row["title"] == "Ön inceleme"


def test_son_durumu_bosluk_olan_durusma_nota_duser(db):
    row = _tek_durusma(db, "   ", "Tanık dinlenecek")
    assert row["title"] == "Tanık dinlenecek"


def test_son_durumsuz_notsuz_durusma_durusma_yazar(db):
    row = _tek_durusma(db, "", None)
    assert row["title"] == "Duruşma"


def test_elle_isaret_degismez(db):
    # Aynı aralıkta son durumlu bir dava varken bile işaretin başlığı kendi başlığıdır
    case = _dava(db, "T-2", "Karar bekleniyor")
    _durusma(db, case, 7, note=None, saat="09:00")
    db.add(models.CalendarEvent(tenant_id=TENANT, title="Müvekkil toplantısı", event_date=date(2026, 10, 7), event_time="14:30"))
    rows = _satirlar(db)
    assert [r["type"] for r in rows] == ["Duruşma", "İşaret"]
    isaret = rows[1]
    assert isaret == {
        "date": date(2026, 10, 7),
        "date_str": "07.10.2026",
        "time": "14:30",
        "type": "İşaret",
        "title": "Müvekkil toplantısı",
        "esas_no": "",
        "court": "",
        "client": "",
        "counter": "",
        "lawyer": "",
        "case_id": None,
    }
    assert rows[0]["title"] == "Karar bekleniyor"


def test_sutun_basliklari_ve_alan_adlari_degismez(db):
    from report_builder import COLUMNS

    assert COLUMNS == [
        "Tarih", "Saat", "Tür", "Açıklama", "Esas No",
        "Mahkeme", "Müvekkil", "Karşı Taraf", "Sorumlu Avukat",
    ]
    row = _tek_durusma(db, "Keşif yapılacak", None)
    assert set(row) == {
        "date", "date_str", "time", "type", "title", "esas_no",
        "court", "client", "counter", "lawyer", "case_id",
    }


def test_excel_aciklama_hucresi_son_durum(db):
    from openpyxl import load_workbook

    from report_builder import COLUMNS, rows_to_excel

    case = _dava(db, "T-3", "İstinaf incelemesinde")
    _durusma(db, case, 12, note="Eski not")
    rows = _satirlar(db)
    data = rows_to_excel(rows, START, END)

    ws = load_workbook(io.BytesIO(data)).active
    header_row = 3
    basliklar = [ws.cell(row=header_row, column=i).value for i in range(1, len(COLUMNS) + 1)]
    assert basliklar == COLUMNS
    aciklama_col = COLUMNS.index("Açıklama") + 1
    assert ws.cell(row=header_row + 1, column=aciklama_col).value == "İstinaf incelemesinde"
