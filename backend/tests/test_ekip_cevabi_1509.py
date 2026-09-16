"""Veri ekibinin 15.09 cevabındaki (Ek-4) kart düzeltmeleri script'i
(`scripts/ekip_cevabi_1509.py`): konu, sabit künye düzeltmesi, aynı dava iki kart,
eski esaslı kartın kapatılması, ana tür, klasör öneki, ilişkili dosya bağı, kapatma.

**TEST VERİSİ KURALI:** ekibin gerçek Ek-4'ü repoya girmez; sayfalar openpyxl ile
SENTETİK üretilir (aynı başlıklarla, araya ilgisiz sütunlar konarak).

Asıl yeni davranış: Ek-4'ün kart numaraları 13.09'da gönderdiğimiz listeden geliyor,
o liste yerel kopyadan üretilmişti — id tutmazsa kart satırdaki FÖY NUMARASI ile
çözülür (`_karti_coz`).
"""
import openpyxl
import pytest

import models
from managers import case_manager
from scripts import ekip_cevabi_1509 as ec
from tests.test_g064_aktarim_cekirdek import _kart
from tests.test_g179_ekip_cevabi_1209 import _belge, _foy, _taraf, db_env  # noqa: F401


@pytest.fixture()
def sabitsiz(monkeypatch):
    """Sabit tablolar (4181/2223, 14321↔14322, 14336) bu testlerin DB'sinde yok."""
    monkeypatch.setattr(ec, "SABIT_DUZELTMELER", ())
    monkeypatch.setattr(ec, "SABIT_ILISKILER", ())
    monkeypatch.setattr(ec, "KAPATILACAK_KARTLAR", ())


def _ek4_yaz(yol, *, konu=(), cift=(), eski_esas=(), tur=(), klasor=(), iliski=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ec.SAYFA_KONU
    ws.append(["Ofis dosya no", "Kart", "Konu (listenizde)", "Büro özel türü", "Müvekkil",
               "Durum", "Önerdiğimiz konu", "Not", "Cevabınız"])
    for kart, deger in konu:
        ws.append(["S3.AXA...", kart, "İtirazın İptali", "Rücu", "AXA", "DERDEST", deger, "", ""])

    ws = wb.create_sheet(ec.SAYFA_CIFT)
    ws.append(["Föysüz kart", "Ofis dosya no", "Klasör no", "Föyün bağlı olduğu kart", "Ofis dosya no ",
               "Klasör no ", "Föy (bizde)", "Mahkeme", "Esas", "Nasıl oluştu", "Önerimiz", "Cevabınız"])
    for eski, yeni, foy in cift:
        ws.append([eski, "", "", yeni, "", "", foy, "X Mahkemesi", "2020/1", "", "", ""])

    ws = wb.create_sheet(ec.SAYFA_ESKI_ESAS)
    ws.append(["Eski esaslı kart", "Ofis dosya no", "Mahkeme", "Esas", "Föyü", "Güncel kart",
               "Ofis dosya no ", "Esas ", "Föy (bizde)", "Bizde eski esas", "Ortak klasör no", "Cevabınız"])
    for eski, guncel, foy in eski_esas:
        ws.append([eski, "", "X Mahkemesi", "2017/1", "föy yok", guncel, "", "2023/1", foy, "", "", ""])

    ws = wb.create_sheet(ec.SAYFA_TUR)
    ws.append(["Kart", "Ofis dosya no", "Ana Tür (listenizde)", "Mahkeme", "Bizde", "Not", "Cevabınız"])
    for kart, bizde in tur:
        ws.append([kart, "D1.X...", "Hukuk", "X Mahkemesi", bizde, "", ""])

    ws = wb.create_sheet(ec.SAYFA_KLASOR)
    ws.append(["Kart", "Ofis dosya no", "Klasör listesi (sizde)", "Föy · DosyaNo (bizde)",
               "Önerimiz", "Cevabınız"])
    for kart, sizde, bizde in klasor:
        ws.append([kart, "S3.AXA...", sizde, bizde, "", ""])

    ws = wb.create_sheet(ec.SAYFA_ILISKI)
    ws.append(["Klasör no", "Kart 1", "Ofis dosya no", "Tür", "Mahkeme", "Esas", "Föy (bizde)",
               "Kart 2", "Ofis dosya no ", "Tür ", "Mahkeme ", "Esas ", "Föy (bizde) ",
               "Aynı TKU", "Önerimiz", "Cevabınız"])
    for a, b, foy_a, foy_b in iliski:
        ws.append(["1.001", a, "", "Hukuk", "", "", foy_a, b, "", "Arabuluculuk", "", "", foy_b, "", "", ""])
    wb.save(yol)
    return yol


def _kos(db_env, tmp_path, *, apply=False, **sayfalar):  # noqa: F811
    return ec.kos(db_env, ek4=_ek4_yaz(tmp_path / "ek4.xlsx", **sayfalar), apply=apply, kim="test")


def _sayim(db_env, model):  # noqa: F811
    db = db_env()
    try:
        return db.query(model).count()
    finally:
        db.close()


# ─── Ek-4 okuma ──────────────────────────────────────────────────────────────

def test_ek4_oku_basliga_gore_iki_foy_sutunu(tmp_path):
    """06'da 'Föy (bizde)' iki kez geçer: birincisi Kart 1'in, ikincisi Kart 2'nin."""
    yol = _ek4_yaz(tmp_path / "ek4.xlsx", konu=[(5336, "Rücuen Alacak")],
                   cift=[(1763, 15336, "H-1557")],
                   iliski=[(13767, 14105, "ARB-1", "H-964 ; H-965")])
    veri = ec.ek4_oku(yol)
    assert veri["konu"] == [(5336, "Rücuen Alacak")]
    assert veri["cift"] == [((1763, ""), (15336, "H-1557"),
                             "Ek-4 › 03: aynı dava iki kartta (DosyaNo + mahkeme + esas aynı)")]
    (a, b, _kanit), = veri["iliski"]
    assert a == (13767, "ARB-1") and b == (14105, "H-964 ; H-965")


def test_ek4_sayfa_yoksa_hata(tmp_path):
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "bos.xlsx")
    with pytest.raises(ValueError, match="sayfası yok"):
        ec.ek4_oku(tmp_path / "bos.xlsx")


@pytest.mark.parametrize("bizde, mevcut, beklenen", [
    ("H-15774 1.21458.00", {"21458"}, [("21458", "1.21458.00")]),
    ("H-15774 1.21458.00", {"1.21458.00"}, []),                 # zaten önekli
    ("id-15819 1.21461.00 · id-15831 1851.001", {"21461", "1851.001"}, [("21461", "1.21461.00")]),
    ("(DosyaNo yok)", {"21476"}, []),                           # numara yok
])
def test_klasor_ciftleri(bizde, mevcut, beklenen):
    assert ec._klasor_ciftleri(bizde, mevcut) == beklenen


# ─── Adım 1: konu ────────────────────────────────────────────────────────────

def test_konu_listeden_kanonik_yazilir_liste_disi_ret(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        db.add(models.CaseSubject(code="RUCUEN_ALACAK", name="Rücuen Alacak"))
        kart = _kart(db, "S3.AXA........0711.HUKUK.00000", "1.1", subject="İtirazın İptali")
        yok = _kart(db, "S3.AXA........0712.HUKUK.00000", "1.2", subject="İtirazın İptali")
        db.commit()
        kart_id, yok_id = kart.id, yok.id
    finally:
        db.close()

    sonuc = _kos(db_env, tmp_path, apply=True,
                 konu=[(kart_id, "rücuen alacak"), (yok_id, "Bilinmeyen Konu")])
    assert sonuc.sayim("konu", "YAPILDI") == 1 and sonuc.sayim("konu", "RET") == 1

    db = db_env()
    try:
        assert db.get(models.Case, kart_id).subject == "Rücuen Alacak"        # kanonik yazım
        assert db.get(models.Case, yok_id).subject == "İtirazın İptali"       # dokunulmadı
    finally:
        db.close()

    ikinci = _kos(db_env, tmp_path, apply=True, konu=[(kart_id, "Rücuen Alacak")])
    assert ikinci.sayim("konu", "YAPILDI") == 0 and ikinci.sayim("konu", "ATLANDI") == 1


# ─── Adım 3: aynı dava iki kart ──────────────────────────────────────────────

def test_cift_id_tutmazsa_foy_ile_cozulur_eski_kart_kalir(db_env, tmp_path, sabitsiz):  # noqa: F811
    """Ek-4'ün ikinci kart numarası yerel kopyadan; canlıda o id yok → föy ile bulunur."""
    db = db_env()
    try:
        eski = _kart(db, "S4.CORPUS.....0001.HUKUK.00000", "27.004.00", file_type="Hukuk",
                     esas_no="2017/218", court="İstanbul Anadolu 15. Asliye Hukuk Mahkemesi")
        yeni = _kart(db, "S4.CORPUS.....0005.HUKUK.00000", "2.557.00", file_type="Hukuk",
                     esas_no="2017/218", court="İstanbul Anadolu 15. Asliye Hukuk Mahkemesi")
        _taraf(db, eski, "Corpus Sigorta A.Ş.")
        _taraf(db, yeni, "Corpus Sigorta A.Ş.")
        _foy(db, yeni, "H-1557")
        db.commit()
        eski_id, yeni_id = eski.id, yeni.id
    finally:
        db.close()

    sonuc = _kos(db_env, tmp_path, apply=True, cift=[(eski_id, 999999, "H-1557")])
    assert sonuc.sayim("cift", "YAPILDI") == 1
    assert "föy H-1557" in sonuc.kalemler[0].aciklama

    db = db_env()
    try:
        kalan, sonen = db.get(models.Case, eski_id), db.get(models.Case, yeni_id)
        assert sonen.deleted_at is not None and kalan.deleted_at is None
        assert [f.sistem_no for f in kalan.foys] == ["H-1557"]
        assert set(kalan.klasor_no_2.split(";")) == {"27.004.00", "2.557.00"}   # listeler birleşti
    finally:
        db.close()

    ikinci = _kos(db_env, tmp_path, apply=True, cift=[(eski_id, 999999, "H-1557")])
    assert ikinci.sayim("cift", "YAPILDI") == 0 and ikinci.sayim("cift", "ATLANDI") == 1


# ─── Adım 4: eski esaslı kart ────────────────────────────────────────────────

def test_eski_esas_kart_kapanir_esas_tarihceye_belge_tasinir(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        guncel = _kart(db, "S3.AXA........2655.IDARE.00000", "1.20755", esas_no="2021/1820",
                       court="İzmir 4. İdare Mahkemesi")
        eski = _kart(db, "S3.AXA........2617.IDARE.00000", "1.20755", esas_no="2017/821",
                     court="İzmir 4. İdare Mahkemesi")
        _foy(db, guncel, "id-9477")
        tombstone = _foy(db, eski, "id-15075")
        tombstone.kapsam_durumu = "SILINDI"
        _belge(db, eski, "karar.pdf")
        db.commit()
        guncel_id, eski_id = guncel.id, eski.id
    finally:
        db.close()

    sonuc = _kos(db_env, tmp_path, apply=True, eski_esas=[(eski_id, guncel_id, "id-9477")])
    assert sonuc.sayim("eski_esas", "YAPILDI") == 1

    db = db_env()
    try:
        assert db.get(models.Case, eski_id).deleted_at is not None
        guncel = db.get(models.Case, guncel_id)
        assert guncel.esas_no == "2021/1820"                       # güncel esas değişmedi
        tarihce = {(r.esas_no, r.stage, r.is_current)
                   for r in db.query(models.CaseEsasNumber).filter_by(case_id=guncel_id)}
        assert ("2017/821", case_manager.ESAS_STAGE_ONCEKI, False) in tarihce
        assert {d.case_id for d in db.query(models.CaseDocument).all()} == {guncel_id}
        assert {f.sistem_no for f in guncel.foys} == {"id-9477", "id-15075"}
    finally:
        db.close()

    ikinci = _kos(db_env, tmp_path, apply=True, eski_esas=[(eski_id, guncel_id, "id-9477")])
    assert ikinci.sayim("eski_esas", "ATLANDI") == 1


def test_eski_esas_canli_foy_varsa_ret(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        guncel = _kart(db, "S3.AXA........2714.HUKUK.00000", "1.20831", esas_no="2023/3")
        eski = _kart(db, "S3.AXA........2715.HUKUK.00000-2", "1.20831", esas_no="2017/632")
        _foy(db, guncel, "H-11728")
        _foy(db, eski, "H-15076")                                   # kapsam işareti YOK → canlı
        db.commit()
        guncel_id, eski_id = guncel.id, eski.id
    finally:
        db.close()

    sonuc = _kos(db_env, tmp_path, apply=True, eski_esas=[(eski_id, guncel_id, "H-11728")])
    assert sonuc.sayim("eski_esas", "RET") == 1
    db = db_env()
    try:
        assert db.get(models.Case, eski_id).deleted_at is None
    finally:
        db.close()


# ─── Adım 5-6: ana tür · klasör ──────────────────────────────────────────────

def test_ana_tur_foy_turuyle_duzelir_celiskili_satir_atlanir(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        db.add(models.FileType(code="ARABULUCULUK", name="Arabuluculuk"))
        kart = _kart(db, "D1.F_BINGUL...0012.ARABU.00000", "1.1", file_type="Hukuk")
        ikinci = _kart(db, "D1.S_TAS......0006.ARABU.00000", "1.2", file_type="Hukuk")
        db.commit()
        kart_id, ikinci_id = kart.id, ikinci.id
    finally:
        db.close()

    sonuc = _kos(db_env, tmp_path, apply=True,
                 tur=[(kart_id, "H-14986: Arabuluculuk"),
                      (ikinci_id, "H-1: Hukuk · H-2: Arabuluculuk")])    # iki farklı tür → okunmaz
    assert sonuc.sayim("tur", "YAPILDI") == 1
    db = db_env()
    try:
        assert db.get(models.Case, kart_id).file_type == "Arabuluculuk"
        assert db.get(models.Case, ikinci_id).file_type == "Hukuk"
    finally:
        db.close()


def test_klasor_oneki_degistirilir_sira_korunur(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        kart = _kart(db, "S3.S_BABA.....0001.IDARE.00000", "21461;1851.001")
        db.commit()
        kart_id = kart.id
    finally:
        db.close()

    sonuc = _kos(db_env, tmp_path, apply=True,
                 klasor=[(kart_id, "21461;1851.001", "id-15819 1.21461.00 · id-15831 1851.001")])
    assert sonuc.sayim("klasor", "YAPILDI") == 1
    db = db_env()
    try:
        assert db.get(models.Case, kart_id).klasor_no_2 == "1.21461.00;1851.001"
    finally:
        db.close()


# ─── Adım 7: ilişkili dosya ──────────────────────────────────────────────────

def test_iliski_tek_yonlu_yazilir_ters_yon_de_taninir(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        a = _kart(db, "D1.H_DEMIREL..0006.HUKUK.00000", "1056.003")
        b = _kart(db, "D1.H_DEMIREL..0005.ARABU.00000", "1056.003")
        db.commit()
        a_id, b_id = a.id, b.id
    finally:
        db.close()

    sonuc = _kos(db_env, tmp_path, apply=True, iliski=[(a_id, b_id, "H-15400", "ARB-15364")])
    assert sonuc.sayim("iliski", "YAPILDI") == 1
    assert _sayim(db_env, models.CaseRelation) == 1

    # Ters sırayla ikinci koşu: bağ zaten var, ikinci satır yazılmaz.
    ikinci = _kos(db_env, tmp_path, apply=True, iliski=[(b_id, a_id, "ARB-15364", "H-15400")])
    assert ikinci.sayim("iliski", "ATLANDI") == 1
    assert _sayim(db_env, models.CaseRelation) == 1

    db = db_env()
    try:
        bag = db.query(models.CaseRelation).one()
        assert (bag.source_case_id, bag.target_case_id) == (a_id, b_id)
        assert bag.relation_type == ec.ILISKI_TURU and bag.created_by == ec.DEGISTIREN
    finally:
        db.close()


# ─── Kuru koşu · kapatma ─────────────────────────────────────────────────────

def test_kuru_kosu_hicbir_seyi_yazmaz(db_env, tmp_path, sabitsiz):  # noqa: F811
    db = db_env()
    try:
        db.add(models.CaseSubject(code="RUCUEN_ALACAK", name="Rücuen Alacak"))
        kart = _kart(db, "S3.AXA........0711.HUKUK.00000", "1.1", subject="İtirazın İptali")
        a = _kart(db, "D1.A.........0001.HUKUK.00000", "2.1")
        b = _kart(db, "D1.A.........0002.ARABU.00000", "2.1")
        db.commit()
        kart_id, a_id, b_id = kart.id, a.id, b.id
    finally:
        db.close()

    sonuc = _kos(db_env, tmp_path, konu=[(kart_id, "Rücuen Alacak")], iliski=[(a_id, b_id, "", "")])
    assert sonuc.sayim("konu", "YAPILDI") == 1 and sonuc.sayim("iliski", "YAPILDI") == 1

    db = db_env()
    try:
        assert db.get(models.Case, kart_id).subject == "İtirazın İptali"
        assert db.query(models.CaseRelation).count() == 0
        assert db.query(models.CaseHistory).count() == 0
    finally:
        db.close()


def test_kapatma_foylu_karti_reddeder(db_env, tmp_path, monkeypatch):  # noqa: F811
    db = db_env()
    try:
        kart = _kart(db, "X1.I_KUTLUK...0002.ICRAA.00000", "9.9")
        _foy(db, kart, "H-1")
        db.commit()
        kart_id = kart.id
    finally:
        db.close()
    monkeypatch.setattr(ec, "SABIT_DUZELTMELER", ())
    monkeypatch.setattr(ec, "SABIT_ILISKILER", ())
    monkeypatch.setattr(ec, "KAPATILACAK_KARTLAR", ((kart_id, "deneme kaydı"),))

    sonuc = _kos(db_env, tmp_path, apply=True)
    assert sonuc.sayim("kapatma", "RET") == 1
    db = db_env()
    try:
        assert db.get(models.Case, kart_id).deleted_at is None
    finally:
        db.close()
