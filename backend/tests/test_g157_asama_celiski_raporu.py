"""G157 — Aşama çelişki raporu üreticisi: yer tutucu sınıfı (S6), E-8 etiketi,
06.09 S5 satır 82 senaryosu, başlık kilidi, G150/G151 uyumu, DB+paket toplama,
ekip cevabının geri okunması, CLI.

**TEST VERİSİ KURALI (A.2):** gerçek paket / cevap dosyası REPOYA GİRMEZ;
xlsx'ler openpyxl ile SENTETİK üretilir.
"""
import collections
from pathlib import Path

import pytest
from sqlalchemy import create_engine, event, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import models
from database import _MIGRATIONS, Base
from scripts import asama_celiski_raporu as cli
from scripts.hukdok_aktarim import CIKIS_GIRDI, CIKIS_TAMAM
from services import asama_celiski_raporu as acr

# 06.09'da ekibe giden dosyanın başlıkları ve sayfa adları — AYNEN (ekip aynı formatı bekliyor).
BASLIKLAR_0609 = [
    "Kart", "Ofis no", "Aşama", "Föyler", "Müvekkiller", "Farklı alanlar", "Değerler (föy=değer)",
    "Ana sayfa değerleri", "Bizim okuma", "TKU aynı", "Hasar no aynı", "Karşı taraf aynı",
    "Yerel mahkeme+esas aynı", "CEVABINIZ", "AÇIKLAMANIZ",
]
OZET_0609 = ["Sınıf", "Ad", "Grup", "Ne görüyoruz", "Sorumuz"]
SAYFALAR_0609 = [
    "OZET", "S1_Yalnız boşluk", "S2_Yazım hatası adayı", "S3_Aynı dava, bambaşka ka",
    "S4_Mahkeme ya da esas far", "S5_Yalnız karar durumu fa",
]


# ─── yardımcılar ─────────────────────────────────────────────────────────────

def _foy(sistem_no, *, muvekkil="Axa Sigorta A.Ş.", kimlik=None, tku="TKU-1", hasar="H-1",
         sheet_asama=None, yerel="İstanbul 8. Tüketici", sheet_esas="2020/143", zincir=(), **satir):
    temel = {"sistem_no": sistem_no, "asama_no": 1, "asama": "Temyiz", "mahkeme": "Danıştay 10. Daire",
             "esas_no": "2022/100", "karar_no": "2023/50", "karar_tarihi": "05.05.2023",
             "karar_durumu": "Onama", "guven": "KESİN"}
    temel.update(satir)
    return acr.FoyAsama(
        sistem_no=sistem_no, satir=temel, sheet_asama=dict(sheet_asama or {}),
        muvekkil=muvekkil, muvekkil_kimligi=kimlik or acr.normalize_party_key(muvekkil),
        karsi_taraf="Ayşe Hasta", tku=tku, hasar_no=hasar, sheet_esas=sheet_esas,
        yerel_mahkeme=yerel, zincir_esaslari=set(zincir),
    )


def _grup(*foyler, asama="TEMYIZ", kart_id=14213, ofis="S2.D_GULBAHAR.0001.IDARE.00000", **extra):
    return acr.AsamaGrubu(kart_id, ofis, asama, list(foyler), **extra)


def _sayfa(yol, sayfa):
    from openpyxl import load_workbook

    wb = load_workbook(yol, read_only=True)
    try:
        return list(wb[sayfa].iter_rows(values_only=True)), wb.sheetnames
    finally:
        wb.close()


# ─── yer tutucu deseni ───────────────────────────────────────────────────────

@pytest.mark.parametrize("alan,deger,beklenen", [
    ("mahkeme", "DANIŞTAY . DAİRE", True),
    ("mahkeme", "DANIŞTAY DAİRE", True),
    ("mahkeme", "Yargıtay . Hukuk Dairesi", True),
    ("mahkeme", "DANIŞTAY 10. DAİRE", False),
    ("mahkeme", "Yargıtay 4. Hukuk Dairesi", False),
    ("mahkeme", "İSTANBUL BİM 8. İDD", False),
    ("mahkeme", "ANKARA 3 İDARİ DA ADAİRESİ", False),       # yazım hatası, yer tutucu değil (lokal ölçüm)
    ("mahkeme", "Ankara 3. İdari Dava Dairesi", False),
    ("mahkeme", "Ankara ?. İdare", True),
    ("esas_no", "2026/", True),
    ("esas_no", "2026", True),
    ("esas_no", "2026/1234", False),
    ("karar_no", "2023/ ", True),
    ("karar_no", "2023/50", False),
    ("karar_durumu", "Onama", False),
    ("esas_no", None, False),
    ("esas_no", "", False),
])
def test_yer_tutucu_deseni(alan, deger, beklenen):
    assert acr.yer_tutucu_mu(alan, deger) is beklenen


# ─── S5 satır 82 senaryosu ───────────────────────────────────────────────────

def test_s5_satir_82_yer_tutucu_farki_sutunlari_doldurur(tmp_path):
    """06.09 S5 satır 82: id-10285/10286 Temyiz — mahkeme ve esas yer tutucu,
    normalize anahtarında eşit; eski üretici `Farklı alanlar`ı boş bırakmıştı.
    Yeni kural: S6, iki alan da listede, xlsx'in üç sütunu dolu."""
    grup = _grup(
        _foy("id-10285", muvekkil="Duru Saygın Gülbahar Dr.", mahkeme="DANIŞTAY . DAİRE", esas_no="2026/",
             karar_no=None, karar_tarihi=None, karar_durumu="Onama",
             sheet_asama={"mahkeme": "DANIŞTAY . DAİRE", "esas_no": "2026/", "karar_durumu": "Onama"}),
        _foy("id-10286", muvekkil="Quick Sigorta A.Ş.", mahkeme="DANIŞTAY DAİRE", esas_no="2026",
             karar_no=None, karar_tarihi=None, karar_durumu="Onama",
             sheet_asama={"mahkeme": "DANIŞTAY DAİRE", "esas_no": "2026", "karar_durumu": "Onama"}),
    )
    sonuc = acr.siniflandir(grup)
    assert sonuc is not None
    assert sonuc.sinif == "S6_YER_TUTUCU"
    assert sonuc.farkli_alanlar == ["mahkeme", "esas_no"]
    assert "yer tutucu" in sonuc.sebep

    yol = acr.xlsx_yaz([sonuc], tmp_path / "rapor.xlsx")
    satirlar, sayfalar = _sayfa(yol, acr.sayfa_adi("S6_YER_TUTUCU"))
    assert sayfalar[:6] == SAYFALAR_0609 and sayfalar[6].startswith("S6_")
    satir = dict(zip(BASLIKLAR_0609, satirlar[1], strict=True))
    assert satir["Farklı alanlar"] == "Mahkeme, Esas"
    assert "id-10285=DANIŞTAY . DAİRE" in satir["Değerler (föy=değer)"]
    assert "id-10286=2026" in satir["Değerler (föy=değer)"]
    assert "id-10285=2026/" in satir["Ana sayfa değerleri"]
    assert satir["Föyler"] == "id-10285; id-10286"
    assert satir["CEVABINIZ"] is None and satir["AÇIKLAMANIZ"] is None


def test_gercek_deger_ile_yer_tutucu_de_s6():
    """`DANIŞTAY 10. DAİRE` ↔ `DANIŞTAY DAİRE`: biri gerçek, biri yer tutucu → yine S6 (daire numarası eksik)."""
    sonuc = acr.siniflandir(_grup(
        _foy("A-1", mahkeme="DANIŞTAY 10. DAİRE"), _foy("A-2", mahkeme="DANIŞTAY DAİRE"),
    ))
    assert sonuc is not None and sonuc.sinif == "S6_YER_TUTUCU"
    assert sonuc.farkli_alanlar == ["mahkeme"]


# ─── E-8 etiketi ─────────────────────────────────────────────────────────────

def test_s5_farkli_muvekkil_e8_etiketi(tmp_path):
    sonuc = acr.siniflandir(_grup(
        _foy("H-1", muvekkil="Axa Sigorta A.Ş.", karar_durumu="Red/Esastan"),
        _foy("H-2", muvekkil="Ali Seyed Resuli Dr", karar_durumu="Kabul"),
        asama="YEREL",
    ))
    assert sonuc is not None
    assert sonuc.sinif == "S5_DURUM"
    assert sonuc.etiket == acr.E8_ETIKETI
    assert sonuc.okuma.startswith("müvekkil yönü farkı (E-8) — hata değil")
    satirlar, _ = _sayfa(acr.xlsx_yaz([sonuc], tmp_path / "r.xlsx"), acr.sayfa_adi("S5_DURUM"))
    satir = dict(zip(BASLIKLAR_0609, satirlar[1], strict=True))
    assert satir["Bizim okuma"].startswith(acr.E8_ETIKETI)
    assert satir["Farklı alanlar"] == "Karar durumu"


def test_s5_ayni_muvekkil_etiketsiz():
    sonuc = acr.siniflandir(_grup(
        _foy("H-1", karar_durumu="Kabul"), _foy("H-2", karar_durumu="Red/Esastan"),
    ))
    assert sonuc is not None and sonuc.sinif == "S5_DURUM" and sonuc.etiket == ""
    assert sonuc.okuma == "künye aynı, yalnız karar durumu farklı"


@pytest.mark.parametrize("a,b", [
    ("Kabul", "Kabul/Kısmen"), ("Red/Esastan", "Red/Husumet"), ("Red/Esastan", "Red/Zamanaşımı"),
    ("Açılmamış Sayılması (HMK 150. Md)", "Red/Esastan"), ("Kaldırma", "Kaldırma/Yeniden Hüküm"),
    ("Onama", "Düzelterek Onama"), ("Onama", "Bilinmeyen Değer"),
])
def test_s5_ayni_yon_farkli_muvekkil_etiketsiz(a, b):
    """Ekip cevabı: aynı yönün ayrıntı farkı (Kabul ⇄ Kabul/Kısmen …) E-8 değil, BELİRSİZ/HATA."""
    sonuc = acr.siniflandir(_grup(
        _foy("H-1", muvekkil="Axa Sigorta A.Ş.", karar_durumu=a),
        _foy("H-2", muvekkil="Ali Seyed Resuli Dr", karar_durumu=b),
    ))
    assert sonuc is not None and sonuc.sinif == "S5_DURUM" and sonuc.etiket == ""


@pytest.mark.parametrize("a,b", [
    ("Kabul/Kısmen", "Red/Esastan"), ("Kabul", "Red/Esastan"), ("Bozma", "Onama"),
    ("Başvuru Ret", "Kaldırma/Yeniden Hüküm"), ("kısmen kabul", "RED/HUSUMET"),
])
def test_s5_zit_yon_farkli_muvekkil_e8(a, b):
    sonuc = acr.siniflandir(_grup(
        _foy("H-1", muvekkil="Axa Sigorta A.Ş.", karar_durumu=a),
        _foy("H-2", muvekkil="Ali Seyed Resuli Dr", karar_durumu=b),
    ))
    assert sonuc is not None and sonuc.etiket == acr.E8_ETIKETI


@pytest.mark.parametrize("deger,yon", [
    ("Kabul", "KABUL"), ("Kabul/Kısmen", "KABUL"), ("Kısmen Kabul", "KABUL"), ("Red/Esastan", "RED"),
    ("Açılmamış Sayılması (HMK 150. Md)", "RED"), ("Karar Verilmesine Yer Olmadığına (HMK 331 Md.)", "RED"),
    ("Onama", "ONAMA"), ("Düzelterek Onama", "ONAMA"), ("Bozma", "BOZMA"), ("Başvuru Ret", "BASVURURET"),
    ("Kaldırma", "KALDIRMA"), ("Kapalı", None), (None, None), ("", None),
])
def test_karar_yonu(deger, yon):
    assert acr.karar_yonu(deger) == yon


def test_e8_kimligi_case_party_id_onceliklidir():
    """Aynı ad, farklı `case_party_id` → yine ayrı müvekkil (E-8)."""
    sonuc = acr.siniflandir(_grup(
        _foy("H-1", kimlik="party:1", karar_durumu="Kabul"),
        _foy("H-2", kimlik="party:2", karar_durumu="Red/Esastan"),
    ))
    assert sonuc is not None and sonuc.etiket == acr.E8_ETIKETI


# ─── G150 / G151 uyumu: çelişki tanımı aktarımın tanımı ──────────────────────

def test_yalniz_bosluk_artik_celiski_degil():
    """G150: kardeş föyde boş hücre imzaya girmez → grup rapora düşmez."""
    assert acr.siniflandir(_grup(_foy("A-1"), _foy("A-2", karar_no=None, karar_durumu=None))) is None


def test_buro_durumu_bos_sayilir_g151():
    assert acr.siniflandir(_grup(_foy("A-1", karar_durumu="Red/Esastan"), _foy("A-2", karar_durumu="Kapalı"))) is None


def test_kalan_gercek_bosluk_s1():
    """Aşama satırlarının HEPSİNDE boş, ana sayfada dolu → S1 (aşama katmanı ana sayfadan künye almaz)."""
    sonuc = acr.siniflandir(_grup(
        _foy("A-1", karar_no=None, sheet_asama={"karar_no": "2023/50"}),
        _foy("A-2", karar_no=None, sheet_asama={"karar_no": None}),
    ))
    assert sonuc is not None and sonuc.sinif == "S1_BOSLUK" and sonuc.farkli_alanlar == ["karar_no"]


def test_tek_foy_ve_ayni_grup_none():
    assert acr.siniflandir(_grup(_foy("A-1"))) is None
    assert acr.siniflandir(_grup(_foy("A-1"), _foy("A-2"))) is None


# ─── S2 / S3 / S4 ────────────────────────────────────────────────────────────

def test_s2_tek_rakam_farki():
    sonuc = acr.siniflandir(_grup(_foy("A-1", karar_no="2023/50"), _foy("A-2", karar_no="2023/56")))
    assert sonuc is not None and sonuc.sinif == "S2_YAZIM" and sonuc.farkli_alanlar == ["karar_no"]


def test_s3_bambaska_karar_yil_uyumsuz():
    sonuc = acr.siniflandir(_grup(
        _foy("A-1", karar_no="2025/4243", karar_tarihi="29.03.2023"),
        _foy("A-2", karar_no="2021/12", karar_tarihi="10.10.2021"),
    ))
    assert sonuc is not None and sonuc.sinif == "S3_BASKA_KARAR"
    assert "yılı uyuşmuyor" in sonuc.sebep
    assert sonuc.farkli_alanlar == ["karar_no", "karar_tarihi"]


def test_s3_ana_sayfa_kunye_celiskisi_notu():
    sonuc = acr.siniflandir(_grup(
        _foy("A-1", karar_no="2025/4243", karar_tarihi="29.03.2023"),
        _foy("A-2", karar_no="2021/12", karar_tarihi="10.10.2021"),
        sheet_kunye_celiskisi=True,
    ))
    assert sonuc is not None and sonuc.sebep.endswith("(ana sayfa karar no/tarihi de çelişiyor)")


def test_s4_mahkeme_yazimi_ve_bayat_foy():
    yazim = acr.siniflandir(_grup(
        _foy("A-1", mahkeme="İSTANBUL 8. İDD"), _foy("A-2", mahkeme="İSTANBUL BİM 8. İDD"),
        asama="ISTINAF",
    ))
    assert yazim is not None and yazim.sinif == "S4_MAHKEME_ESAS" and "yazımı/eksik" in yazim.sebep
    bayat = acr.siniflandir(_grup(
        _foy("A-1", esas_no="2019/7", zincir={"20197"}),
        _foy("A-2", esas_no="2022/900", zincir={"20197", "2022900"}),
    ))
    assert bayat is not None and bayat.sinif == "S4_MAHKEME_ESAS" and "bayat föy" in bayat.sebep


def test_farkli_tur_sayisi_s4():
    sonuc = acr.siniflandir(_grup(_foy("A-1"), _foy("A-2"), tur_notu="çok tur: föyler farklı sayıda tur anlatıyor"))
    assert sonuc is not None and sonuc.sinif == "S4_MAHKEME_ESAS" and sonuc.farkli_alanlar == []


# ─── başlık kilidi ───────────────────────────────────────────────────────────

def test_basliklar_0609_ile_birebir(tmp_path):
    assert list(acr.BASLIKLAR) == BASLIKLAR_0609
    assert list(acr.OZET_BASLIKLARI) == OZET_0609
    yol = acr.xlsx_yaz([], tmp_path / "bos.xlsx")
    ozet, sayfalar = _sayfa(yol, "OZET")
    assert sayfalar[:6] == SAYFALAR_0609
    assert list(ozet[0]) == OZET_0609
    assert [r[0] for r in ozet[1:7]] == ["1", "2", "3", "4", "5", "6"]
    assert ozet[7][1:3] == ("Toplam", 0)
    for sayfa in sayfalar[1:]:
        satirlar, _ = _sayfa(yol, sayfa)
        assert list(satirlar[0]) == BASLIKLAR_0609


def test_kanit_sutunlari_evet_hayir(tmp_path):
    sonuc = acr.siniflandir(_grup(
        _foy("A-1", karar_no="2023/50", hasar="H-1"), _foy("A-2", karar_no="2023/56", hasar="H-2"),
    ))
    satirlar, _ = _sayfa(acr.xlsx_yaz([sonuc], tmp_path / "k.xlsx"), acr.sayfa_adi("S2_YAZIM"))
    satir = dict(zip(BASLIKLAR_0609, satirlar[1], strict=True))
    assert (satir["TKU aynı"], satir["Hasar no aynı"], satir["Karşı taraf aynı"],
            satir["Yerel mahkeme+esas aynı"]) == ("evet", "hayır", "evet", "evet")


# ─── DB + paket → gruplar ────────────────────────────────────────────────────

def _index_ops(table):
    return [sql for op in _MIGRATIONS if op[0] == "index" and op[1] == table for sql in op[2]]


@pytest.fixture()
def db_env():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)

    @event.listens_for(engine, "connect")
    def _fk_ac(dbapi_connection, _record):
        dbapi_connection.isolation_level = None
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    @event.listens_for(engine, "begin")
    def _begin(conn):
        conn.exec_driver_sql("BEGIN")

    Base.metadata.create_all(engine)
    with engine.begin() as conn:
        for sql in _index_ops("case_foys"):
            conn.execute(text(sql))
    yield sessionmaker(bind=engine, autocommit=False, autoflush=False)
    engine.dispose()


SHEET_BASLIKLARI = ["SistemNo", "Dosya No", "Klasör No", "Hasar No", "Müvekkil", "Karşı Taraf",
                    "Yerel Mahkeme", "Esas", "Karar No", "Temyiz Mahkemesi", "Temyiz_Esas_No",
                    "Yargıtay Onama Durumu"]
ASAMA_BASLIKLARI = ["SistemNo", "AsamaNo", "Aşama", "Mahkeme", "Esas No", "Karar No", "Karar Tarihi",
                    "Karar Durumu", "Güven"]


def _paket(yol, sheet, asama):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(SHEET_BASLIKLARI)
    for s in sheet:
        ws.append([s.get(b) for b in SHEET_BASLIKLARI])
    wa = wb.create_sheet("Karar_Asamalari")
    wa.append(ASAMA_BASLIKLARI)
    for s in asama:
        wa.append([s.get(b) for b in ASAMA_BASLIKLARI])
    wb.save(yol)
    wb.close()
    return Path(yol)


def _sheet(sn, dosya, muvekkil, **extra):
    temel = {"SistemNo": sn, "Dosya No": dosya, "Klasör No": "TKU-1", "Hasar No": "H-1", "Müvekkil": muvekkil,
             "Karşı Taraf": "Ayşe Hasta", "Yerel Mahkeme": "İstanbul 8. Tüketici", "Esas": "2020/143",
             "Temyiz Mahkemesi": "DANIŞTAY DAİRE", "Temyiz_Esas_No": "2026", "Yargıtay Onama Durumu": "Onama"}
    temel.update(extra)
    return temel


def _asama(sn, no, ad, **extra):
    temel = {"SistemNo": sn, "AsamaNo": no, "Aşama": ad, "Mahkeme": "İstanbul 8. Tüketici", "Esas No": "2020/143",
             "Karar No": "2021/5", "Karar Tarihi": "05.05.2021", "Karar Durumu": "Red/Esastan", "Güven": "KESİN"}
    temel.update(extra)
    return temel


@pytest.fixture()
def zemin(db_env):
    """İki kart: 14213 (satır 82 ikizi, iki föy) + kapsam dışı üçüncü föy; ikinci kart tek föy."""
    db = db_env()
    try:
        k1 = models.Case(tracking_no="S2.D_GULBAHAR.0001.IDARE.00000", status="DERDEST", klasor_no_2="D-1")
        k2 = models.Case(tracking_no="S3.AXA........0151.HUKUK.00000", status="DERDEST", klasor_no_2="D-2")
        db.add_all([k1, k2])
        db.flush()
        db.add_all([
            models.CaseFoy(sistem_no="id-10285", case_id=k1.id, tku_no="TKU-1", hasar_no="H-1",
                           ham_veri={"Müvekkil": "Duru Saygın Gülbahar Dr."}),
            models.CaseFoy(sistem_no="id-10286", case_id=k1.id, tku_no="TKU-1", hasar_no="H-1",
                           ham_veri={"Müvekkil": "Quick Sigorta A.Ş."}),
            models.CaseFoy(sistem_no="id-9999", case_id=k1.id, tku_no="TKU-1", kapsam_durumu="KAPSAM_DISI"),
            models.CaseFoy(sistem_no="H-407", case_id=k2.id, tku_no="TKU-2"),
        ])
        db.commit()
        return db_env, k1.id
    finally:
        db.close()


def test_gruplari_topla_ve_rapor_uret(zemin, tmp_path):
    fabrika, kart_id = zemin
    paket = _paket(tmp_path / "paket.xlsx", [
        _sheet("id-10285", "D-1", "Duru Saygın Gülbahar Dr.", **{"Temyiz Mahkemesi": "DANIŞTAY . DAİRE",
                                                                  "Temyiz_Esas_No": "2026/"}),
        _sheet("id-10286", "D-1", "Quick Sigorta A.Ş."),
        _sheet("id-9999", "D-1", "Kapsam Dışı"),
        _sheet("H-407", "D-2", "Axa Sigorta A.Ş."),
    ], [
        _asama("id-10285", 1, "Yerel"), _asama("id-10285", 2, "Önceki", **{"Esas No": "2019/1"}),
        _asama("id-10285", 3, "Temyiz", Mahkeme="DANIŞTAY . DAİRE", **{"Esas No": "2026/", "Karar No": None,
                                                                          "Karar Tarihi": None,
                                                                          "Karar Durumu": "Onama"}),
        _asama("id-10286", 1, "Yerel", **{"Karar Durumu": "Kapalı"}),          # G151: boş sayılır
        _asama("id-10286", 3, "Temyiz", Mahkeme="DANIŞTAY DAİRE", **{"Esas No": "2026", "Karar No": None,
                                                                       "Karar Tarihi": None,
                                                                       "Karar Durumu": "Onama"}),
        _asama("id-9999", 3, "Temyiz", Mahkeme="Danıştay 10. Daire", **{"Esas No": "2026/77"}),
        _asama("H-407", 1, "Yerel"),
    ])
    db = fabrika()
    try:
        gruplar = acr.gruplari_topla(db, paket)
        assert {(g.asama, len(g.foyler)) for g in gruplar} == {("YEREL", 2), ("TEMYIZ", 2)}
        temyiz = next(g for g in gruplar if g.asama == "TEMYIZ")
        assert [f.sistem_no for f in temyiz.foyler] == ["id-10285", "id-10286"]     # kapsam dışı yok
        assert temyiz.foyler[0].muvekkil == "Duru Saygın Gülbahar Dr."
        assert temyiz.foyler[0].sheet_asama["mahkeme"] == "DANIŞTAY . DAİRE"
        assert "20191" in temyiz.foyler[0].zincir_esaslari                          # Önceki esas zincirde
        assert temyiz.foyler[0].muvekkil_kimligi != temyiz.foyler[1].muvekkil_kimligi

        yol, celiskiler = acr.rapor_uret(db, paket, tmp_path / "rapor", tarih="2026-09-10")
    finally:
        db.close()
    assert yol.name == "HUKDOK_ASAMA_CELISKILERI_2026-09-10.xlsx"
    assert [(c.sinif, c.asama, c.kart_id) for c in celiskiler] == [("S6_YER_TUTUCU", "TEMYIZ", kart_id)]
    assert celiskiler[0].farkli_alanlar == ["mahkeme", "esas_no"]
    satirlar, _ = _sayfa(yol, acr.sayfa_adi("S6_YER_TUTUCU"))
    satir = dict(zip(BASLIKLAR_0609, satirlar[1], strict=True))
    assert satir["Kart"] == kart_id and satir["Ofis no"] == "S2.D_GULBAHAR.0001.IDARE.00000"
    assert satir["Farklı alanlar"] == "Mahkeme, Esas" and satir["Aşama"] == "Temyiz"
    assert "id-10286=DANIŞTAY DAİRE" in satir["Ana sayfa değerleri"]
    assert satir["Müvekkiller"] == "Duru Saygın Gülbahar Dr.; Quick Sigorta A.Ş."


def test_yalniz_kartlar_filtresi_celiski_listesinden(zemin, tmp_path):
    """`celiskileri_bul` çıktısı (KART kümesi) girdi olarak: yalnız o kartlar toplanır."""
    from scripts.hukdok_aktarim import Celiski

    fabrika, kart_id = zemin
    paket = _paket(tmp_path / "p.xlsx", [
        _sheet("id-10285", "D-1", "A"), _sheet("id-10286", "D-1", "B"),
    ], [
        _asama("id-10285", 1, "Yerel", **{"Karar No": "2021/5"}),
        _asama("id-10286", 1, "Yerel", **{"Karar No": "2021/6"}),
    ])
    idler = acr.celiski_kart_idleri([
        Celiski("KART", str(kart_id), "karar_no", "x"), Celiski("TKU", "TKU-1", "karar_no", "y"),
        Celiski("KART", "bozuk", "karar_no", "z"),
    ])
    assert idler == {kart_id}
    db = fabrika()
    try:
        assert len(acr.gruplari_topla(db, paket, yalniz_kartlar=idler)) == 1
        assert acr.gruplari_topla(db, paket, yalniz_kartlar={kart_id + 100}) == []
    finally:
        db.close()


def test_cok_tur_farkli_sayida_satir(zemin, tmp_path):
    fabrika, _ = zemin
    paket = _paket(tmp_path / "p.xlsx", [
        _sheet("id-10285", "D-1", "A"), _sheet("id-10286", "D-1", "B"),
    ], [
        _asama("id-10285", 1, "Yerel"), _asama("id-10285", 2, "Yerel", **{"Esas No": "2023/9",
                                                                          "Karar Tarihi": "01.01.2024"}),
        _asama("id-10286", 1, "Yerel"),
    ])
    db = fabrika()
    try:
        celiskiler = acr.siniflandir_hepsi(acr.gruplari_topla(db, paket))
    finally:
        db.close()
    assert [(c.sinif, c.sebep) for c in celiskiler] == [
        ("S4_MAHKEME_ESAS", "çok tur: föyler farklı sayıda tur anlatıyor"),
    ]


# ─── cevabı geri okuma ───────────────────────────────────────────────────────

def _cevapli(yol, sayfalar):
    """Sentetik CEVAPLI xlsx: {sayfa adı: [(kart, aşama, föyler, farklı, okuma, cevap, açıklama), ...]}."""
    from openpyxl import Workbook

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "OZET"
    ws0.append(OZET_0609)
    for ad, satirlar in sayfalar.items():
        ws = wb.create_sheet(ad)
        ws.append(BASLIKLAR_0609)
        for kart, asama, foyler, farkli, okuma, cevap, aciklama in satirlar:
            if kart is None and cevap is None:
                ws.append([None] * len(BASLIKLAR_0609))                  # tamamen boş satır
                continue
            ws.append([kart, "S3.X.0001", asama, foyler, "A; B", farkli, "v", "v", okuma,
                       "evet", "evet", "evet", "evet", cevap, aciklama])
    wb.save(yol)
    wb.close()
    return Path(yol)


@pytest.mark.parametrize("ham,beklenen", [
    ("HATA", acr.CEVAP_HATA), (" hata ", acr.CEVAP_HATA), ("BELİRSİZ", acr.CEVAP_BELIRSIZ),
    ("Belirsiz", acr.CEVAP_BELIRSIZ), ("SEBEBİ VAR", acr.CEVAP_SEBEBI_VAR), ("sebebi var (E-8)", acr.CEVAP_SEBEBI_VAR),
    (None, acr.CEVAP_BOS), ("", acr.CEVAP_BOS), ("belki", acr.CEVAP_TANINMAYAN),
])
def test_cevap_anahtari(ham, beklenen):
    assert acr.cevap_anahtari(ham) == beklenen


def test_cevaplari_oku_sayim_ve_beklenti(tmp_path):
    yol = _cevapli(tmp_path / "cevapli.xlsx", collections.OrderedDict([
        ("S2_Yazım hatası adayı", [
            (11, "Yerel", "id-1; id-2", "Karar no", "aynı mahkeme+esas", "HATA", "Karar No: doğru 2014/1133"),
            (12, "Yerel", "H-1; H-2", "Karar no, Karar tarihi", "aynı mahkeme+esas", "BELİRSİZ", ""),
        ]),
        ("S5_Yalnız karar durumu fa", [
            (5, "Yerel", "H-3; H-4", "Karar durumu", f"{acr.E8_ETIKETI}; künye aynı", "SEBEBİ VAR", "E-8"),
            (6, "Yerel", "H-5; H-6", "Karar durumu", f"{acr.E8_ETIKETI}; künye aynı", "HATA", "Kabul"),
            (7, "Yerel", "H-7; H-8", "Karar durumu", "künye aynı", "", ""),
            (None, None, None, None, None, None, None),                      # tamamen boş satır
        ]),
        ("NOTLAR", [(1, "x", "x", "x", "x", "HATA", "sayılmaz")]),          # sınıf sayfası değil
    ]))
    ozet = acr.cevaplari_oku(yol)
    assert set(ozet.sayim) == {"S2_YAZIM", "S5_DURUM"}
    assert ozet.sayim["S2_YAZIM"] == {acr.CEVAP_HATA: 1, acr.CEVAP_BELIRSIZ: 1}
    assert ozet.sayim["S5_DURUM"] == {acr.CEVAP_SEBEBI_VAR: 1, acr.CEVAP_HATA: 1, acr.CEVAP_BOS: 1}
    assert ozet.toplam[acr.CEVAP_HATA] == 2
    assert (ozet.e8_etiketli, ozet.e8_sebebi_var) == (2, 1)
    assert [(b.sinif, b.kart, b.farkli_alanlar, b.aciklama) for b in ozet.beklentiler] == [
        ("S2_YAZIM", "11", "Karar no", "Karar No: doğru 2014/1133"),
        ("S5_DURUM", "6", "Karar durumu", "Kabul"),
    ]
    metin = acr.cevap_ozeti_metni(ozet)
    assert "HATA=2" in metin and "Düzeltme_Logu beklentisi (HATA): 2 kalem" in metin
    csv_yolu = acr.beklenti_csv_yaz(ozet.beklentiler, tmp_path / "b.csv")
    icerik = csv_yolu.read_text(encoding="utf-8-sig").splitlines()
    assert icerik[0] == "sinif;kart;ofis_no;asama;foyler;farkli_alanlar;aciklama"
    # föy listesi ';' ayracı taşıdığı için csv modülü tırnaklar (Excel doğru açar)
    assert icerik[1] == 'S2_YAZIM;11;S3.X.0001;Yerel;"id-1; id-2";Karar no;Karar No: doğru 2014/1133'


def test_uretilen_dosya_geri_okunur(tmp_path):
    """Yazıcı → okuyucu tur: üretilen xlsx'te CEVABINIZ boş → hepsi (boş)."""
    sonuc = acr.siniflandir(_grup(_foy("A-1", karar_durumu="Kabul"), _foy("A-2", karar_durumu="Red/Esastan")))
    yol = acr.xlsx_yaz([sonuc], tmp_path / "r.xlsx")
    ozet = acr.cevaplari_oku(yol)
    assert ozet.sayim["S5_DURUM"] == {acr.CEVAP_BOS: 1}
    assert ozet.beklentiler == []


def test_cevap_dosyasi_yoksa_hata(tmp_path):
    from scripts.hukdok_aktarim import AktarimHatasi

    with pytest.raises(AktarimHatasi):
        acr.cevaplari_oku(tmp_path / "yok.xlsx")


# ─── CLI ─────────────────────────────────────────────────────────────────────

def test_cli_uretim_ve_geri_okuma(zemin, tmp_path, capsys):
    fabrika, _ = zemin
    paket = _paket(tmp_path / "p.xlsx", [
        _sheet("id-10285", "D-1", "A"), _sheet("id-10286", "D-1", "B"),
    ], [
        _asama("id-10285", 1, "Yerel", **{"Karar No": "2021/5"}),
        _asama("id-10286", 1, "Yerel", **{"Karar No": "2021/6"}),
    ])
    kod = cli.main(["--paket", str(paket), "--cikti-dizini", str(tmp_path / "out"), "--tarih", "2026-09-10"],
                   session_factory=fabrika)
    assert kod == CIKIS_TAMAM
    cikti = capsys.readouterr().out
    assert "S2 Yazım hatası adayı" in cikti and "HUKDOK_ASAMA_CELISKILERI_2026-09-10.xlsx" in cikti
    uretilen = tmp_path / "out" / "HUKDOK_ASAMA_CELISKILERI_2026-09-10.xlsx"
    assert uretilen.exists()

    kod = cli.main(["--cevap", str(uretilen), "--beklenti-csv", str(tmp_path / "b.csv")])
    assert kod == CIKIS_TAMAM
    assert "(boş)=1" in capsys.readouterr().out
    assert (tmp_path / "b.csv").exists()

    assert cli.main(["--cevap", str(tmp_path / "yok.xlsx")]) == CIKIS_GIRDI
    assert cli.main(["--paket", str(tmp_path / "yok.xlsx"), "--cikti-dizini", str(tmp_path)],
                    session_factory=fabrika) == CIKIS_GIRDI


def test_cli_eksik_arguman(capsys):
    with pytest.raises(SystemExit):
        cli.main([])
