"""G149 — Föysüz + DERDEST kart listesi (`scripts/foysuz_derdest_kartlar.py`).

Script SALT OKUNURDUR: hiçbir tabloya yazmaz. Testler üç şeyi kilitler:

* seçim kuralı — yalnız `deleted_at IS NULL` + `DERDEST` + `case_foys`'ta satırı
  olmayan kartlar (föylü, MAHZEN, soft-silinmiş kart listede YOK)
* sayfa ayrımı — İcra/Tahkim/Vergi `KAPSAM_DISI_TUR`, kalan (tür boş dahil)
  `MALPRAKTIS_ADAYI`; tür karşılaştırması Türkçe harfe dayanıklı
* çıktı biçimi — kolon başlıkları, müvekkil/karşı taraf `; ` birleşimi,
  `OZET` sayıları, dosya adı sözleşmesi
"""
import re
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import models
from database import Base
from scripts import foysuz_derdest_kartlar as fdk


@pytest.fixture()
def SessionFactory():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    yield sessionmaker(bind=engine, autocommit=False, autoflush=False)
    engine.dispose()


def _kart(db, tracking_no, status="DERDEST", file_type="Hukuk", **alanlar):
    kart = models.Case(tracking_no=tracking_no, status=status, file_type=file_type, **alanlar)
    db.add(kart)
    db.flush()
    return kart


def _taraf(db, kart, name, party_type, role="Davacı"):
    db.add(models.CaseParty(case_id=kart.id, name=name, party_type=party_type, role=role))


@pytest.fixture()
def dolu_db(SessionFactory):
    """Görev tanımındaki fixture: föylü DERDEST, föysüz DERDEST Hukuk, föysüz
    DERDEST İcra, föysüz MAHZEN (+ soft-silinmiş föysüz DERDEST, tür boş DERDEST)."""
    db = SessionFactory()
    foylu = _kart(db, "D1.FOYLU.....0001.HUKUK.00000", court="Bursa 3. Asliye Hukuk")
    db.add(models.CaseFoy(sistem_no="SSTMN-1", case_id=foylu.id))

    hukuk = _kart(
        db, "D1.HUKUK.....0002.HUKUK.00000", klasor_no_2="2005.017", judicial_unit="Asliye Hukuk",
        esas_no="2024/123", court="İstanbul 6. Tüketici Mahkemesi", opening_date=date(2024, 3, 1),
        acceptance_date=date(2024, 2, 10), responsible_lawyer_name="Av. Ayşe Demir",
        updated_at=datetime(2026, 9, 1, 9, 30, tzinfo=timezone.utc),
    )
    _taraf(db, hukuk, "Murat Özcan Dr.", "CLIENT")
    _taraf(db, hukuk, "Ak Sigorta A.Ş.", "CLIENT")
    _taraf(db, hukuk, "Songül Kaya", "COUNTER", role="Davalı")
    _taraf(db, hukuk, "Sağlık Bakanlığı", "THIRD", role="Diğer Davalı")

    icra = _kart(db, "D1.ICRA......0003.ICRA.00000", file_type="İCRA", klasor_no_2="2006.004")
    _taraf(db, icra, "Nadir Yıldırım Dr.", "CLIENT")

    _kart(db, "D1.MAHZEN....0004.HUKUK.00000", status="MAHZEN")
    _kart(db, "D1.SILINMIS..0005.HUKUK.00000", deleted_at=datetime(2026, 8, 1, tzinfo=timezone.utc))
    _kart(db, "D1.TURSUZ....0006.HUKUK.00000", file_type=None)
    db.commit()
    db.close()
    return SessionFactory


def _sayfa(ws):
    satirlar = list(ws.iter_rows(values_only=True))
    return list(satirlar[0]), satirlar[1:]


# ── seçim kuralı + sayfa ayrımı ───────────────────────────────────────────────

def test_yalniz_foysuz_derdest_kartlar_dogru_sayfalarda(dolu_db, tmp_path):
    hedef = tmp_path / "rapor.xlsx"
    ozet = fdk.raporu_uret(dolu_db, hedef)

    assert ozet["toplam"] == 3
    assert ozet["malpraktis_adayi"] == 2
    assert ozet["kapsam_disi_tur"] == 1
    assert ozet["dosya"] == hedef

    wb = load_workbook(hedef)
    assert wb.sheetnames == ["MALPRAKTIS_ADAYI", "KAPSAM_DISI_TUR", "OZET"]

    basliklar, aday = _sayfa(wb["MALPRAKTIS_ADAYI"])
    assert basliklar == list(fdk.BASLIKLAR)
    assert {s[0] for s in aday} == {"D1.HUKUK.....0002.HUKUK.00000", "D1.TURSUZ....0006.HUKUK.00000"}

    _, kapsam_disi = _sayfa(wb["KAPSAM_DISI_TUR"])
    assert [s[0] for s in kapsam_disi] == ["D1.ICRA......0003.ICRA.00000"]

    tum_kartlar = {s[0] for s in aday} | {s[0] for s in kapsam_disi}
    for olmamali in ("D1.FOYLU.....0001.HUKUK.00000", "D1.MAHZEN....0004.HUKUK.00000",
                     "D1.SILINMIS..0005.HUKUK.00000"):
        assert olmamali not in tum_kartlar


def test_kapsam_isaretli_foy_de_foydur(SessionFactory, tmp_path):
    """Kapsam dışı işaretli föyü olan kart teslimde bir kez görünmüştür;
    'föysüz' listesine düşmemeli (ekibin listesinde zaten var)."""
    db = SessionFactory()
    kart = _kart(db, "D1.KAPSAM....0007.HUKUK.00000")
    db.add(models.CaseFoy(sistem_no="SSTMN-7", case_id=kart.id, kapsam_durumu="KAPSAM_DISI"))
    db.commit()
    db.close()

    ozet = fdk.raporu_uret(SessionFactory, tmp_path / "r.xlsx")
    assert ozet["toplam"] == 0


@pytest.mark.parametrize("tur", ["İcra", "İCRA", "icra", "Tahkim", "TAHKİM", "Vergi", " vergi "])
def test_kapsam_disi_turler_turkce_harfe_dayanikli(tur):
    assert fdk.kapsam_disi_mi(tur) is True


@pytest.mark.parametrize("tur", ["Hukuk", "İdare", "Ceza", "Arabuluculuk", "Danışmanlık", "Hukuk Dava", None, ""])
def test_diger_turler_malpraktis_adayidir(tur):
    assert fdk.kapsam_disi_mi(tur) is False


# ── satır içeriği ─────────────────────────────────────────────────────────────

def test_satir_kolonlari_dogru_degerleri_tasiyor(dolu_db, tmp_path):
    hedef = tmp_path / "rapor.xlsx"
    fdk.raporu_uret(dolu_db, hedef)
    basliklar, aday = _sayfa(load_workbook(hedef)["MALPRAKTIS_ADAYI"])
    satir = dict(zip(basliklar, next(s for s in aday if s[0] == "D1.HUKUK.....0002.HUKUK.00000"), strict=True))

    assert satir["DosyaNo"] == "2005.017"
    # Müvekkiller alfabetik, `; ` ile; THIRD taraflar hiçbir kolona girmez
    assert satir["Müvekkil(ler)"] == "Ak Sigorta A.Ş.; Murat Özcan Dr."
    assert satir["Karşı Taraf"] == "Songül Kaya"
    assert satir["Ana Tür"] == "Hukuk"
    assert satir["Yargı Birimi"] == "Asliye Hukuk"
    assert satir["Esas"] == "2024/123"
    assert satir["Mahkeme"] == "İstanbul 6. Tüketici Mahkemesi"
    assert satir["Dava Tarihi"] == datetime(2024, 3, 1)
    assert satir["İş Kabul"] == datetime(2024, 2, 10)
    assert satir["Sorumlu Avukat"] == "Av. Ayşe Demir"
    # sqlite tz bilgisini düşürür (naive 09:30 döner); tz dönüşümü aşağıda ayrı test
    assert satir["Son Güncelleme"] == datetime(2026, 9, 1, 9, 30)


def test_tzli_zaman_tr_saatine_cevrilip_tzsiz_yazilir():
    """Postgres tz'li döner; Excel saat dilimi bilmez → TR saatine çevrilip tz atılır."""
    utc = datetime(2026, 9, 1, 9, 30, 15, 123456, tzinfo=timezone.utc)
    assert fdk._excel_zamani(utc) == datetime(2026, 9, 1, 12, 30, 15)
    assert fdk._excel_zamani(datetime(2026, 9, 1, 9, 30)) == datetime(2026, 9, 1, 9, 30)
    assert fdk._excel_zamani(date(2026, 9, 1)) == date(2026, 9, 1)
    assert fdk._excel_zamani(None) is None


def test_ozet_sayfasi_tur_ve_sayfa_sayilarini_veriyor(dolu_db, tmp_path):
    hedef = tmp_path / "rapor.xlsx"
    fdk.raporu_uret(dolu_db, hedef)
    basliklar, satirlar = _sayfa(load_workbook(hedef)["OZET"])
    assert basliklar == ["Ana Tür", "Sayfa", "Kart"]
    sayilar = {(s[0], s[1]): s[2] for s in satirlar}
    assert sayilar[("Hukuk", "MALPRAKTIS_ADAYI")] == 1
    assert sayilar[("(boş)", "MALPRAKTIS_ADAYI")] == 1
    assert sayilar[("İCRA", "KAPSAM_DISI_TUR")] == 1
    assert sayilar[("TOPLAM", "MALPRAKTIS_ADAYI")] == 2
    assert sayilar[("TOPLAM", "KAPSAM_DISI_TUR")] == 1
    assert sayilar[("TOPLAM", None)] == 3


# ── dosya adı + salt okunurluk ────────────────────────────────────────────────

def test_varsayilan_dosya_adi_sozlesmesi(tmp_path):
    assert fdk.varsayilan_dosya_adi(date(2026, 9, 8)) == "HUKDOK_DERDEST_KARTLAR_2026-09-08.xlsx"
    assert re.fullmatch(r"HUKDOK_DERDEST_KARTLAR_\d{4}-\d{2}-\d{2}\.xlsx", fdk.varsayilan_dosya_adi())
    # --out dizinse varsayılan ad içine düşer, dosyaysa aynen
    assert fdk.cikti_yolu(str(tmp_path)).name == fdk.varsayilan_dosya_adi()
    assert fdk.cikti_yolu(str(tmp_path / "x.xlsx")) == tmp_path / "x.xlsx"
    assert fdk.cikti_yolu(None) == Path(fdk.varsayilan_dosya_adi())


def test_script_salt_okunur():
    """Kabul kriteri: kaynakta commit(/add(/delete(/update( geçmez."""
    kaynak = Path(fdk.__file__).read_text(encoding="utf-8")
    for yasak in ("commit(", "add(", "delete(", "update("):
        assert yasak not in kaynak, yasak


def test_bos_veritabaninda_uc_sayfa_yine_yazilir(SessionFactory, tmp_path):
    hedef = tmp_path / "alt" / "bos.xlsx"
    ozet = fdk.raporu_uret(SessionFactory, hedef)
    assert ozet["toplam"] == 0
    wb = load_workbook(hedef)
    assert wb.sheetnames == ["MALPRAKTIS_ADAYI", "KAPSAM_DISI_TUR", "OZET"]
    _, satirlar = _sayfa(wb["OZET"])
    assert satirlar[-1] == ("TOPLAM", None, 0)
