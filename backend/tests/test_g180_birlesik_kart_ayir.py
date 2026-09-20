"""G180 — Birleşik kartları geri ayırma (`scripts/birlesik_kart_ayir.py`): farklı
tür/esas föy grubu yeni karta, `AYRISTIRILAN` ilişkisi, tarihçe, idempotency.

**TEST VERİSİ KURALI (A.2 dersi):** Ek-3 sentetik (openpyxl), ham satırlar
elde yazılmış sözlükler (ekibin başlıklarıyla).
"""
import openpyxl
import pytest
from sqlalchemy import create_engine, event, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import models
from database import _MIGRATIONS, Base
from scripts import birlesik_kart_ayir as bka
from tests.test_g064_aktarim_cekirdek import _kart


@pytest.fixture()
def db_env():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _fk_ac(dbapi_connection, _record):
        dbapi_connection.isolation_level = None
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    @event.listens_for(engine, "begin")
    def _begin(conn):
        conn.exec_driver_sql("BEGIN")

    Base.metadata.create_all(engine)
    with engine.begin() as conn:
        for op in _MIGRATIONS:
            if op[0] == "index" and op[1] in ("case_foys", "case_esas_numbers"):
                for sql in op[2]:
                    conn.execute(text(sql))
    yield sessionmaker(bind=engine, autocommit=False, autoflush=False)
    engine.dispose()


def _ham(sistem_no, dosya_no, tur, esas, mahkeme, *, muvekkil="Quick Sigorta A.Ş.", karsi="Ayşe Yılmaz",
         tip="Sigorta", durum="Aktif", konu="Tazminat", tarih="2026-03-01"):
    return {"SistemNo": sistem_no, "DosyaNo": dosya_no, "Klasör No": "TKU-9", "Müvekkil": muvekkil,
            "Müvekkil Tipi": tip, "Karşı Taraf": karsi, "Yerel Mahkeme": mahkeme, "Esas": esas,
            "Dava Tarihi": tarih, "Ana Tür": tur, "Durum": durum, "Dava Konusu": konu}


def _foy(db, case, sistem_no, ham=None, onceki=None):
    f = models.CaseFoy(sistem_no=sistem_no, case_id=case.id, ham_veri=ham, onceki_tracking_no=onceki)
    db.add(f)
    db.flush()
    return f


def _birlesik_kart(db_env):
    """Ek-3 › 03 deseni: Hukuk kartı altında arabuluculuk föyü + dava föyü (aynı DosyaNo)."""
    db = db_env()
    try:
        k = _kart(db, "S4.QUICK......0453.HUKUK.00000", "2.455.00", file_type="Hukuk", esas_no="2024/528",
                  court="Ankara 2. Asliye Hukuk Mahkemesi")
        db.add(models.CaseParty(case_id=k.id, name="Quick Sigorta A.Ş.", role="Müvekkil", party_type="CLIENT"))
        db.add(models.CaseDocument(case_id=k.id, original_filename="d.pdf", stored_filename="d.pdf"))
        _foy(db, k, "ARB-16361", _ham("ARB-16361", "2.455.00", "ARABULUCULUK", "2025/48948",
                                       "ANKARA ARABULUCULUK BÜROSU"))
        _foy(db, k, "H-16477", _ham("H-16477", "2.455.00", "HUKUK", "2024/528", "Ankara 2. Asliye Hukuk Mahkemesi"))
        db.commit()
        return k.id
    finally:
        db.close()


# ─── Gruplama ────────────────────────────────────────────────────────────────

def test_foy_satiri_ham_veriden_alanlari_taniyor():
    foy = models.CaseFoy(sistem_no="ARB-1", case_id=1,
                         ham_veri=_ham("ARB-1", "2.1.00", "ARABULUCULUK", "2026/12 ", "X Bürosu"))
    satir = bka.foy_satiri(foy)
    assert satir is not None
    assert satir.degerler["ana_tur"] == "ARABULUCULUK" and satir.degerler["dosya_no"] == "2.1.00"
    assert bka.grup_anahtari(satir) == ("Arabuluculuk", bka.esas_anahtari("2026/12"))
    assert bka.foy_satiri(models.CaseFoy(sistem_no="X", case_id=1, ham_veri=None)) is None


def test_kendi_grubu_kartin_tur_ve_esasiyla_secilir():
    kart = models.Case(tracking_no="K", file_type="Hukuk", esas_no="2024/528")
    f1 = models.CaseFoy(sistem_no="ARB-1", case_id=1, ham_veri=_ham("ARB-1", "2.1", "ARABULUCULUK", "2025/1", "B"))
    f2 = models.CaseFoy(sistem_no="H-1", case_id=1, ham_veri=_ham("H-1", "2.1", "HUKUK", "2024/528", "M"))
    f3 = models.CaseFoy(sistem_no="H-2", case_id=1, ham_veri=_ham("H-2", "2.1", "HUKUK", "2025/9", "M"))
    kendi, gruplar, hamsiz = bka.gruplari_bul(kart, [f1, f2, f3])
    assert kendi == ("Hukuk", bka.esas_anahtari("2024/528")) and len(gruplar) == 3 and hamsiz == []
    # kartın esası hiçbir grupta yoksa: kartın türündeki en büyük grup
    kart.esas_no = "2019/1"
    kendi, _, _ = bka.gruplari_bul(kart, [f1, f2, f3])
    assert kendi[0] == "Hukuk"


# ─── Uçtan uca ───────────────────────────────────────────────────────────────

def test_arabuluculuk_foyu_yeni_karta_iliski_ve_tarihce(db_env):
    kart_id = _birlesik_kart(db_env)

    sonuc = bka.kos(db_env, kart_idler=[kart_id], apply=True, kim="test")
    assert sonuc.sayim("YAPILDI") == 1 and sonuc.yeni_kart_sayisi == 1

    db = db_env()
    try:
        yeni_id, tracking, sistem_nolar = sonuc.kalemler[0].yeni_kartlar[0]
        yeni = db.get(models.Case, yeni_id)
        assert sistem_nolar == ["ARB-16361"]
        assert yeni.file_type == "Arabuluculuk" and yeni.esas_no == "2025/48948"
        assert yeni.klasor_no_2 == "2.455.00" and ".ARABU." in tracking and yeni.deleted_at is None
        assert yeni.notes.startswith(f"#{kart_id} ")
        # föy taşındı, müvekkil bağı yeni karttaki CLIENT'a
        foy = db.query(models.CaseFoy).filter_by(sistem_no="ARB-16361").one()
        assert foy.case_id == yeni_id
        assert foy.case_party_id is not None and db.get(models.CaseParty, foy.case_party_id).case_id == yeni_id
        assert {p.party_type for p in yeni.parties} == {"CLIENT", "COUNTER"}
        # dava föyü ve belge kalan kartta
        assert db.query(models.CaseFoy).filter_by(sistem_no="H-16477").one().case_id == kart_id
        assert db.query(models.CaseDocument).filter_by(case_id=kart_id).count() == 1
        # ilişki + tarihçe
        iliski = db.query(models.CaseRelation).one()
        assert (iliski.source_case_id, iliski.target_case_id, iliski.relation_type) == (kart_id, yeni_id, "AYRISTIRILAN")
        assert db.query(models.CaseHistory).filter_by(field_name="kart_ayirma").count() == 2
        # esas tarihçesi tek yazma yolundan
        assert db.query(models.CaseEsasNumber).filter_by(case_id=yeni_id, is_current=True).count() == 1
    finally:
        db.close()

    tekrar = bka.kos(db_env, kart_idler=[kart_id], apply=True, kim="test")
    assert tekrar.sayim("ATLANDI") == 1 and tekrar.yeni_kart_sayisi == 0
    db = db_env()
    try:
        assert db.query(models.CaseRelation).count() == 1
    finally:
        db.close()


def test_ayni_tur_farkli_esas_da_ayrilir(db_env):
    """14246 deseni: iki Hukuk föyü, farklı esas/mahkeme → kartın esasındaki kalır, öteki yeni karta."""
    db = db_env()
    try:
        k = _kart(db, "S7.F_YESILTAS.0001.HUKUK.00000", "6.5034.00", file_type="Hukuk", esas_no="2020/25")
        _foy(db, k, "H-11742", _ham("H-11742", "6.5034.00", "HUKUK", "2020/25", "İstanbul Anadolu 2. Tüketici Mahkemesi",
                                     muvekkil="Sompo Sigorta A.Ş."))
        _foy(db, k, "H-11743", _ham("H-11743", "6.5034.00", "HUKUK", "2024/24", "İstanbul Anadolu 9. Tüketici Mahkemesi",
                                     muvekkil="Sompo Sigorta A.Ş."))
        db.commit()
        k_id = k.id
    finally:
        db.close()
    sonuc = bka.kos(db_env, kart_idler=[k_id], apply=True, kim="test")
    assert sonuc.sayim("YAPILDI") == 1
    db = db_env()
    try:
        yeni_id = sonuc.kalemler[0].yeni_kartlar[0][0]
        assert db.query(models.CaseFoy).filter_by(sistem_no="H-11743").one().case_id == yeni_id
        assert db.query(models.CaseFoy).filter_by(sistem_no="H-11742").one().case_id == k_id
        assert db.get(models.Case, yeni_id).esas_no == "2024/24"
    finally:
        db.close()


def test_tek_gruplu_ve_hamsiz_kartlar(db_env):
    db = db_env()
    try:
        tek = _kart(db, "T", "1.1", file_type="Hukuk", esas_no="2024/1")
        _foy(db, tek, "H-1", _ham("H-1", "1.1", "HUKUK", "2024/1", "M"))
        _foy(db, tek, "H-2", _ham("H-2", "1.1", "HUKUK", "2024/1", "M"))
        hamsiz = _kart(db, "HS", "1.2", file_type="Hukuk", esas_no="2024/2")
        _foy(db, hamsiz, "H-3", _ham("H-3", "1.2", "HUKUK", "2024/2", "M"))
        _foy(db, hamsiz, "ARB-3", None)
        sonen = _kart(db, "SN", "1.3", file_type="Hukuk", esas_no="2024/3")
        _foy(db, sonen, "H-4", _ham("H-4", "1.3", "HUKUK", "2024/3", "M"))
        _foy(db, sonen, "ARB-4", _ham("ARB-4", "1.3", "ARABULUCULUK", "2025/1", "B"), onceki="ESKI.KART")
        db.commit()
        idler = [tek.id, hamsiz.id, sonen.id, 999]
    finally:
        db.close()
    sonuc = bka.kos(db_env, kart_idler=idler, apply=True, kim="test")
    assert [k.sonuc for k in sonuc.kalemler] == ["ATLANDI", "RET", "RET", "RET"]
    assert "ham satırı olmayan föy: ARB-3" in sonuc.kalemler[1].aciklama
    assert "sönen kart izi" in sonuc.kalemler[2].aciklama
    db = db_env()
    try:
        assert db.query(models.Case).count() == 3 and db.query(models.CaseRelation).count() == 0
    finally:
        db.close()


def test_kuru_kosu_yazmaz(db_env):
    kart_id = _birlesik_kart(db_env)
    sonuc = bka.kos(db_env, kart_idler=[kart_id], apply=False, kim="test")
    assert sonuc.sayim("YAPILDI") == 1 and sonuc.yeni_kart_sayisi == 1
    db = db_env()
    try:
        assert db.query(models.Case).count() == 1 and db.query(models.CaseRelation).count() == 0
        assert db.query(models.CaseFoy).filter_by(sistem_no="ARB-16361").one().case_id == kart_id
    finally:
        db.close()
    print(bka.ozet_metni(sonuc, apply=False))


# ─── Ek-3 okuma ──────────────────────────────────────────────────────────────

def test_ek3_kartlari_iki_sayfadan(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = bka.EK3_BIRLESTIRME_SAYFASI
    ws.append(["Kart", "Ofis dosya no", "Kart türü", "Kart mahkemesi", "Kart esası", "Föy sayısı", "Föylerin künyesi",
               "Durum", "Cevabınız"])
    ws.append([2468, "", "", "", "", 2, "", "", ""])
    ws.append([4177, "", "", "", "", 2, "", "", ""])
    ws2 = wb.create_sheet(bka.EK3_TASIMA_SAYFASI)
    ws2.append(["SistemNo", "DosyaNo (bizde)", "Föy türü", "Föy mahkemesi", "Föy esası", "Bağlı olduğu kart",
                "Ofis dosya no", "Kart türü", "Kart mahkemesi", "Kart esası", "Karttaki öteki föyler",
                "Önerdiğimiz kart", "Ofis dosya no ", "Açıklama", "Durum", "Cevabınız"])
    ws2.append(["H-16799", "", "", "", "", 4957, "", "", "", "", "", None, "", "", "SORU", ""])     # öneri boş → liste
    ws2.append(["H-16477", "", "", "", "", 14328, "", "", "", "", "", 14326, "", "", "", ""])      # öneri dolu → G179
    ws2.append(["H-16925", "", "", "", "", 4177, "", "", "", "", "", None, "", "", "SORU", ""])    # zaten listede
    yol = tmp_path / "ek3.xlsx"
    wb.save(yol)
    assert bka.ek3_kartlari(yol) == [2468, 4177, 4957]


# ─── Müvekkil ayrımı (Ek-6 › 01, kart 14334) ─────────────────────────────────

def _iki_muvekkilli_kart(db_env):
    """14334 deseni: aynı arabuluculuk + aynı dava esası, İKİ müvekkil, dört föy."""
    db = db_env()
    try:
        k = _kart(db, "D1.D_ESINLER..0002.ARABU.00000", "1110.003;1993.001", file_type="Arabuluculuk",
                  esas_no="2026/720", court="Ankara Arabuluculuk Bürosu")
        for sistem_no, dosya_no, tur, esas, mahkeme, muvekkil in (
                ("ARB-16767", "1110.003", "ARABULUCULUK", "2026/720", "Ankara Arabuluculuk Bürosu", "Deniz Esinler Dr."),
                ("ARB-16779", "1993.001", "ARABULUCULUK", "2026/720", "Ankara Arabuluculuk Bürosu", "Aylin Ayrim Dr"),
                ("H-16856", "1110.003", "HUKUK", "2026/91", "Ankara 11. Tüketici Mahkemesi", "Deniz Esinler Dr."),
                ("H-16857", "1993.001", "HUKUK", "2026/91", "Ankara 11. Tüketici Mahkemesi", "Aylin Ayrim Dr")):
            _foy(db, k, sistem_no, _ham(sistem_no, dosya_no, tur, esas, mahkeme,
                                        muvekkil=muvekkil, tip="Doktor"))
        db.commit()
        return k.id
    finally:
        db.close()


def test_muvekkil_ayrimi_dort_foyu_dort_karta_boler(db_env):
    kart_id = _iki_muvekkilli_kart(db_env)
    sonuc = bka.kos(db_env, kart_idler=[kart_id], apply=True, kim="test", muvekkil_ayrimi=True)
    assert sonuc.sayim("YAPILDI") == 1 and sonuc.yeni_kart_sayisi == 3

    db = db_env()
    try:
        # Kartın künyesi kendi müvekkilinde kaldı: isim bloğu D_ESINLER olan ARB grubu
        kalan_foyler = {f.sistem_no for f in db.query(models.CaseFoy).filter_by(case_id=kart_id)}
        assert kalan_foyler == {"ARB-16767"}
        kartlar = {}
        for yeni_id, tracking, sistem_nolar in sonuc.kalemler[0].yeni_kartlar:
            assert len(sistem_nolar) == 1
            kartlar[sistem_nolar[0]] = db.get(models.Case, yeni_id)
            assert tracking.split(".")[1] in ("A_AYRIM", "D_ESINLER")
        assert set(kartlar) == {"ARB-16779", "H-16856", "H-16857"}
        # Klasör no müvekkil başına: kartın iki numaralı listesi yeni karta TAŞINMAZ
        assert kartlar["ARB-16779"].klasor_no_2 == "1993.001"
        assert kartlar["H-16856"].klasor_no_2 == "1110.003"
        assert kartlar["H-16857"].esas_no == "2026/91" and kartlar["H-16857"].file_type == "Hukuk"
        assert db.query(models.CaseRelation).filter_by(source_case_id=kart_id,
                                                       relation_type="AYRISTIRILAN").count() == 3
    finally:
        db.close()

    tekrar = bka.kos(db_env, kart_idler=[kart_id], apply=True, kim="test", muvekkil_ayrimi=True)
    assert tekrar.sayim("ATLANDI") == 1 and tekrar.yeni_kart_sayisi == 0


def test_muvekkil_ayrimi_kapaliyken_davranis_degismez(db_env):
    """Bayraksız: aynı esastaki iki müvekkil TEK grup — dava grubu iki DosyaNo'ya
    bölündüğü için kart RET (bayrağın neden gerektiğinin ölçüsü)."""
    kart_id = _iki_muvekkilli_kart(db_env)
    sonuc = bka.kos(db_env, kart_idler=[kart_id], apply=True, kim="test")
    assert sonuc.sayim("RET") == 1 and "DosyaNo'ya bölünüyor" in sonuc.kalemler[0].aciklama
    db = db_env()
    try:
        assert db.query(models.Case).count() == 1
        assert db.query(models.CaseFoy).filter_by(case_id=kart_id).count() == 4
    finally:
        db.close()
