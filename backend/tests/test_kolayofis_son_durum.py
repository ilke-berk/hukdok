"""KolayOfis raporundan dosya son durumu tazeleme (`scripts/kolayofis_son_durum.py`):
başlıksız sütun, yazım düzeltmesi, yer tutucu/boş hücre, föy→kart çözümü, çelişki,
panel listesi düzeltmesi, kuru koşu ve idempotentlik.

**TEST VERİSİ KURALI:** ekibin gerçek raporu repoya girmez; sayfa openpyxl ile SENTETİK
üretilir (17.09 çıktısının başlık düzeniyle: son durum sütunu BAŞLIKSIZ).
"""
import openpyxl
import pytest

import models
from scripts import kolayofis_son_durum as ko
from tests.test_g064_aktarim_cekirdek import _kart
from tests.test_g179_ekip_cevabi_1209 import _foy, db_env  # noqa: F401

# 17.09 çıktısının başlıkları; 16. sütun (son durum) BAŞLIKSIZ gelir.
BASLIKLAR = ["Dosya - Föy Bilgileri", "SistemNo", "DosyaNo", "Klasör No", "Tarafımız", "Müvekkil",
             "Karşı Taraf", "Yerel Mahkeme", "Esas", "Dava Tarihi", "Ana Tür", "Durum", "Dava Konusu",
             "Dava Türü Alt Kırılımı", "Buro Özel Türü", None, "Yerel Mahkeme Karar Durumu", "Dosya İlgilisi"]


def _rapor_yaz(yol, satirlar, *, basliklar=None):
    """satirlar: (SistemNo, son durum) — son durum None ise hücre boş bırakılır."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(basliklar if basliklar is not None else BASLIKLAR)
    for sistem_no, son_durum in satirlar:
        satir = [None] * len(BASLIKLAR)
        satir[1] = sistem_no
        satir[15] = son_durum
        ws.append(satir)
    wb.save(yol)
    return yol


@pytest.fixture()
def rapor(tmp_path):
    def yaz(satirlar, **kw):
        return _rapor_yaz(tmp_path / "rapor.xlsx", satirlar, **kw)
    return yaz


def _liste(db, *adlar):
    for i, ad in enumerate(adlar):
        db.add(models.FileStatus(code=ko._karar_kodu(ad).replace("_", "-"), name=ad, active=True, sequence=i))
    db.flush()


# ─── Rapor okuma ─────────────────────────────────────────────────────────────

def test_baslıksiz_sutun_komsulardan_bulunur(rapor):
    yol = rapor([("H-1", "Bilirkişide")])
    assert ko.raporu_oku(yol) == [("H-1", "Bilirkişide")]


def test_baslikli_sutun_da_okunur(rapor):
    basliklar = list(BASLIKLAR)
    basliklar[15] = "Son Durum"
    yol = rapor([("H-1", "Temyizde")], basliklar=basliklar)
    assert ko.raporu_oku(yol) == [("H-1", "Temyizde")]


def test_sutun_bulunamazsa_hata(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["SistemNo", "Müvekkil"])
    wb.save(tmp_path / "eksik.xlsx")
    with pytest.raises(ValueError, match="son durum sütunu"):
        ko.raporu_oku(tmp_path / "eksik.xlsx")


@pytest.mark.parametrize("ham, beklenen", [
    ("Delliller Toplanıyor", "Deliller Toplanıyor"),      # ekibin başlık listesindeki yazım hatası
    ("Bekletici Mesele/ceza-hukuk Dosyası", "Bekletici Mesele/Ceza-Hukuk Dosyası"),
    ("Islah ", "Islah"),                                  # sondaki boşluk
    ("  Bilirkişide  ", "Bilirkişide"),
    ("Lütfen Seçiniz", None),                             # yer tutucu → yazılmaz
    ("Kapalı", None),                                     # 22.09: karar durumu havuzundan sızma → yazılmaz
    ("", None),
    (None, None),
])
def test_deger_normalizasyonu(ham, beklenen):
    assert ko._degeri_normalize(ham) == beklenen


def test_yazilmayan_satirlar_rapordan_dusurulur(rapor):
    yol = rapor([("H-1", "Delliller Toplanıyor"), ("H-2", "Lütfen Seçiniz"), ("H-3", None), (None, "Tanık")])
    assert ko.raporu_oku(yol) == [("H-1", "Deliller Toplanıyor")]


# ─── Son durum tazeleme ──────────────────────────────────────────────────────

def test_kart_tazelenir_ve_tarihce_dusuer(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Bilirkişide", "Temyizde")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="Bilirkişide")
    _foy(db, kart, "H-1")
    db.commit()
    kart_id = kart.id
    db.close()

    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Temyizde")]), apply=True, kim="test")
    assert sonuc.sayim("son_durum", "YAPILDI") == 1

    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Temyizde"
    tarihce = db.query(models.CaseHistory).filter_by(field_name="dosya_son_durumu").one()
    assert (tarihce.old_value, tarihce.new_value) == ("Bilirkişide", "Temyizde")
    assert tarihce.changed_by == ko.DEGISTIREN and "H-1" in tarihce.source
    db.close()


def test_bos_kutu_dolar(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Ön İnceleme")
    kart = _kart(db, "T1", "1.1")
    _foy(db, kart, "H-1")
    db.commit()
    kart_id = kart.id
    db.close()

    ko.kos(db_env, rapor=rapor([("H-1", "Ön İnceleme")]), apply=True)
    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Ön İnceleme"
    db.close()


def test_yer_tutucu_dolu_kutuyu_bosaltmaz(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Bilirkişide")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="Bilirkişide")
    _foy(db, kart, "H-1")
    db.commit()
    kart_id = kart.id
    db.close()

    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Lütfen Seçiniz"), ("H-1", None)]), apply=True)
    assert sonuc.sayim("son_durum", "YAPILDI") == 0

    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Bilirkişide"
    assert db.query(models.CaseHistory).filter_by(field_name="dosya_son_durumu").count() == 0
    db.close()


def test_celiskili_kart_atlanir(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Bilirkişide", "Temyizde", "Kesin Lehe")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="Bilirkişide")
    _foy(db, kart, "H-1")
    _foy(db, kart, "H-2")
    db.commit()
    kart_id = kart.id
    db.close()

    # İki dava-seviyesi değer: seviye kuralı da çözemez (22.09) — gerçek çelişki.
    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Temyizde"), ("H-2", "Bilirkişide")]), apply=True)
    assert sonuc.sayim("son_durum", "CELISKI") == 1
    assert sonuc.sayim("son_durum", "YAPILDI") == 0
    kalem = next(k for k in sonuc.kalemler if k.sonuc == "CELISKI")
    assert "H-1" in kalem.aciklama and "H-2" in kalem.aciklama

    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Bilirkişide"
    db.close()


def test_ayni_degeri_soyleyen_coklu_foy_celiski_degil(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Temyizde")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="Bilirkişide")
    _foy(db, kart, "H-1")
    _foy(db, kart, "H-2")
    db.commit()
    kart_id = kart.id
    db.close()

    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Temyizde"), ("H-2", "Temyizde")]), apply=True)
    assert sonuc.sayim("son_durum", "CELISKI") == 0

    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Temyizde"
    db.close()


def test_bizde_olmayan_foy_ret(db_env, rapor):  # noqa: F811
    sonuc = ko.kos(db_env, rapor=rapor([("H-YOK", "Temyizde")]), apply=True)
    assert sonuc.sayim("son_durum", "RET") == 1


def test_silinmis_kartin_foyu_ret(db_env, rapor):  # noqa: F811
    from sqlalchemy import func

    db = db_env()
    _liste(db, "Temyizde")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="Bilirkişide")
    _foy(db, kart, "H-1")
    kart.deleted_at = func.now()
    db.commit()
    db.close()

    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Temyizde")]), apply=True)
    assert sonuc.sayim("son_durum", "RET") == 1
    assert sonuc.sayim("son_durum", "YAPILDI") == 0


# ─── Panel listesi ve kart yazım birliği ─────────────────────────────────────

def test_liste_duzeltmesi(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Bilirkişide", "İstinafda")          # 20.09 koşusunun bıraktığı (yanlış) yazım
    db.commit()
    db.close()

    ko.kos(db_env, rapor=rapor([]), apply=True)
    db = db_env()
    adlar = {s.name for s in db.query(models.FileStatus).all()}
    assert "İstinafta" in adlar and "İstinafda" not in adlar
    assert {"Soruşturma", "Derdest", "İnfaz"} <= adlar
    assert "Kapalı" not in adlar                    # 22.09: sızma değeri artık listeye EKLENMEZ
    kodlar = [s.code for s in db.query(models.FileStatus).all()]
    assert len(kodlar) == len(set(kodlar))          # kod tekilliği korunur
    db.close()


def test_listedeki_kapali_silinmez_pasife_cekilir(db_env, rapor):  # noqa: F811
    """23.09: `Kapalı` son durum değil — satır SİLİNMEZ (eski tarihçe okunur), seçilemez olur."""
    db = db_env()
    _liste(db, "Kapalı")
    db.commit()
    db.close()
    ko.kos(db_env, rapor=rapor([]), apply=True)
    db = db_env()
    satir = db.query(models.FileStatus).filter_by(name="Kapalı").one()
    assert satir.active is False
    db.close()
    ikinci = ko.kos(db_env, rapor=rapor([]), apply=True)
    assert ikinci.sayim("liste", "YAPILDI") == 0


def test_liste_dogru_yazim_varsa_eski_satir_silinmez(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "İstinafta", "İstinafda")
    db.commit()
    db.close()

    sonuc = ko.kos(db_env, rapor=rapor([]), apply=True)
    assert sonuc.sayim("liste", "ATLANDI") >= 1
    db = db_env()
    assert {"İstinafta", "İstinafda"} <= {s.name for s in db.query(models.FileStatus).all()}
    db.close()


def test_kart_yazim_birligi(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Bekletici Mesele/Ceza-Hukuk Dosyası")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="Bekletici Mesele/ceza-hukuk Dosyası")
    db.commit()
    kart_id = kart.id
    db.close()

    ko.kos(db_env, rapor=rapor([]), apply=True)
    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Bekletici Mesele/Ceza-Hukuk Dosyası"
    assert db.query(models.CaseHistory).filter_by(field_name="dosya_son_durumu").count() == 1
    db.close()


def test_istinafda_kartlari_istinafta_olur(db_env, rapor):  # noqa: F811
    """22.09: ekip 'İstinafda' yazımını geri aldı — 20.09'un 114 + paketin 382 kartı tarihçeli döner."""
    db = db_env()
    _liste(db, "İstinafda")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="İstinafda")
    db.commit()
    kart_id = kart.id
    db.close()

    sonuc = ko.kos(db_env, rapor=rapor([]), apply=True, kim="test")
    assert sonuc.sayim("kart_yazim", "YAPILDI") == 1
    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "İstinafta"
    tarihce = db.query(models.CaseHistory).filter_by(field_name="dosya_son_durumu").one()
    assert (tarihce.old_value, tarihce.new_value) == ("İstinafda", "İstinafta")
    assert {s.name for s in db.query(models.FileStatus).all()} >= {"İstinafta"}
    db.close()


# ─── Föy seviyesi kuralı (22.09) ─────────────────────────────────────────────

def _celiskili_kart(db_env, *foyler, son_durum="Lexis Rapor Gönderildi"):  # noqa: F811
    db = db_env()
    _liste(db, "Bilirkişide", "Temyizde", "Kesin Lehe", "Lexis Rapor Gönderildi")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu=son_durum)
    for sistem_no in foyler:
        _foy(db, kart, sistem_no)
    db.commit()
    kart_id = kart.id
    db.close()
    return kart_id


def test_tek_dava_seviyesi_deger_karti_alir(db_env, rapor):  # noqa: F811
    kart_id = _celiskili_kart(db_env, "H-1", "H-2", "H-3")
    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Lexis Rapor Gönderildi"), ("H-2", "Bilirkişide"),
                                         ("H-3", "Kesin Lehe")]), apply=True, kim="test")
    assert sonuc.sayim("son_durum", "YAPILDI") == 1 and sonuc.sayim("son_durum", "CELISKI") == 0
    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Bilirkişide"
    tarihce = db.query(models.CaseHistory).filter_by(field_name="dosya_son_durumu").one()
    assert "kademe kuralı" in tarihce.source and "H-2" in tarihce.source
    db.close()


@pytest.mark.parametrize("satirlar, gerekce", [
    ([("H-1", "Bilirkişide"), ("H-2", "Temyizde")], "aynı kademede"),                     # gerçek çelişki
    ([("H-1", "Lexis Rapor Gönderildi"), ("H-2", "Dava Açılması Bekleniyor")], "aynı kademede"),  # iki hizmet aşaması
    ([("H-1", "Bilirkişide"), ("H-2", "Uydurma Değer")], "sınıflanmamış değer"),
])
def test_cozulemeyen_karisimlar_celiski(db_env, rapor, satirlar, gerekce):  # noqa: F811
    kart_id = _celiskili_kart(db_env, "H-1", "H-2", son_durum="Ön İnceleme")
    sonuc = ko.kos(db_env, rapor=rapor(satirlar), apply=True)
    assert sonuc.sayim("son_durum", "CELISKI") == 1 and sonuc.sayim("son_durum", "YAPILDI") == 0
    kalem = next(k for k in sonuc.kalemler if k.sonuc == "CELISKI")
    assert gerekce in kalem.aciklama and "H-1" in kalem.aciklama
    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Ön İnceleme"
    db.close()


def test_kapali_foy_degeri_yazilmaz_ama_kalani_engellemez(db_env, rapor):  # noqa: F811
    kart_id = _celiskili_kart(db_env, "H-1", "H-2", son_durum="Kapalı")
    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Kapalı"), ("H-2", "Kesin Lehe")]), apply=True)
    assert sonuc.sayim("son_durum", "YAPILDI") == 1
    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Kesin Lehe"
    db.close()


def test_inceleme_kuyrugundaki_kart_bekler(db_env, rapor, monkeypatch):  # noqa: F811
    kart_id = _celiskili_kart(db_env, "H-1", "H-2", son_durum="Ön İnceleme")
    monkeypatch.setattr(ko, "INCELEME_KUYRUGU_FOYLERI", frozenset({"H-2"}))
    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Temyizde"), ("H-2", "Temyizde")]), apply=True)
    assert sonuc.sayim("son_durum", "BEKLIYOR") == 1 and sonuc.sayim("son_durum", "YAPILDI") == 0
    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Ön İnceleme"     # tekil değerde bile yazılmaz
    db.close()
    assert "bekliyor" in ko.ozet_metni(sonuc, apply=True) and "BEKLIYOR" in ko.ozet_metni(sonuc, apply=True)


@pytest.mark.parametrize("degerler, beklenen", [
    ({"Temyizde": ["H-1"]}, "Temyizde"),
    ({"Temyizde": ["H-1"], "Lexis Rapor Gönderildi": ["H-2"], "Kesin Aleyhe": ["H-3"]}, "Temyizde"),
    ({"Temyizde": ["H-1"], "İstinafta": ["H-2"]}, None),
    ({"Temyizde": ["H-1"], "İade": ["H-2"]}, "Temyizde"),                       # 23.09: İade föy seviyesi
    ({"Lexis Rapor Gönderildi": ["H-1"], "Kesin Lehe": ["H-2"]}, "Kesin Lehe"),  # sonuç hizmete baskın
    ({"Lexis Rapor Gönderildi": ["H-1"], "İstifa": ["H-2"]}, "İstifa"),          # büro sonu hizmete baskın
    ({"İstifa": ["H-1"], "Kesin Aleyhe": ["H-2"]}, "Kesin Aleyhe"),              # sonuç büro sonuna baskın
    ({"Kesin Lehe": ["H-1"], "Kesin Aleyhe": ["H-2"]}, "Kesin Aleyhe"),          # karışık sonuç: en ağırı
    ({"Kesin Lehe": ["H-1"], "Sulh İle Kapatma": ["H-2"]}, "Sulh İle Kapatma"),
    ({"Kesin Aleyhe": ["H-1"], "Sulh İle Kapatma": ["H-2"], "Kesin Lehe": ["H-3"]}, "Kesin Aleyhe"),
    ({"İnfaz": ["H-1"], "Kesin Lehe": ["H-2"]}, None),                            # ağırlık listesi dışı
    ({"İstifa": ["H-1"], "İade": ["H-2"]}, None),                                 # iki büro sonu
    ({"Davaya Dönüştü": ["H-1"], "Dava Açılması Bekleniyor": ["H-2"]}, "Davaya Dönüştü"),
])
def test_kart_degeri_kurali(degerler, beklenen):
    assert ko.kart_degeri(degerler)[0] == beklenen


# ─── Kuru koşu ve idempotentlik ──────────────────────────────────────────────

def test_kuru_kosu_yazmaz(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Bilirkişide", "Temyizde")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="Bilirkişide")
    _foy(db, kart, "H-1")
    db.commit()
    kart_id = kart.id
    db.close()

    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Temyizde")]), apply=False)
    assert sonuc.sayim("son_durum", "YAPILDI") == 1      # rapor yine basılır

    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Bilirkişide"
    assert db.query(models.CaseHistory).count() == 0
    assert {s.name for s in db.query(models.FileStatus).all()} == {"Bilirkişide", "Temyizde"}
    db.close()


def test_ikinci_kosu_degisiklik_uretmez(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Bilirkişide", "İstinafta")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="Bilirkişide")
    _foy(db, kart, "H-1")
    db.commit()
    db.close()

    yol = rapor([("H-1", "Delliller Toplanıyor")])
    ko.kos(db_env, rapor=yol, apply=True)
    ikinci = ko.kos(db_env, rapor=yol, apply=True)
    assert ikinci.sayim("son_durum", "YAPILDI") == 0
    assert ikinci.sayim("liste", "YAPILDI") == 0
    assert ikinci.sayim("kart_yazim", "YAPILDI") == 0

    db = db_env()
    assert db.query(models.CaseHistory).filter_by(field_name="dosya_son_durumu").count() == 1
    db.close()


def test_ozet_metni_celiskiyi_gosterir(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Bilirkişide", "Temyizde", "Kesin Lehe")
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="Bilirkişide")
    _foy(db, kart, "H-1")
    _foy(db, kart, "H-2")
    db.commit()
    db.close()

    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Temyizde"), ("H-2", "Bilirkişide")]), apply=False)
    metin = ko.ozet_metni(sonuc, apply=False)
    assert "KURU KOŞU" in metin and "CELISKI" in metin


def test_oncelik_sozlugu_ekibin_41_degeri():
    assert len(ko.ONCELIK) == 41 and set(ko.ONCELIK.values()) == {1, 2, 3, 4}
    assert "Kapalı" not in ko.ONCELIK and set(ko.AGIRLIK) <= set(ko.ONCELIK)


# ─── Ekip cevabı 23.09: ek katmanı, Kapalı boşaltma, İnfaz İcradan ────────────

def _ek_yaz(yol, foyler, t3=()):
    """foyler: (SistemNo, bugünkü değer); t3: (föy listesi, değer, durum)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FOYLER"
    ws.append(["Kart ID", "SistemNo", "Ek-2 Son Durum", "Bugünkü Son Durum", "Değişti", "Seviye", "Durum"])
    for sistem_no, deger in foyler:
        ws.append([1, sistem_no, None, deger, None, None, "Aktif"])
    ws3 = wb.create_sheet("T3_GUNCEL")
    ws3.append(["Kart ID", "Föyler", "Bugünkü değer", "Durum"])
    for satir in t3:
        ws3.append(["1", *satir])
    wb.save(yol)
    return yol


def test_ek_oku_foyler_ve_kapanan_t3(tmp_path):
    yol = _ek_yaz(tmp_path / "ek.xlsx", [("H-1", "Kesin Lehe"), ("H-2", "Kapalı")],
                  t3=[("H-5, H-6", "İstinafta", "KAPANDI — bütün föyler aynı değerde"),
                      ("H-7, H-8", "Bilirkişide ↔ Temyizde", "AÇIK — hâlâ okumada")])
    assert ko.ek_oku(yol) == {"H-1": "Kesin Lehe", "H-2": None, "H-5": "İstinafta", "H-6": "İstinafta"}


def test_ek_eksik_sayfa_hata(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "KARTLAR"
    wb.save(tmp_path / "ek.xlsx")
    with pytest.raises(ValueError, match="T3_GUNCEL"):
        ko.ek_oku(tmp_path / "ek.xlsx")


def test_satirlari_birlestir_ek_kazanir_bos_dusurur():
    birlesik = dict(ko.satirlari_birlestir([("H-1", "Bilirkişide"), ("H-2", "İstinafta"), ("H-3", "Temyizde")],
                                           {"H-1": "Tanık", "H-2": None, "H-9": "Kesin Lehe"}))
    assert birlesik["H-1"] == "Tanık" and "H-2" not in birlesik
    assert birlesik["H-3"] == "Temyizde" and birlesik["H-9"] == "Kesin Lehe"


def test_mail_degerleri_yalniz_ek_ile_girer(db_env, rapor, tmp_path):  # noqa: F811
    db = db_env()
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="Bilirkişide")
    _foy(db, kart, "H-16736")
    db.commit()
    kart_id = kart.id
    db.close()
    assert ko.kos(db_env, rapor=rapor([]), apply=True).sayim("son_durum", "YAPILDI") == 0
    ko.kos(db_env, rapor=rapor([]), ek=_ek_yaz(tmp_path / "ek.xlsx", []), apply=True)
    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Bilirkişi Tazminat Raporu Alındı"     # mail §1 (#5009)
    db.close()


def test_ek_katmani_master_degerini_ezer(db_env, rapor, tmp_path):  # noqa: F811
    kart_id = _celiskili_kart(db_env, "H-1", "H-2", son_durum="Ön İnceleme")
    ek = _ek_yaz(tmp_path / "ek.xlsx", [("H-2", "Bilirkişide")])
    sonuc = ko.kos(db_env, rapor=rapor([("H-1", "Bilirkişide"), ("H-2", "Temyizde")]), ek=ek, apply=True)
    assert sonuc.sayim("son_durum", "CELISKI") == 0
    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "Bilirkişide"
    tarihce = db.query(models.CaseHistory).filter_by(field_name="dosya_son_durumu").one()
    assert "ekip cevabı 23.09" in tarihce.source
    db.close()


def test_kapali_kartlar_bosaltilir(db_env, rapor):  # noqa: F811
    db = db_env()
    _liste(db, "Kapalı")
    foysuz = _kart(db, "T1", "1.1", dosya_son_durumu="Kapalı")
    diger = _kart(db, "T2", "1.2", dosya_son_durumu="Kesin Lehe")
    db.commit()
    foysuz_id, diger_id = foysuz.id, diger.id
    db.close()

    sonuc = ko.kos(db_env, rapor=rapor([]), apply=True, kim="test")
    assert sonuc.sayim("kart_bosalt", "YAPILDI") == 1
    db = db_env()
    assert db.get(models.Case, foysuz_id).dosya_son_durumu is None
    assert db.get(models.Case, diger_id).dosya_son_durumu == "Kesin Lehe"
    tarihce = db.query(models.CaseHistory).filter_by(field_name="dosya_son_durumu").one()
    assert (tarihce.old_value, tarihce.new_value) == ("Kapalı", None) and "§5" in tarihce.source
    db.close()
    assert ko.kos(db_env, rapor=rapor([]), apply=True).sayim("kart_bosalt", "YAPILDI") == 0


def test_infaz_icradan_duz_infaz_olur(db_env, rapor):  # noqa: F811
    db = db_env()
    kart = _kart(db, "T1", "1.1", dosya_son_durumu="İnfaz İcradan")
    db.commit()
    kart_id = kart.id
    db.close()
    ko.kos(db_env, rapor=rapor([]), apply=True)
    db = db_env()
    assert db.get(models.Case, kart_id).dosya_son_durumu == "İnfaz"
    db.close()
