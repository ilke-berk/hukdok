"""Ekip cevabı 23.09 (`scripts/ekip_cevabi_2309.py`): TKU birleştirmesiyle sönen kartın
canlandırılması, föy taşıma (klasör no + taraf temizliği), Aktif/Arşiv kuralı (KORUNDU),
kuru koşu ve idempotency.

**TEST VERİSİ KURALI:** ham satırlar elde yazılmış sözlükler (ekibin başlıklarıyla);
rapor/ek sentetik (openpyxl).
"""
import openpyxl
from sqlalchemy import func

import models
from scripts import ekip_cevabi_2309 as ec
from tests.test_g064_aktarim_cekirdek import _kart
from tests.test_g180_birlesik_kart_ayir import _foy, _ham, db_env  # noqa: F401

MAHKEME = "İstanbul 7. Tüketici Mahkemesi"


def _birlesmis_kart(db_env):  # noqa: F811
    """#3418 deseni: H-961 (2021/390) kartına sönen 0204 kartının föyü H-16815 (2025/618) katılmış."""
    db = db_env()
    try:
        sonen = _kart(db, "S7.SOMPO......0204.HUKUK.00000", "6.5170.00", file_type="Hukuk", esas_no="2025/618",
                      court=MAHKEME)
        sonen.deleted_at = func.now()
        sonen.delete_reason = "TKU kart birleştirmesi: #1 ile birleştirildi"
        sonen.active = False
        k = _kart(db, "S7.SOMPO......0101.HUKUK.00000-2", "6.5064.00;6.5170.00", file_type="Hukuk",
                  esas_no="2025/618", court=MAHKEME)
        db.add(models.CaseParty(case_id=k.id, name="Sompo Sigorta A.Ş.", role="Müvekkil", party_type="CLIENT"))
        db.add(models.CaseDocument(case_id=k.id, original_filename="d.pdf", stored_filename="d.pdf"))
        _foy(db, k, "H-961", _ham("H-961", "6.5064.00", "HUKUK", "2021/390", MAHKEME,
                                   muvekkil="Sompo Sigorta A.Ş.", karsi="Ayşe Yılmaz"))
        _foy(db, k, "H-16815", _ham("H-16815", "6.5170.00", "HUKUK", "2025/618", MAHKEME,
                                     muvekkil="Sompo Sigorta A.Ş.", karsi="Ayşe Yılmaz"),
             onceki="S7.SOMPO......0204.HUKUK.00000")
        db.commit()
        return k.id, sonen.id
    finally:
        db.close()


def _tesaduf_kartlari(db_env):  # noqa: F811
    """#14148 deseni: H-3844 başka davanın föyü; davası H-3845'in kartında."""
    db = db_env()
    try:
        kaynak = _kart(db, "S1.M_TURKER...0002.HUKUK.00000", "3.832.00;1193.002.00", file_type="Hukuk",
                       esas_no="2020/611", court="Antalya 2. Tüketici Mahkemesi")
        hedef = _kart(db, "S2.K_YILDIZ...0002.HUKUK.00000", "1641.002.00;9.819.00", file_type="Hukuk",
                      esas_no="2025/243", court="İstanbul Anadolu 10. Asliye Ticaret Mahkemesi")
        db.add_all([models.CaseParty(case_id=kaynak.id, name="Mehmet Türker Dr", role="Müvekkil", party_type="CLIENT"),
                    models.CaseParty(case_id=kaynak.id, name="Özgül Kaçar", role="Karşı Taraf", party_type="COUNTER"),
                    models.CaseParty(case_id=kaynak.id, name="Egehan Coşgun", role="Karşı Taraf", party_type="COUNTER")])
        _foy(db, kaynak, "H-10597", _ham("H-10597", "1193.002.00", "HUKUK", "2020/611", "Antalya 2. Tüketici Mahkemesi",
                                         muvekkil="Mehmet Türker Dr", karsi="Özgül Kaçar", tip="Doktor"))
        _foy(db, kaynak, "H-3844", _ham("H-3844", "3.832.00", "HUKUK", "2020/611",
                                        "İstanbul Anadolu 10. Asliye Ticaret Mahkemesi",
                                        muvekkil="Mehmet Türker Dr", karsi="Egehan Coşgun", tip="Doktor"))
        _foy(db, hedef, "H-3845", _ham("H-3845", "1641.002.00", "HUKUK", "2025/243",
                                       "İstanbul Anadolu 10. Asliye Ticaret Mahkemesi",
                                       muvekkil="Ak Sigorta A.Ş.", karsi="Egehan Coşgun"))
        db.commit()
        return kaynak.id, hedef.id
    finally:
        db.close()


# ─── Canlandırma ─────────────────────────────────────────────────────────────

def test_sonen_kart_canlanir_foy_doner(db_env):  # noqa: F811
    kart_id, sonen_id = _birlesmis_kart(db_env)
    sonuc = ec.kos(db_env, apply=True, kim="test")
    assert sonuc.sayim("canlandirma", "YAPILDI") == 1

    db = db_env()
    kart, sonen = db.get(models.Case, kart_id), db.get(models.Case, sonen_id)
    foy = db.query(models.CaseFoy).filter_by(sistem_no="H-16815").one()
    assert sonen.deleted_at is None and sonen.active is True and sonen.delete_reason is None
    assert foy.case_id == sonen_id and foy.onceki_tracking_no is None
    # Esaslar kendi föylerine döner; klasör no föyün DosyaNo'suyla ayrılır.
    assert (kart.esas_no, sonen.esas_no) == ("2021/390", "2025/618")
    assert (kart.klasor_no_2, sonen.klasor_no_2) == ("6.5064.00", "6.5170.00")
    # Belge kalan kartta; ilişki ILGILI; tarihçe iki kartta.
    assert db.query(models.CaseDocument).filter_by(case_id=kart_id).count() == 1
    iliski = db.query(models.CaseRelation).one()
    assert (iliski.source_case_id, iliski.target_case_id, iliski.relation_type) == (kart_id, sonen_id, "ILGILI")
    assert {"foy_tasima", "klasor_no_2"} <= {h.field_name for h in db.query(models.CaseHistory).filter_by(case_id=kart_id)}
    assert {"kapatma", "foy_tasima", "onceki_tracking_no"} <= {
        h.field_name for h in db.query(models.CaseHistory).filter_by(case_id=sonen_id)}
    # Müvekkil sönen kartta da taraf olarak yazılı, föy bağı kurulu.
    muvekkil = db.query(models.CaseParty).filter_by(case_id=sonen_id, name="Sompo Sigorta A.Ş.").one()
    assert foy.case_party_id == muvekkil.id
    db.close()


def test_sonen_kart_yoksa_ret(db_env):  # noqa: F811
    kart_id, sonen_id = _birlesmis_kart(db_env)
    db = db_env()
    db.get(models.Case, sonen_id).tracking_no = "BASKA"
    db.commit()
    db.close()
    sonuc = ec.kos(db_env, apply=True)
    assert sonuc.sayim("canlandirma", "RET") == 1
    db = db_env()
    assert db.query(models.CaseFoy).filter_by(sistem_no="H-16815").one().case_id == kart_id
    db.close()


# ─── Föy taşıma ──────────────────────────────────────────────────────────────

def test_foy_tasima_klasor_ve_taraf(db_env):  # noqa: F811
    kaynak_id, hedef_id = _tesaduf_kartlari(db_env)
    sonuc = ec.kos(db_env, apply=True, kim="test")
    assert sonuc.sayim("tasima", "YAPILDI") == 1

    db = db_env()
    assert db.query(models.CaseFoy).filter_by(sistem_no="H-3844").one().case_id == hedef_id
    kaynak, hedef = db.get(models.Case, kaynak_id), db.get(models.Case, hedef_id)
    assert kaynak.klasor_no_2 == "1193.002.00"
    assert hedef.klasor_no_2 == "1641.002.00;9.819.00;3.832.00"
    assert kaynak.esas_no == "2020/611"                    # kalan föyün esası; dokunulmaz
    # Kaynakta yalnız taşınan föye özgü taraf silindi; ortak müvekkil kaldı.
    assert {p.name for p in db.query(models.CaseParty).filter_by(case_id=kaynak_id)} == {
        "Mehmet Türker Dr", "Özgül Kaçar"}
    assert "Mehmet Türker Dr" in {p.name for p in db.query(models.CaseParty).filter_by(case_id=hedef_id)}
    db.close()


def test_belgeye_bagli_taraf_silinmez(db_env):  # noqa: F811
    kaynak_id, _ = _tesaduf_kartlari(db_env)
    db = db_env()
    egehan = db.query(models.CaseParty).filter_by(case_id=kaynak_id, name="Egehan Coşgun").one()
    db.add(models.CaseDocument(case_id=kaynak_id, original_filename="d.pdf", stored_filename="d.pdf",
                               case_party_id=egehan.id))
    db.commit()
    db.close()
    sonuc = ec.kos(db_env, apply=True)
    assert sonuc.sayim("taraf", "ATLANDI") == 1
    db = db_env()
    assert db.query(models.CaseParty).filter_by(case_id=kaynak_id, name="Egehan Coşgun").count() == 1
    db.close()


# ─── Aktif/Arşiv kuralı ──────────────────────────────────────────────────────

def _rapor(yol, durumlar):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Dosya - Föy Bilgileri", "SistemNo", "DosyaNo", "Durum", "Son Durum"])
    for sistem_no, durum in durumlar:
        ws.append([None, sistem_no, None, durum, None])
    wb.save(yol)
    return yol


def _ek(yol, durumlar):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FOYLER"
    ws.append(["Kart ID", "SistemNo", "Ek-2 Son Durum", "Bugünkü Son Durum", "Değişti", "Seviye", "Durum"])
    for sistem_no, durum in durumlar:
        ws.append([1, sistem_no, None, None, None, None, durum])
    wb.save(yol)
    return yol


def _durum_kartlari(db_env):  # noqa: F811
    db = db_env()
    try:
        biri_aktif, hepsi_arsiv, danis, elle = (_kart(db, f"A{i}", f"1.{i}") for i in range(1, 5))
        biri_aktif.status, danis.status = "MAHZEN", "DANIŞ"
        for kart, foyler in ((biri_aktif, ("F-1", "F-2")), (hepsi_arsiv, ("F-3", "F-4")), (danis, ("F-5",)),
                             (elle, ("F-6",))):
            for sistem_no in foyler:
                _foy(db, kart, sistem_no)
        db.add(models.CaseHistory(case_id=elle.id, field_name="status", old_value="MAHZEN", new_value="DERDEST",
                                  changed_by="Nurten Meral", source="auto-status"))
        db.commit()
        return biri_aktif.id, hepsi_arsiv.id, danis.id, elle.id
    finally:
        db.close()


def test_aktif_arsiv_kurali(db_env, tmp_path):  # noqa: F811
    biri_aktif, hepsi_arsiv, danis, elle = _durum_kartlari(db_env)
    rapor = _rapor(tmp_path / "r.xlsx", [("F-1", "Arşiv"), ("F-2", "Arşiv"), ("F-3", "Arşiv"), ("F-4", "Aktif"),
                                          ("F-5", "Arşiv"), ("F-6", "Arşiv")])
    ek = _ek(tmp_path / "e.xlsx", [("F-2", "Aktif"), ("F-4", "Arşiv")])     # ek master'ı ezer
    sonuc = ec.kos(db_env, rapor=rapor, ek=ek, apply=True, kim="test")
    assert sonuc.sayim("status", "YAPILDI") == 2 and sonuc.sayim("status", "KORUNDU") == 1

    db = db_env()
    assert db.get(models.Case, biri_aktif).status == "DERDEST"
    assert db.get(models.Case, hepsi_arsiv).status == "MAHZEN"
    assert db.get(models.Case, danis).status == "DANIŞ"
    assert db.get(models.Case, elle).status == "DERDEST"       # elle girilmiş: KORUNDU
    db.close()


# ─── Kuru koşu ve idempotency ────────────────────────────────────────────────

def test_kuru_kosu_yazmaz(db_env):  # noqa: F811
    kart_id, sonen_id = _birlesmis_kart(db_env)
    sonuc = ec.kos(db_env, apply=False)
    assert sonuc.sayim("canlandirma", "YAPILDI") == 1
    db = db_env()
    assert db.get(models.Case, sonen_id).deleted_at is not None
    assert db.query(models.CaseFoy).filter_by(sistem_no="H-16815").one().case_id == kart_id
    assert db.query(models.CaseHistory).count() == 0
    db.close()
    assert "KURU KOŞU" in ec.ozet_metni(sonuc, apply=False)


def test_ikinci_kosu_degisiklik_uretmez(db_env, tmp_path):  # noqa: F811
    _birlesmis_kart(db_env)
    _tesaduf_kartlari(db_env)
    ec.kos(db_env, apply=True)
    ikinci = ec.kos(db_env, apply=True)
    for adim in ("canlandirma", "tasima", "taraf", "status"):
        assert ikinci.sayim(adim, "YAPILDI") == 0
    assert ikinci.sayim("canlandirma", "ATLANDI") == 1 and ikinci.sayim("tasima", "ATLANDI") == 1
