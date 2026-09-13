"""G179 — Veri ekibinin 12.09 cevabındaki kart düzeltmeleri script'i
(`scripts/ekip_cevabi_1209.py`): çift kart birleştirme (+ ret'te yalnız föy
taşıma), föy taşıma, kart düzeltmeleri, kopya/test kartlarını kapatma; kuru
koşu geri alır, ikinci koşu 0.

**TEST VERİSİ KURALI (A.2 dersi):** ekibin gerçek Ek-3'ü repoya girmez; Ek-3
sayfaları openpyxl ile SENTETİK üretilir (aynı başlıklar).
"""
import openpyxl
import pytest
from sqlalchemy import create_engine, event, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import models
from database import _MIGRATIONS, Base
from scripts import ekip_cevabi_1209 as ec
from tests.test_g064_aktarim_cekirdek import _kart


@pytest.fixture()
def db_env():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _fk_ac(dbapi_connection, _record):
        dbapi_connection.isolation_level = None
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    @event.listens_for(engine, "begin")
    def _begin(conn):
        conn.exec_driver_sql("BEGIN")

    Base.metadata.create_all(engine)
    with engine.begin() as conn:
        for op in _MIGRATIONS:
            if op[0] == "index" and op[1] in ("case_foys", "case_esas_numbers"):
                for sql in op[2]:
                    conn.execute(text(sql))
    yield sessionmaker(bind=engine, autocommit=False, autoflush=False)
    engine.dispose()


def _taraf(db, case, ad, tur="CLIENT"):
    p = models.CaseParty(case_id=case.id, name=ad, role="Müvekkil" if tur == "CLIENT" else "Karşı Taraf",
                         party_type=tur)
    db.add(p)
    db.flush()
    return p


def _foy(db, case, sistem_no, party=None):
    f = models.CaseFoy(sistem_no=sistem_no, case_id=case.id, case_party_id=party.id if party else None)
    db.add(f)
    db.flush()
    return f


def _belge(db, case, ad, party=None):
    d = models.CaseDocument(case_id=case.id, original_filename=ad, stored_filename=ad,
                            case_party_id=party.id if party else None)
    db.add(d)
    db.flush()
    return d


def _ek3_yaz(yol, ciftler, tasimalar):
    """Ek-3'ün iki sayfası, ekibin başlıklarıyla (araya ilgisiz sütunlar da konur)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ec.EK3_CIFT_SAYFASI
    ws.append(["SistemNo", "DosyaNo (bizde)", "Mahkeme", "Esas", "30.07'deki kart", "Ofis dosya no",
               "Klasör no", "Bugünkü föyleri", "Föyün bağlı olduğu kart", "Ofis dosya no ", "Klasör no ",
               "Kart açılış", "Cevabınız"])
    for sn, kalan, mukerrer in ciftler:
        ws.append([sn, "2.500.00", "X", "2025/1", kalan, "", "2.500", "föy yok", mukerrer, "", "2.500.00", "", ""])
    ws2 = wb.create_sheet(ec.EK3_TASIMA_SAYFASI)
    ws2.append(["SistemNo", "DosyaNo (bizde)", "Föy türü", "Föy mahkemesi", "Föy esası", "Bağlı olduğu kart",
                "Ofis dosya no", "Kart türü", "Kart mahkemesi", "Kart esası", "Karttaki öteki föyler",
                "Önerdiğimiz kart", "Ofis dosya no ", "Açıklama", "Durum", "Cevabınız"])
    for sn, hedef in tasimalar:
        ws2.append([sn, "", "Hukuk", "", "", 1, "", "", "", "", "", hedef, "", "", "", ""])
    wb.save(yol)
    return yol


def _kos(db_env, **kw):
    kw.setdefault("sabit_ciftler", ())
    kw.setdefault("duzeltmeler", ())
    kw.setdefault("kapatmalar", ())
    kw.setdefault("ciftler", ())
    kw.setdefault("tasimalar", ())
    return ec.kos(db_env, ek3=None, kim="test", **kw)


# ─── Ek-3 okuma ──────────────────────────────────────────────────────────────

def test_ek3_oku_iki_sayfa_baslikla(tmp_path):
    """Başlıklar ada göre bulunur (kolon sırası değil); 'Önerdiğimiz kart' boş satır taşıma değildir."""
    yol = _ek3_yaz(tmp_path / "ek3.xlsx", [("H-1", 10, 20), ("H-2", 11, 21)], [("F-1", 30), ("F-2", None)])
    ciftler, tasimalar = ec.ek3_oku(yol)
    assert ciftler == [("H-1", 10, 20), ("H-2", 11, 21)]
    assert tasimalar == [("F-1", 30)]


def test_ek3_sayfa_yoksa_hata(tmp_path):
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "bos.xlsx")
    with pytest.raises(ValueError, match="sayfası yok"):
        ec.ek3_oku(tmp_path / "bos.xlsx")


# ─── Adım 1: çift kart ───────────────────────────────────────────────────────

def test_cift_birlesir_foy_kalana_gecer_ikinci_kosu_atlar(db_env):
    """Ek-3 › 01 çekirdek vakası: 30.07 kartı (klasör '2.500', föysüz, belgeli) kalır;
    05.09 kartı ('2.500.00', föylü) birleşir → föy kalanda, mükerrer soft delete."""
    db = db_env()
    try:
        eski = _kart(db, "S4.QUICK......0469.HUKUK.00000", "2.500", file_type="Hukuk", esas_no="2025/393",
                     court="İskenderun 2. Tüketici Mahkemesi")
        yeni = _kart(db, "S4.QUICK......0538.HUKUK.00000", "2.500.00", file_type="Hukuk", esas_no="2025/393",
                     court="İskenderun 2. Tüketici Mahkemesi")
        _taraf(db, eski, "Quick Sigorta A.Ş.")
        _taraf(db, yeni, "Quick Sigorta A.Ş.")
        _belge(db, eski, "dilekce.pdf")
        _foy(db, yeni, "H-16708")
        db.commit()
        eski_id, yeni_id = eski.id, yeni.id
    finally:
        db.close()

    sonuc = _kos(db_env, ciftler=[("H-16708", eski_id, yeni_id)], apply=True)
    assert sonuc.sayim("cift", "YAPILDI") == 1

    db = db_env()
    try:
        assert db.query(models.CaseFoy).filter_by(sistem_no="H-16708").one().case_id == eski_id
        yeni = db.get(models.Case, yeni_id)
        assert yeni.deleted_at is not None and f"#{eski_id} " in yeni.delete_reason
        assert db.get(models.Case, eski_id).klasor_no_2 == "2.500;2.500.00"
    finally:
        db.close()

    tekrar = _kos(db_env, ciftler=[("H-16708", eski_id, yeni_id)], apply=True)
    assert tekrar.sayim("cift", "ATLANDI") == 1 and tekrar.sayim("cift", "YAPILDI") == 0


def test_cift_reddedilirse_yalniz_foy_tasinir(db_env):
    """H-15496 deseni: 'föyün kartı' iki müvekkilli Anadolu kartı → ön koşul reddeder;
    kart BİRLEŞMEZ, yalnız föy 30.07 kartına geçer, öteki föy yerinde kalır."""
    db = db_env()
    try:
        sompo = _kart(db, "S7.SOMPO......0156.HUKUK.00000", "6.5110", file_type="Hukuk", esas_no="2023/43",
                      court="Kocaeli 1. Tüketici Mahkemesi")
        anadolu = _kart(db, "S2.ANADOLU....0921.HUKUK.00000", "9.1662;6.5110.00", file_type="Hukuk",
                        esas_no="2023/43", court="Kocaeli 1. Tüketici Mahkemesi")
        _taraf(db, sompo, "Sompo Sigorta A.Ş.")
        _taraf(db, anadolu, "Anadolu Anonim Türk Sigorta Şirketi")
        sompo_taraf = _taraf(db, anadolu, "Sompo Sigorta A.Ş.")
        _foy(db, anadolu, "H-15496", party=sompo_taraf)
        _foy(db, anadolu, "H-15400")
        db.commit()
        sompo_id, anadolu_id = sompo.id, anadolu.id
    finally:
        db.close()

    sonuc = _kos(db_env, ciftler=[("H-15496", sompo_id, anadolu_id)], apply=True)
    assert sonuc.sayim("cift", "YAPILDI") == 1
    assert "kart birleşmedi" in sonuc.kalemler[0].aciklama

    db = db_env()
    try:
        assert db.get(models.Case, anadolu_id).deleted_at is None
        foy = db.query(models.CaseFoy).filter_by(sistem_no="H-15496").one()
        assert foy.case_id == sompo_id
        # taraf bağı: hedef karttaki aynı adlı CLIENT (Sompo) satırına yeniden bağlandı
        assert db.get(models.CaseParty, foy.case_party_id).case_id == sompo_id
        assert db.query(models.CaseFoy).filter_by(sistem_no="H-15400").one().case_id == anadolu_id
        notlar = db.query(models.CaseHistory).filter_by(field_name="foy_tasima").all()
        assert {n.case_id for n in notlar} == {sompo_id, anadolu_id}
    finally:
        db.close()

    tekrar = _kos(db_env, ciftler=[("H-15496", sompo_id, anadolu_id)], apply=True)
    assert tekrar.sayim("cift", "ATLANDI") == 1


def test_sabit_cift_ret_kalir(db_env):
    """Sabit çiftte (§3 / Ek-2 › 03) ret → föy taşıma yedeği YOK, satır RET."""
    db = db_env()
    try:
        a = _kart(db, "D1.M_MERSIN...0017.HUKUK.00000", "1519.012", file_type="Hukuk", esas_no="2025/483")
        b = _kart(db, "D1.M_MERSIN...0020.HUKUK.00000-2", "1519.012", file_type="Hukuk", esas_no="2025/")
        db.commit()
        a_id, b_id = a.id, b.id
    finally:
        db.close()
    sonuc = _kos(db_env, sabit_ciftler=[(a_id, b_id, "Ek-2 › 03")], apply=True)
    assert sonuc.sayim("cift", "RET") == 1 and "esas no farklı" in sonuc.kalemler[0].aciklama
    db = db_env()
    try:
        assert db.get(models.Case, b_id).deleted_at is None
    finally:
        db.close()


# ─── Adım 2: föy taşıma ──────────────────────────────────────────────────────

def test_foy_tasima_hedefe_gecer_tarihce_iki_kartta(db_env):
    db = db_env()
    try:
        arabu = _kart(db, "S4.QUICK......0453.ARABU.00000", "2.455.00", file_type="Arabuluculuk")
        dava = _kart(db, "S4.QUICK......0463.HUKUK.00000", "2.455", file_type="Hukuk", esas_no="2024/528")
        q = _taraf(db, arabu, "Quick Sigorta A.Ş.")
        _taraf(db, dava, "Quıck Sigorta A.ş")          # yazım farklı, aynı anahtar
        _foy(db, arabu, "H-16477", party=q)
        db.commit()
        arabu_id, dava_id = arabu.id, dava.id
    finally:
        db.close()

    sonuc = _kos(db_env, tasimalar=[("H-16477", dava_id), ("YOK-1", dava_id), ("H-16477", 999)], apply=True)
    assert sonuc.sayim("foy", "YAPILDI") == 1 and sonuc.sayim("foy", "RET") == 2

    db = db_env()
    try:
        foy = db.query(models.CaseFoy).filter_by(sistem_no="H-16477").one()
        assert foy.case_id == dava_id
        assert db.get(models.CaseParty, foy.case_party_id).case_id == dava_id
        assert db.query(models.CaseHistory).filter_by(field_name="foy_tasima", case_id=arabu_id).count() == 1
        assert db.query(models.CaseHistory).filter_by(field_name="foy_tasima", case_id=dava_id).count() == 1
    finally:
        db.close()

    tekrar = _kos(db_env, tasimalar=[("H-16477", dava_id)], apply=True)
    assert tekrar.sayim("foy", "ATLANDI") == 1


# ─── Adım 3: kart düzeltmeleri ───────────────────────────────────────────────

def test_duzeltmeler_muvekkil_klasor_esas_konu(db_env):
    db = db_env()
    try:
        db.add_all([models.CaseSubject(code="HAKEM", name="Hakem Kararına İtiraz"),
                    models.ServiceType(code="VT", name="Vekaletsiz Takip"),
                    models.Client(name="AK SİGORTA A.Ş.")])
        k = _kart(db, "S3.AXA........2963.HUKUK.00000", "3.1400.00", file_type="Hukuk", esas_no="2020/113")
        _taraf(db, k, "Axa Sigorta A.Ş.")
        _taraf(db, k, "Hulusi Erçetin", tur="COUNTER")
        c = _kart(db, "S4.CORPUS.....0002.HUKUK.00000", "27.002.00", file_type="Hukuk")
        e = _kart(db, "D1.D_SARITAS..0003.IDARE.00000", "633.002.00", file_type="İdare", esas_no="2015/1027")
        db.commit()
        k_id, c_id, e_id = k.id, c.id, e.id
    finally:
        db.close()

    duzeltmeler = [
        (k_id, "muvekkil", "Ak Sigorta A.Ş.", "Ek-2 › 06"),
        (c_id, "klasor_ekle", "2.555.00", "Ek-2 › 06"),
        (e_id, "esas_no", "2015/11027", "karar künyesi"),
        (k_id, "subject", "hakem kararına itiraz", "§3 (yazım listeden)"),
        (k_id, "hizmet_turu", "Vekaletsiz Takip", "§3"),
        (k_id, "subject", "Olmayan Konu", "listede yok → RET"),
        (999, "klasor", "1", "kart yok → RET"),
    ]
    sonuc = _kos(db_env, duzeltmeler=duzeltmeler, apply=True)
    assert sonuc.sayim("duzeltme", "YAPILDI") == 5 and sonuc.sayim("duzeltme", "RET") == 2

    db = db_env()
    try:
        k = db.get(models.Case, k_id)
        client = [p for p in k.parties if p.party_type == "CLIENT"][0]
        assert client.name == "Ak Sigorta A.Ş." and client.client_id is not None
        assert k.subject == "Hakem Kararına İtiraz" and k.hizmet_turu == "Vekaletsiz Takip"
        assert db.get(models.Case, c_id).klasor_no_2 == "27.002.00;2.555.00"
        e = db.get(models.Case, e_id)
        assert e.esas_no == "2015/11027"
        guncel = [r for r in db.query(models.CaseEsasNumber).filter_by(case_id=e_id).all() if r.is_current]
        assert [r.esas_no for r in guncel] == ["2015/11027"]
        kaynaklar = {h.source for h in db.query(models.CaseHistory).filter_by(case_id=k_id).all()}
        assert any(s.startswith("ekip cevabı 12.09.2026 (test)") for s in kaynaklar)
    finally:
        db.close()

    tekrar = _kos(db_env, duzeltmeler=duzeltmeler[:5], apply=True)
    assert tekrar.sayim("duzeltme", "ATLANDI") == 5 and tekrar.sayim("duzeltme", "YAPILDI") == 0


def test_duzeltme_esas_ve_mahkeme_bosaltma(db_env):
    """4181 deseni: arabuluculuk kartında dava dosyasının esas/mahkemesi → boşaltılır,
    esas tarihçesi silinmez (E8: değer boşsa kolon temizlenir, güncel işaret kalkar)."""
    db = db_env()
    try:
        k = _kart(db, "X1.MAYADENT...0003.ARABU.00000", "1735.001.00", file_type="Arabuluculuk",
                  esas_no="2023/202", court="İstanbul 1. Tüketici Mahkemesi")
        db.commit()
        k_id = k.id
    finally:
        db.close()
    sonuc = _kos(db_env, duzeltmeler=[(k_id, "klasor", "1735.002", "x"), (k_id, "court", None, "x"),
                                      (k_id, "esas_no", None, "x")], apply=True)
    assert sonuc.sayim("duzeltme", "YAPILDI") == 3
    db = db_env()
    try:
        k = db.get(models.Case, k_id)
        assert (k.klasor_no_2, k.court, k.esas_no) == ("1735.002", None, None)
    finally:
        db.close()


# ─── Adım 4: kapatma ─────────────────────────────────────────────────────────

def test_kapatma_belgeler_gercek_karta_tasinir_kopya_soft_delete(db_env):
    db = db_env()
    try:
        kopya = _kart(db, "HK-837834", None, file_type="Hukuk", esas_no="2020/1759")
        gercek = _kart(db, "S3.L_ARSLAN...0001.HUKUK.00000", "1.20001.00", file_type="Hukuk", esas_no="2023/108")
        t = _taraf(db, kopya, "Selçuk Gülen", tur="COUNTER")
        _belge(db, kopya, "istinaf-karari.pdf", party=t)
        _belge(db, kopya, "tebligat.pdf")
        smoke = _kart(db, "SMOKE-1785412361", None, file_type="Hukuk")
        _belge(db, smoke, "smoke.pdf")
        db.commit()
        kopya_id, gercek_id, smoke_id = kopya.id, gercek.id, smoke.id
    finally:
        db.close()

    kapatmalar = [(kopya_id, gercek_id, "§3 hatalı kopya"), (smoke_id, None, "§3 deneme verisi"), (999, None, "yok")]
    sonuc = _kos(db_env, kapatmalar=kapatmalar, apply=True)
    assert sonuc.sayim("kapatma", "YAPILDI") == 2 and sonuc.sayim("kapatma", "RET") == 1

    db = db_env()
    try:
        belgeler = db.query(models.CaseDocument).filter_by(case_id=gercek_id).all()
        assert len(belgeler) == 2 and all(d.case_party_id is None for d in belgeler)
        kopya = db.get(models.Case, kopya_id)
        assert kopya.deleted_at is not None and kopya.active is False and "hatalı kopya" in kopya.delete_reason
        assert db.query(models.CaseDocument).filter_by(case_id=kopya_id).count() == 0
        smoke = db.get(models.Case, smoke_id)
        assert smoke.deleted_at is not None
        assert db.query(models.CaseDocument).filter_by(case_id=smoke_id).count() == 1   # hedefsiz: belge kartta kalır
        assert db.query(models.CaseHistory).filter_by(case_id=gercek_id, field_name="belge_tasima").count() == 1
    finally:
        db.close()

    tekrar = _kos(db_env, kapatmalar=kapatmalar[:2], apply=True)
    assert tekrar.sayim("kapatma", "ATLANDI") == 2


# ─── Kuru koşu ───────────────────────────────────────────────────────────────

def test_kuru_kosu_hicbir_sey_yazmaz(db_env, tmp_path):
    db = db_env()
    try:
        a = _kart(db, "A", "2.1", file_type="Hukuk", esas_no="2025/1")
        b = _kart(db, "B", "2.1.00", file_type="Hukuk", esas_no="2025/1")
        _foy(db, b, "H-1")
        smoke = _kart(db, "SMOKE-1", None, file_type="Hukuk")
        db.commit()
        a_id, b_id, smoke_id = a.id, b.id, smoke.id
    finally:
        db.close()
    yol = _ek3_yaz(tmp_path / "ek3.xlsx", [("H-1", a_id, b_id)], [])
    sonuc = ec.kos(db_env, ek3=yol, apply=False, kim="test", sabit_ciftler=(), duzeltmeler=(),
                   kapatmalar=[(smoke_id, None, "deneme")])
    assert sonuc.sayim("cift", "YAPILDI") == 1 and sonuc.sayim("kapatma", "YAPILDI") == 1
    db = db_env()
    try:
        assert db.get(models.Case, b_id).deleted_at is None
        assert db.get(models.Case, smoke_id).deleted_at is None
        assert db.query(models.CaseFoy).filter_by(sistem_no="H-1").one().case_id == b_id
        assert db.query(models.CaseHistory).count() == 0
    finally:
        db.close()
    print(ec.ozet_metni(sonuc, apply=False))
