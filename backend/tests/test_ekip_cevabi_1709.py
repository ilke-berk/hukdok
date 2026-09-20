"""Veri ekibinin 17.09 cevabındaki (Ek-6) kart düzeltmeleri script'i
(`scripts/ekip_cevabi_1709.py`): kapatma, alan düzeltmesi, 15 mükerrer çift, klasör
numarasının tek yazıma çekilmesi, müvekkil ayrımıyla ayırma, ilişkili dosya bağı.

**TEST VERİSİ KURALI:** ekibin gerçek Ek-6'sı repoya girmez; sayfalar openpyxl ile
SENTETİK üretilir (aynı başlıklarla). Sabit tablolar PROD kart id'leri taşıdığı için
testte monkeypatch ile bu DB'nin id'lerine çevrilir.
"""
import openpyxl
import pytest

import models
from scripts import ekip_cevabi_1709 as ec
from tests.test_g064_aktarim_cekirdek import _kart
from tests.test_g179_ekip_cevabi_1209 import _belge, _foy, _taraf, db_env  # noqa: F401
from tests.test_g180_birlesik_kart_ayir import _ham


@pytest.fixture()
def sabitsiz(monkeypatch):
    """Prod id'li tabloların hepsi boşalır; her test kendi kalemini koyar."""
    for ad in ("KAPATMALAR", "ALAN_DUZELTMELERI", "BIRLESTIRMELER", "KLASOR_TEKLESTIRME",
               "KLASOR_CIKAR", "AYIRMALAR", "ILISKILER"):
        monkeypatch.setattr(ec, ad, ())
    return monkeypatch


def _ek6_yaz(yol, *, cok_kart=(), sorular=(), duzeltmeler=()):
    """Ek-6'nın üç sayfası, ekibin başlıklarıyla (araya ilgisiz sütunlar da konur)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ec.EK6_SORU_SAYFASI
    ws.append(["Sorunuz", "Kart", "Bizdeki kayıt", "Cevabımız"])
    for kart in sorular:
        ws.append([f"{kart} · soru", kart, "", "cevap"])

    ws = wb.create_sheet(ec.EK6_COK_KART_SAYFASI)
    ws.append(["Sınıf", "Mahkeme", "Esas", "Kart", "Ofis dosya no", "Durum", "Bizim föyümüz",
               "Klasör listesi", "Notumuz"])
    for sinif, kart, foy in cok_kart:
        ws.append([sinif, "X Mahkemesi", "2024/1", kart, "", "DERDEST", foy, "", ""])

    ws = wb.create_sheet(ec.EK6_DUZELTME_SAYFASI)
    ws.append(["Konu", "Kart", "Bugünkü değer", "Talebimiz / bilgi"])
    for kart in duzeltmeler:
        ws.append(["konu", kart, "", "talep"])
    wb.save(yol)
    return yol


# ─── Ek-6 doğrulama kapısı ───────────────────────────────────────────────────

def test_ek6_kapisi_kalan_foylu_sonen_foysuz_bekler(sabitsiz, tmp_path):
    sabitsiz.setattr(ec, "BIRLESTIRMELER", ((10, 20, "A", "kanıt"),))
    temiz = _ek6_yaz(tmp_path / "temiz.xlsx", cok_kart=[("A", 10, "H-1"), ("A", 20, "—")])
    assert ec.ek6_dogrula(temiz) == []

    ters = _ek6_yaz(tmp_path / "ters.xlsx", cok_kart=[("A", 10, "—"), ("A", 20, "H-1")])
    sorunlar = ec.ek6_dogrula(ters)
    assert len(sorunlar) == 2 and "çift ters olabilir" in sorunlar[0]

    eksik = _ek6_yaz(tmp_path / "eksik.xlsx", cok_kart=[("A", 10, "H-1")])
    assert ec.ek6_dogrula(eksik) == ["A: kart 20 Ek-6 › 02'de yok"]


def test_kapi_tutmazsa_hicbir_sey_yazilmaz(sabitsiz, db_env, tmp_path):  # noqa: F811
    db = db_env()
    try:
        bos = _kart(db, "X1.BOS.0001.IDARE.00000", None, file_type="İdare")
        db.commit()
        sabitsiz.setattr(ec, "KAPATMALAR", ((bos.id, "sebep"),))
    finally:
        db.close()
    ek6 = _ek6_yaz(tmp_path / "ek6.xlsx", sorular=[9999])      # kart listede yok

    with pytest.raises(ValueError, match="Ek-6 sabit tablolarla uyuşmuyor"):
        ec.kos(db_env, ek6=ek6, apply=True, kim="test")

    db = db_env()
    try:
        assert db.query(models.Case).filter(models.Case.deleted_at.isnot(None)).count() == 0
    finally:
        db.close()


# ─── Kapatma + alan düzeltmesi ───────────────────────────────────────────────

def test_kapatma_ve_alan_duzeltmeleri(sabitsiz, db_env):  # noqa: F811
    db = db_env()
    try:
        db.add(models.CaseSubject(code="RUCUEN", name="Rücuen Alacak"))
        bos = _kart(db, "X1.XXXX.0001.IDARE.00000", None, file_type="İdare")
        konu = _kart(db, "S1.AK.0001.HUKUK.00000", "1.1", file_type="Hukuk", subject="Alacak")
        karar = _kart(db, "S1.AK.0002.HUKUK.00000", "1.2", file_type="Hukuk", karar_turu="RED")
        dolu = _kart(db, "S1.AK.0003.HUKUK.00000", "1.3", file_type="Hukuk")
        _foy(db, dolu, "H-1")
        db.commit()
        idler = (bos.id, konu.id, karar.id, dolu.id)
    finally:
        db.close()
    sabitsiz.setattr(ec, "KAPATMALAR", ((idler[0], "›01: karşılığı yok"), (idler[3], "›01: karşılığı yok")))
    sabitsiz.setattr(ec, "ALAN_DUZELTMELERI", (
        (idler[1], "subject", "Alacak", "Rücuen Alacak", "›04"),
        (idler[2], "karar_turu", "RED", None, "›04"),
    ))

    sonuc = ec.kos(db_env, apply=True, kim="test")

    assert sonuc.sayim("kapatma", "YAPILDI") == 1 and sonuc.sayim("kapatma", "RET") == 1
    assert sonuc.sayim("alan", "YAPILDI") == 2
    db = db_env()
    try:
        assert db.get(models.Case, idler[0]).deleted_at is not None
        assert db.get(models.Case, idler[0]).active is False
        assert db.get(models.Case, idler[3]).deleted_at is None      # föyü var → kapatılmadı
        assert db.get(models.Case, idler[1]).subject == "Rücuen Alacak"
        assert db.get(models.Case, idler[2]).karar_turu is None
        imzalar = {h.source.split(":")[0] for h in db.query(models.CaseHistory)}
        assert imzalar == {f"{ec.KAYNAK} (test)"}
    finally:
        db.close()

    tekrar = ec.kos(db_env, apply=True, kim="test")
    assert tekrar.sayim("kapatma", "YAPILDI") == 0 and tekrar.sayim("alan", "YAPILDI") == 0
    assert tekrar.sayim("alan", "ATLANDI") == 2


def test_alan_kapali_listede_yoksa_yazilmaz(sabitsiz, db_env):  # noqa: F811
    db = db_env()
    try:
        kart = _kart(db, "S1.AK.0009.HUKUK.00000", "1.9", file_type="Hukuk", subject="Alacak")
        db.commit()
        kart_id = kart.id
    finally:
        db.close()
    sabitsiz.setattr(ec, "ALAN_DUZELTMELERI", ((kart_id, "subject", "Alacak", "Havuzda Yok", "›04"),))

    sonuc = ec.kos(db_env, apply=True, kim="test")

    assert sonuc.sayim("alan", "RET") == 1 and "kapalı listede yok" in sonuc.kalemler[0].aciklama
    db = db_env()
    try:
        assert db.get(models.Case, kart_id).subject == "Alacak"
    finally:
        db.close()


# ─── Mükerrer + klasör ───────────────────────────────────────────────────────

def _a_sinifi_cift(db_env):  # noqa: F811
    """A sınıfı deseni: aynı numaranın iki yazımı, sönen kartta belge var."""
    db = db_env()
    try:
        kalan = _kart(db, "D1.C_HARMANCI.0002.IDARE.00000", "1976.001", file_type="İdare",
                      esas_no="2023/907", court="Hatay 4. İdare Mahkemesi")
        sonen = _kart(db, "D1.C_HARMANCI.0001.IDARE.00000", "1.976.001", file_type="İdare",
                      esas_no="2023/907", court="Hatay 4 İdare Mahkemesi")
        _taraf(db, kalan, "Cem Harmancı Dr.")
        _foy(db, kalan, "id-16645")
        _belge(db, sonen, "dilekce.pdf")
        db.commit()
        return kalan.id, sonen.id
    finally:
        db.close()


def test_birlestirme_belgeyi_tasir_klasor_tek_yazima_iner(sabitsiz, db_env):  # noqa: F811
    kalan_id, sonen_id = _a_sinifi_cift(db_env)
    sabitsiz.setattr(ec, "BIRLESTIRMELER", ((kalan_id, sonen_id, "A", "klasör 1976.001 ↔ 1.976.001"),))
    sabitsiz.setattr(ec, "KLASOR_TEKLESTIRME", ((kalan_id, "1976.001", "›02 A"),))

    sonuc = ec.kos(db_env, apply=True, kim="test")

    assert sonuc.sayim("cift", "YAPILDI") == 1 and sonuc.sayim("klasor", "YAPILDI") == 1
    db = db_env()
    try:
        kalan = db.get(models.Case, kalan_id)
        assert kalan.klasor_no_2 == "1976.001"          # varyant yazım düştü
        assert db.query(models.CaseDocument).filter_by(case_id=kalan_id).count() == 1
        sonen = db.get(models.Case, sonen_id)
        assert sonen.deleted_at is not None and sonen.active is False
        assert ec.KAYNAK in sonen.delete_reason
    finally:
        db.close()

    tekrar = ec.kos(db_env, apply=True, kim="test")
    assert tekrar.sayim("cift", "ATLANDI") == 1 and tekrar.sayim("klasor", "ATLANDI") == 1


def test_teklestirme_baska_numaraya_dokunmaz(sabitsiz, db_env):  # noqa: F811
    db = db_env()
    try:
        kart = _kart(db, "D1.K.0001.IDARE.00000", "2.433.01;9.1110.00;2433.01", file_type="İdare")
        db.commit()
        kart_id = kart.id
    finally:
        db.close()
    sabitsiz.setattr(ec, "KLASOR_TEKLESTIRME", ((kart_id, "2.433.01", "›02 A"),))

    ec.kos(db_env, apply=True, kim="test")

    db = db_env()
    try:
        assert db.get(models.Case, kart_id).klasor_no_2 == "2.433.01;9.1110.00"
    finally:
        db.close()


def test_klasor_cikarma_kalan_parcayi_arar(sabitsiz, db_env):  # noqa: F811
    db = db_env()
    try:
        kart = _kart(db, "H1.H_ALP.0003.CEZAA.00000", "221.002.00;212.001.00", file_type="Ceza")
        eksik = _kart(db, "H1.H_ALP.0004.CEZAA.00000", "212.001.00", file_type="Ceza")
        db.commit()
        kart_id, eksik_id = kart.id, eksik.id
    finally:
        db.close()
    sabitsiz.setattr(ec, "KLASOR_CIKAR", (
        (kart_id, "212.001.00", "221.002.00", "›02 D"),
        (eksik_id, "212.001.00", "221.002.00", "›02 D"),
    ))

    sonuc = ec.kos(db_env, apply=True, kim="test")

    assert sonuc.sayim("klasor", "YAPILDI") == 1 and sonuc.sayim("klasor", "RET") == 1
    db = db_env()
    try:
        assert db.get(models.Case, kart_id).klasor_no_2 == "221.002.00"
        assert db.get(models.Case, eksik_id).klasor_no_2 == "212.001.00"   # dokunulmadı
    finally:
        db.close()


# ─── Ayırma + ilişki ─────────────────────────────────────────────────────────

def test_ayirma_muvekkil_basina_kart_acar(sabitsiz, db_env):  # noqa: F811
    db = db_env()
    try:
        kart = _kart(db, "D1.D_ESINLER..0002.ARABU.00000", "1110.003;1993.001",
                     file_type="Arabuluculuk", esas_no="2026/720", court="Ankara Arabuluculuk Bürosu")
        for sistem_no, dosya_no, tur, esas, mahkeme, muvekkil in (
                ("ARB-16767", "1110.003", "ARABULUCULUK", "2026/720", "Ankara Arabuluculuk Bürosu",
                 "Deniz Esinler Dr."),
                ("ARB-16779", "1993.001", "ARABULUCULUK", "2026/720", "Ankara Arabuluculuk Bürosu",
                 "Aylin Ayrim Dr"),
                ("H-16856", "1110.003", "HUKUK", "2026/91", "Ankara 11. Tüketici Mahkemesi",
                 "Deniz Esinler Dr."),
                ("H-16857", "1993.001", "HUKUK", "2026/91", "Ankara 11. Tüketici Mahkemesi",
                 "Aylin Ayrim Dr")):
            foy = _foy(db, kart, sistem_no)
            foy.ham_veri = _ham(sistem_no, dosya_no, tur, esas, mahkeme, muvekkil=muvekkil, tip="Doktor")
        db.commit()
        kart_id = kart.id
    finally:
        db.close()
    sabitsiz.setattr(ec, "AYIRMALAR", ((kart_id, "›01: dört föy dört ayrı kartta"),))

    sonuc = ec.kos(db_env, apply=True, kim="test")

    assert sonuc.sayim("ayirma", "YAPILDI") == 1
    db = db_env()
    try:
        assert db.query(models.Case).filter(models.Case.deleted_at.is_(None)).count() == 4
        assert db.query(models.CaseFoy).filter_by(case_id=kart_id).count() == 1
    finally:
        db.close()


def test_iliski_tek_satir_ve_kuru_kosu(sabitsiz, db_env):  # noqa: F811
    db = db_env()
    try:
        a = _kart(db, "D1.M_ORAL.0001.HUKUK.00000", "2002.001", file_type="Hukuk", esas_no="2026/246")
        b = _kart(db, "S4.QUICK.0589.HUKUK.00000", "2.365.00", file_type="Hukuk", esas_no="2026/246")
        db.commit()
        a_id, b_id = a.id, b.id
    finally:
        db.close()
    sabitsiz.setattr(ec, "ILISKILER", ((a_id, b_id, "›02 D: aynı davanın iki müvekkil kartı"),))

    kuru = ec.kos(db_env, apply=False, kim="test")
    assert kuru.sayim("iliski", "YAPILDI") == 1
    db = db_env()
    try:
        assert db.query(models.CaseRelation).count() == 0      # kuru koşu yazmadı
    finally:
        db.close()

    ec.kos(db_env, apply=True, kim="test")
    ec.kos(db_env, apply=True, kim="test")
    db = db_env()
    try:
        iliski = db.query(models.CaseRelation).one()
        assert (iliski.source_case_id, iliski.target_case_id, iliski.relation_type) == (a_id, b_id, "ILGILI")
    finally:
        db.close()


def test_ozet_metni_adimlari_sayar(sabitsiz, db_env):  # noqa: F811
    metin = ec.ozet_metni(ec.kos(db_env, apply=False, kim="test"), apply=False)
    assert "KURU KOŞU" in metin and all(ad in metin for ad in ec.ADIM_ADLARI.values())
