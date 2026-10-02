"""G249 — Hizmet kaydı: aktarım föy başına satır yazar + geriye dönük doldurma +
kart birleştir/ayır satırları taşır.

Kullanıcı kararı 01.10.2026 (G248): hizmet türü kartın değil KART × MÜVEKKİL
çiftinin özelliğidir. Bu dosya G249'un kilitlediği davranışları ölçer:

  1. aktarım: her föy (kapsamda + müvekkil bağı + hizmet) kartında BİR föy
     kaynaklı hizmet satırı taşır (`case_hizmetleri.foydan_yaz`, anahtar
     `foy_id`, `source` = aktarım imzası); föyün müvekkili/hizmeti değişince
     satır YERİNDE güncellenir; kapsam dışına çıkınca silinir, işaret kalkınca
     geri gelir; müvekkil bağı yoksa / hizmet listede yoksa satır yazılmaz +
     satır raporuna `UYARI` (HATA değil), "listede yok" koşu özetinde ayrı kalem,
  2. `hizmet_turu` kart alanı uzlaşısından çıktı: kardeş föylerin farklı hizmeti
     çelişki değil, kart özeti (`cases.hizmet_turu`) satırların birleşimi; paket
     elle girilmiş hizmet satırına dokunmaz,
  3. `scripts/hizmet_kayitlari_doldur.py`: varsayılan kuru koşu, `apply` tek
     transaction, ikinci koşu 0 değişiklik,
  4. `mukerrer_kart_birlestir` tarafla birlikte hizmet satırlarını taşır (elle
     satır çakışması birleşir); `birlesik_kart_ayir` föyle birlikte satırını taşır.

**TEST VERİSİ KURALI (A.2 dersi):** gerçek teslim paketi REPOYA GİRMEZ; bütün
testler openpyxl ile SENTETİK mini paket üretir (test_g120 düzeni).
"""
from pathlib import Path

import pytest
from sqlalchemy import create_engine, event, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import models
from database import _MIGRATIONS, Base
from managers import case_hizmetleri, case_manager, seed_data
from scripts import birlesik_kart_ayir as bka
from scripts import hizmet_kayitlari_doldur as hkd
from scripts import hukdok_aktarim
from scripts import mukerrer_kart_birlestir as mb
from scripts.hukdok_aktarim import CIKIS_TAMAM, aktarimi_kos
from services import ofis_no

BASLIKLAR = ["SistemNo", "TKU", "Dosya No", "Müvekkil", "Müvekkil Tipi", "Hizmet Türü"]
SILINEN_BASLIKLAR = ["SistemNo", "Dosya No", "Silinme Gerekçesi", "Tarih"]
IMZA = "HUKDOK_TESLIM_t.xlsx"


def _paket_yaz(yol, satirlar, *, basliklar=None, silinen=None):
    """Sentetik mini teslim paketi: ana sayfa + isteğe bağlı `Silinen_Föyler`."""
    from openpyxl import Workbook

    kullanilan = list(basliklar if basliklar is not None else BASLIKLAR)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(kullanilan)
    for satir in satirlar:
        ws.append([satir.get(baslik) for baslik in kullanilan])
    if silinen is not None:
        w = wb.create_sheet("Silinen_Föyler")
        w.append(SILINEN_BASLIKLAR)
        for sistem_no in silinen:
            w.append([sistem_no, "D-1", "mükerrer açılış", "12.08.2026"])
    wb.save(yol)
    wb.close()
    return Path(yol)


def _satir(sistem_no, dosya_no, muvekkil=None, hizmet=None, **extra):
    temel = {"SistemNo": sistem_no, "Dosya No": dosya_no, "TKU": "TKU-249",
             "Müvekkil": muvekkil, "Hizmet Türü": hizmet}
    temel.update(extra)
    return temel


# ═══════════════════════════════════════════════════════════════════════════
# Ortam — sqlite (test_g064 reçetesi: FK + ÇALIŞAN SAVEPOINT) + hizmet listesi
# ═══════════════════════════════════════════════════════════════════════════

def _index_sqlleri(*tablolar):
    return [sql for op in _MIGRATIONS if op[0] == "index" and op[1] in tablolar
            for sql in op[2] if not sql.lstrip().upper().startswith(("UPDATE", "ALTER", "DO "))]


@pytest.fixture()
def db_env():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _fk_ac(dbapi_connection, _record):
        dbapi_connection.isolation_level = None      # pysqlite BEGIN yaymasın
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    @event.listens_for(engine, "begin")
    def _begin(conn):
        conn.exec_driver_sql("BEGIN")

    Base.metadata.create_all(engine)
    with engine.begin() as conn:
        # Migrasyonun index SQL'leri OLDUĞU GİBİ (G049 dersi): `case_hizmetleri`nin
        # iki kısmi UNIQUE'i (föy başına tek satır, elle satır tekrarı yok) dahil.
        for sql in _index_sqlleri("case_foys", "case_hizmetleri", "case_esas_numbers"):
            conn.execute(text(sql))
    fabrika = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    db = fabrika()
    try:
        for sira, (kod, ad) in enumerate(seed_data.SERVICE_TYPES):
            db.add(models.ServiceType(code=kod, name=ad, active=True, sequence=sira))
        # `birlesik_kart_ayir` yeni kartın numarasını `services/ofis_no`dan alır (G239).
        for kod, ad, anahtarlar in ofis_no.VARSAYILAN_SIGORTA_KODLARI:
            db.add(models.SigortaKisaKodu(kod=kod, ad=ad, eslesme_anahtarlari=list(anahtarlar), aktif=True))
        db.commit()
    finally:
        db.close()
    yield fabrika
    engine.dispose()


def _kart(db, tracking, klasor, **extra):
    case = models.Case(tracking_no=tracking, status="DERDEST", klasor_no_2=klasor, **extra)
    db.add(case)
    db.flush()
    return case


def _taraf(db, case, ad, party_type="CLIENT", role="Müvekkil"):
    taraf = models.CaseParty(case_id=case.id, name=ad, role=role, party_type=party_type)
    db.add(taraf)
    db.flush()
    return taraf


def _foy(db, case, taraf, sistem_no, hizmet, **extra):
    foy = models.CaseFoy(
        sistem_no=sistem_no, case_id=case.id, case_party_id=taraf.id if taraf is not None else None,
        hizmet_turu=hizmet, source=IMZA, **extra,
    )
    db.add(foy)
    db.flush()
    return foy


@pytest.fixture()
def iki_kart(db_env):
    db = db_env()
    try:
        for i in (1, 2):
            _kart(db, f"HA.G249.{i}", f"D-{i}")
        db.commit()
    finally:
        db.close()
    return db_env


def _satirlar(fabrika):
    """Hizmet satırları, yazılma sırasıyla: (kart klasörü, müvekkil adı, hizmet,
    föyün SistemNo'su — elle satırda None —, satır `source`u)."""
    db = fabrika()
    try:
        cikti = []
        for h in db.query(models.CaseHizmeti).order_by(models.CaseHizmeti.id):
            foy = db.get(models.CaseFoy, h.foy_id) if h.foy_id is not None else None
            cikti.append((
                db.get(models.Case, h.case_id).klasor_no_2,
                db.get(models.CaseParty, h.case_party_id).name,
                h.hizmet_turu,
                foy.sistem_no if foy is not None else None,
                h.source,
            ))
        return cikti
    finally:
        db.close()


def _ozetler(fabrika):
    db = fabrika()
    try:
        return {c.klasor_no_2: c.hizmet_turu for c in db.query(models.Case)}
    finally:
        db.close()


def _hizmet_tarihcesi(fabrika):
    db = fabrika()
    try:
        return [(h.old_value, h.new_value) for h in
                db.query(models.CaseHistory)
                .filter(models.CaseHistory.field_name == case_hizmetleri.TARIHCE_ALANI)
                .order_by(models.CaseHistory.id)]
    finally:
        db.close()


def _kos(fabrika, tmp_path, satirlar, **kw):
    paket = _paket_yaz(tmp_path / "t.xlsx", satirlar,
                       basliklar=kw.pop("basliklar", None), silinen=kw.pop("silinen", None))
    return aktarimi_kos(fabrika, girdi=paket, rapor_dizini=tmp_path / "rapor", **kw)


# ═══════════════════════════════════════════════════════════════════════════
# 1. Kayıt kilitleri (DB yok)
# ═══════════════════════════════════════════════════════════════════════════

def test_hizmet_turu_kart_alani_degil_duzeltme_logu_da_tanimaz():
    """Kabul 2: `hizmet_turu` kart alanı kümesinde yok → uzlaşı/çelişki hesabına,
    kart yazımına ve `Düzeltme_Logu` alan haritasına giremez. Föy yolu yerinde."""
    assert "hizmet_turu" not in hukdok_aktarim.KART_ALANLARI
    assert "hizmet_turu" not in hukdok_aktarim.KART_TURETILEN
    assert "hizmet_turu" not in hukdok_aktarim.DUZELTME_ALAN_HARITASI.values()
    satir = hukdok_aktarim.HamSatir(satir_no=2, degerler={"hizmet_turu": "lexis rapor"})
    assert "hizmet_turu" not in hukdok_aktarim.kart_degerleri(satir)
    assert hukdok_aktarim.foy_degerleri(satir)["hizmet_turu"] == "Lexis Rapor"
    assert hukdok_aktarim.HIZMET_UYARI_TURU == "UYARI"


# ═══════════════════════════════════════════════════════════════════════════
# 2. Aktarım — föy başına hizmet satırı
# ═══════════════════════════════════════════════════════════════════════════

def test_foy_basina_satir_imza_ozet_ve_ikinci_kosu_sifir(iki_kart, tmp_path):
    """Kabul 1 + 2: föy satırı `foydan_yaz` ile yazılır (müvekkil tarafına, föy
    bağıyla, aktarım imzasıyla); kart özeti satırdan türetilir; aynı girdiyle
    ikinci koşu satır da tarihçe de üretmez."""
    satirlar = [_satir("H-1", "D-1", "Ak Sigorta A.Ş.", "LEXIS RAPOR")]

    ilk = _kos(iki_kart, tmp_path, satirlar)

    assert ilk.cikis_kodu == CIKIS_TAMAM and ilk.rapor_satirlari == []
    assert (ilk.hizmet_eklenen, ilk.hizmet_guncellenen, ilk.hizmet_silinen) == (1, 0, 0)
    assert ilk.alan_degisikligi == 0                  # hizmet kart hücresi sayılmaz
    assert ilk.kart_degisen == 1
    assert _satirlar(iki_kart) == [("D-1", "Ak Sigorta A.Ş.", "Lexis Rapor", "H-1", IMZA)]
    assert _ozetler(iki_kart) == {"D-1": "Lexis Rapor", "D-2": None}
    assert _hizmet_tarihcesi(iki_kart) == []           # ilk yazım tarihçesiz
    db = iki_kart()
    try:
        tarihce = db.query(models.CaseHistory).count()
        assert db.query(models.CaseHistory).filter_by(field_name="hizmet_turu").count() == 0
    finally:
        db.close()

    ikinci = _kos(iki_kart, tmp_path, satirlar)

    assert (ikinci.hizmet_eklenen, ikinci.hizmet_guncellenen, ikinci.hizmet_silinen) == (0, 0, 0)
    assert ikinci.alan_degisikligi == 0 and ikinci.kart_degisen == 0
    assert len(_satirlar(iki_kart)) == 1
    db = iki_kart()
    try:
        assert db.query(models.CaseHistory).count() == tarihce
    finally:
        db.close()


def test_kardes_foylerin_farkli_hizmeti_celiski_degil_ozet_birlesim(iki_kart, tmp_path):
    """Kabul 2: aynı kartın iki föyü farklı hizmet taşıyor → kardeş çelişkisi
    raporunda `hizmet_turu` satırı YOK, iki satır yazılır, özet birleşimdir."""
    sonuc = _kos(iki_kart, tmp_path, [
        _satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Vekaletli Takip"),
        _satir("H-2", "D-1", "Dr. Ali Veli", "Lexis Rapor"),
    ])

    assert sonuc.celiskiler == []
    assert not [y for y in sonuc.raporlar if "kardes-foy-celiskileri" in y.name]
    assert sonuc.hizmet_eklenen == 2 and sonuc.kart_degisen == 1
    assert {(s[1], s[2], s[3]) for s in _satirlar(iki_kart)} == {
        ("Ak Sigorta A.Ş.", "Vekaletli Takip", "H-1"),
        ("Dr. Ali Veli", "Lexis Rapor", "H-2"),
    }
    assert _ozetler(iki_kart)["D-1"] == "Lexis Rapor ; Vekaletli Takip"


def test_hizmet_ve_muvekkil_degisince_satir_yerinde_guncellenir(iki_kart, tmp_path):
    """Kabul 1: upsert anahtarı `foy_id` — föyün hizmeti ya da müvekkili değişince
    AYNI satır güncellenir (yeni satır açılmaz), tarihçeye düşer, özet yenilenir."""
    _kos(iki_kart, tmp_path, [_satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Lexis Rapor")])
    db = iki_kart()
    try:
        satir_id = db.query(models.CaseHizmeti).one().id
    finally:
        db.close()

    hizmet = _kos(iki_kart, tmp_path, [_satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Vekaletli Takip")])

    assert (hizmet.hizmet_eklenen, hizmet.hizmet_guncellenen) == (0, 1)
    assert hizmet.kart_degisen == 1 and hizmet.alan_degisikligi == 0
    assert _satirlar(iki_kart) == [("D-1", "Ak Sigorta A.Ş.", "Vekaletli Takip", "H-1", IMZA)]
    assert _ozetler(iki_kart)["D-1"] == "Vekaletli Takip"
    assert _hizmet_tarihcesi(iki_kart) == [
        ("Ak Sigorta A.Ş. — Lexis Rapor", "Ak Sigorta A.Ş. — Vekaletli Takip"),
    ]

    muvekkil = _kos(iki_kart, tmp_path, [_satir("H-1", "D-1", "Dr. Ali Veli", "Vekaletli Takip")])

    assert muvekkil.foy_muvekkil_degisen == 1 and muvekkil.hizmet_guncellenen == 1
    assert _satirlar(iki_kart) == [("D-1", "Dr. Ali Veli", "Vekaletli Takip", "H-1", IMZA)]
    assert _hizmet_tarihcesi(iki_kart)[-1] == (
        "Ak Sigorta A.Ş. — Vekaletli Takip", "Dr. Ali Veli — Vekaletli Takip")
    db = iki_kart()
    try:
        assert db.query(models.CaseHizmeti).one().id == satir_id      # yerinde
    finally:
        db.close()


def test_muvekkil_bagi_yok_uyari_hata_degil(iki_kart, tmp_path):
    """Kabul 1: föyün `case_party_id`'si NULL (Müvekkil hücresi boş) → satır
    yazılmaz, satır raporuna UYARI düşer; çıkış kodu TAMAM, her koşuda yeniden."""
    satirlar = [_satir("H-1", "D-1", None, "Lexis Rapor")]

    sonuc = _kos(iki_kart, tmp_path, satirlar)

    assert sonuc.cikis_kodu == CIKIS_TAMAM and sonuc.hatalar == []
    assert [(r.satir_no, r.sistem_no, r.dosya_no, r.tur) for r in sonuc.rapor_satirlari] == [
        (2, "H-1", "D-1", "UYARI")]
    assert "hizmet satırı yazılmadı" in sonuc.rapor_satirlari[0].sebep
    assert "müvekkil bağı yok" in sonuc.rapor_satirlari[0].sebep
    assert (sonuc.hizmet_muvekkilsiz, sonuc.hizmet_listede_yok, sonuc.hizmet_eklenen) == (1, 0, 0)
    assert _satirlar(iki_kart) == [] and _ozetler(iki_kart)["D-1"] is None
    rapor = [y for y in sonuc.raporlar if "satir-raporu" in y.name]
    assert rapor and ";UYARI;" in rapor[0].read_text(encoding="utf-8-sig")

    tekrar = _kos(iki_kart, tmp_path, satirlar)
    assert tekrar.hizmet_muvekkilsiz == 1 and len(tekrar.rapor_satirlari) == 1


def test_listede_olmayan_hizmet_uyari_ayri_kalem_mevcut_satir_kalir(iki_kart, tmp_path):
    """Kabul 1 (G256 senaryosu): hizmet admin panelinden yeniden adlandırıldı,
    paket ESKİ adı taşıyor → satır yazılmaz/değişmez, UYARI düşer, "listede yok"
    koşu özetinde AYRI kalemdir; mevcut satır yeni adıyla yerinde kalır."""
    satirlar = [_satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Lexis Rapor"),
                _satir("H-2", "D-2", "Dr. Ali Veli", "Lexis Rapor")]
    _kos(iki_kart, tmp_path, satirlar[:1])
    db = iki_kart()
    try:
        # yeniden adlandırma (liste + satır + özet) — panelin yaptığının düz hâli
        db.query(models.ServiceType).filter_by(name="Lexis Rapor").update({"name": "Lexis Raporu"})
        db.query(models.CaseHizmeti).update({"hizmet_turu": "Lexis Raporu"})
        db.query(models.Case).filter_by(klasor_no_2="D-1").update({"hizmet_turu": "Lexis Raporu"})
        db.commit()
    finally:
        db.close()

    sonuc = _kos(iki_kart, tmp_path, satirlar)

    assert sonuc.cikis_kodu == CIKIS_TAMAM and sonuc.hatalar == []
    assert sonuc.hizmet_listede_yok == 2 and sonuc.hizmet_muvekkilsiz == 0
    assert (sonuc.hizmet_eklenen, sonuc.hizmet_guncellenen, sonuc.hizmet_silinen) == (0, 0, 0)
    uyarilar = {r.sistem_no: r.sebep for r in sonuc.rapor_satirlari if r.tur == "UYARI"}
    assert set(uyarilar) == {"H-1", "H-2"} and len(sonuc.rapor_satirlari) == 2
    assert all("listede yok" in s and "Lexis Rapor" in s for s in uyarilar.values())
    # mevcut satır yeni adıyla durur; listede olmayan adla yeni satır açılmadı
    assert [(s[0], s[2], s[3]) for s in _satirlar(iki_kart)] == [("D-1", "Lexis Raporu", "H-1")]
    assert _ozetler(iki_kart) == {"D-1": "Lexis Raporu", "D-2": None}
    ozet = hukdok_aktarim.ozet_metni(sonuc)
    assert "hizmet uyarısı    : 2 listede yok, 0 müvekkil bağı yok" in ozet
    assert "hizmet satırı     : 0 eklendi, 0 güncellendi, 0 silindi" in ozet


def test_hizmet_sutunu_olmayan_paket_sessiz_ve_satiri_korur(iki_kart, tmp_path):
    """None sözleşmesi: paket `Hizmet Türü` sütununu taşımıyorsa föyün hizmeti
    ve satırı korunur; hiç hizmeti olmayan föy UYARI üretmez."""
    _kos(iki_kart, tmp_path, [_satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Lexis Rapor")])

    dar = _kos(iki_kart, tmp_path, [
        _satir("H-1", "D-1", "Ak Sigorta A.Ş."), _satir("H-2", "D-2", "Dr. Ali Veli"),
    ], basliklar=["SistemNo", "TKU", "Dosya No", "Müvekkil"])

    assert dar.rapor_satirlari == []
    assert (dar.hizmet_eklenen, dar.hizmet_guncellenen, dar.hizmet_silinen) == (0, 0, 0)
    assert (dar.hizmet_listede_yok, dar.hizmet_muvekkilsiz) == (0, 0)
    assert _satirlar(iki_kart) == [("D-1", "Ak Sigorta A.Ş.", "Lexis Rapor", "H-1", IMZA)]
    assert _ozetler(iki_kart) == {"D-1": "Lexis Rapor", "D-2": None}


def test_kapsam_disi_satiri_siler_isaret_kalkinca_geri_gelir(iki_kart, tmp_path):
    """Kabul 1: föy kapsam dışı işaretlenince satırı AYNI koşuda silinir
    (tarihçeli), işaret kalkınca AYNI koşuda geri gelir."""
    _kos(iki_kart, tmp_path, [
        _satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Lexis Rapor"),
        _satir("H-2", "D-1", "Ak Sigorta A.Ş.", "Vekaletli Takip"),
    ])
    assert _ozetler(iki_kart)["D-1"] == "Lexis Rapor ; Vekaletli Takip"

    # H-1 yalnız kapsam sayfasında (ana sayfada yok) — yine de dolaşılır.
    silme = _kos(iki_kart, tmp_path, [_satir("H-2", "D-1", "Ak Sigorta A.Ş.", "Vekaletli Takip")],
                 silinen=["H-1"])

    assert silme.kapsam_isaretlenen == 1 and silme.hizmet_silinen == 1
    assert silme.rapor_satirlari == []                 # kapsam dışı UYARI değil
    assert [s[3] for s in _satirlar(iki_kart)] == ["H-2"]
    assert _ozetler(iki_kart)["D-1"] == "Vekaletli Takip"
    assert _hizmet_tarihcesi(iki_kart) == [("Ak Sigorta A.Ş. — Lexis Rapor", "")]

    ayni = _kos(iki_kart, tmp_path, [_satir("H-2", "D-1", "Ak Sigorta A.Ş.", "Vekaletli Takip")],
                silinen=["H-1"])
    assert (ayni.hizmet_silinen, ayni.hizmet_eklenen) == (0, 0)   # idempotent

    geri = _kos(iki_kart, tmp_path, [
        _satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Lexis Rapor"),
        _satir("H-2", "D-1", "Ak Sigorta A.Ş.", "Vekaletli Takip"),
    ])

    assert geri.kapsam_geri_alinan == 1 and geri.hizmet_eklenen == 1
    assert sorted(s[3] for s in _satirlar(iki_kart)) == ["H-1", "H-2"]
    assert _ozetler(iki_kart)["D-1"] == "Lexis Rapor ; Vekaletli Takip"


def test_ana_sayfada_da_gecen_kapsam_disi_foy_satir_almaz(iki_kart, tmp_path):
    """İlk kez gelen ve aynı pakette kapsam sayfasında da olan föy: kimliği
    yazılır ama hizmet satırı AÇILMAZ (kapsam dışı föy kart düzeyine yazmaz)."""
    sonuc = _kos(iki_kart, tmp_path,
                 [_satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Lexis Rapor")], silinen=["H-1"])

    assert sonuc.foy_yeni == 1 and sonuc.kapsam_isaretlenen == 1
    assert (sonuc.hizmet_eklenen, sonuc.hizmet_silinen) == (0, 0)
    assert sonuc.rapor_satirlari == []
    assert _satirlar(iki_kart) == [] and _ozetler(iki_kart)["D-1"] is None


def test_kuru_kosu_hizmet_satiri_yazmaz_envanter_denk(iki_kart, tmp_path):
    """Kabul 2: kuru koşu hizmet satırını farkta gösterir, yazmaz; belge
    envanteri DENK kalır."""
    db = iki_kart()
    try:
        kart = db.query(models.Case).filter_by(klasor_no_2="D-1").one()
        db.add(models.CaseDocument(case_id=kart.id, original_filename="d.pdf", stored_filename="d.pdf"))
        db.commit()
    finally:
        db.close()

    sonuc = _kos(iki_kart, tmp_path,
                 [_satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Lexis Rapor")], dry_run=True)

    assert sonuc.dry_run and not sonuc.yazildi
    assert sonuc.hizmet_eklenen == 1 and sonuc.envanter_farki == {}
    assert sonuc.cikis_kodu == CIKIS_TAMAM
    db = iki_kart()
    try:
        assert db.query(models.CaseHizmeti).count() == 0
        assert db.query(models.CaseFoy).count() == 0
        assert db.query(models.CaseDocument).count() == 1
    finally:
        db.close()
    assert _ozetler(iki_kart)["D-1"] is None


def test_paket_elle_satira_dokunmaz_ozet_birlesim(iki_kart, tmp_path):
    """Kabul 2: kart özeti yalnız `ozeti_yenile`den — kullanıcının elle girdiği
    hizmet satırı paketle silinmez/ezilmez, özet elle + föy satırlarının birleşimi."""
    db = iki_kart()
    try:
        kart = db.query(models.Case).filter_by(klasor_no_2="D-1").one()
        taraf = _taraf(db, kart, "Ak Sigorta A.Ş.")
        case_hizmetleri.elle_ekle(db, kart, taraf.id, "Vekaletli Takip", changed_by="Avukat", source="panel")
        db.commit()
    finally:
        db.close()
    assert _ozetler(iki_kart)["D-1"] == "Vekaletli Takip"

    sonuc = _kos(iki_kart, tmp_path, [_satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Lexis Rapor")])

    assert sonuc.alan_degisikligi == 0 and sonuc.hizmet_eklenen == 1
    assert [(s[2], s[3], s[4]) for s in _satirlar(iki_kart)] == [
        ("Vekaletli Takip", None, "panel"), ("Lexis Rapor", "H-1", IMZA)]
    assert _ozetler(iki_kart)["D-1"] == "Lexis Rapor ; Vekaletli Takip"
    db = iki_kart()
    try:
        assert db.query(models.CaseHistory).filter_by(field_name="hizmet_turu").count() == 0
    finally:
        db.close()


# ═══════════════════════════════════════════════════════════════════════════
# 3. Geriye dönük doldurma — scripts/hizmet_kayitlari_doldur.py
# ═══════════════════════════════════════════════════════════════════════════

@pytest.fixture()
def doldurma_zemini(db_env):
    """Aktarılmış ama hizmet satırı yazılmamış föyler (G249 öncesi durum):
    D-1 iki föy/iki hizmet (özeti eski tek değer), D-2 tek föy (özeti zaten doğru),
    D-3 satır alamayan dört föy (kapsam dışı / müvekkilsiz / listede yok / hizmetsiz)."""
    db = db_env()
    try:
        k1 = _kart(db, "HA.G249.1", "D-1", hizmet_turu="Lexis Rapor")
        ak = _taraf(db, k1, "Ak Sigorta A.Ş.")
        ali = _taraf(db, k1, "Dr. Ali Veli")
        _foy(db, k1, ak, "H-1", "Lexis Rapor")
        _foy(db, k1, ali, "H-2", "Vekaletli Takip")
        k2 = _kart(db, "HA.G249.2", "D-2", hizmet_turu="Danışmanlık")
        _foy(db, k2, _taraf(db, k2, "Quick Sigorta A.Ş."), "H-3", "Danışmanlık")
        k3 = _kart(db, "HA.G249.3", "D-3")
        veli = _taraf(db, k3, "Veli Kaya")
        _foy(db, k3, veli, "H-4", "Lexis Rapor", kapsam_durumu="SILINDI")
        _foy(db, k3, None, "H-5", "Lexis Rapor")
        _foy(db, k3, veli, "H-6", "Rapor ; Takip")
        _foy(db, k3, veli, "H-7", None)
        db.commit()
    finally:
        db.close()
    return db_env


def test_doldurma_kuru_kosu_sayar_yazmaz(doldurma_zemini):
    sonuc = hkd.doldur(doldurma_zemini)

    assert not sonuc.apply and not sonuc.yazildi and sonuc.cikis_kodu == 0
    assert sonuc.foy_toplam == 7
    assert (sonuc.eklenen, sonuc.guncellenen, sonuc.silinen, sonuc.degismeyen) == (3, 0, 0, 0)
    assert sonuc.atlanan == 4
    assert {sebep: [f[0] for f in foyler] for sebep, foyler in sonuc.atlananlar.items()} == {
        case_hizmetleri.SEBEP_KAPSAM_DISI: ["H-4"],
        case_hizmetleri.SEBEP_TARAF_YOK: ["H-5"],
        case_hizmetleri.SEBEP_LISTEDE_YOK: ["H-6"],
        case_hizmetleri.SEBEP_HIZMET_BOS: ["H-7"],
    }
    assert sonuc.satir_alan_kart == 2
    assert sonuc.ozeti_degisen_kart == 1               # yalnız D-1: tek değer → birleşim
    assert sonuc.dagilim == {"Lexis Rapor": 1, "Vekaletli Takip": 1, "Danışmanlık": 1}
    # hiçbir şey yazılmadı
    assert _satirlar(doldurma_zemini) == []
    assert _ozetler(doldurma_zemini) == {"D-1": "Lexis Rapor", "D-2": "Danışmanlık", "D-3": None}
    metin = hkd.ozet_metni(sonuc)
    assert "KURU KOŞU" in metin and "yazılacak satır" in metin
    assert "hizmet türü listede yok" in metin and "H-6" in metin
    assert "özeti değişen kart  : 1" in metin


def test_doldurma_apply_yazar_ikinci_kosu_sifir(doldurma_zemini):
    sonuc = hkd.doldur(doldurma_zemini, apply=True)

    assert sonuc.yazildi and sonuc.eklenen == 3 and sonuc.degisiklik == 3
    assert _satirlar(doldurma_zemini) == [
        ("D-1", "Ak Sigorta A.Ş.", "Lexis Rapor", "H-1", IMZA),      # source = föyün teslim imzası
        ("D-1", "Dr. Ali Veli", "Vekaletli Takip", "H-2", IMZA),
        ("D-2", "Quick Sigorta A.Ş.", "Danışmanlık", "H-3", IMZA),
    ]
    assert _ozetler(doldurma_zemini) == {
        "D-1": "Lexis Rapor ; Vekaletli Takip", "D-2": "Danışmanlık", "D-3": None}
    assert _hizmet_tarihcesi(doldurma_zemini) == []    # ilk yazım tarihçesiz
    assert "YAZILDI" in hkd.ozet_metni(sonuc)

    ikinci = hkd.doldur(doldurma_zemini, apply=True)

    assert ikinci.degisiklik == 0 and ikinci.degismeyen == 3
    assert ikinci.ozeti_degisen_kart == 0 and ikinci.satir_alan_kart == 0
    assert ikinci.atlanan == 4                         # atlananlar her koşuda raporlanır
    assert len(_satirlar(doldurma_zemini)) == 3


def test_doldurmadan_sonra_aktarim_satirlari_degismedi_bulur(doldurma_zemini, tmp_path):
    """Doldurma ile aktarım AYNI yazıcıyı kullanır: doldurulmuş föyleri taşıyan
    paket hiçbir hizmet satırı eklemez/değiştirmez."""
    hkd.doldur(doldurma_zemini, apply=True)

    sonuc = _kos(doldurma_zemini, tmp_path, [
        _satir("H-1", "D-1", "Ak Sigorta A.Ş.", "Lexis Rapor"),
        _satir("H-2", "D-1", "Dr. Ali Veli", "Vekaletli Takip"),
        _satir("H-3", "D-2", "Quick Sigorta A.Ş.", "Danışmanlık"),
    ])

    assert (sonuc.hizmet_eklenen, sonuc.hizmet_guncellenen, sonuc.hizmet_silinen) == (0, 0, 0)
    assert len(_satirlar(doldurma_zemini)) == 3


# ═══════════════════════════════════════════════════════════════════════════
# 4. Kart birleştirme — scripts/mukerrer_kart_birlestir.py
# ═══════════════════════════════════════════════════════════════════════════

def _mukerrer_kart(db, tracking, klasor, muvekkiller):
    kart = _kart(db, tracking, klasor, court="Yalova 2. Asliye Hukuk Mahkemesi", file_type="Hukuk")
    case_manager.sync_current_esas(db, kart, "2019/82", court=kart.court, source="test")
    taraflar = {ad: _taraf(db, kart, ad) for ad in muvekkiller}
    return kart, taraflar


def test_birlestirme_eslesen_tarafin_hizmet_satirlarini_tasir_elle_cakisma_birlesir(db_env):
    """Kabul 4: mükerrerdeki müvekkil kalan kartta zaten var → o tarafın hizmet
    satırları kalandaki tarafa geçer (föy satırı hep; elle satır çakışırsa
    birleşir), taraf satırı RESTRICT'e takılmadan silinir, iki özet yenilenir."""
    db = db_env()
    try:
        kalan, kt = _mukerrer_kart(db, "K.1", "3.655.00", ["Ak Sigorta A.Ş."])
        case_hizmetleri.elle_ekle(db, kalan, kt["Ak Sigorta A.Ş."].id, "Danışmanlık", source="panel")
        muk, mt = _mukerrer_kart(db, "M.1", "3.658.00", ["AK SİGORTA A.Ş."])
        m_taraf = mt["AK SİGORTA A.Ş."]
        foy = _foy(db, muk, m_taraf, "H-9", "Lexis Rapor")
        case_hizmetleri.foydan_yaz(db, foy)
        case_hizmetleri.elle_ekle(db, muk, m_taraf.id, "Danışmanlık", source="panel")        # çakışır
        case_hizmetleri.elle_ekle(db, muk, m_taraf.id, "Vekaletli Takip", source="panel")    # taşınır
        db.commit()
        kalan_id, muk_id, kalan_taraf_id = kalan.id, muk.id, kt["Ak Sigorta A.Ş."].id
    finally:
        db.close()

    sonuc = mb.ciftleri_birlestir(db_env, [(kalan_id, muk_id)], apply=True, kim="test")

    assert sonuc[0].ret is None
    assert sonuc[0].tasinan["hizmet"] == 2 and sonuc[0].tasinan["hizmet_birlesen"] == 1
    assert sonuc[0].tasinan["taraf_tekil"] == 1
    db = db_env()
    try:
        satirlar = db.query(models.CaseHizmeti).order_by(models.CaseHizmeti.hizmet_turu).all()
        assert [(s.case_id, s.case_party_id, s.hizmet_turu, s.foy_id is not None) for s in satirlar] == [
            (kalan_id, kalan_taraf_id, "Danışmanlık", False),
            (kalan_id, kalan_taraf_id, "Lexis Rapor", True),
            (kalan_id, kalan_taraf_id, "Vekaletli Takip", False),
        ]
        foy = db.query(models.CaseFoy).filter_by(sistem_no="H-9").one()
        assert (foy.case_id, foy.case_party_id) == (kalan_id, kalan_taraf_id)
        assert db.get(models.Case, kalan_id).hizmet_turu == "Danışmanlık ; Lexis Rapor ; Vekaletli Takip"
        muk = db.get(models.Case, muk_id)
        assert muk.deleted_at is not None and muk.hizmet_turu is None
        assert db.query(models.CaseParty).filter_by(case_id=muk_id).count() == 0
    finally:
        db.close()


def test_birlestirme_tasinan_tarafin_satiri_kartla_birlikte_gider(db_env):
    """Kabul 4: mükerrerdeki müvekkil kalan kartta YOK (müvekkil ayrımı) → taraf
    satırı taşınır, hizmet satırı aynı tarafta kalır, yalnız kartı değişir."""
    db = db_env()
    try:
        kalan, _kt = _mukerrer_kart(db, "K.2", "1.1.00", ["Ak Sigorta A.Ş."])
        muk, mt = _mukerrer_kart(db, "M.2", "1.2.00", ["Dr. Ali Veli"])
        foy = _foy(db, muk, mt["Dr. Ali Veli"], "H-10", "Lexis Rapor")
        case_hizmetleri.foydan_yaz(db, foy)
        db.commit()
        kalan_id, muk_id, taraf_id = kalan.id, muk.id, mt["Dr. Ali Veli"].id
        with db.begin_nested():
            sonuc = mb.birlestir(db, kalan, muk, kim="test", muvekkil_ayrimi=True)
        db.commit()
    finally:
        db.close()

    assert sonuc.ret is None
    assert sonuc.tasinan["hizmet"] == 1 and "hizmet_birlesen" not in sonuc.tasinan
    db = db_env()
    try:
        satir = db.query(models.CaseHizmeti).one()
        assert (satir.case_id, satir.case_party_id, satir.hizmet_turu) == (kalan_id, taraf_id, "Lexis Rapor")
        assert db.get(models.CaseParty, taraf_id).case_id == kalan_id
        assert db.get(models.Case, kalan_id).hizmet_turu == "Lexis Rapor"
        assert db.get(models.Case, muk_id).hizmet_turu is None
    finally:
        db.close()


def test_birlestirme_ozeti_duz_alan_gibi_kopyalamaz_satirsiz_kartta_eski_deger_korunur(db_env):
    """`hizmet_turu` türetilmiş özettir: kalan kartta hizmet satırı varsa özet
    satırlardan gelir (mükerrerin eski tek değeri kopyalanmaz); iki kartta da
    satır yoksa (doldurma öncesi) mükerrerin eski değeri kaybolmaz."""
    db = db_env()
    try:
        # (a) iki kartta da satır yok → eski tek değer tamamlanır
        kalan_a, _ = _mukerrer_kart(db, "K.3", "2.1.00", ["Ak Sigorta A.Ş."])
        muk_a, _ = _mukerrer_kart(db, "M.3", "2.2.00", ["Ak Sigorta A.Ş."])
        muk_a.hizmet_turu = "Lexis Rapor"
        # (b) mükerrerin satırı var, kalanın eski tek değeri var → özet satırlardan
        kalan_b, _ = _mukerrer_kart(db, "K.4", "3.1.00", ["Quick Sigorta A.Ş."])
        kalan_b.hizmet_turu = "Danışmanlık"
        muk_b, mt_b = _mukerrer_kart(db, "M.4", "3.2.00", ["Quick Sigorta A.Ş."])
        case_hizmetleri.foydan_yaz(db, _foy(db, muk_b, mt_b["Quick Sigorta A.Ş."], "H-11", "Vekaletli Takip"))
        db.commit()
        idler = (kalan_a.id, muk_a.id, kalan_b.id, muk_b.id)
    finally:
        db.close()

    sonuc = mb.ciftleri_birlestir(db_env, [(idler[0], idler[1]), (idler[2], idler[3])], apply=True, kim="test")

    assert [s.ret for s in sonuc] == [None, None]
    assert sonuc[0].tasinan.get("alan_tamamlanan") == 1 and "hizmet" not in sonuc[0].tasinan
    assert sonuc[1].tasinan["hizmet"] == 1
    db = db_env()
    try:
        assert db.get(models.Case, idler[0]).hizmet_turu == "Lexis Rapor"
        assert db.get(models.Case, idler[2]).hizmet_turu == "Vekaletli Takip"
    finally:
        db.close()


# ═══════════════════════════════════════════════════════════════════════════
# 5. Kart ayırma — scripts/birlesik_kart_ayir.py
# ═══════════════════════════════════════════════════════════════════════════

def _ham(sistem_no, tur, esas, mahkeme, hizmet):
    return {"SistemNo": sistem_no, "DosyaNo": "2.455.00", "Klasör No": "TKU-9",
            "Müvekkil": "Quick Sigorta A.Ş.", "Müvekkil Tipi": "Sigorta", "Karşı Taraf": "Ayşe Yılmaz",
            "Yerel Mahkeme": mahkeme, "Esas": esas, "Dava Tarihi": "2026-03-01", "Ana Tür": tur,
            "Durum": "Aktif", "Dava Konusu": "Tazminat", "Hizmet Türü": hizmet}


def test_ayirma_foyun_hizmet_satirini_foyle_birlikte_tasir(db_env):
    """Kabul 4: ayrılan föyün hizmet satırı yeni karta, yeni kartın müvekkil
    tarafına YERİNDE geçer (aynı satır); kalan föyün ve elle satırın yeri
    değişmez; iki kartın özeti ve tarihçesi güncellenir."""
    db = db_env()
    try:
        kart = _kart(db, "S4.QUICK......0453.HUKUK.00000", "2.455.00", file_type="Hukuk",
                     esas_no="2024/528", court="Ankara 2. Asliye Hukuk Mahkemesi")
        taraf = _taraf(db, kart, "Quick Sigorta A.Ş.")
        arb = _foy(db, kart, taraf, "ARB-16361", "Vekaletsiz Takip",
                   ham_veri=_ham("ARB-16361", "ARABULUCULUK", "2025/48948", "ANKARA ARABULUCULUK BÜROSU",
                                 "Vekaletsiz Takip"))
        dava = _foy(db, kart, taraf, "H-16477", "Vekaletli Takip",
                    ham_veri=_ham("H-16477", "HUKUK", "2024/528", "Ankara 2. Asliye Hukuk Mahkemesi",
                                  "Vekaletli Takip"))
        case_hizmetleri.foydan_yaz(db, arb)
        case_hizmetleri.foydan_yaz(db, dava)
        case_hizmetleri.elle_ekle(db, kart, taraf.id, "Danışmanlık", source="panel")
        db.commit()
        kart_id, taraf_id = kart.id, taraf.id
        arb_satir_id = db.query(models.CaseHizmeti).filter_by(foy_id=arb.id).one().id
        assert kart.hizmet_turu == "Danışmanlık ; Vekaletli Takip ; Vekaletsiz Takip"
    finally:
        db.close()

    sonuc = bka.kos(db_env, kart_idler=[kart_id], apply=True, kim="test")

    assert sonuc.sayim("YAPILDI") == 1 and sonuc.yeni_kart_sayisi == 1
    yeni_id = sonuc.kalemler[0].yeni_kartlar[0][0]
    db = db_env()
    try:
        tasinan = db.get(models.CaseHizmeti, arb_satir_id)             # aynı satır, yerinde
        foy = db.query(models.CaseFoy).filter_by(sistem_no="ARB-16361").one()
        assert tasinan.foy_id == foy.id and tasinan.hizmet_turu == "Vekaletsiz Takip"
        assert (tasinan.case_id, tasinan.case_party_id) == (yeni_id, foy.case_party_id)
        assert db.get(models.CaseParty, tasinan.case_party_id).case_id == yeni_id
        kalanlar = {(s.hizmet_turu, s.foy_id is not None, s.case_party_id) for s in
                    db.query(models.CaseHizmeti).filter_by(case_id=kart_id)}
        assert kalanlar == {("Vekaletli Takip", True, taraf_id), ("Danışmanlık", False, taraf_id)}
        assert db.get(models.Case, yeni_id).hizmet_turu == "Vekaletsiz Takip"
        assert db.get(models.Case, kart_id).hizmet_turu == "Danışmanlık ; Vekaletli Takip"
        tarihce = {(h.case_id, h.old_value, h.new_value) for h in
                   db.query(models.CaseHistory).filter_by(field_name=case_hizmetleri.TARIHCE_ALANI)
                   if h.source and h.source.startswith("birlesik_kart_ayir")}
        assert tarihce == {
            (kart_id, "Quick Sigorta A.Ş. — Vekaletsiz Takip", ""),
            (yeni_id, "", "Quick Sigorta A.Ş. — Vekaletsiz Takip"),
        }
    finally:
        db.close()
