"""Aşama çelişki raporu üreticisi (G157) — kardeş föylerin `Karar_Asamalari`
satırlarını sınıflandırır, ekibe giden xlsx'i üretir, cevabı geri okur.

06.09'da ekibe giden `HUKDOK_ASAMA_CELISKILERI_2026-09-06.xlsx` (532 grup, 5
sınıf, sarı CEVABINIZ) tek seferlik scratch betiklerle üretilmişti
(`kanit_cikar.py` + `rapor_uret.py`, repoya girmemişti). Bu modül o üreticinin
repodaki kalıcı hâlidir; CLI `scripts/asama_celiski_raporu.py`.

Sözleşme (ekip aynı formatı bekliyor):

* Sınıf sayfaları `S<no>_<ad[:22]>`, sütun başlıkları `BASLIKLAR` — 06.09
  dosyasıyla BİREBİR; `OZET` sayfası `OZET_BASLIKLARI`. Ekip yalnız sarı
  `CEVABINIZ` (HATA / SEBEBİ VAR / BELİRSİZ) ve `AÇIKLAMANIZ` sütunlarını doldurur.
* Çelişki tanımı aktarımın tanımıdır: `hukdok_aktarim._asama_imzasi` (G150:
  boş hücre imzaya girmez; G151: büro durumu `Kapalı`/`Derdest` boş sayılır).
  İmzası aynı grup rapora DÜŞMEZ — o aşama zaten birleşik satır olarak yazılır.
  Bu yüzden 06.09'un S1 sınıfı (302 grup "yalnız boşluk") artık yalnız KALAN
  gerçek boşluğu taşır: aşama satırlarının HEPSİNDE boş olup ana sayfada dolu
  görünen künye alanı (aşama katmanı ana sayfadan künye doldurmaz).
* **S6 · Daire/numara eksik (yer tutucu)** — 06.09'un S5 satır 82 dersi:
  `DANIŞTAY . DAİRE` ≈ `DANIŞTAY DAİRE`, `2026/` ≈ `2026` normalize anahtarında
  eşit sayılıp `Farklı alanlar` boş gitmişti. Yer tutucu deseni (`DAİRE` öncesi
  numara yok, `YYYY/` numarasız, `?`) taşıyan künye farkı ayrı sınıftır ve
  ham yazımı farklı olan her yer tutucu alan listeye girer.
* **E-8 etiketi** — S5'te föyler ayrı müvekkile aitse (`case_foys.case_party_id`
  ya da ana sayfa `Müvekkil`) VE karar durumları zıt yönlüyse (`karar_yonu`:
  Kabul ⇄ Red, Onama ⇄ Bozma, Başvuru Ret ⇄ Kaldırma) "Bizim okuma" `E8_ETIKETI`
  ile başlar: sözlük E-8 kuralı karar durumunu müvekkil yönünden yazar, fark hata
  değildir (ekip cevabı 06.09 §2-b, 29 grup; kullanıcı kararı 08.09: föy düzeyi
  model YOK). Yön koşulunun gerekçesi `_KARAR_YONLERI` yorumunda.

Gerçek paket ve cevap dosyası REPOYA GİRMEZ (A.2); testler sentetik xlsx üretir.
"""
from __future__ import annotations

import collections
import difflib
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Set, Tuple, cast

from sqlalchemy.orm import Session

import models
from managers.reference_lists import tr_upper
from party_check import normalize_party_key
from scripts import hukdok_aktarim as ha

logger = logging.getLogger("AsamaCeliskiRaporu")

# ─── Sınıflar ────────────────────────────────────────────────────────────────

SINIFLAR: "collections.OrderedDict[str, Tuple[str, str, str]]" = collections.OrderedDict([
    ("S1_BOSLUK", ("1", "Yalnız boşluk",
                   "Künye alanı kardeş föylerin aşama satırlarının hepsinde boş, ana sayfada dolu; "
                   "dolu değerler birbiriyle çelişmiyor.")),
    ("S2_YAZIM", ("2", "Yazım hatası adayı",
                  "Aynı mahkeme + aynı esas; karar numarası ya da tarihte tek rakam / yıl farkı.")),
    ("S3_BASKA_KARAR", ("3", "Aynı dava, bambaşka karar künyesi",
                        "Aynı mahkeme + aynı esas; karar numarası ve/veya tarihi bambaşka.")),
    ("S4_MAHKEME_ESAS", ("4", "Mahkeme ya da esas farklı",
                         "Aynı kartın kardeş föyleri aynı aşamada farklı mahkeme ya da esas numarası taşıyor.")),
    ("S5_DURUM", ("5", "Yalnız karar durumu farklı",
                  "Mahkeme, esas, karar no ve tarih aynı; yalnız karar durumu farklı.")),
    ("S6_YER_TUTUCU", ("6", "Daire/numara eksik (yer tutucu)",
                       "Mahkeme ya da esas numarası yer tutucu (daire numarasız 'DAİRE', numarasız 'YYYY/', '?'); "
                       "kardeş föyler yer tutucuyu farklı yazmış ya da biri gerçek değeri taşıyor.")),
])

SORULAR: Dict[str, str] = {
    "S1_BOSLUK": "Ana sayfadaki değer aşama satırına taşınabilir mi?",
    "S2_YAZIM": "Hangi föy doğru?",
    "S3_BASKA_KARAR": "HATA mı, ek karar mı? Ek kararsa hangisi asıl karar?",
    "S4_MAHKEME_ESAS": "Yazım ve bayat föyler: HATA teyidi. Çok tur olanlar: iki turu da tüm kardeş föylere yazar mısınız?",
    "S5_DURUM": "Fark müvekkil yönünden mi (SEBEBİ VAR — E-8 etiketli satırlar), yoksa HATA mı?",
    "S6_YER_TUTUCU": "Daire numarası / esas numarası tamamlanabilir mi?",
}

E8_ETIKETI = "müvekkil yönü farkı (E-8) — hata değil"

# Sütun başlıkları — 06.09 dosyasıyla BİREBİR (ekip aynı formatı bekliyor; test kilitler).
BASLIKLAR: Tuple[str, ...] = (
    "Kart", "Ofis no", "Aşama", "Föyler", "Müvekkiller", "Farklı alanlar", "Değerler (föy=değer)",
    "Ana sayfa değerleri", "Bizim okuma", "TKU aynı", "Hasar no aynı", "Karşı taraf aynı",
    "Yerel mahkeme+esas aynı", "CEVABINIZ", "AÇIKLAMANIZ",
)
OZET_BASLIKLARI: Tuple[str, ...] = ("Sınıf", "Ad", "Grup", "Ne görüyoruz", "Sorumuz")
OZET_SAYFASI = "OZET"
CEVAP_NOTU = "CEVABINIZ sütunu: HATA / SEBEBİ VAR / BELİRSİZ · AÇIKLAMANIZ: doğru değer ya da sebep"
SARI = "FFF3BF"
SUTUN_GENISLIKLERI = (7, 30, 10, 26, 40, 18, 60, 60, 40, 8, 8, 8, 10, 14, 40)
DOSYA_ONEKI = "HUKDOK_ASAMA_CELISKILERI"

KUNYE_ALANLARI: Tuple[str, ...] = ("mahkeme", "esas_no", "karar_no", "karar_tarihi")
ALANLAR: Tuple[str, ...] = KUNYE_ALANLARI + ("karar_durumu",)
ALAN_AD: Dict[str, str] = {
    "mahkeme": "Mahkeme", "esas_no": "Esas", "karar_no": "Karar no",
    "karar_tarihi": "Karar tarihi", "karar_durumu": "Karar durumu",
}
ASAMA_AD: Dict[str, str] = {
    "YEREL": "Yerel", "ISTINAF": "İstinaf", "TEMYIZ": "Temyiz", "KARAR_DUZELTME": "Karar Düzeltme",
}
# Ana sayfanın (Sheet) aynı aşamayı anlatan sütunları — "Ana sayfa değerleri" sütunu
# ve S1 boşluk ölçümü buradan okur (orijinal başlıklar, `HamSatir.ham`).
SHEET_SUTUNLARI: Dict[str, Dict[str, str]] = {
    "YEREL": {"mahkeme": "Yerel Mahkeme", "esas_no": "Esas", "karar_no": "Karar No",
              "karar_tarihi": "Yerel Mahkeme Karar Tarihi", "karar_durumu": "Yerel Mahkeme Karar Durumu"},
    "ISTINAF": {"mahkeme": "İstinaf Mahkemesi", "esas_no": "İstinaf Mahkeme Esas",
                "karar_no": "İstinaf Mahkeme Karar No", "karar_tarihi": "İstinaf Mahkeme Karar Tar.",
                "karar_durumu": "İstinaf Karar Durumu"},
    "TEMYIZ": {"mahkeme": "Temyiz Mahkemesi", "esas_no": "Temyiz_Esas_No",
               "karar_tarihi": "Temyiz Karar Tarihi", "karar_durumu": "Yargıtay Onama Durumu"},
    "KARAR_DUZELTME": {"karar_durumu": "Karar Düzeltme Kararı Durumu"},
}
KANIT_ANAHTARLARI: Tuple[str, ...] = ("TKU", "hasar_no", "karsi_taraf", "yerel_mahkeme_esas")

# Cevap sütunu anahtarları (`_baslik_anahtari` sonrası).
CEVAP_HATA = "HATA"
CEVAP_BELIRSIZ = "BELIRSIZ"
CEVAP_SEBEBI_VAR = "SEBEBIVAR"
CEVAP_BOS = "(boş)"
CEVAP_TANINMAYAN = "(tanınmayan)"
CEVAPLAR: Tuple[str, ...] = (CEVAP_HATA, CEVAP_BELIRSIZ, CEVAP_SEBEBI_VAR, CEVAP_BOS, CEVAP_TANINMAYAN)


# ─── Veri yapıları ───────────────────────────────────────────────────────────

@dataclass
class FoyAsama:
    """Bir föyün bir aşamadaki satırı + raporun kanıt sütunları için bağlamı."""
    sistem_no: str
    satir: Dict[str, Any]                          # Karar_Asamalari alanları (ASAMA_SUTUNLARI anahtarları)
    sheet_asama: Dict[str, Any] = field(default_factory=dict)   # ana sayfanın aynı aşama sütunları
    muvekkil: Optional[str] = None
    muvekkil_kimligi: Optional[str] = None         # E-8: case_party_id ya da müvekkil anahtarı
    karsi_taraf: Optional[str] = None
    sigortali: Optional[str] = None
    tku: Optional[str] = None
    hasar_no: Optional[str] = None
    sheet_esas: Optional[str] = None
    eski_dosya_no: Optional[str] = None
    yerel_mahkeme: Optional[str] = None
    zincir_esaslari: Set[str] = field(default_factory=set)   # föyün tüm aşamalarındaki esas anahtarları


@dataclass
class AsamaGrubu:
    """Girdi: aynı kartın aynı aşamadaki kardeş föyleri."""
    kart_id: int
    ofis_no: str
    asama: str                                     # YEREL | ISTINAF | TEMYIZ | KARAR_DUZELTME
    foyler: List[FoyAsama]
    tur_notu: str = ""                             # "2. tur" gibi (çok turlu aşama)
    sheet_kunye_celiskisi: bool = False            # `celiskileri_bul`: ana sayfa karar no/tarihi de çelişiyor


@dataclass
class CeliskiGrubu:
    """Çıktı: raporun bir satırı."""
    kart_id: int
    ofis_no: str
    asama: str
    sinif: str
    sebep: str
    farkli_alanlar: List[str]
    kanit: Dict[str, bool]
    foyler: List[FoyAsama]
    etiket: str = ""                               # E-8 ön etiketi (S5)

    @property
    def asama_adi(self) -> str:
        return ASAMA_AD.get(self.asama, self.asama)

    @property
    def okuma(self) -> str:
        return f"{self.etiket}; {self.sebep}" if self.etiket else self.sebep


@dataclass
class DuzeltmeBeklentisi:
    """Ekip HATA dediyse sonraki paketin `Düzeltme_Logu`'nda beklediğimiz kalem."""
    sinif: str
    kart: str
    ofis_no: str
    asama: str
    foyler: str
    farkli_alanlar: str
    aciklama: str


@dataclass
class CevapOzeti:
    sayim: Dict[str, "collections.Counter[str]"] = field(default_factory=dict)   # sınıf → cevap → adet
    beklentiler: List[DuzeltmeBeklentisi] = field(default_factory=list)
    e8_etiketli: int = 0
    e8_sebebi_var: int = 0

    @property
    def toplam(self) -> "collections.Counter[str]":
        toplam: "collections.Counter[str]" = collections.Counter()
        for sayac in self.sayim.values():
            toplam.update(sayac)
        return toplam


# ─── Normalizasyon ───────────────────────────────────────────────────────────

def _metin(deger: Any) -> Optional[str]:
    return ha._metin(deger)


def _ham(satir: Mapping[str, Any], alan: str) -> Optional[str]:
    """Ham yazım (boşluk normalize); tarih alanı ISO'nun ilk 10 karakteri."""
    deger = satir.get(alan)
    if deger in (None, ""):
        return None
    if alan in ("karar_tarihi", "teblig_tarihi", "basvuru_tarihi"):
        return ha._tarih_yumusak(deger, alan) or None
    return _metin(deger)


def _sade(deger: str) -> str:
    """TR büyük harf + aksan sadeleştirme, boşluk/noktalama KORUNUR (yer tutucu deseni için)."""
    ayrik = unicodedata.normalize("NFD", tr_upper(deger))
    return "".join(ch for ch in ayrik if not unicodedata.combining(ch))


# "DAİRE" öncesinde herhangi bir yerde rakam varsa daire numaralıdır ("DANIŞTAY 10. DAİRE",
# "ANKARA 3 İDARİ DA ADAİRESİ" — ikincisi yazım hatası, yer tutucu değil → S4).
_DAIRE_NUMARALI = re.compile(r"\d.*DAIRE")
_YIL_NUMARASIZ = re.compile(r"^\d{4}\s*/?\s*$")


def yer_tutucu_mu(alan: str, deger: Optional[str]) -> bool:
    """Değer yer tutucu mu? (`DAİRE` öncesi numara yok, `YYYY/` numarasız, `?`).

    06.09 S5 satır 82: `DANIŞTAY . DAİRE` / `DANIŞTAY DAİRE` ve `2026/` / `2026`
    — ekibin §2.6 "daire numarası eksik" ailesi.
    """
    metin = _metin(deger)
    if not metin:
        return False
    if "?" in metin:
        return True
    sade = _sade(metin)
    if alan == "mahkeme":
        return "DAIRE" in sade and _DAIRE_NUMARALI.search(sade) is None
    if alan in ("esas_no", "karar_no"):
        return _YIL_NUMARASIZ.match(sade) is not None
    return False


def _imza(satir: Mapping[str, Any]) -> Dict[str, str]:
    """Aktarımın çelişki anahtarı (`_asama_imzasi`): tek tanım, tek yer."""
    return ha._asama_imzasi(ha.HamSatir(satir_no=0, degerler=dict(satir)))


def _anahtar(alan: str, deger: Optional[str]) -> Optional[str]:
    """Sınıflandırma anahtarı (yazım duyarsız): tarih ISO, diğerleri başlık anahtarı."""
    if deger is None:
        return None
    if alan == "karar_tarihi":
        return deger[:10]
    return ha._baslik_anahtari(deger) or None


def _yakin(a: str, b: str) -> bool:
    """Tek rakam farkı ya da rakam sırası karışmış."""
    if len(a) != len(b):
        return False
    return sum(x != y for x, y in zip(a, b, strict=True)) <= 1 or sorted(a) == sorted(b)


# ─── Sınıflandırma ───────────────────────────────────────────────────────────

def _degerler(foyler: Sequence[FoyAsama], alan: str) -> Dict[str, str]:
    """İmza anahtarı → ilk ham yazım (yalnız DOLU değerler)."""
    sonuc: Dict[str, str] = {}
    for foy in foyler:
        imza = _imza(foy.satir)
        if alan in imza and imza[alan] not in sonuc:
            sonuc[imza[alan]] = _ham(foy.satir, alan) or imza[alan]
    return sonuc


def _sebep_s4(grup: AsamaGrubu, kunye: Sequence[str]) -> str:
    sebepler: Set[str] = set()
    foyler = grup.foyler
    if "mahkeme" in kunye:
        degerler = sorted({_anahtar("mahkeme", _ham(f.satir, "mahkeme")) or "" for f in foyler} - {""})
        if degerler and all(
            difflib.SequenceMatcher(None, degerler[0], b).ratio() > 0.8
            or degerler[0].startswith(b) or b.startswith(degerler[0])
            for b in degerler[1:]
        ):
            sebepler.add("mahkeme adı yazımı/eksik")
        elif grup.asama != "YEREL" and any(
            _anahtar("mahkeme", _ham(f.satir, "mahkeme")) == _anahtar("mahkeme", f.yerel_mahkeme)
            and f.yerel_mahkeme for f in foyler
        ):
            sebepler.add("üst aşama satırında yerel mahkeme (aşama yanlış dosyalanmış)")
        else:
            sebepler.add("mahkeme bambaşka")
    if "esas_no" in kunye:
        esaslar = {f.sistem_no: _anahtar("esas_no", _ham(f.satir, "esas_no")) for f in foyler}
        farkli = sorted({v for v in esaslar.values() if v})
        if any(a.startswith(b) or b.startswith(a) for a in farkli for b in farkli if a != b):
            sebepler.add("esas eksik yazılmış")
        elif any(
            esaslar[a.sistem_no] and esaslar[a.sistem_no] != esaslar[b.sistem_no]
            and esaslar[a.sistem_no] in b.zincir_esaslari
            for a in foyler for b in foyler if a is not b
        ):
            sebepler.add("bayat föy: bir föy eski esasta kalmış")
        elif grup.asama != "YEREL" and len(
            {_anahtar("karar_tarihi", _ham(f.satir, "karar_tarihi")) for f in foyler} - {None}
        ) > 1:
            sebepler.add("çok tur adayı: aynı aşamada iki ayrı karar (bozma sonrası?)")
        else:
            sebepler.add("esas bambaşka")
    if grup.tur_notu:
        sebepler.add(grup.tur_notu)
    return "; ".join(sorted(sebepler))


def _kanit(foyler: Sequence[FoyAsama]) -> Dict[str, bool]:
    def ayni(degerler: Iterable[Optional[Any]]) -> bool:
        kume = set(degerler)
        return len(kume) == 1 and None not in kume

    return {
        "TKU": ayni(f.tku for f in foyler),
        "hasar_no": ayni(f.hasar_no for f in foyler),
        "karsi_taraf": ayni(normalize_party_key(f.karsi_taraf) or None if f.karsi_taraf else None for f in foyler),
        "yerel_mahkeme_esas": ayni(
            (ha._baslik_anahtari(f.yerel_mahkeme or ""), ha._baslik_anahtari(f.sheet_esas or ""))
            if f.sheet_esas else None
            for f in foyler
        ),
    }


# Karar durumunun YÖNÜ (E-8 etiketi için). Ekibin 06.09 cevabında S5'in 82
# grubunun 82'sinde föyler ayrı müvekkile aitti (kardeş föy tanımı gereği:
# hekim föyü + sigorta föyü) — "müvekkil farklı" tek başına ayırıcı değil.
# Ayırıcı olan sonucun yönü: SEBEBİ VAR dediği 26 grubun 25'i zıt yönlü çift
# (Kabul ⇄ Red 20, Onama ⇄ Bozma 4, Başvuru Ret ⇄ Kaldırma 1); aynı yönün
# ayrıntı farkı (Kabul ⇄ Kabul/Kısmen, Red/Esastan ⇄ Red/Husumet) BELİRSİZ ya
# da HATA. Anahtar `_baslik_anahtari` ile (aksan/boşluk duyarsız), ön ek eşleşmesi.
_KARAR_YONLERI: Tuple[Tuple[str, str], ...] = (
    ("KABUL", "KABUL"), ("KISMENKABUL", "KABUL"),
    ("RED", "RED"), ("RET", "RED"), ("ACILMAMIS", "RED"), ("KARARVERILMESINEYER", "RED"),
    ("DUZELTEREKONAMA", "ONAMA"), ("ONAMA", "ONAMA"), ("BOZMA", "BOZMA"),
    ("BASVURURET", "BASVURURET"), ("BASVURUNUN", "BASVURURET"), ("KALDIRMA", "KALDIRMA"),
)


def karar_yonu(deger: Optional[str]) -> Optional[str]:
    """Karar durumunun yönü (KABUL / RED / ONAMA / BOZMA / BASVURURET / KALDIRMA); tanınmayan → None."""
    anahtar = ha._baslik_anahtari(_metin(deger) or "")
    if not anahtar:
        return None
    for onek, yon in _KARAR_YONLERI:
        if anahtar.startswith(onek):
            return yon
    return None


def _e8_etiketi(foyler: Sequence[FoyAsama]) -> str:
    """E-8: föyler ayrı müvekkile ait VE karar durumları zıt yönlü."""
    kimlikler = {f.muvekkil_kimligi for f in foyler if f.muvekkil_kimligi}
    if len(kimlikler) < 2:
        return ""
    yonler = {karar_yonu(_ham(f.satir, "karar_durumu")) for f in foyler} - {None}
    return E8_ETIKETI if len(yonler) > 1 else ""


def _kalan_bosluk(grup: AsamaGrubu) -> List[str]:
    """Aşama satırlarının HEPSİNDE boş, ana sayfada (en az bir föyde) dolu künye alanları."""
    bosluk: List[str] = []
    for alan in ALANLAR:
        if any(alan in _imza(f.satir) for f in grup.foyler):
            continue
        if any(_metin(f.sheet_asama.get(alan)) for f in grup.foyler):
            bosluk.append(alan)
    return bosluk


def siniflandir(grup: AsamaGrubu) -> Optional[CeliskiGrubu]:
    """Bir (kart, aşama) grubunu sınıflandırır; çelişki yoksa None.

    Çelişki = aktarımın çelişkisi (`_asama_imzasi`: iki farklı DOLU değer).
    Aşama satırları birleşiyorsa yalnız "kalan boşluk" (S1) bakılır.
    """
    if len(grup.foyler) < 2:
        return None
    gercek = [alan for alan in ALANLAR if len(_degerler(grup.foyler, alan)) > 1]
    kanit = _kanit(grup.foyler)

    if not gercek:
        if grup.tur_notu:
            return CeliskiGrubu(
                grup.kart_id, grup.ofis_no, grup.asama, "S4_MAHKEME_ESAS", grup.tur_notu,
                [], kanit, list(grup.foyler),
            )
        bosluk = _kalan_bosluk(grup)
        if not bosluk:
            return None
        return CeliskiGrubu(
            grup.kart_id, grup.ofis_no, grup.asama, "S1_BOSLUK",
            "aşama satırlarında boş, ana sayfada (Sheet) dolu — aşama katmanı ana sayfadan künye almaz",
            bosluk, kanit, list(grup.foyler),
        )

    kunye = [alan for alan in gercek if alan != "karar_durumu"]

    # S6 — yer tutucu: farklı künye alanlarının HEPSİNDE en az bir yer tutucu değer.
    if kunye and all(any(yer_tutucu_mu(alan, v) for v in _degerler(grup.foyler, alan).values()) for alan in kunye):
        farkli = list(gercek)
        for alan in KUNYE_ALANLARI:          # ham yazımı farklı öteki yer tutucular da listeye
            hamlar = {_ham(f.satir, alan) for f in grup.foyler} - {None}
            if alan not in farkli and len(hamlar) > 1 and any(yer_tutucu_mu(alan, v) for v in hamlar):
                farkli.append(alan)
        farkli.sort(key=ALANLAR.index)
        return CeliskiGrubu(
            grup.kart_id, grup.ofis_no, grup.asama, "S6_YER_TUTUCU",
            "Daire/numara eksik (yer tutucu)" + (f"; {grup.tur_notu}" if grup.tur_notu else ""),
            farkli, kanit, list(grup.foyler),
        )

    if not kunye:
        return CeliskiGrubu(
            grup.kart_id, grup.ofis_no, grup.asama, "S5_DURUM",
            "künye aynı, yalnız karar durumu farklı", gercek, kanit, list(grup.foyler),
            etiket=_e8_etiketi(grup.foyler),
        )

    if "mahkeme" in kunye or "esas_no" in kunye:
        return CeliskiGrubu(
            grup.kart_id, grup.ofis_no, grup.asama, "S4_MAHKEME_ESAS", _sebep_s4(grup, kunye),
            gercek, kanit, list(grup.foyler),
        )

    # Aynı mahkeme + esas: karar no / tarih farkı → yazım adayı ya da bambaşka karar
    yazim = True
    for alan in kunye:
        degerler = sorted({_anahtar(alan, v) or "" for v in _degerler(grup.foyler, alan).values()} - {""})
        if len(degerler) != 2:
            yazim = False
            break
        a, b = degerler
        if alan == "karar_tarihi":
            if not (a[5:] == b[5:] or a[:7] == b[:7]):
                yazim = False
        elif not _yakin(a, b):
            yazim = False
    ek = " (ana sayfa karar no/tarihi de çelişiyor)" if grup.sheet_kunye_celiskisi else ""
    if yazim:
        return CeliskiGrubu(
            grup.kart_id, grup.ofis_no, grup.asama, "S2_YAZIM",
            "aynı mahkeme+esas; karar no/tarihte tek rakam ya da yıl farkı" + ek,
            gercek, kanit, list(grup.foyler),
        )
    yil_uyumsuz = False
    for f in grup.foyler:
        karar_no = _anahtar("karar_no", _ham(f.satir, "karar_no")) or ""
        tarih = _anahtar("karar_tarihi", _ham(f.satir, "karar_tarihi")) or ""
        if karar_no and tarih and karar_no[:4] != tarih[:4]:
            yil_uyumsuz = True
    sebep = "aynı mahkeme+esas; karar no ve/veya tarih bambaşka"
    if yil_uyumsuz:
        sebep += " (bir föyde karar no yılı ile karar tarihi yılı uyuşmuyor)"
    return CeliskiGrubu(
        grup.kart_id, grup.ofis_no, grup.asama, "S3_BASKA_KARAR", sebep + ek,
        gercek, kanit, list(grup.foyler),
    )


def siniflandir_hepsi(gruplar: Iterable[AsamaGrubu]) -> List[CeliskiGrubu]:
    sonuc = [c for c in (siniflandir(g) for g in gruplar) if c is not None]
    sonuc.sort(key=lambda c: (c.sinif, c.kart_id, c.asama))
    return sonuc


def sinif_sayimi(celiskiler: Sequence[CeliskiGrubu]) -> "collections.OrderedDict[str, int]":
    sayim: "collections.OrderedDict[str, int]" = collections.OrderedDict((k, 0) for k in SINIFLAR)
    for c in celiskiler:
        sayim[c.sinif] = sayim.get(c.sinif, 0) + 1
    return sayim


# ─── DB + paketten gruplar ───────────────────────────────────────────────────

def celiski_kart_idleri(celiskiler: Iterable[ha.Celiski]) -> Set[int]:
    """`celiskileri_bul` / `AktarimSonucu.celiskiler` çıktısından KART kümesi id'leri."""
    idler: Set[int] = set()
    for c in celiskiler:
        if c.kume == "KART":
            try:
                idler.add(int(c.kume_anahtari))
            except ValueError:
                continue
    return idler


def _muvekkil_kimligi(foy: models.CaseFoy, muvekkil: Optional[str]) -> Optional[str]:
    if foy.case_party_id is not None:
        return f"party:{foy.case_party_id}"
    anahtar = normalize_party_key(muvekkil) if muvekkil else ""
    return anahtar or None


def gruplari_topla(db: Session, paket: Path, *, yalniz_kartlar: Optional[Set[int]] = None) -> List[AsamaGrubu]:
    """Paketin `Karar_Asamalari` satırlarını DB'deki föy→kart bağıyla (kart, aşama) gruplarına ayırır.

    Kapsam dışı föyler (`kapsam_durumu` dolu) katılmaz (G113). `yalniz_kartlar`
    verilirse (örn. `celiski_kart_idleri(sonuc.celiskiler)`) yalnız o kartlar.
    Çok turlu aşamada föyler farklı sayıda satır anlatıyorsa grup `tur_notu` ile
    işaretlenir; aynı sayıdaysa tur tur ayrı grup.
    """
    satirlar, _ = ha.xlsx_oku(paket)
    sheet: Dict[str, ha.HamSatir] = {}
    for s in satirlar:
        sn = _metin(s.degerler.get("sistem_no"))
        if sn:
            sheet[sn] = s
    asama_satirlari = ha.asama_satirlarini_oku(paket)

    # cast: eski stil Column() modellerinde mypy instance alanını Column[...] görür
    # (Mapped[] geçişine kadar) — köprü tek noktada.
    foys: Dict[str, models.CaseFoy] = {
        cast(str, f.sistem_no): f for f in db.query(models.CaseFoy).all()
    }
    kart_ids = {cast(int, f.case_id) for f in foys.values()}
    cases: Dict[int, models.Case] = {
        cast(int, c.id): c for c in db.query(models.Case).filter(models.Case.id.in_(kart_ids)).all()
    } if kart_ids else {}

    # Ana sayfa künye çelişkileri (kart düzeyi, `celiskileri_bul` yalnız çağrılır)
    kunye_kayitlari = [
        {
            "sistem_no": sn, "case_id": cast(int, foys[sn].case_id),
            "tku_no": _metin(s.degerler.get("tku_no")),
            "karar_no": _metin(s.degerler.get("karar_no")),
            "karar_tarihi": ha._tarih_yumusak(s.degerler.get("karar_tarihi"), "karar_tarihi"),
        }
        for sn, s in sheet.items() if sn in foys and not foys[sn].kapsam_durumu
    ]
    sheet_celiskili = celiski_kart_idleri(ha.celiskileri_bul(kunye_kayitlari))

    zincir: Dict[str, List[ha.HamSatir]] = collections.defaultdict(list)
    grup: Dict[Tuple[int, str], Dict[str, List[ha.HamSatir]]] = collections.defaultdict(
        lambda: collections.defaultdict(list)
    )
    for s in asama_satirlari:
        sn = _metin(s.degerler.get("sistem_no")) or ""
        foy = foys.get(sn)
        if foy is None or foy.kapsam_durumu:
            continue
        case_id = cast(int, foy.case_id)
        if yalniz_kartlar is not None and case_id not in yalniz_kartlar:
            continue
        zincir[sn].append(s)
        stage = ha.ASAMA_ESLEMESI.get(ha._baslik_anahtari(_metin(s.degerler.get("asama")) or ""))
        if stage:
            grup[(case_id, stage)][sn].append(s)

    def foy_baglami(sn: str, stage: str, satir: ha.HamSatir) -> FoyAsama:
        foy = foys[sn]
        sh = sheet.get(sn)
        ham: Dict[str, Any] = dict(foy.ham_veri or {})
        if sh is not None:
            ham.update(sh.ham)
        degerler: Dict[str, Any] = sh.degerler if sh is not None else {}
        muvekkil = _metin(degerler.get("muvekkil")) or _metin(ham.get("Müvekkil"))
        esaslar = {_anahtar("esas_no", _ham(r.degerler, "esas_no")) for r in zincir[sn]}
        esaslar |= {ha._baslik_anahtari(ham.get("Esas") or ""), ha._baslik_anahtari(ham.get("Eski Dosya No") or "")}
        return FoyAsama(
            sistem_no=sn,
            satir=dict(satir.degerler),
            sheet_asama={alan: ham.get(kol) for alan, kol in SHEET_SUTUNLARI.get(stage, {}).items()},
            muvekkil=muvekkil,
            muvekkil_kimligi=_muvekkil_kimligi(foy, muvekkil),
            karsi_taraf=_metin(degerler.get("karsi_taraf")) or _metin(ham.get("Karşı Taraf")),
            sigortali=_metin(degerler.get("sigortali")) or _metin(ham.get("Sigortalı")),
            tku=cast(Optional[str], foy.tku_no),
            hasar_no=cast(Optional[str], foy.hasar_no),
            sheet_esas=_metin(ham.get("Esas")),
            eski_dosya_no=_metin(ham.get("Eski Dosya No")),
            yerel_mahkeme=_metin(ham.get("Yerel Mahkeme")),
            zincir_esaslari={e for e in esaslar if e},
        )

    gruplar: List[AsamaGrubu] = []
    for (case_id, stage), foyler in sorted(grup.items()):
        if len(foyler) < 2:
            continue
        case = cases.get(case_id)
        ofis_no = cast(str, case.tracking_no) if case is not None else ""
        sirali = {sn: sorted(r, key=ha._asama_sira) for sn, r in sorted(foyler.items())}
        uzunluklar = {len(r) for r in sirali.values()}
        ortak = AsamaGrubu(case_id, ofis_no, stage, [], sheet_kunye_celiskisi=case_id in sheet_celiskili)
        if len(uzunluklar) != 1:
            ortak.tur_notu = "çok tur: föyler farklı sayıda tur anlatıyor"
            ortak.foyler = [foy_baglami(sn, stage, r[0]) for sn, r in sirali.items()]
            gruplar.append(ortak)
            continue
        for konum in range(uzunluklar.pop()):
            g = AsamaGrubu(
                case_id, ofis_no, stage,
                [foy_baglami(sn, stage, r[konum]) for sn, r in sirali.items()],
                tur_notu=f"{konum + 1}. tur" if konum else "",
                sheet_kunye_celiskisi=ortak.sheet_kunye_celiskisi,
            )
            gruplar.append(g)
    return gruplar


# ─── xlsx yazımı ─────────────────────────────────────────────────────────────

def sayfa_adi(sinif: str) -> str:
    no, ad, _ = SINIFLAR[sinif]
    return f"S{no}_{ad[:22]}".replace("/", "-")


def _d(deger: Any) -> str:
    metin = _metin(deger)
    return metin if metin else "—"


def satir_hucreleri(c: CeliskiGrubu) -> List[Any]:
    """Rapor satırı — `BASLIKLAR` sırasıyla (son iki hücre ekibin, boş)."""
    degerler = "\n".join(
        ALAN_AD.get(a, a) + ": " + " | ".join(f"{f.sistem_no}={_d(_ham(f.satir, a))}" for f in c.foyler)
        for a in c.farkli_alanlar
    )
    sheet = "\n".join(
        ALAN_AD.get(a, a) + ": " + " | ".join(f"{f.sistem_no}={_d(f.sheet_asama.get(a))}" for f in c.foyler)
        for a in c.farkli_alanlar
    )
    return [
        c.kart_id, c.ofis_no, c.asama_adi,
        "; ".join(f.sistem_no for f in c.foyler),
        "; ".join(_d(f.muvekkil) for f in c.foyler),
        ", ".join(ALAN_AD.get(a, a) for a in c.farkli_alanlar),
        degerler, sheet, c.okuma,
        *("evet" if c.kanit.get(k) else "hayır" for k in KANIT_ANAHTARLARI),
        "", "",
    ]


def xlsx_yaz(celiskiler: Sequence[CeliskiGrubu], yol: Path) -> Path:
    """OZET + sınıf sayfaları (06.09 formatı; boş sınıfın sayfası da yazılır)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sayim = sinif_sayimi(celiskiler)
    by: Dict[str, List[CeliskiGrubu]] = collections.defaultdict(list)
    for c in celiskiler:
        by[c.sinif].append(c)

    bold = Font(bold=True)
    sari = PatternFill("solid", fgColor=SARI)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    ws0 = wb.active
    ws0.title = OZET_SAYFASI
    ws0.append(list(OZET_BASLIKLARI))
    for sinif, (no, ad, tanim) in SINIFLAR.items():
        ws0.append([no, ad, sayim[sinif], tanim, SORULAR[sinif]])
    ws0.append(["", "Toplam", len(celiskiler), "", ""])
    ws0.append([])
    ws0.append([CEVAP_NOTU])
    for harf, genislik in zip("ABCDE", (6, 30, 8, 80, 60), strict=True):
        ws0.column_dimensions[harf].width = genislik
    for hucre in ws0[1]:
        hucre.font = bold

    for sinif in SINIFLAR:
        ws = wb.create_sheet(sayfa_adi(sinif))
        ws.append(list(BASLIKLAR))
        for hucre in ws[1]:
            hucre.font = bold
            hucre.alignment = wrap
        for c in by.get(sinif, []):
            ws.append(satir_hucreleri(c))
        for i, genislik in enumerate(SUTUN_GENISLIKLERI, start=1):
            ws.column_dimensions[get_column_letter(i)].width = genislik
        for satir in ws.iter_rows(min_row=2):
            for hucre in satir:
                hucre.alignment = wrap
            satir[13].fill = sari
            satir[14].fill = sari
        ws.freeze_panes = "A2"

    yol = Path(yol)
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb.save(yol)
    wb.close()
    return yol


def rapor_uret(db: Session, paket: Path, cikti_dizini: Path, *, tarih: str,
               yalniz_kartlar: Optional[Set[int]] = None) -> Tuple[Path, List[CeliskiGrubu]]:
    """Paket + DB → `<cikti_dizini>/HUKDOK_ASAMA_CELISKILERI_<tarih>.xlsx`."""
    gruplar = gruplari_topla(db, paket, yalniz_kartlar=yalniz_kartlar)
    celiskiler = siniflandir_hepsi(gruplar)
    yol = xlsx_yaz(celiskiler, Path(cikti_dizini) / f"{DOSYA_ONEKI}_{tarih}.xlsx")
    logger.info(f"Aşama çelişki raporu: {len(gruplar)} grup okundu, {len(celiskiler)} çelişki → {yol}")
    return yol, celiskiler


# ─── Cevabı geri okuma ───────────────────────────────────────────────────────

def cevap_anahtari(deger: Any) -> str:
    """CEVABINIZ hücresi → HATA | BELIRSIZ | SEBEBIVAR | (boş) | (tanınmayan)."""
    metin = _metin(deger)
    if not metin:
        return CEVAP_BOS
    anahtar = ha._baslik_anahtari(metin)
    for aday in (CEVAP_HATA, CEVAP_BELIRSIZ, CEVAP_SEBEBI_VAR):
        if anahtar.startswith(aday):
            return aday
    return CEVAP_TANINMAYAN


def _sutun_indeksleri(baslik: Sequence[Any]) -> Dict[str, int]:
    anahtarlar = {ha._baslik_anahtari(b): i for i, b in enumerate(baslik) if b is not None}
    return {ad: anahtarlar[ha._baslik_anahtari(ad)] for ad in BASLIKLAR if ha._baslik_anahtari(ad) in anahtarlar}


def _hucre(satir: Sequence[Any], indeks: Mapping[str, int], ad: str) -> str:
    i = indeks.get(ad)
    if i is None or i >= len(satir):
        return ""
    return _metin(satir[i]) or ""


def cevaplari_oku(yol: Path) -> CevapOzeti:
    """Ekibin CEVAPLI xlsx'i → sınıf başına cevap sayımı + HATA satırlarından
    `Düzeltme_Logu` beklentisi listesi. Sınıf sayfası adıyla (`S<no>_`) tanınır;
    sütunlar başlıkla bulunur (aksan/boşluk duyarsız)."""
    from openpyxl import load_workbook

    yol = Path(yol)
    if not yol.exists():
        raise ha.AktarimHatasi(f"Cevap dosyası yok: {yol}")
    sinif_adlari = {sayfa_adi(s)[:3]: s for s in SINIFLAR}     # "S1_" → sınıf
    ozet = CevapOzeti()
    wb = load_workbook(yol, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            sinif = sinif_adlari.get(ws.title[:3])
            if sinif is None:
                continue
            satirlar = list(ws.iter_rows(values_only=True))
            if not satirlar:
                continue
            indeks = _sutun_indeksleri(satirlar[0])
            if "CEVABINIZ" not in indeks:
                logger.warning(f"{ws.title}: CEVABINIZ sütunu yok — sayfa atlandı")
                continue
            sayac: "collections.Counter[str]" = collections.Counter()
            for satir in satirlar[1:]:
                if not any(v not in (None, "") for v in satir):
                    continue
                cevap = cevap_anahtari(_hucre(satir, indeks, "CEVABINIZ"))
                sayac[cevap] += 1
                okuma = _hucre(satir, indeks, "Bizim okuma")
                if okuma.startswith(E8_ETIKETI):
                    ozet.e8_etiketli += 1
                    if cevap == CEVAP_SEBEBI_VAR:
                        ozet.e8_sebebi_var += 1
                if cevap == CEVAP_HATA:
                    ozet.beklentiler.append(DuzeltmeBeklentisi(
                        sinif=sinif, kart=_hucre(satir, indeks, "Kart"),
                        ofis_no=_hucre(satir, indeks, "Ofis no"), asama=_hucre(satir, indeks, "Aşama"),
                        foyler=_hucre(satir, indeks, "Föyler"),
                        farkli_alanlar=_hucre(satir, indeks, "Farklı alanlar"),
                        aciklama=_hucre(satir, indeks, "AÇIKLAMANIZ"),
                    ))
            ozet.sayim[sinif] = sayac
    finally:
        wb.close()
    return ozet


def cevap_ozeti_metni(ozet: CevapOzeti, *, beklenti_siniri: int = 20) -> str:
    satirlar = ["=" * 70, "Aşama çelişki cevabı — geri okuma (G157)", "=" * 70]
    for sinif, sayac in ozet.sayim.items():
        no, ad, _ = SINIFLAR[sinif]
        dagilim = " · ".join(f"{c}={sayac.get(c, 0)}" for c in CEVAPLAR if sayac.get(c, 0))
        satirlar.append(f"  S{no} {ad:<34} {sum(sayac.values()):>4}  {dagilim}")
    toplam = ozet.toplam
    satirlar.append(
        f"  {'Toplam':<37} {sum(toplam.values()):>4}  "
        + " · ".join(f"{c}={toplam.get(c, 0)}" for c in CEVAPLAR if toplam.get(c, 0))
    )
    satirlar.append(f"  E-8 etiketli satır: {ozet.e8_etiketli} (SEBEBİ VAR: {ozet.e8_sebebi_var})")
    satirlar.append(f"  Düzeltme_Logu beklentisi (HATA): {len(ozet.beklentiler)} kalem")
    for b in ozet.beklentiler[:beklenti_siniri]:
        satirlar.append(f"    - kart {b.kart} · {b.asama} · {b.foyler} · {b.farkli_alanlar}: {b.aciklama[:80]}")
    if len(ozet.beklentiler) > beklenti_siniri:
        satirlar.append(f"    … {len(ozet.beklentiler) - beklenti_siniri} kalem daha")
    satirlar.append("=" * 70)
    return "\n".join(satirlar)


def beklenti_csv_yaz(beklentiler: Sequence[DuzeltmeBeklentisi], yol: Path) -> Path:
    """HATA satırları → CSV (UTF-8 BOM, ';' — `hukdok_aktarim._csv_yaz` deseni)."""
    import csv

    yol = Path(yol)
    yol.parent.mkdir(parents=True, exist_ok=True)
    with open(yol, "w", newline="", encoding="utf-8-sig") as dosya:
        yazici = csv.writer(dosya, delimiter=";")
        yazici.writerow(["sinif", "kart", "ofis_no", "asama", "foyler", "farkli_alanlar", "aciklama"])
        for b in beklentiler:
            yazici.writerow([b.sinif, b.kart, b.ofis_no, b.asama, b.foyler, b.farkli_alanlar, b.aciklama])
    return yol
