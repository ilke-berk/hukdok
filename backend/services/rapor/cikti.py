"""Rapor çıktı üretimi — Excel (xlsx) ve CSV, akışlı (G131, plan K3).

Bellek disiplini: satırlar `motor.satirlari_akit` (`yield_per`) iteratöründen
gelir, TEK geçişte dosyaya yazılır; xlsx `openpyxl.Workbook(write_only=True)`
(satırlar sayfa temp dosyasına akar, `save` zip'ler), csv doğrudan dosya akışı.
Hiçbir yerde tüm satır listesi tutulmaz (backend 2g limit, OOM dersleri).

Dosya-önce kuralı: çıktı ÖNCE `hedef` yoluna yazılır, boyut + sha256 dosyadan
hesaplanır; route aynı dosyayı `FileResponse` ile verir → saklanan dosya ile
indirilen dosya AYNI bayttır, `report_runs.sha256` bunu kanıtlar.

Hücre kuralları:
- `tarih` tipinde kolonlar Excel'de GERÇEK tarih (motor ISO string verir, burada
  `date`/`datetime`'a çevrilir; biçim `DD.MM.YYYY` / `DD.MM.YYYY HH:MM`), CSV'de
  `GG.AA.YYYY` metni (Türkçe Excel doğrudan tanır). Saat dilimli zaman Türkiye saatine
  çevrilir, tz atılır (openpyxl tz'li datetime kabul etmez).
- `para` → `#,##0.00`; `sayi` olduğu gibi; `mantik` → Evet/Hayır. CSV'de `para` "1234,50",
  ondalıklı `sayi` virgüllü (Türkçe Excel `;` ayraçlı dosyada `,` ondalık bekler).
- Başlık satırı bordo zemin + beyaz kalın (`report_builder.rows_to_excel` stili,
  `4A1530`), dondurulmuş, otomatik filtreli.
- CSV: `utf-8-sig` + `;` (`services/teslim_cevap._csv_yaz` deseni) + formül
  enjeksiyonu koruması: `= + - @` ile başlayan METİN hücreye `'` öneki
  (sayısal hücreler etkilenmez; negatif sayı sayı olarak kalır).
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator

from schemas_rapor import SAAT_DILIMI, KolonBasligi

BASLIK_RENGI = "4A1530"           # report_builder.rows_to_excel ile aynı bordo
SAYFA_ADI = "Rapor"
TARIH_BICIMI = "DD.MM.YYYY"
ZAMAN_BICIMI = "DD.MM.YYYY HH:MM"
PARA_BICIMI = "#,##0.00"
CSV_AYRAC = ";"
CSV_KODLAMA = "utf-8-sig"
_ENJEKSIYON_ONEKLERI = ("=", "+", "-", "@")
_SHA_PARCA = 1024 * 1024

MEDYA_TURLERI: dict[str, str] = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
}


@dataclass(frozen=True)
class CiktiOzeti:
    satir_sayisi: int
    dosya_boyutu: int
    sha256: str


# ─── Değer çevirimleri ───────────────────────────────────────────────────────

def _tarihe_cevir(deger: Any) -> Any:
    """Motorun ISO string'ini `date`/`datetime`'a çevirir; çevrilemezse olduğu gibi.
    Saat dilimli değer (Postgres `timestamptz`; motor `+03:00` ile verir) Türkiye saatine çevrilip
    tz'siz bırakılır — openpyxl tz'li datetime'ı REDDEDER ("Excel does not support timezones",
    12.09 lokal Postgres bulgusu: `uploaded_at` kolonlu Excel export 500 veriyordu)."""
    if isinstance(deger, dt.datetime):
        return _yerel_saat(deger)
    if isinstance(deger, dt.date):
        return deger
    if not isinstance(deger, str) or not deger:
        return deger
    try:
        if len(deger) == 10:
            return dt.date.fromisoformat(deger)
        return _yerel_saat(dt.datetime.fromisoformat(deger))
    except ValueError:
        return deger


def _yerel_saat(zaman: dt.datetime) -> dt.datetime:
    if zaman.tzinfo is None:
        return zaman
    return zaman.astimezone(SAAT_DILIMI).replace(tzinfo=None)


def _sayi_metni(deger: Any, tip: str) -> Any:
    """CSV sayı hücresi: Türkçe Excel `;` ayraçlı dosyada ondalık VİRGÜL bekler (12.09). `para` daima
    iki basamak ("1234,50"); `sayi` tam sayıysa olduğu gibi, ondalıklıysa virgüllü. Sayı olmayan değer dokunulmaz."""
    if isinstance(deger, bool) or not isinstance(deger, (int, float)):
        return deger
    if tip == "para":
        return f"{deger:.2f}".replace(".", ",")
    return str(deger).replace(".", ",") if isinstance(deger, float) else deger


def _mantik_metni(deger: Any) -> Any:
    if isinstance(deger, bool):
        return "Evet" if deger else "Hayır"
    return deger


def _tarih_metni(deger: Any) -> Any:
    tarih = _tarihe_cevir(deger)
    if isinstance(tarih, dt.datetime):
        return tarih.strftime("%d.%m.%Y %H:%M:%S")
    if isinstance(tarih, dt.date):
        return tarih.strftime("%d.%m.%Y")
    return tarih


def csv_hucresi_koru(deger: Any) -> Any:
    """CSV formül enjeksiyonu: `=`, `+`, `-`, `@` ile başlayan metne `'` öneki.
    Yalnız `str` değerler; sayılar (negatif dahil) dokunulmaz."""
    if isinstance(deger, str) and deger.startswith(_ENJEKSIYON_ONEKLERI):
        return "'" + deger
    return deger


# ─── Excel ───────────────────────────────────────────────────────────────────

def xlsx_uret(kolonlar: list[KolonBasligi], satir_iter: Iterable[dict[str, Any]], hedef: Path) -> int:
    """Satır iteratörünü `hedef`e write-only xlsx olarak yazar; yazılan veri satır
    sayısını döner. Tarih hücreleri gerçek tarih, para `#,##0.00`."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(SAYFA_ADI)
    for i, kolon in enumerate(kolonlar, start=1):
        genislik = 14 if kolon.tip in ("tarih", "sayi", "para", "mantik") else 28
        ws.column_dimensions[get_column_letter(i)].width = max(genislik, min(len(kolon.etiket) + 2, 40))
    ws.freeze_panes = "A2"

    baslik_zemin = PatternFill(start_color=BASLIK_RENGI, end_color=BASLIK_RENGI, fill_type="solid")
    baslik_yazi = Font(bold=True, color="FFFFFF")
    baslik_hiza = Alignment(horizontal="left", vertical="center")
    basliklar = []
    for kolon in kolonlar:
        hucre = WriteOnlyCell(ws, value=kolon.etiket)
        hucre.fill = baslik_zemin
        hucre.font = baslik_yazi
        hucre.alignment = baslik_hiza
        basliklar.append(hucre)
    ws.append(basliklar)

    anahtarlar = [k.anahtar for k in kolonlar]
    tipler = [k.tip for k in kolonlar]
    sayac = 0
    for satir in satir_iter:
        hucreler: list[Any] = []
        for anahtar, tip in zip(anahtarlar, tipler, strict=True):
            deger = satir.get(anahtar)
            if deger is None:
                hucreler.append(None)
                continue
            if tip == "tarih":
                tarih = _tarihe_cevir(deger)
                if isinstance(tarih, (dt.datetime, dt.date)):
                    hucre = WriteOnlyCell(ws, value=tarih)
                    hucre.number_format = ZAMAN_BICIMI if isinstance(tarih, dt.datetime) else TARIH_BICIMI
                    hucreler.append(hucre)
                else:
                    hucreler.append(tarih)
            elif tip == "para" and isinstance(deger, (int, float)) and not isinstance(deger, bool):
                hucre = WriteOnlyCell(ws, value=deger)
                hucre.number_format = PARA_BICIMI
                hucreler.append(hucre)
            elif tip == "mantik":
                hucreler.append(_mantik_metni(deger))
            else:
                hucreler.append(deger)
        ws.append(hucreler)
        sayac += 1

    son_sutun = get_column_letter(max(1, len(kolonlar)))
    ws.auto_filter.ref = f"A1:{son_sutun}{sayac + 1}"
    wb.save(str(hedef))
    return sayac


# ─── CSV ─────────────────────────────────────────────────────────────────────

def _csv_degerleri(kolonlar: list[KolonBasligi], satir: dict[str, Any]) -> list[Any]:
    degerler: list[Any] = []
    for kolon in kolonlar:
        deger = satir.get(kolon.anahtar)
        if deger is None:
            degerler.append("")
            continue
        if kolon.tip == "tarih":
            deger = _tarih_metni(deger)
        elif kolon.tip == "mantik":
            deger = _mantik_metni(deger)
        elif kolon.tip in ("para", "sayi") and isinstance(deger, (int, float)) and not isinstance(deger, bool):
            # Sayıdan üretilen metin enjeksiyon korumasına GİRMEZ: "-250,00" negatif sayıdır, formül değil
            degerler.append(_sayi_metni(deger, kolon.tip))
            continue
        degerler.append(csv_hucresi_koru(deger))
    return degerler


def csv_uret(kolonlar: list[KolonBasligi], satir_iter: Iterable[dict[str, Any]], hedef: Path) -> int:
    """Satır iteratörünü `hedef`e `utf-8-sig` + `;` CSV olarak yazar (enjeksiyon
    korumalı); yazılan veri satır sayısını döner."""
    sayac = 0
    with open(hedef, "w", newline="", encoding=CSV_KODLAMA) as dosya:
        yazici = csv.writer(dosya, delimiter=CSV_AYRAC)
        yazici.writerow([csv_hucresi_koru(k.etiket) for k in kolonlar])
        for satir in satir_iter:
            yazici.writerow(_csv_degerleri(kolonlar, satir))
            sayac += 1
    return sayac


# ─── Ortak giriş ─────────────────────────────────────────────────────────────

_URETICILER: dict[str, Callable[[list[KolonBasligi], Iterable[dict[str, Any]], Path], int]] = {
    "xlsx": xlsx_uret,
    "csv": csv_uret,
}


def dosya_ozeti(yol: Path) -> tuple[int, str]:
    """(boyut, sha256) — dosya parça parça okunur."""
    ozet = hashlib.sha256()
    boyut = 0
    with open(yol, "rb") as dosya:
        for parca in iter(lambda: dosya.read(_SHA_PARCA), b""):
            ozet.update(parca)
            boyut += len(parca)
    return boyut, ozet.hexdigest()


def ciktiyi_yaz(format: str, kolonlar: list[KolonBasligi], satir_iter: Iterator[dict[str, Any]],
                hedef: Path) -> CiktiOzeti:
    """Formata göre üreticiyi seçer, dosyayı yazar, boyut + sha256'yı dosyadan hesaplar."""
    uretici = _URETICILER.get(format)
    if uretici is None:
        raise ValueError(f"tanınmayan çıktı formatı: {format}")
    hedef.parent.mkdir(parents=True, exist_ok=True)
    satir_sayisi = uretici(kolonlar, satir_iter, hedef)
    boyut, sha = dosya_ozeti(hedef)
    return CiktiOzeti(satir_sayisi=satir_sayisi, dosya_boyutu=boyut, sha256=sha)


__all__ = [
    "BASLIK_RENGI", "CSV_AYRAC", "CSV_KODLAMA", "CiktiOzeti", "MEDYA_TURLERI", "ciktiyi_yaz",
    "csv_hucresi_koru", "csv_uret", "dosya_ozeti", "xlsx_uret",
]
