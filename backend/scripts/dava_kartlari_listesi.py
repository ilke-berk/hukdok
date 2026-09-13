# -*- coding: utf-8 -*-
"""Dava kartları listesi — veri ekibine giden "HUKDOK_DAVA_KARTLARI_<tarih>.xlsx" (11.09 biçimi).

11.09.2026'da ekibe ilk kez gönderilen listenin üreticisi (o gün oturum geçici dosyasındaydı; ekip
12.09 cevabı §8'de "bundan sonra ofisin bütün dosyaları için bu tek liste üzerinden çalışacağız" dedi
→ repoya alındı, G181). Üç sayfa: `Dava Kartlari` (kart başına satır, 89 sütun 11.09 ile birebir +
sonda `İlişkili kartlar`), `Foyler` (föy başına satır), `Aciklama` (sayımlar + dışa aktarım SAATİ —
ekip 12.09 §8: "saatini de yazarsanız karşılaştırmalarımızı ona göre damgalarız"). Salt okunur.

    docker compose exec -T backend python scripts/dava_kartlari_listesi.py --out /app/calibration-data/_g179/HUKDOK_DAVA_KARTLARI_2026-09-13.xlsx
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import sys
from collections import defaultdict
from typing import Any, Dict, List
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import create_engine, text

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

TR = ZoneInfo("Europe/Istanbul")


def _join(vals) -> str:
    seen, out = set(), []
    for v in vals:
        if v is None:
            continue
        s = str(v).strip()
        if s and s not in seen:
            seen.add(s)
            out.append(s)
    return " ; ".join(out)


def _naive(ts):
    return ts.replace(tzinfo=None) if ts else None


HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
BODY_FONT = Font(name="Arial", size=10)


def _write_sheet(ws, cols, data):
    ws.append([c[0] for c in cols])
    for cell in ws[1]:
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for rec in data:
        ws.append([fn(rec) for _, fn in cols])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            if isinstance(cell.value, (dt.date, dt.datetime)):
                cell.number_format = "DD.MM.YYYY"
    for i, (name, _) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, len(name) + 2), 40)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def liste_uret(database_url: str, out: str) -> Dict[str, Any]:
    engine = create_engine(database_url)

    def rows(sql, **p):
        with engine.connect() as c:
            r = c.execute(text(sql), p)
            return [dict(zip(r.keys(), row, strict=True)) for row in r.fetchall()]

    simdi = dt.datetime.now(TR)
    cases = rows("""
        SELECT id, tracking_no, esas_no, status, active, file_type, sub_type, sub_type_extra, service_type,
               subject, court, judicial_unit, opening_date, acceptance_date, atama_tarihi, arsiv_tarihi,
               responsible_lawyer_name, uyap_lawyer_name, bureau_type, klasor_no_2, tku_no, sistem_no,
               hasar_dosya_no, hukuk_no,
               maddi_tazminat, manevi_tazminat, dava_degeri, islah_tutari, para_birimi,
               case_stage, dosya_son_durumu, yerel_karar_durumu, karar_tarihi, karar_turu, karar_lehine, karar_no,
               karar_teblig_tarihi, karar_aciklama, hukmedilen_maddi, hukmedilen_manevi, hukmedilen_toplam,
               istinaf_mahkemesi, istinaf_esas_no, istinaf_basvuru_tarihi, istinaf_karar_no, istinaf_karar_tarihi,
               istinaf_karar_durumu, istinaf_teblig_tarihi, istinaf_basvuran_taraf,
               temyiz_mahkemesi, temyiz_esas_no, temyiz_basvuru_tarihi, temyiz_karar_no, temyiz_karar_tarihi,
               temyiz_karar_durumu, temyiz_eden_durumu, temyiz_teblig_tarihi,
               karar_duzeltme_durumu, karar_duzeltme_esas_no, karar_duzeltme_karar_no, karar_duzeltme_tarihi,
               karar_duzeltme_teblig_tarihi, yeni_esas_no, kesinlesme_tarihi, infaz_tarihi,
               arabuluculuk_no, arabuluculuk_karar_tarihi,
               muvekkil_tipi, hizmet_turu, olay_turu, hukumdeki_rol,
               tibbi_surec, tibbi_olay, iddia_edilen_kusur, hastada_olusan_zarar, uygulanan_yontem,
               missing_required_bucket, notes, created_at, updated_at
        FROM cases WHERE deleted_at IS NULL ORDER BY id
    """)
    parties: Dict[int, Dict[str, List[str]]] = defaultdict(lambda: defaultdict(list))
    for p in rows("SELECT case_id, name, role, party_type FROM case_parties ORDER BY id"):
        parties[p["case_id"]][p["party_type"]].append(f'{p["name"]} ({p["role"]})' if p["role"] else p["name"])
    lawyers: Dict[int, List[str]] = defaultdict(list)
    for lw in rows("SELECT case_id, name FROM case_lawyers ORDER BY id"):
        lawyers[lw["case_id"]].append(lw["name"])
    esas_hist: Dict[int, List[str]] = defaultdict(list)
    for e in rows("SELECT case_id, esas_no, stage, court, is_current FROM case_esas_numbers ORDER BY case_id, id"):
        esas_hist[e["case_id"]].append(f'{e["esas_no"]} [{e["stage"]}{" · güncel" if e["is_current"] else ""}]')
    foys = rows("""
        SELECT f.id, f.sistem_no, f.case_id, f.tku_no, f.hasar_no, f.mko_id, f.muvekkil_no, f.muvekkil_tipi,
               f.hizmet_turu, f.durum, f.onceki_tracking_no, f.source, f.kapsam_durumu, f.kapsam_gerekcesi,
               f.kapsam_tarihi, f.updated_at, cp.name AS muvekkil, c.tracking_no, c.deleted_at
        FROM case_foys f JOIN cases c ON c.id = f.case_id
        LEFT JOIN case_parties cp ON cp.id = f.case_party_id
        ORDER BY f.sistem_no
    """)
    foy_by_case: Dict[int, List[dict]] = defaultdict(list)
    for f in foys:
        foy_by_case[f["case_id"]].append(f)
    doc_cnt = {d["case_id"]: d["n"] for d in rows(
        "SELECT case_id, COUNT(*) AS n FROM case_documents WHERE deleted_at IS NULL AND case_id IS NOT NULL GROUP BY case_id")}
    stage_cnt = {d["case_id"]: d["n"] for d in rows("SELECT case_id, COUNT(*) AS n FROM case_stage_decisions GROUP BY case_id")}
    # İlişkili kartlar (G180: AYRISTIRILAN + elle bağlar; öneri reddi hariç) — her iki yönde
    tracking = {c["id"]: c["tracking_no"] for c in cases}
    iliski: Dict[int, List[str]] = defaultdict(list)
    for r in rows("SELECT source_case_id, target_case_id, relation_type FROM case_relations WHERE relation_type <> 'ONERI_RED' ORDER BY id"):
        a, b, t = r["source_case_id"], r["target_case_id"], r["relation_type"]
        if b in tracking:
            iliski[a].append(f"#{b} {tracking[b]} [{t}]")
        if a in tracking:
            iliski[b].append(f"#{a} {tracking[a]} [{t}]")

    kart_cols = [
        ("Kart ID (HukDok)", lambda c: c["id"]),
        ("Ofis Dosya No (HukDok)", lambda c: c["tracking_no"]),
        ("SistemNo (föyler)", lambda c: _join(f["sistem_no"] for f in foy_by_case[c["id"]])),
        ("Föy sayısı", lambda c: len(foy_by_case[c["id"]])),
        ("TKU (Klasör No)", lambda c: _join([c["tku_no"]] + [f["tku_no"] for f in foy_by_case[c["id"]]])),
        ("Dosya No listesi (klasor_no_2)", lambda c: c["klasor_no_2"]),
        ("MüvekkilNo (föyler)", lambda c: _join(f["muvekkil_no"] for f in foy_by_case[c["id"]])),
        ("Müvekkil(ler)", lambda c: _join(parties[c["id"]]["CLIENT"])),
        ("Karşı Taraf", lambda c: _join(parties[c["id"]]["COUNTER"])),
        ("Diğer taraflar", lambda c: _join(parties[c["id"]]["THIRD"])),
        ("Durum", lambda c: c["status"]),
        ("Aktif mi", lambda c: "Evet" if c["active"] else "Hayır"),
        ("Dosya Son Durumu", lambda c: c["dosya_son_durumu"]),
        ("Aşama (case_stage)", lambda c: c["case_stage"]),
        ("Ana Tür", lambda c: c["file_type"]),
        ("Yargı Birimi", lambda c: c["judicial_unit"]),
        ("Uzmanlık Alanı", lambda c: c["sub_type"]),
        ("Ek Alt Kırılım", lambda c: c["sub_type_extra"]),
        ("Dava Konusu", lambda c: c["subject"]),
        ("Yerel Mahkeme", lambda c: c["court"]),
        ("Esas (güncel)", lambda c: c["esas_no"]),
        ("Esas tarihçesi", lambda c: _join(esas_hist[c["id"]])),
        ("Eski/Yeni Esas No", lambda c: c["yeni_esas_no"]),
        ("Dava Tarihi", lambda c: c["opening_date"]),
        ("İş Kabul Tarihi", lambda c: c["acceptance_date"]),
        ("Atama Tarihi", lambda c: c["atama_tarihi"]),
        ("Arşiv Tarihi", lambda c: c["arsiv_tarihi"]),
        ("Sorumlu Avukat (kart kutusu)", lambda c: c["responsible_lawyer_name"]),
        ("Sorumlu Avukatlar (ilişki listesi)", lambda c: _join(lawyers[c["id"]])),
        ("UYAP Avukatı", lambda c: c["uyap_lawyer_name"]),
        ("Buro Özel Türü", lambda c: c["bureau_type"]),
        ("Hizmet Türü", lambda c: c["hizmet_turu"]),
        ("Müvekkil Tipi", lambda c: c["muvekkil_tipi"]),
        ("Hizmet Türü (eski alan service_type)", lambda c: c["service_type"]),
        ("Hasar Dosya No", lambda c: c["hasar_dosya_no"]),
        ("Hukuk No", lambda c: c["hukuk_no"]),
        ("Sistem No (kart)", lambda c: c["sistem_no"]),
        ("Maddi Tazminat", lambda c: c["maddi_tazminat"]),
        ("Manevi Tazminat", lambda c: c["manevi_tazminat"]),
        ("Dava Değeri", lambda c: c["dava_degeri"]),
        ("Islah Tutarı", lambda c: c["islah_tutari"]),
        ("Para Birimi", lambda c: c["para_birimi"]),
        ("Yerel Karar Durumu", lambda c: c["yerel_karar_durumu"]),
        ("Yerel Karar Tarihi", lambda c: c["karar_tarihi"]),
        ("Yerel Karar No", lambda c: c["karar_no"]),
        ("Yerel Karar Türü", lambda c: c["karar_turu"]),
        ("Yerel Karar Lehine", lambda c: c["karar_lehine"]),
        ("Yerel Tebliğ Tarihi", lambda c: c["karar_teblig_tarihi"]),
        ("Yerel Karar Açıklaması", lambda c: c["karar_aciklama"]),
        ("Hükmedilen Maddi", lambda c: c["hukmedilen_maddi"]),
        ("Hükmedilen Manevi", lambda c: c["hukmedilen_manevi"]),
        ("Hükmedilen Toplam", lambda c: c["hukmedilen_toplam"]),
        ("İstinaf Mahkemesi", lambda c: c["istinaf_mahkemesi"]),
        ("İstinaf Esas", lambda c: c["istinaf_esas_no"]),
        ("İstinaf Başvuru Tarihi", lambda c: c["istinaf_basvuru_tarihi"]),
        ("İstinaf Başvuran Taraf", lambda c: c["istinaf_basvuran_taraf"]),
        ("İstinaf Karar No", lambda c: c["istinaf_karar_no"]),
        ("İstinaf Karar Tarihi", lambda c: c["istinaf_karar_tarihi"]),
        ("İstinaf Karar Durumu", lambda c: c["istinaf_karar_durumu"]),
        ("İstinaf Tebliğ Tarihi", lambda c: c["istinaf_teblig_tarihi"]),
        ("Temyiz Mahkemesi", lambda c: c["temyiz_mahkemesi"]),
        ("Temyiz Esas", lambda c: c["temyiz_esas_no"]),
        ("Temyiz Başvuru Tarihi", lambda c: c["temyiz_basvuru_tarihi"]),
        ("Temyiz Karar No", lambda c: c["temyiz_karar_no"]),
        ("Temyiz Karar Tarihi", lambda c: c["temyiz_karar_tarihi"]),
        ("Temyiz Karar Durumu", lambda c: c["temyiz_karar_durumu"]),
        ("Temyiz Eden", lambda c: c["temyiz_eden_durumu"]),
        ("Temyiz Tebliğ Tarihi", lambda c: c["temyiz_teblig_tarihi"]),
        ("Karar Düzeltme Durumu", lambda c: c["karar_duzeltme_durumu"]),
        ("Karar Düzeltme Esas", lambda c: c["karar_duzeltme_esas_no"]),
        ("Karar Düzeltme Karar No", lambda c: c["karar_duzeltme_karar_no"]),
        ("Karar Düzeltme Tarihi", lambda c: c["karar_duzeltme_tarihi"]),
        ("Kesinleşme Tarihi", lambda c: c["kesinlesme_tarihi"]),
        ("İnfaz Tarihi", lambda c: c["infaz_tarihi"]),
        ("Arabuluculuk No", lambda c: c["arabuluculuk_no"]),
        ("Arabuluculuk Karar Tarihi", lambda c: c["arabuluculuk_karar_tarihi"]),
        ("Olay Türü", lambda c: c["olay_turu"]),
        ("Hükümdeki Rol", lambda c: c["hukumdeki_rol"]),
        ("Tıbbi Süreç", lambda c: c["tibbi_surec"]),
        ("Tıbbi Olay", lambda c: c["tibbi_olay"]),
        ("İddia Edilen Kusur", lambda c: c["iddia_edilen_kusur"]),
        ("Hastada Oluşan Zarar", lambda c: c["hastada_olusan_zarar"]),
        ("Uygulanan Yöntem", lambda c: c["uygulanan_yontem"]),
        ("Aşama karar satırı sayısı", lambda c: stage_cnt.get(c["id"], 0)),
        ("Belge sayısı (HukDok)", lambda c: doc_cnt.get(c["id"], 0)),
        ("Eksik zorunlu alan kovası", lambda c: c["missing_required_bucket"]),
        ("Notlar", lambda c: c["notes"]),
        ("Kart oluşturma", lambda c: _naive(c["created_at"])),
        ("Son güncelleme", lambda c: _naive(c["updated_at"])),
        ("İlişkili kartlar", lambda c: _join(iliski[c["id"]])),          # 13.09: G180 ayrılan çiftler
    ]
    foy_cols = [
        ("SistemNo", lambda f: f["sistem_no"]),
        ("Kart ID (HukDok)", lambda f: f["case_id"]),
        ("Ofis Dosya No (HukDok)", lambda f: f["tracking_no"]),
        ("Önceki Ofis Dosya No (birleştirme öncesi)", lambda f: f["onceki_tracking_no"]),
        ("TKU (Klasör No)", lambda f: f["tku_no"]),
        ("MüvekkilNo", lambda f: f["muvekkil_no"]),
        ("Müvekkil (bağlı taraf)", lambda f: f["muvekkil"]),
        ("Müvekkil Tipi", lambda f: f["muvekkil_tipi"]),
        ("Hizmet Türü", lambda f: f["hizmet_turu"]),
        ("Hasar No", lambda f: f["hasar_no"]),
        ("Föy Bilgisi (MKO)", lambda f: f["mko_id"]),
        ("Durum", lambda f: f["durum"]),
        ("Kapsam Durumu", lambda f: f["kapsam_durumu"]),
        ("Kapsam Gerekçesi", lambda f: f["kapsam_gerekcesi"]),
        ("Kapsam Tarihi", lambda f: f["kapsam_tarihi"]),
        ("Yazan Teslim", lambda f: f["source"]),
        ("Kart silinmiş mi", lambda f: "Evet" if f["deleted_at"] else ""),
    ]

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Dava Kartlari"
    _write_sheet(ws1, kart_cols, cases)
    ws2 = wb.create_sheet("Foyler")
    _write_sheet(ws2, foy_cols, foys)

    n_kart, n_foy = len(cases), len(foys)
    n_foylu = sum(1 for c in cases if foy_by_case[c["id"]])
    n_cok_foy = sum(1 for c in cases if len(foy_by_case[c["id"]]) > 1)
    n_birlesen = sum(1 for f in foys if f["onceki_tracking_no"])
    n_iliski = sum(len(v) for v in iliski.values()) // 2
    status_dagilim: Dict[str, int] = defaultdict(int)
    for c in cases:
        status_dagilim[c["status"] or "(boş)"] += 1
    lines = [
        f"HukDok dava kartları — lokal veritabanı fotoğrafı, {simdi.strftime('%d.%m.%Y %H:%M')} (Türkiye saati)",
        "",
        "Bu dosya HukDok'ta o an kayıtlı dava KARTLARININ tamamıdır (silinmiş kartlar hariç). Kıyaslama içindir; sisteme geri yüklenmez.",
        "'Kart ID (HukDok)' sistemdeki kart kimliğidir (cases.id) — kalıcıdır, yeniden verilmez.",
        "",
        "SAYFALAR",
        "• Dava Kartlari — kart (dava) başına BİR satır. Bir kartın birden çok föyü olabilir; föy anahtarları (SistemNo, TKU, MüvekkilNo) ' ; ' ile birleşik yazılır.",
        "• Foyler — föy (SistemNo) başına BİR satır; hangi HukDok kartına bağlı olduğu ve birleştirme öncesi ofis dosya numarası görünür.",
        "",
        "SAYILAR",
        f"• Kart: {n_kart}  (föylü {n_foylu}, föysüz {n_kart - n_foylu}, birden çok föylü {n_cok_foy})",
        f"• Föy: {n_foy}  (birleştirmeyle kart değiştirmiş föy: {n_birlesen})",
        f"• İlişkili kart çifti (case_relations, son sütun): {n_iliski}",
        "• Durum dağılımı: " + " · ".join(f"{k}: {v}" for k, v in sorted(status_dagilim.items(), key=lambda kv: -kv[1])),
        "",
        "KART MODELİ (12.09 cevabınız §8 · 13.09 kararı)",
        "• Aynı davanın müvekkil başına föyleri tek kartta toplanır (esas + mahkeme + tür aynı).",
        "• Arabuluculuk ile ardından açılan dava, soruşturma ile ceza davası, aynı türde farklı esaslı davalar AYRI karttır; aralarındaki bağ 'İlişkili kartlar' sütununda [AYRISTIRILAN] etiketiyle görünür.",
        "• Aynı DosyaNo iki kartta görünebilir (arabuluculuk kartı + dava kartı) — 'Dosya No listesi (klasor_no_2)' paylaşılır.",
        "",
        "TESLİM PAKETİNİZDEN FARKLAR",
        "• Satır birimi: sizde föy (SistemNo), bizde kart (dava).",
        "• 'Ofis Dosya No' HukDok'un kendi numarasıdır; sizin DosyaNo'nuz 'Dosya No listesi (klasor_no_2)' sütunundadır.",
        "• Esas numarası tarihçelidir: 'Esas (güncel)' + 'Esas tarihçesi' (aşama etiketli).",
        "• Sorumlu avukat iki yerde: kart kutusu (tek ad) + ilişki listesi (paketteki tüm adlar).",
        "• Yargı aşamaları kartta tek slot olarak fotoğraflanır (Yerel/İstinaf/Temyiz/Karar Düzeltme); çok turlu aşamalar 'Aşama karar satırı sayısı'nda görünür.",
        "• Hükmedilen tutarlar ve klinik tasnif alanları paketten geldiği gibi kart alanına yazılır; kardeş föyler çelişiyorsa boş kalır.",
    ]
    ws3 = wb.create_sheet("Aciklama")
    for i, line in enumerate(lines, 1):
        ws3.cell(row=i, column=1, value=line).font = Font(name="Arial", size=10, bold=(i == 1 or line.isupper()))
    ws3.column_dimensions["A"].width = 160
    wb.save(out)
    return {"out": out, "zaman": simdi.isoformat(timespec="minutes"), "kart": n_kart, "foy": n_foy,
            "foylu": n_foylu, "cok_foy": n_cok_foy, "birlesen_foy": n_birlesen, "iliski": n_iliski,
            "durum": dict(status_dagilim)}


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--out", default=f"/tmp/HUKDOK_DAVA_KARTLARI_{dt.date.today().isoformat()}.xlsx")
    args = parser.parse_args(argv)
    ozet = liste_uret(os.environ["DATABASE_URL"], args.out)
    print(json.dumps(ozet, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
