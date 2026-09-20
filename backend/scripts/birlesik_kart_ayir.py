#!/usr/bin/env python3
"""Birleşik kartları geri ayırır (G180) — farklı tür/esas taşıyan föyler ayrı karta,
iki kart arasında "ilişkili dosya" bağı (`case_relations.AYRISTIRILAN`).

Kaynak: veri ekibinin 12.09 cevabı §8 + Ek-3 › 03 (`03_BIRLESTIRME_KARTLARI`, 30 kart) ve
Ek-3 › 02'nin "Önerdiğimiz kart" BOŞ satırları (21 SORU). Kullanıcı kararı 13.09 (iki kez):
"ekibin önerisi gibi olsun" — arabuluculuk ile ardından açılan dava, soruşturma ile ceza
davası, aynı türde farklı esaslı davalar AYRI kart; TKU tek başına birleştirme ölçütü değil.

Bu kartlar TKU birleştirmesinden (scripts/tku_kart_birlestir.py) GELMİYOR — föylerinde
`onceki_tracking_no` yok. Bizim kart açıcımız (scripts/kartsiz_foy_kart_ac.py) aynı
DosyaNo'daki ARB + HUKUK föylerini bilerek tek karta koyar (MİCRO iki aşamaya aynı
DosyaNo verir); aktarımın DosyaNo köprüsü de öyle. Bu script o kararı ekibin modeline
çevirir: kartın KENDİ grubu kalır, öteki her grup için YENİ kart açılır.

Kural:
* Grup anahtarı = (dosya türü, esas anahtarı) — föyün `ham_veri`sinden (`Ana Tür`,
  `Esas`; `case_relations_auto.esas_anahtari`). `ham_veri`siz föy gruplanamaz → kart RET.
* Kartın kendi grubu: (kart.file_type, esas_anahtari(kart.esas_no)) ile birebir; yoksa
  kartın türündeki en büyük grup; o da yoksa en büyük grup. Tek grup → ATLANDI.
* `--muvekkil-ayrimi` (Ek-6 › 01, 17.09.2026 — kart 14334): grup anahtarına ÜÇÜNCÜ boyut
  olarak föyün müvekkili eklenir, böylece aynı tür + aynı esastaki iki müvekkil de ayrılır
  (ekibin sorusu: "dört föy dört ayrı kartta olmalı — müvekkil başına bir arabuluculuk, bir
  dava kartı"). Aynı (tür, esas) birden çok gruba bölününce kartta KALAN grup, kartın ofis
  numarasındaki isim bloğuyla seçilir (`kartsiz_foy_kart_ac.isim_blogu`) — kartın künyesi
  kendi müvekkilinde kalsın diye. Bu modda yeni kartın klasör numarası GRUBUN kendi
  DosyaNo'sudur (müvekkil başına ayrı klasör); kapalıyken kartın numarası paylaşılır.
  Bayrak yalnız adıyla verilen kartta açılır — genel kullanımda KAPALIDIR.
* Yeni kart `kartsiz_foy_kart_ac.kart_adaylari` + `ofis_numarasi` ile (künye grubun asıl
  föyünden, ofis numarası mevcut kural), klasör no = kartın DosyaNo'su (paylaşılır — sonraki
  aktarım köprüde iki kart görür, esas + tür ikinci anahtarı ayırır: tasarlanmış ikiz yolu).
  Taraflar föyün ham satırından `hukdok_aktarim._taraflari_yaz`, föy ↔ müvekkil bağı
  `_foy_muvekkilini_bagla`; esas `case_manager.sync_current_esas` (tek yazma yolu).
  Föyün `onceki_tracking_no`su varsa (sönen kart) elle karar — RET (veride 0).
* Belgeler KALAN kartta kalır (hangi föyün belgesi olduğu bilinmiyor — belge koruma şartı).
* İlişki: `source=kalan`, `target=yeni`, `relation_type=AYRISTIRILAN`; aynı çift için ikinci
  satır açılmaz. İki kartta `case_history` `kart_ayirma`; `refresh_missing_required`.
* Tek transaction; `--apply` yoksa geri alınır. İkinci koşu 0 (tek gruplu kart ATLANDI).

    docker compose exec -T backend python scripts/birlesik_kart_ayir.py \\
        --ek3 /app/calibration-data/_g179/HUKDOK_CEVAP_EKI_3_2026-09-12.xlsx        # kuru koşu
    docker compose exec -T backend python scripts/birlesik_kart_ayir.py --ek3 … --apply --kim ilke
    docker compose exec -T backend python scripts/birlesik_kart_ayir.py --kart 2468 --kart 4177   # liste elle
    docker compose exec -T backend python scripts/birlesik_kart_ayir.py --kart 14334 --muvekkil-ayrimi
"""
from __future__ import annotations

import argparse
import logging
import os
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import models
from managers import case_manager
from party_check import normalize_party_key
from scripts import hukdok_aktarim as ha
from scripts import kartsiz_foy_kart_ac as kfa
from scripts.ekip_cevabi_1209 import _kolon, _sayfa, _tam_sayi
from services.case_relations_auto import esas_anahtari

logger = logging.getLogger("BirlesikKartAyir")
DEGISTIREN = "birlesik_kart_ayir"
KAYNAK = "ekip cevabı 12.09.2026 Ek-3 › 02/03 — ekibin kart modeli (kullanıcı kararı 13.09)"
ILISKI_TURU = "AYRISTIRILAN"
EK3_BIRLESTIRME_SAYFASI = "03_BIRLESTIRME_KARTLARI"
EK3_TASIMA_SAYFASI = "02_BASKA_ASAMA_KARTI"

# (file_type, esas anahtarı) — müvekkil ayrımı açıkken üçüncü boyut müvekkil anahtarı.
# Arite tek koşu içinde sabittir (bayrak koşu başına seçilir), karşılaştırmalar güvenli.
GrupAnahtari = Tuple[str, ...]


@dataclass
class AyirmaKalemi:
    kart_id: int
    sonuc: str                       # YAPILDI | ATLANDI | RET
    aciklama: str = ""
    yeni_kartlar: List[Tuple[int, str, List[str]]] = field(default_factory=list)   # (id, tracking_no, sistem_nolar)


@dataclass
class Sonuc:
    kalemler: List[AyirmaKalemi] = field(default_factory=list)

    def ekle(self, kalem: AyirmaKalemi) -> None:
        self.kalemler.append(kalem)
        (logger.info if kalem.sonuc == "YAPILDI" else logger.warning)(
            f"#{kalem.kart_id}: {kalem.sonuc} {kalem.aciklama}".rstrip())

    def sayim(self, sonuc: str) -> int:
        return sum(1 for k in self.kalemler if k.sonuc == sonuc)

    @property
    def yeni_kart_sayisi(self) -> int:
        return sum(len(k.yeni_kartlar) for k in self.kalemler)


# ─── Ek-3 okuma ──────────────────────────────────────────────────────────────

def ek3_kartlari(yol: Path) -> List[int]:
    """Ek-3 › 03 `Kart` sütunu + Ek-3 › 02'de 'Önerdiğimiz kart' boş satırların 'Bağlı olduğu kart'ı."""
    import openpyxl
    wb = openpyxl.load_workbook(yol, data_only=True, read_only=True)
    idler: List[int] = []
    ws = _sayfa(wb, EK3_BIRLESTIRME_SAYFASI)
    satirlar = list(ws.iter_rows(values_only=True))
    i_kart = _kolon(satirlar[0], "Kart")
    for satir in satirlar[1:]:
        kart = _tam_sayi(satir[i_kart])
        if kart and kart not in idler:
            idler.append(kart)
    ws = _sayfa(wb, EK3_TASIMA_SAYFASI)
    satirlar = list(ws.iter_rows(values_only=True))
    i_bagli = _kolon(satirlar[0], "Bağlı olduğu kart")
    i_oneri = _kolon(satirlar[0], "Önerdiğimiz kart")
    for satir in satirlar[1:]:
        bagli, oneri = _tam_sayi(satir[i_bagli]), _tam_sayi(satir[i_oneri])
        if bagli and not oneri and bagli not in idler:
            idler.append(bagli)
    wb.close()
    return sorted(idler)


# ─── Föy → ham satır → grup ──────────────────────────────────────────────────

def foy_satiri(foy: models.CaseFoy) -> Optional[ha.HamSatir]:
    """`case_foys.ham_veri` (G125: orijinal başlıklar) → HamSatir; ham yoksa None."""
    ham = foy.ham_veri if isinstance(foy.ham_veri, dict) else None
    if not ham:
        return None
    basliklar = list(ham.keys())
    degerler_ham = list(ham.values())
    indeksler, _ = ha._sutun_indeksleri(basliklar)
    degerler = {alan: degerler_ham[i] for alan, i in indeksler.items()}
    degerler.setdefault("sistem_no", foy.sistem_no)
    return ha.HamSatir(satir_no=0, degerler=degerler, ham=dict(ham))


def grup_anahtari(satir: ha.HamSatir, *, muvekkil_ayrimi: bool = False) -> GrupAnahtari:
    tur = ha._esleme(ha.ANA_TUR_ESLEMESI)(satir.degerler.get("ana_tur"), "file_type") or ""
    anahtar = (str(tur), esas_anahtari(ha._metin(satir.degerler.get("esas"))))
    if not muvekkil_ayrimi:
        return anahtar
    return anahtar + (normalize_party_key(kfa._ilk(satir.degerler.get("muvekkil")) or ""),)


def _grup_isim_blogu(grup: Sequence[Tuple[models.CaseFoy, ha.HamSatir]]) -> Optional[str]:
    """Grubun ofis numarası isim bloğu (tek DosyaNo'ya düşmüyorsa None)."""
    adaylar = kfa.kart_adaylari([s for _, s in grup])
    if len(adaylar) != 1 or not adaylar[0].muvekkil:
        return None
    return kfa.isim_blogu(adaylar[0])


def gruplari_bul(kart: models.Case, foyler: Sequence[models.CaseFoy], *, muvekkil_ayrimi: bool = False
                 ) -> Tuple[Optional[GrupAnahtari], Dict[GrupAnahtari, List[Tuple[models.CaseFoy, ha.HamSatir]]], List[str]]:
    """(kartın kendi grubu, {anahtar: [(föy, satır)]}, ham_veri'siz föyler)."""
    gruplar: Dict[GrupAnahtari, List[Tuple[models.CaseFoy, ha.HamSatir]]] = {}
    hamsiz: List[str] = []
    for foy in foyler:
        satir = foy_satiri(foy)
        if satir is None:
            hamsiz.append(foy.sistem_no)
            continue
        gruplar.setdefault(grup_anahtari(satir, muvekkil_ayrimi=muvekkil_ayrimi), []).append((foy, satir))
    if not gruplar:
        return None, gruplar, hamsiz
    def en_buyuk(adaylar: Sequence[GrupAnahtari]) -> GrupAnahtari:
        return max(adaylar, key=lambda k: (len(gruplar[k]), k))

    kendi_temel = (kart.file_type or "", esas_anahtari(kart.esas_no))
    tam = [k for k in gruplar if tuple(k[:2]) == kendi_temel]
    if len(tam) == 1:
        return tam[0], gruplar, hamsiz
    if tam:
        # Müvekkil ayrımı: kartın künyesi kendi müvekkilinde kalsın — kartta kalacak
        # grup, ofis numarasının isim bloğuyla seçilir (tek eşleşme yoksa en büyük grup).
        blok = (kart.tracking_no or "").split(".")[1] if "." in (kart.tracking_no or "") else None
        eslesen = [k for k in tam if blok and _grup_isim_blogu(gruplar[k]) == blok]
        return (eslesen[0] if len(eslesen) == 1 else en_buyuk(tam)), gruplar, hamsiz
    ayni_tur = [k for k in gruplar if k[0] == kendi_temel[0]]
    return en_buyuk(ayni_tur or list(gruplar)), gruplar, hamsiz


# ─── Yeni kart ───────────────────────────────────────────────────────────────

def _tarihce(db, case_id: int, alan: str, eski, yeni, kim: str, not_: str) -> None:
    db.add(models.CaseHistory(
        case_id=case_id, field_name=alan,
        old_value=None if eski is None else str(eski), new_value=None if yeni is None else str(yeni),
        changed_by=DEGISTIREN, source=f"{KAYNAK} ({kim}): {not_}"[:300],
    ))


def _iliski_var(db, a: int, b: int) -> bool:
    q = db.query(models.CaseRelation).filter(
        ((models.CaseRelation.source_case_id == a) & (models.CaseRelation.target_case_id == b))
        | ((models.CaseRelation.source_case_id == b) & (models.CaseRelation.target_case_id == a)))
    return db.query(q.exists()).scalar()


def yeni_kart_ac(db, kalan: models.Case, grup: Sequence[Tuple[models.CaseFoy, ha.HamSatir]],
                 kullanilan: Dict[str, int], *, kim: str,
                 muvekkil_ayrimi: bool = False) -> Tuple[Optional[models.Case], str]:
    """Grubun föyleri için kart açar, föyleri taşır; (kart, hata_metni)."""
    satirlar = [s for _, s in grup]
    adaylar = kfa.kart_adaylari(satirlar)
    if len(adaylar) != 1:
        return None, f"grup {len(adaylar)} DosyaNo'ya bölünüyor — elle"
    aday = adaylar[0]
    if not aday.muvekkil:
        return None, "Müvekkil boş — ofis numarası üretilemez"
    aday.tracking_no = kfa.ofis_numarasi(db, aday, kullanilan)
    source = f"{DEGISTIREN} ({kim})"
    # `kart_adaylari` tarihi ISO metin verir; Date kolonu date ister (sqlite katı, Postgres toleranslı)
    opening_date = date.fromisoformat(aday.opening_date) if aday.opening_date else None
    # Klasör no: aynı dosyanın iki aşaması ayrılırken kartın numarası paylaşılır
    # (tasarlanmış ikiz yolu); müvekkil ayrımında her müvekkilin kendi DosyaNo'su
    # vardır ve kartın numarası yeni karta TAŞINMAZ (ekibin Ek-6 › 01 ricası).
    klasor_no_2 = (aday.dosya_no or kalan.klasor_no_2) if muvekkil_ayrimi else (kalan.klasor_no_2 or aday.dosya_no)
    yeni = models.Case(
        tracking_no=aday.tracking_no, status=aday.status, file_type=aday.file_type,
        subject=aday.subject, court=aday.court, opening_date=opening_date,
        klasor_no_2=klasor_no_2, tenant_id=kalan.tenant_id,
        responsible_lawyer_name=kalan.responsible_lawyer_name,
        notes=f"#{kalan.id} {kalan.tracking_no} kartından ayrıldı ({KAYNAK}): {', '.join(aday.sistem_nolar)}",
    )
    db.add(yeni)
    db.flush()
    if aday.esas_no:
        case_manager.sync_current_esas(db, yeni, aday.esas_no, court=aday.court, source=source)
    for _foy, satir in grup:
        ha._taraflari_yaz(db, yeni, satir, source)
    db.flush()
    for foy, satir in grup:
        foy.case_id = yeni.id
        foy.case_party_id = None
        db.flush()
        ha._foy_muvekkilini_bagla(db, yeni, satir, sistem_no=foy.sistem_no, source=source)
    db.flush()
    return yeni, ""


def karti_ayir(db, kart_id: int, kullanilan: Dict[str, int], *, kim: str,
               muvekkil_ayrimi: bool = False) -> AyirmaKalemi:
    kart = db.get(models.Case, kart_id)
    if kart is None or kart.deleted_at is not None:
        return AyirmaKalemi(kart_id, "RET", "kart yok ya da silinmiş")
    foyler = (db.query(models.CaseFoy).filter(models.CaseFoy.case_id == kart_id)
              .order_by(models.CaseFoy.sistem_no).all())
    if len(foyler) < 2:
        return AyirmaKalemi(kart_id, "ATLANDI", f"{len(foyler)} föy — ayrılacak grup yok")
    kendi, gruplar, hamsiz = gruplari_bul(kart, foyler, muvekkil_ayrimi=muvekkil_ayrimi)
    if hamsiz:
        return AyirmaKalemi(kart_id, "RET", f"ham satırı olmayan föy: {', '.join(hamsiz)} — gruplanamaz")
    if len(gruplar) < 2:
        return AyirmaKalemi(kart_id, "ATLANDI", f"tek grup {kendi} — ayrılacak grup yok")
    sonen = [f.sistem_no for f in foyler if f.onceki_tracking_no]
    if sonen:
        return AyirmaKalemi(kart_id, "RET", f"sönen kart izi taşıyan föy: {', '.join(sonen)} — elle karar")
    kalem = AyirmaKalemi(kart_id, "YAPILDI", f"kendi grubu {kendi}")
    for anahtar, grup in sorted(gruplar.items()):
        if anahtar == kendi:
            continue
        with db.begin_nested():
            yeni, hata = yeni_kart_ac(db, kart, grup, kullanilan, kim=kim, muvekkil_ayrimi=muvekkil_ayrimi)
            if yeni is None:
                kalem.sonuc, kalem.aciklama = "RET", f"grup {anahtar}: {hata}"
                return kalem
            if not _iliski_var(db, kart.id, yeni.id):
                db.add(models.CaseRelation(source_case_id=kart.id, target_case_id=yeni.id,
                                           relation_type=ILISKI_TURU, created_by=kim,
                                           note=f"{KAYNAK}: {anahtar[0]} {anahtar[1] or '(esassız)'}"
                                                f"{' · ' + anahtar[2] if len(anahtar) > 2 else ''} ayrı kart"))
            sistem_nolar = [f.sistem_no for f, _ in grup]
            _tarihce(db, kart.id, "kart_ayirma", ", ".join(sistem_nolar), f"→ #{yeni.id} {yeni.tracking_no}", kim,
                     f"grup {anahtar}")
            _tarihce(db, yeni.id, "kart_ayirma", f"#{kart.id} {kart.tracking_no}", ", ".join(sistem_nolar), kim,
                     f"grup {anahtar}")
            case_manager.refresh_missing_required(db, yeni)
            kalem.yeni_kartlar.append((yeni.id, yeni.tracking_no, sistem_nolar))
    case_manager.refresh_missing_required(db, kart)
    return kalem


def kos(session_factory, *, kart_idler: Sequence[int], apply: bool = False, kim: str = DEGISTIREN,
        muvekkil_ayrimi: bool = False) -> Sonuc:
    sonuc = Sonuc()
    kullanilan: Dict[str, int] = {}
    db = session_factory()
    try:
        for kart_id in kart_idler:
            sonuc.ekle(karti_ayir(db, kart_id, kullanilan, kim=kim, muvekkil_ayrimi=muvekkil_ayrimi))
        if apply:
            db.commit()
        else:
            db.rollback()
    finally:
        db.close()
    return sonuc


def ozet_metni(sonuc: Sonuc, *, apply: bool) -> str:
    satirlar = ["Birleşik kart ayırma — " + ("YAZILDI" if apply else "KURU KOŞU (geri alındı)"),
                f"  kart: {len(sonuc.kalemler)} · ayrıldı: {sonuc.sayim('YAPILDI')} · atlandı: {sonuc.sayim('ATLANDI')} · "
                f"ret: {sonuc.sayim('RET')} · yeni kart: {sonuc.yeni_kart_sayisi}"]
    for k in sonuc.kalemler:
        satirlar.append(f"    #{k.kart_id}: {k.sonuc} {k.aciklama}".rstrip())
        for yeni_id, tracking, sistem_nolar in k.yeni_kartlar:
            satirlar.append(f"        → #{yeni_id} {tracking}: {', '.join(sistem_nolar)}")
    return "\n".join(satirlar)


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--ek3", default=None, help="HUKDOK_CEVAP_EKI_3_2026-09-12.xlsx (kart listesi)")
    parser.add_argument("--kart", type=int, action="append", default=[], help="kart id (tekrarlanabilir)")
    parser.add_argument("--apply", action="store_true", help="yaz (varsayılan kuru koşu)")
    parser.add_argument("--kim", default=DEGISTIREN, help="tarihçe imzası")
    parser.add_argument("--muvekkil-ayrimi", action="store_true",
                        help="aynı tür + esastaki farklı müvekkilleri de ayır (yalnız adıyla verilen kartta)")
    args = parser.parse_args(argv)
    from logging_setup import configure_logging    # log sözleşmesi: yapılandırma tek yerde
    configure_logging()
    idler = list(args.kart)
    if args.ek3:
        idler = sorted(set(idler) | set(ek3_kartlari(Path(args.ek3))))
    if not idler:
        parser.error("--ek3 ya da --kart gerekli")
    from database import SessionLocal
    sonuc = kos(SessionLocal, kart_idler=idler, apply=args.apply, kim=args.kim,
                muvekkil_ayrimi=args.muvekkil_ayrimi)
    print(ozet_metni(sonuc, apply=args.apply))
    return 0


if __name__ == "__main__":
    sys.exit(main())
