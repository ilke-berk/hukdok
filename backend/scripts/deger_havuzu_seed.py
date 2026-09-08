#!/usr/bin/env python3
"""Teslim paketinin değer havuzlarından kapalı listeleri kurar (G124).

Kullanıcı kararı (05.09.2026): "menüleri listedeki bilgilerden oluşturalım —
onların hatalısını geçirelim, lokal migrasyon bitince hepsini elden geçiririz."
Bu script `Sheet` sayfasının sütunlarındaki ATOMİK değerleri (çok değerli
hücreler " ; " ile ayrılır, services.multi_value) ilgili liste tablosuna
YENİ satır olarak ekler; mevcut adlara (büyük/küçük harf duyarsız) dokunmaz,
hiçbir satırı silmez ya da yeniden adlandırmaz. Yazımlar OLDUĞU GİBİ alınır
("YARGITAY .....HD", "Karar Aaleyhe" dahil) — temizlik yönetim panelinin işi.

    docker compose exec -T backend python scripts/deger_havuzu_seed.py \\
        --input /tmp/paket.xlsx                # kuru koşu: yalnız sayar
    docker compose exec -T backend python scripts/deger_havuzu_seed.py \\
        --input /tmp/paket.xlsx --apply        # yazar

İdempotent: ikinci koşu 0 yeni satır. Gerçek paket repoya GİRMEZ (A.2);
testler sentetik paketle koşar (tests/test_g124_deger_havuzlari.py).

G151 — iki ek:

* **Büro durumu elemesi:** dört karar listesine (`*_decisions`) paketten
  `Kapalı`/`Derdest` GİRMEZ (`hukdok_aktarim.BURO_DURUMLARI`, büyük/küçük
  harf ve aksan toleranslı) — bunlar mahkeme kararı değil büro dosya
  durumudur; yerel havuzdan çıkarıldılar ve paket seed'i geri sokmamalı.
  Başka BÜYÜK/BOZUK yazım ("Karar Aaleyhe", "YARGITAY .....HD") yine
  GEÇER — kullanıcı kararı "hatalısını geçirelim" sürüyor, temizlik panelden.
* **`--kaldir AD` (kuru koşu varsayılan):** mevcut kurulumdaki bir liste
  satırını kaldırır — yalnız hiçbir kart kolonunda (`reference_lists.DEPENDENCIES`)
  ve hiçbir aşama satırında (`case_stage_decisions.karar_durumu`) kullanılmıyorsa;
  kullanılıyorsa SİLMEZ, kullanım sayısıyla raporlar. `--apply` olmadan
  hiçbir şey yazılmaz.

    docker compose exec -T backend python scripts/deger_havuzu_seed.py \\
        --liste local_decisions --kaldir "Kapalı" --kaldir "Derdest"   # kuru koşu
    docker compose exec -T backend python scripts/deger_havuzu_seed.py \\
        --liste local_decisions --kaldir "Kapalı" --kaldir "Derdest" --apply
"""
from __future__ import annotations

import argparse
import collections
import os
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from sqlalchemy import func

import models
from managers.reference_lists import DEPENDENCIES
from managers.seed_data import _karar_kodu
from managers.stage_decisions import STAGE_DECISION_LISTS
from scripts.hukdok_aktarim import YER_TUTUCULAR, _baslik_anahtari, _buro_durumu_mu, _metin
from services.multi_value import split_values


@dataclass(frozen=True)
class Havuz:
    liste_adi: str
    model: type
    basliklar: Tuple[str, ...]      # kabul edilen sütun başlıkları (aday)
    coklu: bool                     # hücre " ; " ile çok değerli mi


# Sütun → liste. Tıbbi beşli çok değerli; karar durumları ve mahkemeler tek.
HAVUZLAR: Tuple[Havuz, ...] = (
    Havuz("medical_processes", models.MedicalProcess, ("Tıbbi Süreç",), True),
    Havuz("medical_events", models.MedicalEvent, ("Tıbbi Olay",), True),
    Havuz("alleged_faults", models.AllegedFault, ("İddia Edilen Kusur",), True),
    Havuz("patient_harms", models.PatientHarm, ("Hastada Oluşan Zarar",), True),
    Havuz("applied_methods", models.AppliedMethod, ("Uygulanan Yöntem",), True),
    Havuz("cassation_courts", models.CassationCourt, ("Temyiz Mahkemesi",), False),
    Havuz("appeal_courts", models.AppealCourt, ("İstinaf Mahkemesi",), False),
    Havuz("defendant_administrations", models.DefendantAdministration, ("Davalı İdare",), True),
    Havuz("local_decisions", models.LocalDecision, ("Yerel Mahkeme Karar Durumu",), False),
    Havuz("appeal_decisions", models.AppealDecision, ("İstinaf Karar Durumu",), False),
    Havuz("cassation_decisions", models.CassationDecision, ("Yargıtay Onama Durumu",), False),
    Havuz("revision_decisions", models.RevisionDecision, ("Karar Düzeltme Kararı Durumu",), False),
    Havuz("currencies", models.Currency, ("Para Birimi TL", "Para Birimi"), False),
)
HAVUZ_HARITASI: Dict[str, Havuz] = {h.liste_adi: h for h in HAVUZLAR}
# Karar listeleri: büro durumu (Kapalı/Derdest) bu dört listeye paketten
# GİRMEZ (G151) — `STAGE_DECISION_LISTS`ten türetilir, üçüncü kopya değil.
KARAR_LISTELERI: frozenset = frozenset(
    h.liste_adi for h in HAVUZLAR if h.model in STAGE_DECISION_LISTS.values()
)


@dataclass
class HavuzSonucu:
    liste_adi: str
    sutun: Optional[str]            # bulunan başlık; None = paket bu sütunu taşımıyor
    mevcut: int = 0
    yeni: int = 0
    ornekler: List[str] = None      # type: ignore[assignment]

    def __post_init__(self) -> None:
        if self.ornekler is None:
            self.ornekler = []


def _yer_tutucu(deger: str) -> bool:
    return deger.strip().upper() in YER_TUTUCULAR


def degerleri_say(satirlar: Sequence[Dict[str, Any]], sutun: str, coklu: bool) -> "collections.Counter[str]":
    """Sütunun atomik değerlerini sıklıkla sayar (yer tutucu ve boş düşer)."""
    sayac: "collections.Counter[str]" = collections.Counter()
    for satir in satirlar:
        ham = _metin(satir.get(sutun))
        if ham is None:
            continue
        parcalar = split_values(ham) if coklu else [ham]
        for parca in parcalar:
            if not _yer_tutucu(parca):
                sayac[parca] += 1
    return sayac


def xlsx_satirlari(yol: Path, sheet: str) -> Tuple[List[str], List[Dict[str, Any]]]:
    import openpyxl

    wb = openpyxl.load_workbook(yol, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        basliklar = [str(h) if h is not None else "" for h in next(it)]
        # strict=False: read_only satırları başlıktan kısa gelebilir (sondaki
        # boş hücreler düşer); eksik sütun "bu satırda yok" demektir.
        satirlar = [dict(zip(basliklar, row, strict=False)) for row in it]
    finally:
        wb.close()
    return basliklar, satirlar


def _sutunu_bul(basliklar: Sequence[str], adaylar: Sequence[str]) -> Optional[str]:
    anahtarlar = {_baslik_anahtari(b): b for b in basliklar}
    for aday in adaylar:
        bulunan = anahtarlar.get(_baslik_anahtari(aday))
        if bulunan is not None:
            return bulunan
    return None


def _tekil_kod(taban: str, kullanilan: set) -> str:
    """Kod çakışmasında ("Karar Aleyhe" / "Karar-Aleyhe" → KARAR_ALEYHE) sonek."""
    kod = taban or "DEGER"
    aday, n = kod, 2
    while aday in kullanilan:
        aday = f"{kod}_{n}"
        n += 1
    kullanilan.add(aday)
    return aday


def havuzu_isle(db, havuz: Havuz, basliklar: Sequence[str],
                satirlar: Sequence[Dict[str, Any]]) -> HavuzSonucu:
    """Tek listeyi paketten besler (flush eder, COMMIT ETMEZ)."""
    sonuc = HavuzSonucu(havuz.liste_adi, _sutunu_bul(basliklar, havuz.basliklar))
    if sonuc.sutun is None:
        return sonuc
    mevcut_satirlar = db.query(havuz.model).all()
    mevcut_adlar = {r.name.casefold() for r in mevcut_satirlar}
    kullanilan_kodlar = {r.code for r in mevcut_satirlar}
    sonuc.mevcut = len(mevcut_satirlar)
    sira = len(mevcut_satirlar)
    karar_listesi = havuz.liste_adi in KARAR_LISTELERI
    for ad, _sayi in degerleri_say(satirlar, sonuc.sutun, havuz.coklu).most_common():
        if ad.casefold() in mevcut_adlar:
            continue
        if karar_listesi and _buro_durumu_mu(ad):
            continue                      # büro durumu karar değil — havuza girmez (G151)
        db.add(havuz.model(
            code=_tekil_kod(_karar_kodu(ad), kullanilan_kodlar), name=ad,
            active=True, sequence=sira,
        ))
        mevcut_adlar.add(ad.casefold())
        sira += 1
        sonuc.yeni += 1
        if len(sonuc.ornekler) < 5:
            sonuc.ornekler.append(ad)
    db.flush()
    return sonuc


def havuzlari_kur(session_factory, *, girdi: Path, sheet: str = "Sheet",
                  apply: bool = False) -> List[HavuzSonucu]:
    basliklar, satirlar = xlsx_satirlari(girdi, sheet)
    db = session_factory()
    try:
        sonuclar = [havuzu_isle(db, havuz, basliklar, satirlar) for havuz in HAVUZLAR]
        if apply:
            db.commit()
        else:
            db.rollback()
        return sonuclar
    finally:
        db.close()


# ─── Satır kaldırma (G151) ───────────────────────────────────────────────────

@dataclass
class KaldirmaSonucu:
    liste_adi: str
    ad: str                         # istenen ad (komut satırından)
    bulunan: Optional[str] = None   # listedeki gerçek yazım; None = satır yok
    kart_kullanimi: int = 0         # DEPENDENCIES kolonlarında bu adı taşıyan kayıt
    asama_kullanimi: int = 0        # case_stage_decisions.karar_durumu (ilgili aşama)
    silindi: bool = False

    @property
    def kullaniliyor(self) -> bool:
        return (self.kart_kullanimi + self.asama_kullanimi) > 0


def satir_kullanimi(db, liste_adi: str, ad: str) -> Tuple[int, int]:
    """(kart kolonu kullanımı, aşama satırı kullanımı) — adı BİREBİR taşıyan
    kayıtlar (kapalı havuz doğrulaması adı birebir yazar; soft-delete'li
    kartlar da sayılır: geri alınabilir kayıt bağ sayılır)."""
    kart = 0
    for dep in DEPENDENCIES.get(liste_adi, []):
        kart += db.query(func.count()).select_from(dep.model).filter(
            getattr(dep.model, dep.column) == ad).scalar() or 0
    model = HAVUZ_HARITASI[liste_adi].model
    stage = next((s for s, m in STAGE_DECISION_LISTS.items() if m is model), None)
    asama = 0
    if stage is not None:
        asama = db.query(func.count()).select_from(models.CaseStageDecision).filter(
            models.CaseStageDecision.stage == stage,
            models.CaseStageDecision.karar_durumu == ad,
        ).scalar() or 0
    return kart, asama


def havuz_satirlarini_kaldir(session_factory, *, liste_adi: str, adlar: Sequence[str],
                             apply: bool = False) -> List[KaldirmaSonucu]:
    """Liste satırlarını kaldırır — kullanılmayanı siler, kullanılanı raporlar.

    Eşleşme `_baslik_anahtari` ile (büyük/küçük harf, aksan toleranslı); silinen
    satırın adı `bulunan`da. `apply=False` (varsayılan) hiçbir şeyi yazmaz,
    yalnız ne olacağını söyler.
    """
    if liste_adi not in HAVUZ_HARITASI:
        raise ValueError(f"bilinmeyen liste: {liste_adi!r} (izinli: {', '.join(sorted(HAVUZ_HARITASI))})")
    model = HAVUZ_HARITASI[liste_adi].model
    db = session_factory()
    try:
        mevcut = {_baslik_anahtari(r.name): r for r in db.query(model).all()}
        sonuclar: List[KaldirmaSonucu] = []
        for ad in adlar:
            sonuc = KaldirmaSonucu(liste_adi, ad)
            satir = mevcut.get(_baslik_anahtari(ad))
            if satir is not None:
                sonuc.bulunan = satir.name
                sonuc.kart_kullanimi, sonuc.asama_kullanimi = satir_kullanimi(db, liste_adi, satir.name)
                if apply and not sonuc.kullaniliyor:
                    db.delete(satir)
                    sonuc.silindi = True
            sonuclar.append(sonuc)
        if apply:
            db.commit()
        else:
            db.rollback()
        return sonuclar
    finally:
        db.close()


def kaldirma_ozeti(sonuclar: Sequence[KaldirmaSonucu], *, apply: bool) -> str:
    satirlar = [f"{'liste':28} {'ad':32} {'kart':>5} {'aşama':>5}  sonuç"]
    for s in sonuclar:
        if s.bulunan is None:
            durum = "listede yok"
        elif s.kullaniliyor:
            durum = "KULLANILIYOR — silinmedi"
        elif s.silindi:
            durum = "SİLİNDİ"
        else:
            durum = "silinebilir (kuru koşu)"
        satirlar.append(f"{s.liste_adi:28} {(s.bulunan or s.ad)[:32]:32} "
                        f"{s.kart_kullanimi:5} {s.asama_kullanimi:5}  {durum}")
    satirlar.append(f"silinen satır: {sum(1 for s in sonuclar if s.silindi)} — "
                    + ("YAZILDI" if apply else "kuru koşu, yazılmadı (--apply ile yazar)"))
    return "\n".join(satirlar)


def ozet_metni(sonuclar: Sequence[HavuzSonucu], *, apply: bool) -> str:
    satirlar = [f"{'liste':28} {'sütun':32} {'mevcut':>6} {'yeni':>5}  örnek"]
    for s in sonuclar:
        sutun = s.sutun or "(pakette yok)"
        satirlar.append(f"{s.liste_adi:28} {sutun[:32]:32} {s.mevcut:6} {s.yeni:5}  "
                        + " | ".join(o[:28] for o in s.ornekler))
    satirlar.append(f"toplam yeni satır: {sum(s.yeni for s in sonuclar)} — "
                    + ("YAZILDI" if apply else "kuru koşu, yazılmadı (--apply ile yazar)"))
    return "\n".join(satirlar)


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Teslim paketi havuzlarından kapalı listeleri kurar (G124)")
    parser.add_argument("--input", default=None, help="teslim paketi (.xlsx); --kaldir ile gerekmez")
    parser.add_argument("--sheet", default="Sheet")
    parser.add_argument("--apply", action="store_true", help="yaz (varsayılan kuru koşu)")
    parser.add_argument("--kaldir", action="append", default=[], metavar="AD",
                        help="listeden satır kaldır (G151; kullanılan satır silinmez, tekrarlanabilir)")
    parser.add_argument("--liste", default="local_decisions", choices=sorted(HAVUZ_HARITASI),
                        help="--kaldir'ın hedef listesi (varsayılan local_decisions)")
    args = parser.parse_args(argv)
    if not args.input and not args.kaldir:
        parser.error("--input ya da --kaldir gerekli")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    from logging_setup import configure_logging

    configure_logging()
    import database

    if args.kaldir:
        kaldirilan = havuz_satirlarini_kaldir(
            database.SessionLocal, liste_adi=args.liste, adlar=args.kaldir, apply=args.apply)
        print(kaldirma_ozeti(kaldirilan, apply=args.apply))
    if args.input:
        sonuclar = havuzlari_kur(database.SessionLocal, girdi=Path(args.input),
                                 sheet=args.sheet, apply=args.apply)
        print(ozet_metni(sonuclar, apply=args.apply))
    return 0


if __name__ == "__main__":
    sys.exit(main())
