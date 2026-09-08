#!/usr/bin/env python3
"""Föysüz + DERDEST kartların xlsx listesi — SALT OKUNUR (G149, veri ekibine ek).

    docker compose exec -T backend python scripts/foysuz_derdest_kartlar.py \\
        --out /tmp/HUKDOK_DERDEST_KARTLAR_2026-09-08.xlsx

Hiçbir tabloya yazmaz. Ürünü tek bir xlsx'tir; kararı veri ekibi verir.

Neden bu liste
--------------
Veri ekibi (06.09 §1.4) teslim paketinde föyü olmayan ama bizde DERDEST duran
kartların listesini istedi: icra/tahkim/vergi olanları ayırıp kalanın
malpraktis olanlarını kendi listelerine alacaklar. 06.09 ad-hoc ölçümü 370 kart
saymıştı (İcra 195, Hukuk 139); o ölçümü üreten script repoda yoktu. Bu script o
ölçümü tekrarlanabilir kılar. 06.09 "föysüz kartlar plan dışı" kararıyla
çelişmez: liste vermek maliyetsizdir, kapsam kararı ekibindir.

Tanım
-----
`cases` içinde `deleted_at IS NULL`, `status = 'DERDEST'` ve `case_foys`'ta HİÇ
föyü olmayan kartlar. Kapsam işaretli (`kapsam_durumu` dolu) föy de föydür:
kart teslimde bir kez görünmüşse "föysüz" değildir — ekibin listesinde zaten
vardır, bu liste onu bir daha göstermez.

Üç sayfa
--------
1. `MALPRAKTIS_ADAYI` — `file_type` ∉ {İcra, Tahkim, Vergi}. Tür boş olan
   kartlar da buradadır (tür bilinmiyorsa dışlamak, ekibin karar vereceği bir
   kartı gizlemek olurdu).
2. `KAPSAM_DISI_TUR` — İcra/Tahkim/Vergi; ekibin "hiç gelmeyecek" dediği türler
   (bilgi amaçlı, listeye alınmayacak).
3. `OZET` — Ana Tür × sayfa sayıları + toplam.

Tür karşılaştırması Türkçe büyük/küçük harfe dayanıklıdır (`_tur_anahtari`):
"İCRA", "İcra" ve "icra" aynı türdür.
"""
from __future__ import annotations

import argparse
import logging
import os
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Sequence, Tuple
from zoneinfo import ZoneInfo

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

logger = logging.getLogger("FoysuzDerdestKartlar")

DOSYA_ADI_ONEKI = "HUKDOK_DERDEST_KARTLAR"
SAYFA_ADAY = "MALPRAKTIS_ADAYI"
SAYFA_KAPSAM_DISI = "KAPSAM_DISI_TUR"
SAYFA_OZET = "OZET"

# Ekibin "hiç gelmeyecek" dediği türler (06.09 §1.4 / plan P3: 4.437 icra/tahkim + vergi).
KAPSAM_DISI_TURLER = frozenset({"icra", "tahkim", "vergi"})

BASLIKLAR: Tuple[str, ...] = (
    "HukuDok Kart No", "DosyaNo", "Müvekkil(ler)", "Ana Tür", "Yargı Birimi", "Esas",
    "Mahkeme", "Dava Tarihi", "İş Kabul", "Karşı Taraf", "Sorumlu Avukat", "Son Güncelleme",
)
OZET_BASLIKLAR: Tuple[str, ...] = ("Ana Tür", "Sayfa", "Kart")

_TR_SAAT = ZoneInfo("Europe/Istanbul")


def _tur_anahtari(file_type: str | None) -> str:
    """Türü Türkçe harflere dayanıklı küçük harfe indirger ("İCRA" → "icra").

    `str.lower()` "İ"yi "i̇" (i + birleşik nokta) yapar, "I"yı "i" yapar; iki
    yön de Türkçe için yanlıştır. Önce iki büyük harfi elle çevirip sonra lower.
    """
    metin = (file_type or "").strip()
    return metin.replace("İ", "i").replace("I", "ı").lower()


def kapsam_disi_mi(file_type: str | None) -> bool:
    return _tur_anahtari(file_type) in KAPSAM_DISI_TURLER


def _taraf_adlari(kart, party_type: str) -> List[str]:
    return sorted({
        (p.name or "").strip()
        for p in (getattr(kart, "parties", None) or [])
        if (p.party_type or "") == party_type and (p.name or "").strip()
    })


def _excel_zamani(deger):
    """Excel saat dilimi bilmez: tz'li datetime'ı TR saatine çevirip tz'siz yaz.

    Tarih (`date`) olduğu gibi gider; None boş hücre olur.
    """
    if isinstance(deger, datetime):
        if deger.tzinfo is not None:
            deger = deger.astimezone(_TR_SAAT).replace(tzinfo=None)
        return deger.replace(microsecond=0)
    if isinstance(deger, date):
        return deger
    return None


def kart_satiri(kart) -> Tuple[object, ...]:
    """Bir kartın xlsx satırı — `BASLIKLAR` ile aynı sırada."""
    return (
        kart.tracking_no,
        kart.klasor_no_2,
        "; ".join(_taraf_adlari(kart, "CLIENT")),
        kart.file_type,
        kart.judicial_unit,
        kart.esas_no,
        kart.court,
        _excel_zamani(kart.opening_date),
        _excel_zamani(kart.acceptance_date),
        "; ".join(_taraf_adlari(kart, "COUNTER")),
        kart.responsible_lawyer_name,
        _excel_zamani(kart.updated_at),
    )


def foysuz_derdest_kartlar(db) -> list:
    """`deleted_at IS NULL` + `status='DERDEST'` + `case_foys`'ta satırı yok.

    Föy varlığı `EXISTS` ile sorulur: kapsam işaretli föy de föydür (docstring).
    Taraflar tek seferde yüklenir (370 kart için kart başına sorgu atılmaz).
    """
    from sqlalchemy import exists
    from sqlalchemy.orm import selectinload

    import models

    foyu_var = exists().where(models.CaseFoy.case_id == models.Case.id)
    return (
        db.query(models.Case)
        .options(selectinload(models.Case.parties))
        .filter(
            models.Case.deleted_at.is_(None),
            models.Case.status == "DERDEST",
            ~foyu_var,
        )
        .order_by(models.Case.file_type, models.Case.tracking_no)
        .all()
    )


def _sayfa_yaz(wb, ad: str, basliklar: Sequence[str], satirlar: Sequence[Sequence[object]]) -> None:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title=ad)
    ws.append(list(basliklar))
    for hucre in ws[1]:
        hucre.font = Font(bold=True)
    for satir in satirlar:
        ws.append(list(satir))
    ws.freeze_panes = "A2"
    for sutun, baslik in enumerate(basliklar, start=1):
        en_uzun = max([len(str(baslik))] + [len(str(s[sutun - 1] or "")) for s in satirlar])
        ws.column_dimensions[get_column_letter(sutun)].width = min(max(en_uzun + 2, 10), 60)
    for satir_hucreleri in ws.iter_rows(min_row=2):
        for hucre in satir_hucreleri:
            if isinstance(hucre.value, datetime):
                hucre.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(hucre.value, date):
                hucre.number_format = "yyyy-mm-dd"


def varsayilan_dosya_adi(gun: date | None = None) -> str:
    return f"{DOSYA_ADI_ONEKI}_{(gun or date.today()).isoformat()}.xlsx"


def cikti_yolu(out: str | None) -> Path:
    """`--out` dosya yolu ya da dizin olabilir; dizinse varsayılan ad içine düşer."""
    if not out:
        return Path(varsayilan_dosya_adi())
    yol = Path(out)
    if yol.is_dir() or out.endswith(("/", os.sep)):
        return yol / varsayilan_dosya_adi()
    return yol


def raporu_uret(SessionFactory, hedef: Path) -> Dict[str, object]:
    """xlsx'i yazar, özet sözlüğü döndürür. Hiçbir tabloya yazmaz."""
    from openpyxl import Workbook

    db = SessionFactory()
    try:
        kartlar = foysuz_derdest_kartlar(db)
        aday: List[Tuple[object, ...]] = []
        kapsam_disi: List[Tuple[object, ...]] = []
        tur_sayimi: Counter = Counter()
        for kart in kartlar:
            sayfa = SAYFA_KAPSAM_DISI if kapsam_disi_mi(kart.file_type) else SAYFA_ADAY
            (kapsam_disi if sayfa == SAYFA_KAPSAM_DISI else aday).append(kart_satiri(kart))
            tur_sayimi[((kart.file_type or "").strip() or "(boş)", sayfa)] += 1
    finally:
        db.close()

    logger.info(f"föysüz DERDEST kart: {len(kartlar)} (aday {len(aday)} · kapsam dışı tür {len(kapsam_disi)})")

    ozet_satirlari: List[Tuple[object, ...]] = [
        (tur, sayfa, adet)
        for (tur, sayfa), adet in sorted(tur_sayimi.items(), key=lambda p: (p[0][1], -p[1], p[0][0]))
    ]
    ozet_satirlari += [
        ("TOPLAM", SAYFA_ADAY, len(aday)),
        ("TOPLAM", SAYFA_KAPSAM_DISI, len(kapsam_disi)),
        ("TOPLAM", None, len(kartlar)),
    ]

    wb = Workbook()
    wb.remove(wb.active)
    _sayfa_yaz(wb, SAYFA_ADAY, BASLIKLAR, aday)
    _sayfa_yaz(wb, SAYFA_KAPSAM_DISI, BASLIKLAR, kapsam_disi)
    _sayfa_yaz(wb, SAYFA_OZET, OZET_BASLIKLAR, ozet_satirlari)
    hedef.parent.mkdir(parents=True, exist_ok=True)
    wb.save(hedef)
    wb.close()

    return {
        "toplam": len(kartlar),
        "malpraktis_adayi": len(aday),
        "kapsam_disi_tur": len(kapsam_disi),
        "turler": {f"{tur} [{sayfa}]": adet for (tur, sayfa), adet in tur_sayimi.items()},
        "dosya": hedef,
    }


def ozet_metni(ozet: Dict[str, object]) -> str:
    satirlar = [
        "=" * 78,
        "Föysüz + DERDEST kart listesi (SALT OKUNUR)",
        "=" * 78,
        f"  föysüz DERDEST kart : {ozet['toplam']}",
        f"  {SAYFA_ADAY:20s}: {ozet['malpraktis_adayi']}",
        f"  {SAYFA_KAPSAM_DISI:20s}: {ozet['kapsam_disi_tur']}",
    ]
    turler = ozet["turler"]
    assert isinstance(turler, dict)
    for etiket, adet in sorted(turler.items(), key=lambda p: -p[1]):
        satirlar.append(f"    {etiket:40s}: {adet}")
    satirlar.append(f"  dosya               : {ozet['dosya']}")
    satirlar.append("=" * 78)
    return "\n".join(satirlar)


def main(argv: "list[str] | None" = None) -> int:
    parser = argparse.ArgumentParser(
        description="Föysüz + DERDEST kartları xlsx'e çıkarır (hiçbir tabloya yazmaz)",
    )
    parser.add_argument("--out", default=None,
                        help=f"Çıktı dosyası ya da dizini (varsayılan ./{DOSYA_ADI_ONEKI}_<tarih>.xlsx)")
    args = parser.parse_args(argv)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    from logging_setup import configure_logging

    configure_logging()

    import database

    ozet = raporu_uret(database.SessionLocal, cikti_yolu(args.out))
    print(ozet_metni(ozet))
    return 0


if __name__ == "__main__":
    sys.exit(main())
