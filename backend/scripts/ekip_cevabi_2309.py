#!/usr/bin/env python3
"""Ekibin 23.09.2026 cevabı (22.09 çelişki mailimize) — kart düzeni ve Aktif/Arşiv kuralı.

Kaynak: `HUKDOK_CELISKI_CEVABI_2026-09-23.xlsx` (Ek-1) + `MASTER_DOSYA_DURUM_RAPORU_2026-09-20.xlsx`
(föy "Durum" sütunu). Son durum tazelemesi AYRI script'tir ve bundan SONRA koşar
(`scripts/kolayofis_son_durum.py --ek …`): ayrılan/taşınan föyün değeri yeni kartına düşsün.

Adımlar (tek transaction; `--apply` yoksa geri alınır):

1. **Sönen kartı canlandırma (#3418, Ek-1 › BÖLÜNMELİ):** H-961 (İstanbul 7. Tüketici
   2021/390) ile H-16815 (aynı mahkeme 2025/618, bekletici mesele) iki ayrı dava; ikincisi
   11.09 TKU birleştirmesiyle (`tku_kart_birlestir`) #3418'e katılmıştı — föyünde
   `onceki_tracking_no` izi var, bu yüzden `birlesik_kart_ayir` RET verir. Doğru geri dönüş
   birleştirmeyi o föy için GERİ ALMAKTIR: silinmiş eski kart (`tracking_no` = föyün izi)
   canlandırılır, föy oraya geçer, iki kart `ILGILI` bağıyla ilişkilenir, #3418'in esası
   kalan föyün esasına döner. Belgeler kalan kartta kalır (belge koruma şartı).
2. **Föy taşıma (#14148, Ek-1 › BÖLÜNMELİ):** H-3844 (İstanbul Anadolu 10. Asliye Ticaret
   2025/243) H-10597 ile tesadüfen aynı esası (2020/611) paylaşıyordu; davası H-3845'in
   kartında — H-3844 oraya taşınır.
   Her iki adımda: taraflar ham satırdan yalnız-ekleme (`_taraflari_yaz`), föy ↔ müvekkil
   bağı, `klasor_no_2`'de föyün DosyaNo'su kaynaktan hedefe geçer, kaynakta yalnız taşınan
   föye özgü ve bağsız taraflar silinir (`ekip_cevabi_2209` deseni), tarihçe iki kartta.
   Kartlar id ile değil FÖY anahtarıyla bulunur (prod/lokal kart id'leri farklı olabilir).
3. **Aktif/Arşiv kuralı (ekip 23.09 §6):** föylerden en az biri Aktif ise kart DERDEST,
   hepsi Arşiv ise MAHZEN. Föy durumu Ek-1 FOYLER › "Durum" (yeni), yoksa master › "Durum".
   DANIŞ karta dokunulmaz; kullanıcı imzalı `status` tarihçesi olan kart `KORUNDU`
   (script/aktarım/migrasyon imzası dışındaki her imza kullanıcıdır).

    docker compose exec -T backend python scripts/ekip_cevabi_2309.py \\
        --rapor /tmp/ek/MASTER_DOSYA_DURUM_RAPORU_2026-09-20.xlsx \\
        --ek /tmp/ek/HUKDOK_CELISKI_CEVABI_2026-09-23.xlsx                   # kuru koşu
    docker compose exec -T backend python scripts/ekip_cevabi_2309.py --rapor … --ek … --apply --kim ilke

İkinci koşu 0 değişiklik: föyler hedef kartta olduğundan ATLANDI, status eşitse ATLANDI.
"""
from __future__ import annotations

import argparse
import logging
import os
import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Set, Tuple

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import models
from managers import case_manager
from party_check import normalize_party_key
from scripts import birlesik_kart_ayir as bka
from scripts import hukdok_aktarim as ha
from scripts.ekip_cevabi_1209 import Sonuc
from scripts.ekip_cevabi_2209 import _foy_taraf_anahtarlari
from scripts.kolayofis_son_durum import _sayfa_satirlari, _sutun_indeksi

logger = logging.getLogger("EkipCevabi2309")
DEGISTIREN = "ekip_cevabi_2309"
KAYNAK = "ekip cevabı 23.09.2026"
ILISKI_TURU = "ILGILI"

# (kalan föy, sönen kartına dönecek föy, kanıt) — Ek-1 › KARTLAR #3418.
CANLANDIRMALAR: Tuple[Tuple[str, str, str], ...] = (
    ("H-961", "H-16815",
     "iki ayrı dava: İstanbul 7. Tüketici 2021/390 (istinafta) ↔ 2025/618 (bekletici mesele); "
     "ikinci dava birinci karara bağlanmadan açılmış, yeniden yargılama turu değil"),
)
# (taşınacak föy, hedef kartın föyü, kanıt) — Ek-1 › KARTLAR #14148.
TASIMALAR: Tuple[Tuple[str, str, str], ...] = (
    ("H-3844", "H-3845",
     "esas tesadüfen aynı (2020/611): H-3844 İstanbul Anadolu 10. Asliye Ticaret 2025/243 — "
     "davası H-3845/H-9344 kartında"),
)

DURUM_KARTA = {"aktif": "DERDEST", "arşiv": "MAHZEN", "arsiv": "MAHZEN"}
KURAL_STATUSLERI = frozenset({"DERDEST", "MAHZEN"})
# `status` tarihçesinde kullanıcı SAYILMAYAN imzalar (aktarım, migrasyon, ekip script'leri).
SCRIPT_IMZALARI = frozenset({"hukdok_aktarim", "sistem", "birlesik_kart_ayir", "kartsiz_foy_kart_ac"})

ADIM_ADLARI = {"canlandirma": "Sönen kart canlandırma", "tasima": "Föy taşıma",
               "taraf": "Taraf temizliği", "status": "Aktif/Arşiv kuralı"}


def _tarihce(db, case_id: int, alan: str, eski, yeni, kim: str, kanit: str) -> None:
    db.add(models.CaseHistory(
        case_id=case_id, field_name=alan,
        old_value=None if eski is None else str(eski), new_value=None if yeni is None else str(yeni),
        changed_by=DEGISTIREN, source=f"{KAYNAK} ({kim}): {kanit}"[:300],
    ))


def _foy(db, sistem_no: str) -> Optional[models.CaseFoy]:
    return db.query(models.CaseFoy).filter(models.CaseFoy.sistem_no == sistem_no).one_or_none()


def _canli_kart(db, foy: Optional[models.CaseFoy]) -> Optional[models.Case]:
    if foy is None:
        return None
    kart = db.get(models.Case, foy.case_id)
    return kart if kart is not None and kart.deleted_at is None else None


def _klasor_listesi(deger: Optional[str]) -> List[str]:
    return [p.strip() for p in (deger or "").split(";") if p.strip()]


# ─── Föyü karta geçirme (adım 1 ve 2 ortak) ──────────────────────────────────

def _foy_gecir(db, foy: models.CaseFoy, hedef: models.Case, *, kim: str, kanit: str) -> models.Case:
    """Föyü hedef karta geçirir; kaynak kartı döndürür."""
    kaynak = db.get(models.Case, foy.case_id)
    satir = bka.foy_satiri(foy)
    source = f"{DEGISTIREN} ({kim})"
    if satir is not None:
        ha._taraflari_yaz(db, hedef, satir, source)
        db.flush()
    foy.case_id = hedef.id
    foy.case_party_id = None
    db.flush()
    if satir is not None:
        ha._foy_muvekkilini_bagla(db, hedef, satir, sistem_no=foy.sistem_no, source=source)
        dosya_no = ha._metin(satir.degerler.get("dosya_no"))
        if dosya_no:
            kalan = [p for p in _klasor_listesi(kaynak.klasor_no_2) if p != dosya_no]
            if kalan != _klasor_listesi(kaynak.klasor_no_2):
                yeni_deger = ";".join(kalan) or None
                _tarihce(db, kaynak.id, "klasor_no_2", kaynak.klasor_no_2, yeni_deger, kim, f"{foy.sistem_no} ayrıldı")
                kaynak.klasor_no_2 = yeni_deger
            hedef_liste = _klasor_listesi(hedef.klasor_no_2)
            if dosya_no not in hedef_liste:
                yeni_deger = ";".join(hedef_liste + [dosya_no])
                _tarihce(db, hedef.id, "klasor_no_2", hedef.klasor_no_2, yeni_deger, kim, f"{foy.sistem_no} geldi")
                hedef.klasor_no_2 = yeni_deger
    _tarihce(db, kaynak.id, "foy_tasima", foy.sistem_no, f"→ #{hedef.id} {hedef.tracking_no}", kim, kanit)
    _tarihce(db, hedef.id, "foy_tasima", f"#{kaynak.id} {kaynak.tracking_no}", foy.sistem_no, kim, kanit)
    db.flush()
    return kaynak


def taraflari_temizle(db, kart: models.Case, giden: Sequence[str], *, kim: str, sonuc: Sonuc) -> None:
    """Kaynak kartta YALNIZ taşınan föye özgü, belge/föy bağı olmayan tarafları siler."""
    kalan_foyler = db.query(models.CaseFoy).filter(models.CaseFoy.case_id == kart.id).all()
    giden_foyler = db.query(models.CaseFoy).filter(models.CaseFoy.sistem_no.in_(list(giden))).all()
    ozgu = _foy_taraf_anahtarlari(giden_foyler) - _foy_taraf_anahtarlari(kalan_foyler)
    if not ozgu:
        sonuc.ekle("taraf", f"#{kart.id}", "ATLANDI", "taşınan föye özgü taraf yok")
        return
    for taraf in db.query(models.CaseParty).filter(models.CaseParty.case_id == kart.id).all():
        if normalize_party_key(taraf.name) not in ozgu:
            continue
        hedef = f"#{kart.id} {taraf.name} ({taraf.role})"
        belge = db.query(models.CaseDocument).filter(models.CaseDocument.case_party_id == taraf.id).count()
        foy_bagi = db.query(models.CaseFoy).filter(models.CaseFoy.case_party_id == taraf.id).count()
        if belge or foy_bagi:
            sonuc.ekle("taraf", hedef, "ATLANDI", f"bağlı: {belge} belge, {foy_bagi} föy — dokunulmadı")
            continue
        _tarihce(db, kart.id, "taraf", f"{taraf.name} ({taraf.role})", None, kim,
                 f"taşınan föyün tarafı: {', '.join(giden)}")
        db.delete(taraf)
        sonuc.ekle("taraf", hedef, "YAPILDI", "kaynak karttan silindi (hedef kartta yazılı)")
    db.flush()


def _iliski_kur(db, a: models.Case, b: models.Case, kim: str, not_: str) -> None:
    if not bka._iliski_var(db, a.id, b.id):
        db.add(models.CaseRelation(source_case_id=a.id, target_case_id=b.id, relation_type=ILISKI_TURU,
                                   created_by=kim, note=f"{KAYNAK}: {not_}"[:500]))
        db.flush()


def _esasi_foye_cek(db, kart: models.Case, foy: models.CaseFoy, kim: str) -> None:
    satir = bka.foy_satiri(foy)
    esas = ha._metin(satir.degerler.get("esas")) if satir is not None else None
    if esas and kart.esas_no != esas:
        case_manager.sync_current_esas(db, kart, esas, court=kart.court, source=f"{DEGISTIREN} ({kim})")


# ─── Adım 1: sönen kartı canlandırma ─────────────────────────────────────────

def kartlari_canlandir(db, kalemler: Sequence[Tuple[str, str, str]], *, kim: str, sonuc: Sonuc) -> None:
    for kalan_sn, giden_sn, kanit in kalemler:
        hedef_ad = f"{kalan_sn} | {giden_sn}"
        kalan_foy, giden_foy = _foy(db, kalan_sn), _foy(db, giden_sn)
        kart = _canli_kart(db, kalan_foy)
        if kart is None or giden_foy is None:
            sonuc.ekle("canlandirma", hedef_ad, "RET", "föy bizde yok ya da kartı silinmiş")
            continue
        if giden_foy.case_id != kart.id:
            sonuc.ekle("canlandirma", hedef_ad, "ATLANDI", f"föyler zaten ayrı kartlarda (#{giden_foy.case_id})")
            continue
        iz = giden_foy.onceki_tracking_no
        sonen = (db.query(models.Case).filter(models.Case.tracking_no == iz, models.Case.deleted_at.isnot(None))
                 .one_or_none()) if iz else None
        if sonen is None:
            sonuc.ekle("canlandirma", hedef_ad, "RET", f"sönen kart bulunamadı (iz: {iz or '—'}) — elle karar")
            continue
        silme_notu = sonen.delete_reason
        sonen.deleted_at = None
        sonen.deleted_by = None
        sonen.delete_reason = None
        sonen.active = True
        _tarihce(db, sonen.id, "kapatma", silme_notu, None, kim, f"canlandırıldı: {kanit}")
        db.flush()
        _foy_gecir(db, giden_foy, sonen, kim=kim, kanit=kanit)
        _tarihce(db, sonen.id, "onceki_tracking_no", iz, None, kim, "föy kendi kartına döndü")
        giden_foy.onceki_tracking_no = None
        _esasi_foye_cek(db, sonen, giden_foy, kim)
        _esasi_foye_cek(db, kart, kalan_foy, kim)
        _iliski_kur(db, kart, sonen, kim, "ayrı dava (TKU birleştirmesi geri alındı)")
        case_manager.refresh_missing_required(db, kart)
        case_manager.refresh_missing_required(db, sonen)
        sonuc.ekle("canlandirma", f"#{kart.id} {kart.tracking_no}", "YAPILDI",
                   f"{giden_sn} → canlanan #{sonen.id} {sonen.tracking_no} — {kanit}")
        taraflari_temizle(db, kart, [giden_sn], kim=kim, sonuc=sonuc)


# ─── Adım 2: föy taşıma ──────────────────────────────────────────────────────

def foyleri_tasi(db, kalemler: Sequence[Tuple[str, str, str]], *, kim: str, sonuc: Sonuc) -> None:
    for giden_sn, hedef_sn, kanit in kalemler:
        hedef_ad = f"{giden_sn} → {hedef_sn}'in kartı"
        giden_foy = _foy(db, giden_sn)
        kaynak, hedef = _canli_kart(db, giden_foy), _canli_kart(db, _foy(db, hedef_sn))
        if giden_foy is None or kaynak is None or hedef is None:
            sonuc.ekle("tasima", hedef_ad, "RET", "föy bizde yok ya da kartı silinmiş")
            continue
        if kaynak.id == hedef.id:
            sonuc.ekle("tasima", hedef_ad, "ATLANDI", f"zaten #{hedef.id}'de")
            continue
        _foy_gecir(db, giden_foy, hedef, kim=kim, kanit=kanit)
        case_manager.refresh_missing_required(db, kaynak)
        case_manager.refresh_missing_required(db, hedef)
        sonuc.ekle("tasima", hedef_ad, "YAPILDI", f"#{kaynak.id} → #{hedef.id} {hedef.tracking_no} — {kanit}")
        taraflari_temizle(db, kaynak, [giden_sn], kim=kim, sonuc=sonuc)


# ─── Adım 3: Aktif/Arşiv kuralı ──────────────────────────────────────────────

def durumlari_oku(rapor: Optional[Path], ek: Optional[Path]) -> Dict[str, str]:
    """Föy → Durum (Aktif/Arşiv); ek (Ek-1 FOYLER) master'ı ezer."""
    import openpyxl

    durumlar: Dict[str, str] = {}
    if rapor is not None:
        wb = openpyxl.load_workbook(rapor, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            basliklar = [str(h) if h is not None else "" for h in next(ws.iter_rows(max_row=1, values_only=True))]
            if _sutun_indeksi(basliklar, ("Durum",)) is None:
                raise ValueError(f"Raporda Durum sütunu yok — başlıklar: {basliklar}")
            for sistem_no, durum in _sayfa_satirlari(wb, ws.title, ("SistemNo", "Sistem No"), ("Durum",)):
                if ha._metin(sistem_no) and ha._metin(durum):
                    durumlar[ha._metin(sistem_no)] = ha._metin(durum)
        finally:
            wb.close()
    if ek is not None:
        wb = openpyxl.load_workbook(ek, read_only=True, data_only=True)
        try:
            for sistem_no, durum in _sayfa_satirlari(wb, "FOYLER", ("SistemNo",), ("Durum",)):
                if ha._metin(sistem_no) and ha._metin(durum):
                    durumlar[ha._metin(sistem_no)] = ha._metin(durum)
        finally:
            wb.close()
    return durumlar


def _kullanici_statusu_var(db, case_id: int) -> Optional[str]:
    for (changed_by,) in (db.query(models.CaseHistory.changed_by)
                          .filter(models.CaseHistory.case_id == case_id,
                                  models.CaseHistory.field_name == "status").all()):
        imza = (changed_by or "").strip()
        if imza in SCRIPT_IMZALARI or imza.startswith("ekip_cevabi") or imza == DEGISTIREN:
            continue
        return imza or "(imzasız)"
    return None


def statuslari_uygula(db, durumlar: Dict[str, str], *, kim: str, sonuc: Sonuc) -> None:
    kart_durumlari: Dict[int, Set[str]] = defaultdict(set)
    parca = 500
    sistem_nolar = list(durumlar)
    for i in range(0, len(sistem_nolar), parca):
        for sistem_no, case_id in (db.query(models.CaseFoy.sistem_no, models.CaseFoy.case_id)
                                   .filter(models.CaseFoy.sistem_no.in_(sistem_nolar[i:i + parca])).all()):
            hedef = DURUM_KARTA.get(durumlar[sistem_no].casefold())
            if hedef:
                kart_durumlari[case_id].add(hedef)
    for case_id, hedefler in sorted(kart_durumlari.items()):
        kart = db.get(models.Case, case_id)
        if kart is None or kart.deleted_at is not None or kart.status not in KURAL_STATUSLERI:
            continue
        yeni = "DERDEST" if "DERDEST" in hedefler else "MAHZEN"
        if kart.status == yeni:
            continue
        hedef_ad = f"#{kart.id} {kart.tracking_no}"
        kullanici = _kullanici_statusu_var(db, kart.id)
        if kullanici:
            sonuc.ekle("status", hedef_ad, "KORUNDU", f"{kart.status} (kural {yeni}) — elle girilmiş status: {kullanici}")
            continue
        aciklama = f"{kart.status} → {yeni} (föyler: {', '.join(sorted(hedefler))})"
        _tarihce(db, kart.id, "status", kart.status, yeni, kim,
                 "§6 kuralı: föylerden biri Aktif ise kart Aktif, hepsi Arşiv ise Arşiv")
        kart.status = yeni
        sonuc.ekle("status", hedef_ad, "YAPILDI", aciklama)
    db.flush()


# ─── Koşu ────────────────────────────────────────────────────────────────────

def kos(session_factory, *, rapor: Optional[Path] = None, ek: Optional[Path] = None, apply: bool = False,
        kim: str = DEGISTIREN) -> Sonuc:
    durumlar = durumlari_oku(rapor, ek)
    sonuc = Sonuc()
    db = session_factory()
    try:
        kartlari_canlandir(db, CANLANDIRMALAR, kim=kim, sonuc=sonuc)
        foyleri_tasi(db, TASIMALAR, kim=kim, sonuc=sonuc)
        statuslari_uygula(db, durumlar, kim=kim, sonuc=sonuc)
        if apply:
            db.commit()
        else:
            db.rollback()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
    return sonuc


def ozet_metni(sonuc: Sonuc, *, apply: bool) -> str:
    satirlar = ["=" * 78, f"Ekip cevabı 23.09 — {'UYGULANDI' if apply else 'KURU KOŞU'}", "=" * 78]
    for adim, ad in ADIM_ADLARI.items():
        y, a, r, k = (sonuc.sayim(adim, s) for s in ("YAPILDI", "ATLANDI", "RET", "KORUNDU"))
        satirlar.append(f"  {ad:22} {y:3} yapıldı · {a:3} atlandı · {r:3} ret · {k:3} korundu")
    satirlar.append("  " + "-" * 74)
    satirlar.extend(f"  {k.sonuc:7} [{k.adim}] {k.hedef}: {k.aciklama}" for k in sonuc.kalemler)
    satirlar.append("=" * 78)
    return "\n".join(satirlar)


def main(argv: Optional[Sequence[str]] = None) -> int:
    ayristirici = argparse.ArgumentParser(description="Ekip cevabı 23.09: kart düzeni + Aktif/Arşiv kuralı")
    ayristirici.add_argument("--rapor", type=Path, default=None, help="MASTER_DOSYA_DURUM_RAPORU_<tarih>.xlsx")
    ayristirici.add_argument("--ek", type=Path, default=None, help="HUKDOK_CELISKI_CEVABI_<tarih>.xlsx")
    ayristirici.add_argument("--apply", action="store_true", help="yazar (yoksa kuru koşu)")
    ayristirici.add_argument("--kim", default=DEGISTIREN, help="tarihçe imzası")
    args = ayristirici.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    from database import SessionLocal
    from logging_setup import configure_logging
    configure_logging()

    sonuc = kos(SessionLocal, rapor=args.rapor, ek=args.ek, apply=args.apply, kim=args.kim)
    print(ozet_metni(sonuc, apply=args.apply))
    return 0


if __name__ == "__main__":
    sys.exit(main())
