#!/usr/bin/env python3
"""Ekibin CEVAPLI mükerrer-kart xlsx'inden aktarım kart haritası üretir (G154).

`HUKDOK_MUKERRER_VE_YENI_KARTLAR_<tarih>_CEVAPLI.xlsx` dosyasının
`MUKERRER_ESLESME` sayfasında her grup bir "FÖY (paket)" satırı + o föyün
düşebileceği "KART (HukuDok)" satırlarından oluşur; ekip föy satırının
`CEVABINIZ (hangi kart / not)` hücresini doldurur. Bu script SALT OKUNUR:
xlsx → CSV (`sistem_no,tracking_no,kaynak_not`), aktarım o CSV'yi
`hukdok_aktarim.py --kart-esleme` ile alır ve föyü verilen karta bağlar.

    docker compose exec -T backend python scripts/cevapli_kart_eslemesi.py \\
        --input /tmp/CEVAPLI.xlsx --output /tmp/kart_eslemesi.csv

Tanınan cevap desenleri (06.09 cevabındaki üç biçim):

* **Kart no** (`S3.AXA........2915.IDARE.00000`, `…HUKUK.00000-2`): ofis dosya
  numarası biçimi (`retag_tracking_nos.generate_tracking_number`: kategori +
  10 karakterlik isim bloğu + 4 haneli sıra + tür + hizmet + isteğe bağlı
  `-N` eki). Hücrede not da olabilir; ilk kart no alınır.
* **`MÜVEKKİL: <ad>`** (öneri yoktu, ekip müvekkili söyledi): grubun KART
  satırları arasında CLIENT anahtarı ada uyan TEK kart varsa onun numarası.
  Sigorta adı marka sözcükleriyle karşılaştırılır (G153: "Quıck Sigorta A.ş"
  = "Quick Sigorta A.Ş."); ikiz hekim kartı sigortayı ORTAK müvekkil olarak
  da taşıdığından önce müvekkili YALNIZ o ad olan kart aranır (`_kokun_karti_mi`
  ilkesi), yoksa adı içeren kart. Tek kart yoksa `tracking_no` boş + uyarı.
* **`BAĞLAMAYIN`** (H-6589: "önce bizde düzeltilecek"): CSV'ye YAZILMAZ, özette
  raporlanır.

Boş ya da tanınmayan cevap: satır CSV'ye boş `tracking_no` ile düşer (insan
gözden geçirsin diye), aktarım boş satırı WARNING ile atlar. Gerçek cevap
dosyası repoya girmez (A.2); testler sentetik xlsx üretir.
"""
from __future__ import annotations

import argparse
import csv
import logging
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Set

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from party_check import normalize_party_key
from scripts import hukdok_aktarim as ha

logger = logging.getLogger("CevapliKartEslemesi")

SAYFA = "MUKERRER_ESLESME"

# Sütun adları (rapor üreticimizin başlıkları; karşılaştırma `_baslik_anahtari`
# ile aksan/boşluk duyarsız). Aday listesi: ilk eşleşen kullanılır.
SUTUN_ADAYLARI: Dict[str, Sequence[str]] = {
    "grup": ("Grup",),
    "satir_turu": ("Satır türü", "Satir turu"),
    "kimlik": ("SistemNo / Kart No", "SistemNo"),
    "muvekkil": ("Müvekkil",),
    "oneri": ("HukuDok öneri kart", "Öneri kart"),
    "cevap": ("CEVABINIZ (hangi kart / not)", "CEVABINIZ"),
}
ZORUNLU_SUTUNLAR = ("satir_turu", "kimlik", "cevap")

FOY_TURU_ONEKI = "FOY"          # "FÖY (paket)" → anahtar "FOYPAKET"
KART_TURU_ONEKI = "KART"        # "KART (HukuDok)"

# Ofis dosya numarası: `S3.AXA........2915.IDARE.00000`, `D1.A_VARAN....0002.IDARE.00000`,
# `X1.A_HEM......0001.HUKUK.00000`, `S3.AXA........2754.HUKUK.00000-2`.
TRACKING_NO_DESENI = re.compile(r"[A-Z]{1,2}\d\.[A-Z0-9_]+\.+\d{4}\.[A-Z]+\.\d{5}(?:-\d+)?")
BAGLAMAYIN_ANAHTARI = "BAGLAMAYIN"
MUVEKKIL_ANAHTARI = "MUVEKKIL"

DURUM_KART = "kart"
DURUM_MUVEKKIL = "muvekkil"
DURUM_BAGLAMAYIN = "baglamayin"
DURUM_COZULEMEDI = "cozulemedi"


@dataclass
class KartSatiri:
    tracking_no: str
    muvekkil_anahtarlari: Set[str] = field(default_factory=set)


@dataclass
class EslemeSatiri:
    sistem_no: str
    tracking_no: str                  # boş = çözülemedi (CSV'ye boş düşer)
    kaynak_not: str
    durum: str                        # kart | muvekkil | baglamayin | cozulemedi
    grup: str = ""

    @property
    def csv_satiri(self) -> Sequence[str]:
        return (self.sistem_no, self.tracking_no, self.kaynak_not)


# ─── xlsx okuma ─────────────────────────────────────────────────────────────

def _sutun_indeksleri(baslik: Sequence[Any]) -> Dict[str, int]:
    anahtarlar = {ha._baslik_anahtari(b): i for i, b in enumerate(baslik) if b is not None}
    bulunan: Dict[str, int] = {}
    for alan, adaylar in SUTUN_ADAYLARI.items():
        for aday in adaylar:
            indeks = anahtarlar.get(ha._baslik_anahtari(aday))
            if indeks is not None:
                bulunan[alan] = indeks
                break
    eksik = [alan for alan in ZORUNLU_SUTUNLAR if alan not in bulunan]
    if eksik:
        raise ha.AktarimHatasi(
            f"'{SAYFA}' sayfasında zorunlu sütun yok: {', '.join(eksik)} "
            f"(bulunan başlıklar: {', '.join(str(b) for b in baslik if b)})"
        )
    return bulunan


def _hucre(satir: Sequence[Any], indeksler: Dict[str, int], alan: str) -> Optional[str]:
    indeks = indeksler.get(alan)
    if indeks is None or indeks >= len(satir):
        return None
    return ha._metin(satir[indeks])


def cevaplari_oku(yol: Path, *, sheet: str = SAYFA) -> List[EslemeSatiri]:
    """xlsx'i okur, her FÖY satırı için bir `EslemeSatiri` üretir (grup sırası)."""
    from openpyxl import load_workbook

    yol = Path(yol)
    if not yol.exists():
        raise ha.AktarimHatasi(f"Cevap dosyası yok: {yol}")
    wb = load_workbook(yol, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ha.AktarimHatasi(
                f"'{sheet}' sayfası yok (sayfalar: {', '.join(wb.sheetnames)})"
            )
        satirlar = list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()
    if not satirlar:
        raise ha.AktarimHatasi(f"'{sheet}' sayfası boş")
    indeksler = _sutun_indeksleri(satirlar[0])

    # Gruplar: FÖY satırı + onu izleyen KART satırları (Grup sütunu varsa onunla,
    # yoksa "bir sonraki FÖY satırına kadar").
    foyler: List[EslemeSatiri] = []
    ham_cevaplar: List[Dict[str, Any]] = []
    for satir in satirlar[1:]:
        tur = ha._baslik_anahtari(_hucre(satir, indeksler, "satir_turu"))
        kimlik = _hucre(satir, indeksler, "kimlik")
        grup = _hucre(satir, indeksler, "grup") or ""
        if not kimlik:
            continue
        if tur.startswith(FOY_TURU_ONEKI):
            ham_cevaplar.append({
                "sistem_no": kimlik, "grup": grup,
                "cevap": _hucre(satir, indeksler, "cevap"),
                "oneri": _hucre(satir, indeksler, "oneri"),
                "kartlar": [],
            })
        elif tur.startswith(KART_TURU_ONEKI) and ham_cevaplar:
            hedef = ham_cevaplar[-1]
            if grup and hedef["grup"] and grup != hedef["grup"]:
                continue              # gruplar karışmış: kart satırı bu föyün değil
            hedef["kartlar"].append(KartSatiri(
                tracking_no=kimlik,
                muvekkil_anahtarlari={
                    normalize_party_key(ad)
                    for ad in ha._taraf_adlari(_hucre(satir, indeksler, "muvekkil"))
                    if normalize_party_key(ad)
                },
            ))
    for ham in ham_cevaplar:
        foyler.append(cevabi_coz(
            ham["sistem_no"], ham["cevap"], oneri=ham["oneri"],
            kartlar=ham["kartlar"], grup=ham["grup"],
        ))
    return foyler


# ─── cevap ayrıştırma ───────────────────────────────────────────────────────

def _muvekkil_adi(cevap: str) -> Optional[str]:
    """`MÜVEKKİL: <ad>` deseni → ad; değilse None."""
    if ":" not in cevap:
        return None
    etiket, ad = cevap.split(":", 1)
    if ha._baslik_anahtari(etiket) != MUVEKKIL_ANAHTARI:
        return None
    ad = " ".join(ad.split())
    return ad or None


def _ad_uyuyor(ad_anahtari: str, kart_anahtari: str) -> bool:
    """Tam anahtar eşitliği ya da (ikisi de sigortaysa) marka kapsaması."""
    if ad_anahtari == kart_anahtari:
        return True
    ad_markasi = ha._sigorta_markasi(ad_anahtari)
    kart_markasi = ha._sigorta_markasi(kart_anahtari)
    return ad_markasi is not None and kart_markasi is not None and ad_markasi <= kart_markasi


def muvekkille_kart_sec(ad: str, kartlar: Sequence[KartSatiri]) -> List[str]:
    """Ada uyan kartlar; tek elemanlıysa çözüm. Önce müvekkili YALNIZ bu ad
    olan kartlar (sigortanın kendi kartı; ikiz hekim kartı sigortayı ortak
    müvekkil olarak da taşır), o küme tek değilse adı içeren kartlar."""
    anahtar = normalize_party_key(ad)
    if not anahtar:
        return []
    kendi = [
        k.tracking_no for k in kartlar
        if k.muvekkil_anahtarlari and all(_ad_uyuyor(anahtar, m) for m in k.muvekkil_anahtarlari)
    ]
    if len(kendi) == 1:
        return kendi
    iceren = [
        k.tracking_no for k in kartlar
        if any(_ad_uyuyor(anahtar, m) for m in k.muvekkil_anahtarlari)
    ]
    return iceren


def cevabi_coz(sistem_no: str, cevap: Optional[str], *, oneri: Optional[str] = None,
               kartlar: Sequence[KartSatiri] = (), grup: str = "") -> EslemeSatiri:
    """TEK föyün cevabını çözer (saf; xlsx'ten bağımsız — testler doğrudan çağırır)."""
    metin = ha._metin(cevap)
    adaylar = {k.tracking_no for k in kartlar}
    if not metin:
        return EslemeSatiri(sistem_no, "", "cevap yok", DURUM_COZULEMEDI, grup)

    anahtar = ha._baslik_anahtari(metin)
    if BAGLAMAYIN_ANAHTARI in anahtar:
        return EslemeSatiri(sistem_no, "", f"BAĞLAMAYIN: {metin}", DURUM_BAGLAMAYIN, grup)

    kart = TRACKING_NO_DESENI.search(metin)
    if kart is not None:
        tracking_no = kart.group(0)
        notlar = ["cevap: kart no"]
        if oneri and oneri == tracking_no:
            notlar.append("öneriyle aynı")
        elif oneri:
            notlar.append(f"öneri farklıydı: {oneri}")
        else:
            notlar.append("öneri yoktu")
        if adaylar and tracking_no not in adaylar:
            notlar.append("grubun aday kartları dışında")
            logger.warning(
                f"{sistem_no}: cevap {tracking_no!r} grubun aday kartları arasında değil "
                f"({', '.join(sorted(adaylar))}) — DB doğrulaması aktarımda"
            )
        return EslemeSatiri(sistem_no, tracking_no, "; ".join(notlar), DURUM_KART, grup)

    ad = _muvekkil_adi(metin)
    if ad is not None:
        secilen = muvekkille_kart_sec(ad, kartlar)
        if len(secilen) == 1:
            return EslemeSatiri(
                sistem_no, secilen[0], f"MÜVEKKİL: {ad} → tek aday kart", DURUM_MUVEKKIL, grup,
            )
        sebep = "aday yok" if not secilen else f"{len(secilen)} aday: {', '.join(secilen)}"
        logger.warning(f"{sistem_no}: MÜVEKKİL {ad!r} tek karta inmedi ({sebep})")
        return EslemeSatiri(
            sistem_no, "", f"MÜVEKKİL: {ad} → çözülemedi ({sebep})", DURUM_COZULEMEDI, grup,
        )

    logger.warning(f"{sistem_no}: tanınmayan cevap {metin!r}")
    return EslemeSatiri(sistem_no, "", f"tanınmayan cevap: {metin}", DURUM_COZULEMEDI, grup)


# ─── CSV + özet ─────────────────────────────────────────────────────────────

def csv_yaz(satirlar: Sequence[EslemeSatiri], yol: Path) -> Path:
    """BAĞLAMAYIN dışındaki satırlar; UTF-8 BOM, ',' ayraç (aktarım sözleşmesi)."""
    yol = Path(yol)
    yol.parent.mkdir(parents=True, exist_ok=True)
    with open(yol, "w", newline="", encoding="utf-8-sig") as dosya:
        yazici = csv.writer(dosya)
        yazici.writerow(ha.KART_ESLEME_BASLIKLARI)
        for satir in satirlar:
            if satir.durum == DURUM_BAGLAMAYIN:
                continue
            yazici.writerow(satir.csv_satiri)
    return yol


def ozet_metni(satirlar: Sequence[EslemeSatiri], yol: Optional[Path]) -> str:
    sayim = {d: sum(1 for s in satirlar if s.durum == d)
             for d in (DURUM_KART, DURUM_MUVEKKIL, DURUM_BAGLAMAYIN, DURUM_COZULEMEDI)}
    metin = [
        "=" * 70,
        "Cevaplı kart eşlemesi (G154)",
        "=" * 70,
        f"  föy satırı         : {len(satirlar)}",
        f"  kart no ile        : {sayim[DURUM_KART]}",
        f"  müvekkil adıyla    : {sayim[DURUM_MUVEKKIL]}",
        f"  BAĞLAMAYIN (atlandı): {sayim[DURUM_BAGLAMAYIN]}",
        f"  çözülemedi (boş)   : {sayim[DURUM_COZULEMEDI]}",
    ]
    for satir in satirlar:
        if satir.durum in (DURUM_BAGLAMAYIN, DURUM_COZULEMEDI):
            metin.append(f"    - {satir.sistem_no}: {satir.kaynak_not}")
    if yol is not None:
        metin.append(f"  csv                : {yol}")
    metin.append("=" * 70)
    return "\n".join(metin)


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="Ekibin CEVAPLI mükerrer-kart xlsx'inden --kart-esleme CSV'si üretir (G154)",
    )
    parser.add_argument("--input", required=True, help="CEVAPLI xlsx")
    parser.add_argument("--output", required=True, help="üretilecek CSV (sistem_no,tracking_no,kaynak_not)")
    parser.add_argument("--sheet", default=SAYFA, help=f"sayfa adı (varsayılan: {SAYFA})")
    args = parser.parse_args(argv)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    from logging_setup import configure_logging

    configure_logging()

    try:
        satirlar = cevaplari_oku(Path(args.input), sheet=args.sheet)
    except ha.AktarimHatasi as exc:
        logger.error(f"Cevap dosyası okunamadı: {exc}")
        return ha.CIKIS_GIRDI
    yol = csv_yaz(satirlar, Path(args.output))
    print(ozet_metni(satirlar, yol))
    return ha.CIKIS_TAMAM


if __name__ == "__main__":
    sys.exit(main())
