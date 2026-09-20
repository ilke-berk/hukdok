#!/usr/bin/env python3
"""Ek-6 (17.09.2026) kart düzeltmeleri — veri ekibinin 17.09 cevabındaki ricalar.

Kaynak: `HUKDOK_CEVAP_EKI_6_2026-09-17.xlsx` (kesim 17.09, dört sayfa). Ekteki bütün
kart numaraları bizim 17.09 canlı listemizden alınmıştır, yani PROD id'leridir — bu
script LOKALDE "koşuldu" sayılmaz (veri-teslim-hatti.md §7 kuralı).

Adımlar (ekin sayfa/satır referansı her kalemin `kanıt`ında):

1. **Kapatma** (›01) — 5546: ekip "bizde karşılığı yok, kapatılmasını onaylıyoruz" dedi.
2. **Alan** (›04) — 5567 Dava Konusu "Alacak" → "Rücuen Alacak" (Ek-4 › 02'deki düzeltme
   birleştirmede kapanan 5295'te kalmıştı); 2553 Yerel Karar Türü "RED" → BOŞ (dosya
   istinaf kaldırma kararıyla yeniden görülüyor, güncel turun satırında yerel karar
   künyesi bulunmamalı; `yerel_karar_durumu` 16.09 turunda boşaltılmıştı).
3. **Mükerrer** (›02) — 15 çift; kalan DAİMA föylü kart, sönen föysüz:
   * A (4): aynı numaranın iki yazımı ("1976.001" ↔ "1.976.001") köprüde ayrı anahtar
     sayıldığı için ikinci kart açılmış. Kök neden aktarımda da kapatıldı
     (`hukdok_aktarim._dosya_no_adaylari` noktasız ikincil anahtar).
   * B (3): föysüz ikinci kart, föylü kartın klasör numaralarından birini tekrar taşıyor.
   * C (6): bir kartta klasör numarası hiç yok; üçünde BELGE var (14357·13, 14341·4,
     3694·1) — kapatmak yerine birleştirme, belgeler kalan karta taşınır.
   * D (2): Başsavcılık 11163 → 14591 (ekip önerdi) ve Bursa 14375 → 14538 (kararı bize
     bıraktılar; aynı müvekkil Quick Sigorta A.Ş., sıfatları farklı yazılmış — müvekkil
     başına kart desenimiz gereği tek kart, 14375'in 2 belgesi taşınır).
   Kapı: mahkeme anahtarı + esas AYNEN kontrol edilir (ekteki yazım farklarını
   `_baslik_anahtari` zaten eritiyor — ölçüldü). Müvekkil kümesi eşitliği bu adımda
   ARANMAZ (`muvekkil_ayrimi=True`): çiftler tek tek adıyla verilmiştir ve ekibin bütün
   sınıflarında kartların müvekkil kapsamı bilerek farklıdır (föysüz kartta eksik,
   föylü kartta "müvekkiller zaten birlikte kayıtlı").
4. **Klasör** (›02) — A sınıfında birleşme iki yazımı ";" ile birleştirir; numara föylü
   karttaki TEK yazıma çekilir (yalnız aynı numaranın varyantı düşer, başka numara
   korunur). Ayrıca 11163'ten gelen "212.001.00" düşürülür: ekip "kayıtlarımızda hiç
   geçmiyor, C-111'in numarası 221.002.00" dedi (eski değer tarihçede kalır).
5. **Ayırma** (›01) — 14334: dört föy dört ayrı karta (müvekkil başına bir arabuluculuk,
   bir dava kartı) — `birlesik_kart_ayir --muvekkil-ayrimi`.
6. **İlişki** (›02 D) — 14571 ↔ 14730: aynı davanın iki müvekkil kartı, ayrı kalır,
   aralarında "ilişkili dosya" bağı (`ILGILI`).

DOKUNULMAYANLAR (ekibin ricası): 4370/ARB-16909 — ekip düzeltilmiş künyeyi ayrıca
bildirecek, ayırma o zaman ve "2.011.00 yeni karta taşınmadan"; 4953 — föy henüz MİCRO
numarası almadı, alan olduğu gibi kalır; 13397 ↔ 14408 (Kocaeli) — işlem gerekmiyor.

`--ek6` verilirse sabit tablolar ekle KARŞILAŞTIRILIR (kart numarası ekte var mı, kalan
föylü/sönen föysüz mü); uyuşmazsa hiçbir şey yazılmadan çıkılır. Beklenen ESKİ değeri
taşımayan kalem RET, zaten düzeltilmiş kalem ATLANDI. Tek transaction; `--apply` yoksa
geri alınır. İkinci koşu 0 değişiklik.

    docker compose exec -T backend python scripts/ekip_cevabi_1709.py \\
        --ek6 /app/data/HUKDOK_CEVAP_EKI_6_2026-09-17.xlsx                     # kuru koşu
    docker compose exec -T backend python scripts/ekip_cevabi_1709.py \\
        --ek6 /app/data/HUKDOK_CEVAP_EKI_6_2026-09-17.xlsx --apply --kim ilke
"""
from __future__ import annotations

import argparse
import logging
import os
import sys
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from sqlalchemy import func

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import models
from managers import case_manager
from scripts import birlesik_kart_ayir as bka
from scripts import hukdok_aktarim as ha
from scripts import mukerrer_kart_birlestir as mb
from scripts.ekip_cevabi_1209 import Sonuc, _canli_kart, _kolon, _sayfa, _tam_sayi

logger = logging.getLogger("EkipCevabi1709")
DEGISTIREN = "ekip_cevabi_1709"
KAYNAK = "ekip cevabı 17.09.2026 (Ek-6)"
ILISKI_TURU = "ILGILI"

EK6_SORU_SAYFASI = "01_SORULARINIZ"
EK6_COK_KART_SAYFASI = "02_AYNI_DAVA_COK_KART"
EK6_DUZELTME_SAYFASI = "04_KUCUK_DUZELTMELER"

# (kart, sebep)
KAPATMALAR: Tuple[Tuple[int, str], ...] = (
    (5546, "›01: ekip onayladı — bizde karşılığı yok; müvekkil, mahkeme, esas ve klasör boş"),
)

# (kart, alan, beklenen eski değer, yeni değer | None, kanıt)
ALAN_DUZELTMELERI: Tuple[Tuple[int, str, str, Optional[str], str], ...] = (
    (5567, "subject", "Alacak", "Rücuen Alacak",
     "›04: Ek-4 › 02 konu düzeltmesi 5295 ile birleştirmede kapanan kartta kalmış"),
    (2553, "karar_turu", "RED", None,
     "›04: istinaf kaldırma kararıyla yeniden görülüyor — güncel turun satırında yerel karar künyesi olmaz"),
)

# Kapalı listeden yazılan alanlar: yeni değer havuzda yoksa YAZILMAZ (havuz farkı doğurmasın)
KAPALI_LISTE_ALANLARI: Dict[str, type] = {"subject": models.CaseSubject}

# (kalan = föylü kart, sönen = föysüz kart, sınıf, kanıt)
BIRLESTIRMELER: Tuple[Tuple[int, int, str, str], ...] = (
    (14482, 4966, "A", "Hatay 4. İdare 2023/907 — klasör 1976.001 ↔ 1.976.001"),
    (14478, 4903, "A", "İstanbul 11. İdare 2025/1426 — klasör 1957.001 ↔ 1.957.001"),
    (14495, 13995, "A", "İzmir Arabuluculuk 2025/8148 — klasör 2.433.01 ↔ 2433.01"),
    (14511, 14330, "A", "Kayseri Arabuluculuk 2025/1334 — klasör 2.473.001 ↔ 2473.001"),
    (791, 792, "B", "İstanbul 4. İdare 2020/1906 — föysüz kart 3.563.00'ı tekrar taşıyor"),
    (746, 745, "B", "Salihli 3. Asliye Hukuk 2020/113 — föysüz kart 3.1400.00'ı tekrar taşıyor"),
    (747, 748, "B", "Van 2. İdare 2020/1164 — föysüz kart 9.1365.00'ı tekrar taşıyor"),
    (941, 687, "C", "Antalya 3. Tüketici 2023/260 — 687 klasörsüz ve belgesiz"),
    (4262, 4052, "C", "Diyarbakır 3. Tüketici 2024/85 — 4052 klasörsüz ve belgesiz"),
    (4441, 3694, "C", "İstanbul 6. Tüketici 2024/249 — 3694 klasörsüz, 1 belgesi taşınır"),
    (14717, 14357, "C", "İstanbul Anadolu 11. Tüketici 2026/254 — 14357 klasörsüz, 13 belgesi taşınır"),
    (14474, 14341, "C", "Şanlıurfa 4. Tüketici 2026/1379 — 14341 klasörsüz, 4 belgesi taşınır"),
    (4141, 3522, "C", "Turhal 2. Asliye Hukuk 2023/122 — 3522 klasörsüz ve belgesiz"),
    (14591, 11163, "D", "İstanbul Anadolu Başsavcılık 2014/170680 — ekip birleştirmeyi önerdi"),
    (14538, 14375, "D", "Bursa 2. İdare 2025/1542 — aynı müvekkil (sıfatlar farklı yazılmış), 2 belge taşınır"),
)

# A sınıfı: birleşmeden sonra numaranın föylü karttaki TEK yazımı kalır
KLASOR_TEKLESTIRME: Tuple[Tuple[int, str, str], ...] = (
    (14482, "1976.001", "›02 A: föylü karttaki yazım"),
    (14478, "1957.001", "›02 A: föylü karttaki yazım"),
    (14495, "2.433.01", "›02 A: föylü karttaki yazım"),
    (14511, "2.473.001", "›02 A: föylü karttaki yazım"),
)

# (kart, düşürülecek parça, kalması beklenen parça, kanıt)
KLASOR_CIKAR: Tuple[Tuple[int, str, str, str], ...] = (
    (14591, "212.001.00", "221.002.00",
     "›02 D: 212.001.00 ekibin hiçbir kaydında geçmiyor; C-111 föyünün numarası 221.002.00"),
)

# (kart, kanıt) — müvekkil ayrımıyla ayrılır
AYIRMALAR: Tuple[Tuple[int, str], ...] = (
    (14334, "›01: dört föy dört ayrı kartta — müvekkil başına bir arabuluculuk, bir dava kartı"),
)

# (a, b, kanıt)
ILISKILER: Tuple[Tuple[int, int, str], ...] = (
    (14571, 14730, "›02 D: Manisa Asliye Ticaret 2026/246 — aynı davanın iki müvekkil kartı"),
)

ADIM_ADLARI = {
    "kapatma": "1. Kapatma (›01)", "alan": "2. Alan düzeltmesi (›04)",
    "cift": "3. Mükerrer kart (›02)", "klasor": "4. Klasör numarası (›02)",
    "ayirma": "5. Ayırma (›01)", "iliski": "6. İlişkili dosya (›02 D)",
}


# ─── Ek-6 doğrulama kapısı ───────────────────────────────────────────────────

def ek6_dogrula(yol: Path) -> List[str]:
    """Sabit tabloları ekle karşılaştırır; uyuşmazlık listesi döner (boş = temiz).

    Ekte kart numarası + "Bizim föyümüz" sütunu var; birleştirmede kalan kartın
    föyü OLMALI, sönenin OLMAMALI (ekin dört sınıfında da desen budur).
    """
    import openpyxl
    wb = openpyxl.load_workbook(yol, data_only=True, read_only=True)
    try:
        satirlar = list(_sayfa(wb, EK6_COK_KART_SAYFASI).iter_rows(values_only=True))
        i_kart = _kolon(satirlar[0], "Kart")
        i_foy = _kolon(satirlar[0], "Bizim föyümüz")
        foylu: Dict[int, bool] = {}
        for satir in satirlar[1:]:
            kart = _tam_sayi(satir[i_kart])
            if kart is not None:
                foylu[kart] = (ha._metin(satir[i_foy]) or "—") not in ("—", "-")
        soru_kartlari = _sayfa_kartlari(wb, EK6_SORU_SAYFASI)
        duzeltme_kartlari = _sayfa_kartlari(wb, EK6_DUZELTME_SAYFASI)
    finally:
        wb.close()

    sorunlar: List[str] = []
    for kalan, sonen, sinif, _kanit in BIRLESTIRMELER:
        for kart, beklenen_foy in ((kalan, True), (sonen, False)):
            if kart not in foylu:
                sorunlar.append(f"{sinif}: kart {kart} Ek-6 › 02'de yok")
            elif foylu[kart] is not beklenen_foy:
                sorunlar.append(f"{sinif}: kart {kart} ekte "
                                f"{'föysüz' if beklenen_foy else 'föylü'} görünüyor — çift ters olabilir")
    for kart, _dogru, _kanit in KLASOR_TEKLESTIRME:
        if kart not in foylu:
            sorunlar.append(f"klasör: kart {kart} Ek-6 › 02'de yok")
    for kart, _cikar, _kalan_parca, _kanit in KLASOR_CIKAR:
        if kart not in foylu:
            sorunlar.append(f"klasör: kart {kart} Ek-6 › 02'de yok")
    for a, b, _kanit in ILISKILER:
        for kart in (a, b):
            if kart not in foylu:
                sorunlar.append(f"ilişki: kart {kart} Ek-6 › 02'de yok")
    for kart, _sebep in KAPATMALAR:
        if kart not in soru_kartlari:
            sorunlar.append(f"kapatma: kart {kart} Ek-6 › 01'de yok")
    for kart, _kanit in AYIRMALAR:
        if kart not in soru_kartlari:
            sorunlar.append(f"ayırma: kart {kart} Ek-6 › 01'de yok")
    for kart, _alan, _eski, _yeni, _kanit in ALAN_DUZELTMELERI:
        if kart not in duzeltme_kartlari:
            sorunlar.append(f"alan: kart {kart} Ek-6 › 04'te yok")
    return sorunlar


def _sayfa_kartlari(wb, ad: str) -> set:
    satirlar = list(_sayfa(wb, ad).iter_rows(values_only=True))
    i_kart = _kolon(satirlar[0], "Kart")
    return {k for k in (_tam_sayi(s[i_kart]) for s in satirlar[1:]) if k is not None}


# ─── Yardımcılar ─────────────────────────────────────────────────────────────

def _tarihce(db, case_id: int, alan: str, eski, yeni, kim: str, kanit: str) -> None:
    db.add(models.CaseHistory(
        case_id=case_id, field_name=alan,
        old_value=None if eski is None else str(eski), new_value=None if yeni is None else str(yeni),
        changed_by=DEGISTIREN, source=f"{KAYNAK} ({kim}): {kanit}"[:300],
    ))


def _parcalar(kart: models.Case) -> List[str]:
    return [p for p in (ha._metin(x) or "" for x in (kart.klasor_no_2 or "").split(";")) if p]


def _ayni_numara(a: str, b: str) -> bool:
    """Aynı numaranın iki yazımı mı (yalnız nokta farkı)?"""
    return ha._noktasiz_anahtar(ha._eslesme_anahtari(a)) == ha._noktasiz_anahtar(ha._eslesme_anahtari(b))


# ─── Adım 1: kapatma ─────────────────────────────────────────────────────────

def kartlari_kapat(db, kalemler: Sequence[Tuple[int, str]], *, kim: str, sonuc: Sonuc) -> None:
    for kart_id, sebep in kalemler:
        hedef = f"#{kart_id}"
        kart = db.get(models.Case, kart_id)
        if kart is None:
            sonuc.ekle("kapatma", hedef, "RET", "kart yok")
            continue
        if kart.deleted_at is not None:
            sonuc.ekle("kapatma", hedef, "ATLANDI", "zaten kapalı")
            continue
        foy = db.query(models.CaseFoy).filter(models.CaseFoy.case_id == kart_id).count()
        if foy or kart.documents:
            sonuc.ekle("kapatma", hedef, "RET", f"{foy} föy / {len(kart.documents)} belge var — insan kararı")
            continue
        kart.deleted_at = func.now()
        kart.deleted_by = kim
        kart.delete_reason = f"{KAYNAK}: {sebep}"[:500]
        kart.active = False
        _tarihce(db, kart_id, "kapatma", kart.tracking_no, None, kim, sebep)
        db.flush()
        case_manager.refresh_missing_required(db, kart)
        sonuc.ekle("kapatma", hedef, "YAPILDI", kart.tracking_no)


# ─── Adım 2: alan düzeltmesi ─────────────────────────────────────────────────

def alanlari_duzelt(db, kalemler, *, kim: str, sonuc: Sonuc) -> None:
    for kart_id, alan, eski_beklenen, yeni, kanit in kalemler:
        hedef = f"#{kart_id} {alan}"
        kart = _canli_kart(db, kart_id)
        if kart is None:
            sonuc.ekle("alan", hedef, "RET", "kart yok/silinmiş")
            continue
        mevcut = getattr(kart, alan)
        if (mevcut or None) == (yeni or None):
            sonuc.ekle("alan", hedef, "ATLANDI", "zaten düzeltilmiş")
            continue
        if (mevcut or "") != eski_beklenen:
            sonuc.ekle("alan", hedef, "RET", f"beklenmeyen değer {mevcut!r} — elle bakılmalı")
            continue
        model = KAPALI_LISTE_ALANLARI.get(alan)
        if yeni is not None and model is not None and \
                db.query(model).filter(model.name == yeni).count() == 0:
            sonuc.ekle("alan", hedef, "RET", f"{yeni!r} kapalı listede yok — havuza önce eklenmeli")
            continue
        setattr(kart, alan, yeni)
        _tarihce(db, kart_id, alan, mevcut, yeni, kim, kanit)
        db.flush()
        case_manager.refresh_missing_required(db, kart)
        sonuc.ekle("alan", hedef, "YAPILDI", f"{mevcut!r} → {yeni!r}")


# ─── Adım 3: mükerrer kart ───────────────────────────────────────────────────

def ciftleri_birlestir(db, kalemler, *, kim: str, sonuc: Sonuc) -> None:
    for kalan_id, sonen_id, sinif, kanit in kalemler:
        hedef = f"{sinif} #{sonen_id} → #{kalan_id}"
        kalan, sonen = db.get(models.Case, kalan_id), db.get(models.Case, sonen_id)
        if kalan is None or sonen is None:
            sonuc.ekle("cift", hedef, "RET", "kart bulunamadı")
            continue
        if sonen.deleted_at is not None and kalan.deleted_at is None:
            sonuc.ekle("cift", hedef, "ATLANDI", "zaten birleştirilmiş")
            continue
        with db.begin_nested():
            b = mb.birlestir(db, kalan, sonen, kim=kim, muvekkil_ayrimi=True,
                             tarihce_alani="ekip_cevabi_1709_birlestirme",
                             sebep_etiketi=f"{KAYNAK} ›02 {sinif}")
        if b.ret:
            sonuc.ekle("cift", hedef, "RET", f"{b.ret} ({kanit})")
        else:
            sonuc.ekle("cift", hedef, "YAPILDI", f"{b.tasinan} — {kanit}")


# ─── Adım 4: klasör numarası ─────────────────────────────────────────────────

def klasorleri_teklestir(db, kalemler, *, kim: str, sonuc: Sonuc) -> None:
    """Aynı numaranın varyant yazımları düşer; başka numaralar korunur."""
    for kart_id, dogru, kanit in kalemler:
        hedef = f"#{kart_id} ={dogru}"
        kart = _canli_kart(db, kart_id)
        if kart is None:
            sonuc.ekle("klasor", hedef, "RET", "kart yok/silinmiş")
            continue
        parcalar = _parcalar(kart)
        if dogru not in parcalar:
            sonuc.ekle("klasor", hedef, "RET", f"doğru yazım listede yok: {parcalar}")
            continue
        kalanlar = [p for p in parcalar if p == dogru or not _ayni_numara(p, dogru)]
        if kalanlar == parcalar:
            sonuc.ekle("klasor", hedef, "ATLANDI", "varyant yazım yok")
            continue
        eski = kart.klasor_no_2
        kart.klasor_no_2 = ";".join(kalanlar)
        _tarihce(db, kart_id, "klasor_no_2", eski, kart.klasor_no_2, kim, kanit)
        db.flush()
        sonuc.ekle("klasor", hedef, "YAPILDI", f"{eski!r} → {kart.klasor_no_2!r}")


def klasorleri_cikar(db, kalemler, *, kim: str, sonuc: Sonuc) -> None:
    for kart_id, cikar, kalan_parca, kanit in kalemler:
        hedef = f"#{kart_id} -{cikar}"
        kart = _canli_kart(db, kart_id)
        if kart is None:
            sonuc.ekle("klasor", hedef, "RET", "kart yok/silinmiş")
            continue
        parcalar = _parcalar(kart)
        if cikar not in parcalar:
            sonuc.ekle("klasor", hedef, "ATLANDI", "numara listede yok")
            continue
        if kalan_parca not in parcalar:
            sonuc.ekle("klasor", hedef, "RET", f"kalması beklenen {kalan_parca} listede yok: {parcalar}")
            continue
        eski = kart.klasor_no_2
        kart.klasor_no_2 = ";".join(p for p in parcalar if p != cikar)
        _tarihce(db, kart_id, "klasor_no_2", eski, kart.klasor_no_2, kim, kanit)
        db.flush()
        sonuc.ekle("klasor", hedef, "YAPILDI", f"{eski!r} → {kart.klasor_no_2!r}")


# ─── Adım 5: ayırma ──────────────────────────────────────────────────────────

def kartlari_ayir(db, kalemler, *, kim: str, sonuc: Sonuc) -> None:
    kullanilan: Dict[str, int] = {}
    for kart_id, kanit in kalemler:
        hedef = f"#{kart_id}"
        kalem = bka.karti_ayir(db, kart_id, kullanilan, kim=kim, muvekkil_ayrimi=True)
        aciklama = "; ".join(f"→ #{i} {t} ({', '.join(s)})" for i, t, s in kalem.yeni_kartlar)
        sonuc.ekle("ayirma", hedef, kalem.sonuc, f"{aciklama or kalem.aciklama} — {kanit}")


# ─── Adım 6: ilişki ──────────────────────────────────────────────────────────

def iliskileri_kur(db, kalemler, *, kim: str, sonuc: Sonuc) -> None:
    for a_id, b_id, kanit in kalemler:
        hedef = f"#{a_id} ↔ #{b_id}"
        a, b = _canli_kart(db, a_id), _canli_kart(db, b_id)
        if a is None or b is None:
            sonuc.ekle("iliski", hedef, "RET", "kart yok/silinmiş")
            continue
        if bka._iliski_var(db, a_id, b_id):
            sonuc.ekle("iliski", hedef, "ATLANDI", "bağ zaten var")
            continue
        db.add(models.CaseRelation(source_case_id=a_id, target_case_id=b_id, relation_type=ILISKI_TURU,
                                   created_by=kim, note=f"{KAYNAK}: {kanit}"[:500]))
        db.flush()
        sonuc.ekle("iliski", hedef, "YAPILDI", f"{a.tracking_no} ↔ {b.tracking_no}")


# ─── Koşu ────────────────────────────────────────────────────────────────────

def kos(session_factory, *, ek6: Optional[Path] = None, apply: bool = False,
        kim: str = DEGISTIREN) -> Sonuc:
    if ek6 is not None:
        sorunlar = ek6_dogrula(ek6)
        if sorunlar:
            raise ValueError("Ek-6 sabit tablolarla uyuşmuyor — hiçbir şey yazılmadı:\n  "
                             + "\n  ".join(sorunlar))
    sonuc = Sonuc()
    db = session_factory()
    try:
        kartlari_kapat(db, KAPATMALAR, kim=kim, sonuc=sonuc)
        alanlari_duzelt(db, ALAN_DUZELTMELERI, kim=kim, sonuc=sonuc)
        ciftleri_birlestir(db, BIRLESTIRMELER, kim=kim, sonuc=sonuc)
        klasorleri_teklestir(db, KLASOR_TEKLESTIRME, kim=kim, sonuc=sonuc)
        klasorleri_cikar(db, KLASOR_CIKAR, kim=kim, sonuc=sonuc)
        kartlari_ayir(db, AYIRMALAR, kim=kim, sonuc=sonuc)
        iliskileri_kur(db, ILISKILER, kim=kim, sonuc=sonuc)
        if apply:
            db.commit()
        else:
            db.rollback()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
    return sonuc


def ozet_metni(sonuc: Sonuc, *, apply: bool, ayrinti: bool = False) -> str:
    satirlar = ["=" * 78,
                f"Ek-6 (17.09.2026) kart düzeltmeleri — {'UYGULANDI' if apply else 'KURU KOŞU'}",
                "=" * 78]
    for adim, ad in ADIM_ADLARI.items():
        y, a, r = (sonuc.sayim(adim, s) for s in ("YAPILDI", "ATLANDI", "RET"))
        satirlar.append(f"  {ad:26} {y:4} yapıldı · {a:4} atlandı · {r:3} ret")
    gosterilecek = sonuc.kalemler if ayrinti else [k for k in sonuc.kalemler if k.sonuc == "RET"]
    if gosterilecek:
        satirlar.append("  " + "-" * 74)
        satirlar.extend(f"  {k.sonuc:7} [{k.adim}] {k.hedef}: {k.aciklama}" for k in gosterilecek)
    satirlar.append("=" * 78)
    return "\n".join(satirlar)


def main(argv: Optional[Sequence[str]] = None) -> int:
    ayristirici = argparse.ArgumentParser(description="Ek-6 (17.09.2026) kart düzeltmeleri")
    ayristirici.add_argument("--ek6", type=Path, default=None,
                             help="HUKDOK_CEVAP_EKI_6_2026-09-17.xlsx (verilirse tablolar ekle doğrulanır)")
    ayristirici.add_argument("--apply", action="store_true", help="yazar (yoksa kuru koşu)")
    ayristirici.add_argument("--kim", default=DEGISTIREN, help="tarihçe imzası")
    ayristirici.add_argument("--ayrinti", action="store_true", help="bütün kalemleri bas (yalnız RET değil)")
    args = ayristirici.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    from database import SessionLocal
    from logging_setup import configure_logging
    configure_logging()

    sonuc = kos(SessionLocal, ek6=args.ek6, apply=args.apply, kim=args.kim)
    print(ozet_metni(sonuc, apply=args.apply, ayrinti=args.ayrinti))
    return 0


if __name__ == "__main__":
    sys.exit(main())
